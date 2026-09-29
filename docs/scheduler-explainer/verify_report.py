"""Check the deliverable's structure, source provenance and illustrative arithmetic.

Visual review of every rendered page remains a separate required step.
"""
from decimal import Decimal
from hashlib import sha256
from pathlib import Path
import subprocess
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from docx import Document
from build_report import OUTPUT, REFERENCE_SHA256, REVISION, ROOT

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def verify():
    assert sha256((ROOT / "reference.docx").read_bytes()).hexdigest() == REFERENCE_SHA256
    original = ROOT.parents[1] / "artifacts/scheduler-explainer/WaterFlex_Scheduler_and_Route_Optimization_Explained.docx"
    if original.exists():
        assert sha256(original.read_bytes()).hexdigest() == REFERENCE_SHA256
    with ZipFile(ROOT / "reference.docx") as reference, ZipFile(OUTPUT) as final:
        assert set(reference.namelist()) == set(final.namelist())
        editable = {"word/document.xml", "word/_rels/document.xml.rels", "docProps/core.xml"}
        for part in reference.namelist():
            if part not in editable:
                assert reference.read(part) == final.read(part), f"Unexpected change to {part}"
        original_xml = ET.fromstring(reference.read("word/document.xml"))
        final_xml = ET.fromstring(final.read("word/document.xml"))
        assert ET.tostring(original_xml.find(f".//{W}sectPr")) == ET.tostring(final_xml.find(f".//{W}sectPr"))
        for part in final.namelist():
            if part.endswith(".xml"):
                xml = ET.fromstring(final.read(part))
                for text in xml.iter(W + "t"):
                    assert "\u2014" not in (text.text or ""), f"Em dash in {part}"
        anchors = {node.get(W + "name") for node in final_xml.iter(W + "bookmarkStart")}
        for node in final_xml.iter(W + "hyperlink"):
            if node.get(W + "anchor"):
                assert node.get(W + "anchor") in anchors

    doc = Document(OUTPUT)
    reference = Document(ROOT / "reference.docx")
    assert [p.text for p in doc.paragraphs if p.style.name == "Heading 1"] == [
        p.text for p in reference.paragraphs if p.style.name == "Heading 1"]
    assert len(doc.tables) == 15
    assert len(doc.sections) == 1
    project_links = 0
    for rel in doc.part.rels.values():
        if rel.reltype.endswith("/hyperlink") and rel.is_external and "ethan-mcgee" in rel.target_ref:
            assert f"/blob/{REVISION}/" in rel.target_ref
            path = rel.target_ref.split(f"/blob/{REVISION}/")[1]
            subprocess.run(["git", "cat-file", "-e", f"{REVISION}:{path}"], cwd=ROOT, check=True)
            project_links += 1
    assert project_links == 16

    def variance(workloads, capacities):
        paid, capacity = sum(workloads), sum(capacities)
        mean = Decimal(paid) / Decimal(capacity)
        return sum(Decimal(c) * (Decimal(w) / Decimal(c) - mean) ** 2 for w, c in zip(workloads, capacities)) / Decimal(capacity)

    assert round(variance([240, 240], [480, 240]), 5) == Decimal("0.05556")
    assert round(variance([360, 120], [480, 240]), 5) == Decimal("0.01389")
    assert variance([320, 160], [480, 240]) == 0
    assert Decimal("970") * Decimal("1.02") == Decimal("989.40")
    assert 105 + 3 * 17 == 156
    assert 45 + 2 * 17 == 79
    assert 156 + 79 == 235
    assert 815 + 587 + 47 == 3780 - 2331
    assert round(2331 / 3780 * 100, 2) == 61.67
    assert round(670 / 1260 * 100, 2) == 53.17
    assert round(401 / 1260 * 100, 2) == 31.83
    assert Decimal("23200.94") - Decimal("19033.00") == Decimal("4167.94")
    assert round(Decimal("4167.94") / Decimal("23200.94") * 100, 2) == Decimal("17.96")
    assert Decimal("74683.33") - Decimal("74422.56") == Decimal("260.77")
    print("PASS reference preservation, package fidelity, section order, 15 tables, navigation, 16 pinned project links, punctuation and arithmetic")


if __name__ == "__main__":
    verify()
