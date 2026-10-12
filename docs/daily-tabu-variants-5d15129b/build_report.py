"""Build the TABU variant report from validated evidence and the selected explainer.

Install requirements.txt first. No network, services or benchmark execution.
"""
from copy import deepcopy
from datetime import datetime
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from statistics import mean
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from charts import SOLVERS, ALTS, LABEL, FLEETS, BUDGETS, WORKLOADS, WORK_LABEL, select, pair, avg, make_charts
from report_data import load, RUN, ANALYSIS, CONFIRM, CONFIRM_ANALYSIS, SEQUENTIAL, CONTENTION, REVISION

HERE = Path(__file__).resolve().parent
REFERENCE = HERE.parent / "WaterFlex_Scheduler_Algorithm_Comparison_Explained.docx"
REFERENCE_SHA = "5366f3017d6f18c15de295ba5fb9c1a014b6222911040f1c5e044845e57f832d"
OUTPUT = HERE / "Daily_Tabu_Variants_5d15129b_Report.docx"
PRESERVE = ("word/styles.xml", "word/stylesWithEffects.xml", "word/numbering.xml", "word/theme/theme1.xml", "word/fontTable.xml")
SHORT = dict(zip(SOLVERS, ("Tabu", "Size 3", "Size 15", "Tabu+K", "K opt", "Sublist", "Late acc.")))


def dollars(value):
    return f"${value:,.2f}" if value >= 0 else f"-${-value:,.2f}"


def signed(value):
    return f"{value:+.2f}"


def wtl(rows):
    return "/".join(str(sum(test(r["saving"]) for r in rows)) for test in (lambda x: x > 0, lambda x: x == 0, lambda x: x < 0))


def own_budget_gain(rows, solver):
    """Same solver, fixture and seed: 60 second accepted cost minus 240 second accepted cost."""
    short = {r["key"][:3]: r for r in select(rows, solver=solver, budget_ms=60000)}
    return [short[r["key"][:3]]["accepted_cost"] - r["accepted_cost"] for r in select(rows, solver=solver, budget_ms=240000)]


def confirmation_pairs(rows, confirm_rows):
    sequential = {r["key"][:-1]: r["raw"]["after"]["costCents"] for r in confirm_rows if r["solver"] == "TABU"}
    parallel = {r["key"][:-1]: r for r in rows if r["solver"] == "TABU_KOPT"}
    out = []
    for r in confirm_rows:
        if r["solver"] == "TABU_KOPT":
            key = r["key"][:-1]
            out.append({"key": key, "sequential": (sequential[key] - r["raw"]["after"]["costCents"]) / 100,
                        "parallel": parallel[key]["saving"]})
    assert len(out) == 64
    return out


def build():
    assert sha256(REFERENCE.read_bytes()).hexdigest() == REFERENCE_SHA
    config, raw_rows, confirm_rows, contention = load()
    rows = pair(raw_rows)
    assets = HERE / "generated"
    make_charts(rows, confirm_rows, contention, assets)
    confirmed = confirmation_pairs(rows, confirm_rows)
    doc = Document(REFERENCE)
    table_properties = deepcopy(doc.tables[0]._tbl.tblPr)
    for child in list(doc.element.body):
        if child.tag != qn("w:sectPr"): doc.element.body.remove(child)
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype.endswith(("/image", "/hyperlink")): del doc.part.rels[rid]
    doc.core_properties.title = "WaterFlex Tabu variant experiment explained"
    doc.core_properties.subject = "How seven daily search configurations behave and why, run 5d15129b"
    doc.core_properties.author = doc.core_properties.last_modified_by = "WaterFlex"
    doc.core_properties.created = doc.core_properties.modified = datetime(2026, 10, 11)
    figures = []

    def p(text, style=None):
        assert "—" not in text and "–" not in text
        para = doc.add_paragraph(text, style); para.paragraph_format.widow_control = True
        return para

    def heading(text):
        h = p(text, "Heading 1"); h.paragraph_format.page_break_before = True

    def sub(text): p(text, "Heading 2")

    def table(headers, data, widths, size=10):
        assert abs(sum(widths) - 6.9) < .01
        t = doc.add_table(rows=1, cols=len(headers)); t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t._tbl.replace(t._tbl.tblPr, deepcopy(table_properties))
        for col, w in zip(t.columns, widths): col.width = Inches(w)
        for i, values in enumerate([headers, *data]):
            assert len(values) == len(headers)
            row = t.rows[0] if i == 0 else t.add_row()
            trpr = row._tr.get_or_add_trPr(); trpr.append(OxmlElement("w:cantSplit"))
            if i == 0: trpr.append(OxmlElement("w:tblHeader"))
            for j, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
                cell.width = Inches(width); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell.text = str(value)
                props = cell._tc.get_or_add_tcPr()
                shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "DCE6EF" if i == 0 else ("FFFFFF" if i % 2 else "F4F6F8")); props.append(shade)
                margins = OxmlElement("w:tcMar")
                for side in ("top", "left", "bottom", "right"):
                    n = OxmlElement("w:" + side); n.set(qn("w:w"), "65"); n.set(qn("w:type"), "dxa"); margins.append(n)
                props.append(margins)
                para = cell.paragraphs[0]; para.paragraph_format.space_after = Pt(2); para.paragraph_format.space_before = Pt(2); para.paragraph_format.line_spacing = 1.0
                if len(headers) > 2 and j > 0: para.alignment = 1
                if len(data) <= 12 and i < len(data): para.paragraph_format.keep_with_next = True  # Short tables stay on one page.
                for run in para.runs: run.font.name = "Calibri"; run.font.size = Pt(size); run.bold = i == 0
        spacer = p(""); spacer.paragraph_format.space_after = Pt(0); spacer.paragraph_format.line_spacing = Pt(3)

    def figure(name, caption):
        figures.append(name)
        caption = f"Figure {len(figures)}. {caption}"
        para = p(""); shape = para.add_run().add_picture(str(assets / f"{name}.png"), width=Inches(6.85))
        shape._inline.docPr.set("descr", caption); para.paragraph_format.keep_with_next = True
        cap = p(caption, "Caption"); cap.paragraph_format.space_after = Pt(7)
        for run in cap.runs:
            run.font.color.rgb = RGBColor(0, 0, 0); run.bold = False; run.italic = True

    def by(solver, **kw):
        return select(rows, solver=solver, **kw)

    tk, la, ko = by("TABU_KOPT"), by("LATE_ACCEPTANCE"), by("KOPT")
    late50 = [r for s in ("KOPT", "SUBLIST", "LATE_ACCEPTANCE") for r in by(s, fleet=50)]
    mid = lambda s: [r for r in by(s) if r["fleet"] in (10, 20)]

    p("WaterFlex Tabu variant experiment explained", "Title")
    p("How seven daily search configurations behave, and why", "Subtitle")
    p("October 11 2026  |  Run 5d15129b with sequential confirmation 3a77f220")
    p(f"This experiment improved the same booked days with seven search configurations: the production comparison control Tabu, three new Tabu variants and three late acceptance methods. Across {len(tk)} matched cases per configuration, Tabu combined with the K opt moves accepted cheaper schedules than Tabu in {sum(r['saving'] > 0 for r in tk)} cases and dearer ones in {sum(r['saving'] < 0 for r in tk)}, saving {dollars(avg(tk, 'saving'))} per case on average. It is the only configuration that is ahead in every fleet size, workload and budget tested. [1]")
    sub("What the comparison tells us")
    table(["Configuration", "Mean accepted cost", "Mean saving vs Tabu", "Wins / ties / losses"],
          [[LABEL[s], dollars(avg(by(s), "accepted_cost")), "Control" if s == "TABU" else dollars(avg(by(s), "saving")),
            "Control" if s == "TABU" else wtl(by(s))] for s in SOLVERS], [1.7, 1.7, 1.8, 1.7])
    p(f"The late acceptance family (K opt, sublist and late acceptance) is the strongest choice for 10 and 20 technicians, saving about {dollars(mean(avg(mid(s), 'saving') for s in ('KOPT', 'SUBLIST', 'LATE_ACCEPTANCE')))} per case there. On 50 technician days it loses about {dollars(-avg(late50, 'saving'))} per case, because it is still improving when its time runs out. Tabu + K opt keeps most of the mid-size gain and avoids the large-fleet loss. Its cost is somewhat less even workload balance than Tabu.")
    p("The report explains each result through the recorded search mechanics: steps taken, candidates examined per step, and when the best schedule appeared. Costs are modeled on synthetic schedules, not payroll savings, and the production default is unchanged.")
    sub("How to read this report")
    table(["Sections", "What you will learn"], [("1 to 4", "What was tested, how local search works, the seven configurations and the two search phases."),
                                                ("5", "Where each configuration saves or loses cost."),
                                                ("6 to 9", "Why: step economics, convergence, tabu size and the Tabu + K opt combination."),
                                                ("10 and 11", "Workload balance and case to case variation."),
                                                ("12 and 13", "Running six cases at once, and the one at a time confirmation."),
                                                ("14 and appendices", "Decisions supported by the evidence, every fleet group, mechanics and sources.")], [1.45, 5.45])

    heading("1 What was tested")
    p("One case is one configuration improving one fixed daily schedule for one search seed and one time budget. In a matched comparison every configuration receives the same starting schedule. The experiment rearranges appointments that are already booked; it does not admit new booking requests. [1, 3]")
    daily = config["daily"]
    table(["Dimension", "Values in this archive"], [("Configurations", "Tabu (control), Tabu size 3, Tabu size 15, Tabu + K opt, K opt, sublist, late acceptance"),
                                                    ("Fleet sizes", "5, 10, 20 and 50 technicians"), ("Workloads", "Clustered, dispersed, mixed skill and tight window"),
                                                    ("Search seeds", ", ".join(map(str, daily["seeds"]))), ("Time budgets", "60, 90, 120 and 240 seconds per case"),
                                                    ("Execution", "Six independent cases at once, one dedicated core and two active processors per JVM")], [1.5, 5.4])
    p(f"The design is 7 configurations x 4 fleets x 4 workloads x 10 seeds x 4 budgets = {len(rows):,} cases. Every configuration has 640 cases. A fleet, workload and budget group holds ten cases per configuration. All {len(rows):,} results are present and paired, report zero violations, and both search phases ended at their time limits in every case.")
    p(f"There are {len({r['raw']['datasetFingerprint'] for r in rows})} fixed synthetic fixtures, one per fleet and workload. Seeds change the random choices of the search on those same schedules. Ten seeds are not ten territories, and the cases are not {len(rows):,} different customer days.")
    sub("The four workload shapes")
    table(["Workload", "How the fixture is built", "What makes it hard"], [
        ("Clustered", "Four travel groups: about 2 to 3 minutes within a group, about 8 to 9 between groups. 30 minute visits.", "Few big wins; gains come from keeping each route inside one group."),
        ("Dispersed", "Irregular one way travel of 3 to 10 minutes per leg with no group structure. 30 minute visits.", "Many routes of similar cost; reordering pays, and the search space is rough."),
        ("Mixed skill", "Shorter legs of 1.5 to 5 minutes. Only every third technician can do the scarce service, which half of their visits need.", "Many moves are illegal, so feasible improvements are sparse."),
        ("Tight window", "The same short legs, 45 minute visits and arrival windows starting every 90 instead of 120 minutes.", "Long jobs packed close together leave little slack for reordering.")], [1.15, 3.1, 2.65])
    p("Every fixture starts roughly half the technicians with six visits and the rest with two, inside two hour arrival promises on a nine hour day. Travel is a deterministic synthetic matrix, not measured road travel.")
    table(["Technicians", "Appointments per case", "Mean starting cost $", "Mean Tabu accepted $", "Tabu improvement"],
          [[f, by("TABU", fleet=f)[0]["raw"]["appointments"], f"{avg(by('TABU', fleet=f), 'start_cost'):,.2f}", f"{avg(by('TABU', fleet=f), 'accepted_cost'):,.2f}",
            f"{(1 - avg(by('TABU', fleet=f), 'accepted_cost') / avg(by('TABU', fleet=f), 'start_cost')) * 100:.1f}%"] for f in FLEETS], [1.2, 1.5, 1.45, 1.5, 1.25])
    p("The last column shows the room for improvement. Five technician days are already close to as good as these searches can make them, so every configuration is within a few dollars there. The larger fleets have far more to gain, and that is where the configurations separate.")

    heading("2 How a local search improves a schedule")
    p("None of these methods builds a schedule from scratch. Each starts from the booked day and repeatedly makes small changes. Picture a dispatcher with a whiteboard: move one visit to another technician, swap two visits, or move a run of visits. Each change is tried on paper, scored, and either kept or discarded. Repeat that thousands of times a second and you have a local search. [3, 4]")
    table(["Term", "Meaning in this report"], [
        ("Move", "One candidate change to the schedule, such as moving a visit or swapping two visits."),
        ("Move evaluation", "Scoring one candidate move. This is the basic unit of work and costs roughly the same for every method."),
        ("Step", "The moment the search actually changes its current schedule by adopting one move."),
        ("Acceptor", "The rule deciding which candidate moves are allowed to be adopted at all."),
        ("Forager", "The rule deciding how many allowed moves to collect before choosing the best of them."),
        ("Selected count limit", "The most candidates examined in one step: 10,000 for every configuration here."),
        ("Accepted count limit", "How many allowed candidates the forager collects before choosing: 1,000 for Tabu, 1 for late acceptance.")], [1.7, 5.2])
    sub("Two ways to spend the same effort")
    p("All seven configurations evaluate a similar number of moves per second, so they do about the same amount of work. They spend it very differently. Tabu looks hard and steps rarely: in each step it examines candidates until it has collected 1,000 allowed ones (or reached 10,000 examined), then adopts the best of them. Late acceptance looks briefly and steps often: it adopts the first candidate that is no worse than the score it had 400 steps earlier.")
    p("Neither is better in general. Careful steps make steady progress in one direction but can settle into a schedule that no single change improves, a local optimum. Quick steps cover more ground and can climb out of local optima by temporarily accepting a worse schedule, but each step is less informed, so they need many more steps to get anywhere. Sections 6 and 7 measure exactly this tradeoff.")

    heading("3 The seven configurations")
    p("Every configuration uses the same Timefold engine, score rules, starting schedules and budgets. They differ in three settings: which kinds of moves they try, the acceptor, and the forager. [3, 4]")
    table(["Configuration", "Moves tried", "Acceptor", "Accepted per step"], [
        ("Tabu (control)", "Change, swap", "Entity tabu 7", "1,000"),
        ("Tabu size 3", "Change, swap", "Entity tabu 3", "1,000"),
        ("Tabu size 15", "Change, swap", "Entity tabu 15", "1,000"),
        ("Tabu + K opt", "Change, swap, sublist, 2 to 3 opt", "Entity tabu 7", "1,000"),
        ("K opt", "Change, swap, sublist, 2 to 3 opt", "Late acceptance 400", "1"),
        ("Sublist", "Change, swap, sublist", "Late acceptance 400", "1"),
        ("Late acceptance", "Change, swap", "Late acceptance 400", "1")], [1.6, 2.4, 1.7, 1.2])
    p("Change and swap moves each have probability weight 45. Where present, sublist and K opt moves each have weight 5, so roughly one candidate in ten is a structural move.")
    sub("The move types")
    p("A change move takes one visit and puts it somewhere else, on the same route or another technician's route. A swap move exchanges two visits. A sublist move picks up a run of two to four consecutive visits and moves it, optionally reversed. A K opt move cuts two or three links in a route and reconnects the pieces another way, which can untangle a route that crosses over itself. Sublist and K opt moves change route structure in ways that would take several single changes, each of which might look worse on its own.")
    sub("The acceptors")
    p("Entity tabu remembers the planning items moved in the last N steps and forbids moving them again until they age out of the list. This stops the search from immediately undoing its own work and pushes it to try other parts of the schedule. A longer list forces more variety; a shorter list lets the search concentrate on the most promising area.")
    p("Late acceptance keeps the score of each of the last 400 steps. A candidate is allowed if it is at least as good as the current schedule or as the schedule 400 steps ago. Early on that bar is loose, so the search can wander through worse schedules; as the history fills with good scores the bar tightens.")
    sub("Why the three new variants were tested")
    p("Tabu size 3 and Tabu size 15 test whether the control's tabu size of 7 is well chosen: 3 should intensify the search and 15 should diversify it. The earlier 1,280 case study found Tabu losing to the late acceptance family overall but winning clearly on 50 technician dispersed days, so Tabu + K opt asks a sharper question. Do the late acceptance family's gains come from its richer moves or from its acceptor? Tabu + K opt keeps Tabu's acceptor and forager and borrows K opt's move set. [5]")

    heading("4 The two search phases and the accepted result")
    p("Each case runs two searches. The reference phase gets two thirds of the budget (40 of 60 seconds, for example) and searches purely for the cheapest schedule. The fairness phase gets the remaining time and searches for a more even distribution of work. It may spend at most 2% above that configuration's own reference cost, rounded down to the cent. A final policy check chooses the accepted schedule. [3]")
    table(["Measure", "How to interpret it"], [
        ("Accepted cost", "The final recorded schedule cost. All comparisons in this report use it unless stated."),
        ("Reference cost", "The cheapest schedule from the cost phase, before fairness. It is lower than accepted cost."),
        ("Saving vs Tabu", "Tabu accepted cost minus the configuration's accepted cost on the matched case. Positive favors the configuration."),
        ("Fairness variance", "Capacity weighted variation of paid time divided by capacity. Lower is more even."),
        ("Best found at", "The share of the reference phase that had elapsed when its best schedule was first found."),
        ("Steps and moves per step", "How often the search changed its schedule, and how many candidates it examined each time.")], [1.75, 5.15])
    p("Modeled cost uses $30 per regular labor hour, $45 per overtime hour and $0.67 per mile, with a 20% travel buffer plus five minutes per leg. New overtime is not allowed. The measured cases recorded zero overtime at every stage.")
    figure("phases", "Left: each point averages 40 cases for one configuration, fleet and workload. Points on the diagonal mean the accepted saving equals the reference saving. Right: mean accepted minus reference cost, 160 cases per bar.")
    p(f"The left panel holds the key to everything that follows. Every point sits on the diagonal: whatever a configuration gained or lost against Tabu in the cost phase, it kept almost exactly after fairness. The right panel shows why. On a given fleet, every configuration spends about the same amount to rebalance work, roughly {dollars(avg(by('TABU', fleet=50), 'fairness_spend'))} per 50 technician day. That is because each uses up most of its own 2% allowance.")
    p("So the cost race is decided in the reference phase. Explaining why a configuration wins or loses therefore means explaining how its cost search behaves, which is what sections 6 to 9 do.")

    heading("5 Where modeled cost improves")
    p("Each cell averages 160 paired differences: four workloads, ten seeds and four budgets, holding fixture and starting schedule fixed. Read the sign before the size. Blue cells mean the configuration accepted cheaper schedules than Tabu; tan cells mean Tabu was cheaper. [1]")
    figure("cost_fleet", "Mean accepted cost saving vs Tabu in dollars per case. Each fleet cell has 160 pairs; the All fleets column has 640. Tabu + K opt is bold.")
    p("Three patterns stand out. First, the late acceptance family (bottom three rows) wins by $13 to $17 at 10 and 20 technicians and loses by $26 to $28 at 50. Its pooled average is therefore small even though it wins most cases. Second, Tabu + K opt is positive in every column, and largest at 50 technicians, exactly where the late acceptance family fails. Third, tabu size matters in opposite directions: size 3 helps from 10 technicians up, while size 15 hurts badly from 20 up.")
    figure("cost_workload", "Mean accepted cost saving vs Tabu by workload. Each cell has 160 pairs: four fleets, ten seeds and four budgets.")
    p("By workload, Tabu + K opt and Tabu size 3 are positive everywhere. The late acceptance family looks negative on mixed skill and tight window work, but that is the 50 technician loss showing through. Appendix E shows that on 5, 10 and 20 technicians it gains on those workloads too. Mixed skill and tight window days have short legs and little slack, so there is simply less travel to save, and a large-fleet loss outweighs the smaller-fleet gains in the pooled number.")
    table(["Late acceptance family", "5 techs", "10 techs", "20 techs", "50 techs"],
          [[WORK_LABEL[w], *[signed(mean(avg(by(s, fleet=f, workload=w), "saving") for s in ("KOPT", "SUBLIST", "LATE_ACCEPTANCE"))) for f in FLEETS]] for w in WORKLOADS], [2.1, 1.2, 1.2, 1.2, 1.2])
    p("This table averages K opt, sublist and late acceptance, which behave almost identically, by workload and fleet. Every workload follows the same shape: gains up to 20 technicians, a loss at 50, except clustered work, where the family still gains at 50.")

    heading("6 Why step style matters")
    figure("steps", "Reference phase means per configuration and fleet, 160 cases per point, on log scales. Left: steps per second. Right: candidate moves examined per step.")
    t50, l50 = by("TABU", fleet=50), by("LATE_ACCEPTANCE", fleet=50)
    p(f"The two families sit about two orders of magnitude apart. Late acceptance takes {avg(la, 'steps_rate'):,.0f} steps per second on average, examining about {avg(la, 'moves_per_step'):.0f} candidates per step. Tabu takes {avg(by('TABU'), 'steps_rate'):.0f} steps per second, examining thousands of candidates each time. The total work is similar: Tabu evaluates {avg(by('TABU'), 'move_rate'):.0f} thousand moves per second and late acceptance {avg(la, 'move_rate'):.0f} thousand.")
    table(["Configuration", "Thousand moves / s", "Steps / s", "Moves per step", "Saving vs Tabu"],
          [[LABEL[s], f"{avg(by(s), 'move_rate'):.1f}", f"{avg(by(s), 'steps_rate'):,.1f}", f"{avg(by(s), 'moves_per_step'):,.0f}", "Control" if s == "TABU" else signed(avg(by(s), "saving"))] for s in SOLVERS], [1.7, 1.35, 1.2, 1.35, 1.3])
    p("Why the gap is so large: the late acceptance forager stops at the first allowed candidate, and early in the search almost anything is allowed. Tabu's forager waits for 1,000 allowed candidates. At 5 to 20 technicians many candidates touch recently moved items and are forbidden, so Tabu usually hits the 10,000 examination cap before collecting 1,000. That is why its moves per step sit between about 8,500 and 10,000 there.")
    p(f"At 50 technicians the picture changes. With 200 visits a tabu list of 7 forbids only a small share of the schedule, so Tabu reaches 1,000 allowed candidates after about {avg(t50, 'moves_per_step'):,.0f} examinations and steps faster ({avg(t50, 'steps_rate'):.0f} per second). Late acceptance steps at {avg(l50, 'steps_rate'):,.0f} per second, but each step is one random allowed move on a schedule with 200 visits. Section 7 shows that this is not enough steps to finish.")
    p("Moves per second fall as fleets grow for every configuration. Larger routes cost more to rescore. Configurations with sublist and K opt moves are a few percent slower per move, because those moves change more of a route.")

    heading("7 Did the search finish in time?")
    p("Each run records when its reference phase first found the schedule it ended with. If that happens early, the search had time to spare and was either finished or stuck. If it happens in the last moments, the search was still improving when the clock stopped and more time would probably have helped. [1]")
    figure("convergence", "Left: mean time of the best reference schedule as a percentage of the reference phase. Right: percentage of cases whose best appeared in the final 10% of the phase. 160 cases per point.")
    p(f"At 50 technicians the late acceptance family found its best schedule at about {avg(late50, 'best_fraction') * 100:.0f}% of the phase on average, and {avg(late50, 'late_best') * 100:.0f}% of those cases were still improving in the final 10%. These searches ran out of time while still improving. The Tabu family's best arrived around two thirds of the way through, and only about a quarter of its cases improved at the very end.")
    p(f"At 5 and 10 technicians the picture flips. Tabu found its best within the first {avg(by('TABU', fleet=5), 'best_fraction') * 100:.0f}% (5 techs) and {avg(by('TABU', fleet=10), 'best_fraction') * 100:.0f}% (10 techs) of the phase and then improved no further. On these small days Tabu settles into a local optimum early and its careful steps cannot climb back out. The late acceptance family keeps finding better schedules later in the phase. Its tolerance for temporarily worse schedules is exactly what lets it escape.")
    figure("budgets", "Mean saving vs Tabu by budget. Left: 5, 10 and 20 technicians (120 pairs per point). Right: 50 technicians (40 pairs per point). Lines join separately run budgets.")
    gains = {s: own_budget_gain(rows, s) for s in SOLVERS}
    p(f"Extra time confirms the diagnosis. At 50 technicians the late acceptance family's deficit shrinks from about {dollars(-mean(avg(by(s, fleet=50, budget_ms=60000), 'saving') for s in ('KOPT', 'SUBLIST', 'LATE_ACCEPTANCE')))} at 60 seconds to about {dollars(-mean(avg(by(s, fleet=50, budget_ms=240000), 'saving') for s in ('KOPT', 'SUBLIST', 'LATE_ACCEPTANCE')))} at 240 seconds. It is catching up and might overtake Tabu with more time; this run cannot tell. On smaller fleets its advantage is flat, because those searches had already finished.")
    table(["Configuration", "Own saving 240 s vs 60 s", "Cheaper / same / dearer"],
          [[LABEL[s], dollars(mean(gains[s])), f"{sum(v > 0 for v in gains[s])}/{sum(v == 0 for v in gains[s])}/{sum(v < 0 for v in gains[s])}"] for s in SOLVERS], [2.0, 2.6, 2.3])
    p("This table compares each configuration with itself, holding fixture and seed fixed, over 160 pairs. The late acceptance family gains about three times as much from four times the budget as the Tabu family does. Its searches are time-limited, while the Tabu family's are mostly done.")
    sub("Follow one measured case")
    example = [select(rows, solver=s, fleet=50, workload="DISPERSED", seed=17, budget_ms=60000)[0] for s in ("TABU", "TABU_KOPT", "KOPT", "LATE_ACCEPTANCE")]
    table(["50 dispersed, seed 17, 60 s", "Start $", "Reference $", "Accepted $", "Steps", "Best at"],
          [[LABEL[r["solver"]], f"{r['start_cost']:,.2f}", f"{r['reference_cost']:,.2f}", f"{r['accepted_cost']:,.2f}", f"{r['raw']['referencePhase']['steps']:,}",
            f"{r['best_fraction'] * 100:.0f}%"] for r in example], [2.0, .95, 1.05, 1.05, .9, .95])
    e = {r["solver"]: r for r in example}
    p(f"All four start from the same {dollars(e['TABU']['start_cost'])} schedule. Late acceptance takes {e['LATE_ACCEPTANCE']['raw']['referencePhase']['steps']:,} steps, about {e['LATE_ACCEPTANCE']['raw']['referencePhase']['steps'] / e['TABU']['raw']['referencePhase']['steps']:.0f} times as many as Tabu, yet its reference is {dollars(e['LATE_ACCEPTANCE']['reference_cost'] - e['TABU']['reference_cost'])} more expensive. Its best schedule appeared at {e['LATE_ACCEPTANCE']['best_fraction'] * 100:.0f}% of the phase, so it was still improving at the deadline. Tabu + K opt takes about as many steps as Tabu and finishes {dollars(e['TABU']['accepted_cost'] - e['TABU_KOPT']['accepted_cost'])} cheaper. This is one observed case, chosen to illustrate the mechanism, not a forecast.")

    heading("8 Why tabu size matters")
    p("The tabu size is the number of recently moved items the search may not touch. Its effect on cost runs entirely through the step economics of section 6. [1, 4]")
    table(["Fleet", "Size", "Steps / s", "Moves per step", "Best at", "Saving vs Tabu"],
          [[f, size, f"{avg(by(s, fleet=f), 'steps_rate'):.1f}", f"{avg(by(s, fleet=f), 'moves_per_step'):,.0f}", f"{avg(by(s, fleet=f), 'best_fraction') * 100:.0f}%",
            "Control" if s == "TABU" else signed(avg(by(s, fleet=f), "saving"))]
           for f in FLEETS for s, size in (("TABU_SIZE_3", 3), ("TABU", 7), ("TABU_SIZE_15", 15))], [.9, .8, 1.15, 1.45, 1.1, 1.5])
    p(f"Size 15 forbids so many items that, from 20 technicians up, the forager almost never collects 1,000 allowed candidates. Every step burns the full 10,000 examinations. At 50 technicians it manages {avg(by('TABU_SIZE_15', fleet=50), 'steps_rate'):.1f} steps per second against Tabu's {avg(t50, 'steps_rate'):.1f}, roughly a third. Fewer steps means less progress, and it loses {dollars(-avg(by('TABU_SIZE_15', fleet=50), 'saving'))} per case there. The extra variety a longer list was meant to buy is swamped by the slowdown.")
    p(f"Size 3 forbids fewer items, so the forager fills up sooner and the search steps more often: {avg(by('TABU_SIZE_3', fleet=20), 'steps_rate'):.0f} against {avg(by('TABU', fleet=20), 'steps_rate'):.0f} steps per second at 20 technicians. That buys more progress from 10 technicians up. On 5 technician days, with only 18 visits, it is slightly worse than Tabu ({signed(avg(by('TABU_SIZE_3', fleet=5), 'saving'))}). A very short list there lets the search drift back to arrangements it just left. That reading is an interpretation consistent with the early best time, not something the archive records directly.")
    p("The lesson is that tabu size is not a free diversity knob. Through the forager it sets how many candidates each step costs, and on large schedules that cost dominates.")

    heading("9 Why Tabu + K opt wins everywhere")
    p("Tabu + K opt pairs Tabu's acceptor and forager with K opt's move set. Comparing three configurations separates what the moves contribute from what the acceptor contributes. [1]")
    table(["Fleet", "Tabu + K opt saving", "K opt saving", "Tabu + K opt steps / s", "K opt steps / s", "Tabu + K opt best at"],
          [[f, signed(avg(by("TABU_KOPT", fleet=f), "saving")), signed(avg(by("KOPT", fleet=f), "saving")), f"{avg(by('TABU_KOPT', fleet=f), 'steps_rate'):.1f}",
            f"{avg(by('KOPT', fleet=f), 'steps_rate'):,.0f}", f"{avg(by('TABU_KOPT', fleet=f), 'best_fraction') * 100:.0f}%"] for f in FLEETS], [.7, 1.35, 1.1, 1.35, 1.15, 1.25])
    p("The moves explain the gains. With Tabu's acceptor unchanged, adding sublist and K opt moves turns Tabu into a winner at every fleet size. The acceptor explains the losses. With the K opt moves unchanged, switching to late acceptance gains more at 10 and 20 technicians but collapses at 50.")
    p("Why the combination works: each Tabu step collects up to 1,000 allowed candidates and adopts the best. About one candidate in ten is a structural sublist or K opt move. When such a move is the best available it is chosen deliberately, so Tabu reliably harvests the restructuring it was missing. Late acceptance adopts the first allowed candidate, so a structural move is taken only when it happens to come first. That costs nothing on small schedules with plenty of steps to spare, but it wastes the large-fleet budget.")
    p(f"Tabu + K opt does not capture all of K opt's mid-size advantage: about {dollars(avg(mid('TABU_KOPT'), 'saving'))} per case at 10 and 20 technicians against K opt's {dollars(avg(mid('KOPT'), 'saving'))}. There, Tabu's early stall in a local optimum (section 7) still limits it. Its gains are also front-loaded, at {dollars(avg(by('TABU_KOPT', budget_ms=60000), 'saving'))} at 60 seconds and {dollars(avg(by('TABU_KOPT', budget_ms=240000), 'saving'))} at 240. The improvement comes from better steps, not more time.")
    p(f"It pays a small evaluation cost: at 50 technicians it evaluates {avg(by('TABU_KOPT', fleet=50), 'move_rate'):.1f} thousand moves per second against Tabu's {avg(t50, 'move_rate'):.1f} thousand, because structural moves rescore more of a route.")

    heading("10 How workload balance changes")
    p("Fairness variance measures how unevenly paid time falls across technicians, relative to each one's capacity. Lower is more even. For teaching arithmetic: two technicians with equal 480 minute days and 240 and 480 paid minutes have utilizations of 0.50 and 1.00, and a variance of 0.0625. Equal 360 minute days give zero. These illustrative values are not experiment observations. [3]")
    figure("fairness", "Each point averages 160 cases for one configuration and fleet; labels give the fleet. Right is cheaper than Tabu; up is more balanced than Tabu. Variance differences are multiplied by 1,000.")
    table(["Configuration", "5 techs", "10 techs", "20 techs", "50 techs", "All"],
          [[LABEL[s], *[f"{avg(by(s, fleet=f), 'fairness'):.4f}" for f in FLEETS], f"{avg(by(s), 'fairness'):.4f}"] for s in SOLVERS], [1.7, 1.0, 1.05, 1.05, 1.05, 1.05])
    p("Tabu produces the most even workloads overall. Most points lie below the zero line: cheaper schedules here usually concentrate work. A cheaper reference tends to pack visits onto fewer, denser routes, and the fairness phase can only partly undo that within its 2% allowance. The late acceptance family pays the most balance for its savings. Tabu + K opt sits between them, closer to Tabu.")
    reasons = {s: sum(r["raw"]["acceptanceReason"] == "COST_REDUCTION" for r in by(s)) for s in SOLVERS}
    p(f"The policy check labels each accepted schedule with a reason. Tabu's result is labeled a fairness improvement in {640 - reasons['TABU']} of 640 cases, and the late acceptance family's in about {640 - round(mean(reasons[s] for s in ('KOPT', 'SUBLIST', 'LATE_ACCEPTANCE')))}. For the rest, the accepted schedule is justified by cost alone. These labels compare against each configuration's own starting point, not against Tabu.")
    p("Tabu size 15 is the one configuration that is more balanced than Tabu at 50 technicians. It pays about $25 per case for that, so it is not an efficient way to buy balance.")

    heading("11 What individual cases and seeds show")
    p("A win counts a lower accepted cost in exact cents, regardless of size. The table covers all 640 matched cases per configuration, including ties and losses. [1]")
    table(["Configuration", "Wins / ties / losses", "Worst saving $", "Median saving $", "Best saving $"],
          [[LABEL[s], wtl(by(s)), signed(min(r["saving"] for r in by(s))), signed(sorted(r["saving"] for r in by(s))[319:321][0] / 2 + sorted(r["saving"] for r in by(s))[319:321][1] / 2),
            signed(max(r["saving"] for r in by(s)))] for s in ALTS], [1.7, 1.55, 1.2, 1.25, 1.2])
    figure("variation", "Left: mean saving per search seed, 64 cases per point. Right: cumulative distribution of all 640 case savings per configuration, clipped to the -$120 to +$80 range. These describe observed variation, not confidence intervals.")
    p("On the left, Tabu + K opt is above zero for every seed, so its headline result does not depend on one lucky seed. The late acceptance family swings more from seed to seed because the 50 technician losses are large.")
    p("On the right, pick a dollar value and read up to a curve: its height is the share of cases saving that much or less. The late acceptance family has a long left tail of 50 technician losses near -$100. Tabu + K opt's curve rises steeply just right of zero: many modest wins and few losses.")

    heading("12 How long it took, and running six at once")
    p("Every case used essentially its whole budget; both phases report TIME_LIMIT in all cases. Six cases ran at once, each on its own physical core with two active processors and a serial garbage collector. [1, 2]")
    p("Running cases side by side still shares memory bandwidth, cache and boost clock. Before the study, a contention check ran 24 identical cases one at a time and then 4, 6 and 10 at once, comparing how many moves per second each solver managed. The bar for treating parallel timings as solo timings is 0.95.")
    figure("contention", "Median parallel over sequential move evaluation rate in the reference phase, per solver. Circles are 20 technician cases, squares 50. Each point has 2 or 3 pairs.")
    table(["Cases at once", "Lowest ratio", "Highest ratio", "Pairs used", "Meets 0.95 bar"],
          [[n, f"{min(g['reference_median_ratio'] for g in contention[n]['groups']):.3f}", f"{max(g['reference_median_ratio'] for g in contention[n]['groups']):.3f}",
            sum(g["pairs"] for g in contention[n]["groups"]), "No"] for n in (4, 6, 10)], [1.5, 1.35, 1.35, 1.3, 1.4])
    p("No setting fully met the bar. Four at once came closest; only Tabu missed (0.94 and 0.91 on three pairs each). The owner chose six at once. Every solver loses about 15% of its search rate there, so a 240 second case does roughly the work of a 200 second solo run. Because all configurations share that slowdown, comparisons between them stay fair. Section 13 checks this directly.")
    p("The earlier 1,280 case study ran ten at once, where each solver kept only about 60% of its solo rate. Its comparisons between solvers remain fair for the same reason, but its budget figures overstate how much search each budget bought.")

    heading("13 Checking the result one case at a time")
    seq, par = [c["sequential"] for c in confirmed], [c["parallel"] for c in confirmed]
    agree = sum((a > 0) == (b > 0) and (a < 0) == (b < 0) for a, b in zip(seq, par))
    p(f"After the main study, the driver picked the configuration with the best accepted cost saving, Tabu + K opt, and reran it beside Tabu one case at a time. This used seeds 17 and 23, every fleet and workload, and the 120 and 240 second budgets: 128 cases with no contention at all. [6]")
    figure("confirmation", "Each point is one of 64 matched cases: Tabu + K opt saving vs Tabu with six at once (horizontal) and one at a time (vertical). The diagonal is perfect agreement.")
    table(["Group", "Pairs", "One at a time $", "Six at once $", "W/T/L one at a time"],
          [["All", 64, signed(mean(seq)), signed(mean(par)), "/".join(str(sum(t(v) for v in seq)) for t in (lambda x: x > 0, lambda x: x == 0, lambda x: x < 0))]]
          + [[f"{b} s budget", 32, signed(mean(c["sequential"] for c in confirmed if c["key"][3] == b * 1000)), signed(mean(c["parallel"] for c in confirmed if c["key"][3] == b * 1000)), ""] for b in (120, 240)]
          + [[f"{f} technicians", 16, signed(mean(c["sequential"] for c in confirmed if c["key"][0] == f)), signed(mean(c["parallel"] for c in confirmed if c["key"][0] == f)), ""] for f in FLEETS],
          [1.7, .8, 1.5, 1.4, 1.5])
    p(f"Run one at a time, Tabu + K opt saves {dollars(mean(seq))} per case and wins {sum(v > 0 for v in seq)} of 64. In {agree} of 64 cases the sign of the saving matches its parallel twin, and it is ahead at every fleet and both budgets. The parallel result stands. These seeds also appeared in the main study, so this rules out a contention artifact; it is not an independent sample of new schedules.")

    heading("14 What should happen next")
    p("The evidence supports three conclusions. Tabu + K opt is a strict improvement over Tabu on these fixtures, with a modest balance cost. The late acceptance family is better still from 10 to 20 technicians but needs more time than these budgets give on 50 technician days. Tabu size 15 should be dropped, and size 3 is a cheap partial improvement.")
    table(["Question", "Next evidence to collect", "Decision informed"], [
        ("Adopt Tabu + K opt?", "Replay real customer days and measured road travel through both; confirm balance is acceptable.", "Whether to change the daily default."),
        ("Choose by fleet size?", "Test K opt up to 20 technicians and Tabu + K opt above, on new seeds and fixtures.", "Whether a size-based rule beats one configuration."),
        ("Combine both ideas?", "Try Tabu + K opt with tabu size 3, which may step more often.", "Whether the two improvements add up."),
        ("Large fleet budgets", "Give late acceptance 50 technician days longer budgets until its deadline improvement stops.", "How much overnight time large fleets need."),
        ("Balance tradeoff", "Set an explicit acceptable variance, or price it in the score.", "Whether cheaper but less even schedules are wanted.")], [1.5, 3.1, 2.3])
    best = {f: max(ALTS, key=lambda s: avg(by(s, fleet=f), "saving")) for f in FLEETS}
    p(f"Picking the best configuration per fleet in hindsight ({', '.join(f'{LABEL[best[f]]} at {f}' for f in FLEETS)}) would average {dollars(mean(avg(by(best[f], fleet=f), 'saving') for f in FLEETS))} per case. That figure is selected on the same data it is measured on, so it is optimistic and needs new fixtures to confirm.")
    sub("How the calculations were checked")
    p("Checks cover the archived configuration and case hashes, every raw result hash in the saved analysis, unique and complete case coverage, fixture identity and starting schedules between matched pairs, the 2% cost ceiling, zero overtime and violations, and an independent recomputation of fairness variance from paid minutes. The contention reports and the confirmation run are loaded and checked the same way. verify_report.py recomputes every table and chart dataset from raw cases. [1, 2, 6]")

    for letter, f in zip("ABCD", FLEETS):
        heading(f"Appendix {letter} The {f} technician cases")
        p("Each row holds ten matched seeds per configuration at the stated budget. Work gives the workload initial. Savings are Tabu accepted cost minus the configuration's accepted cost in dollars per case, with wins, ties and losses below in exact cents; each entry totals ten. [1]")
        data = []
        for w in WORKLOADS:
            for b in BUDGETS:
                groups = [by(s, fleet=f, workload=w, budget_ms=b * 1000) for s in SOLVERS]
                # Integer cents first, so an exact half cent rounds the same way as in verify_report.py.
                data.append([w[0], b, f"{mean(r['raw']['after']['costCents'] for r in groups[0]) / 100:,.2f}", *[f"{avg(g, 'saving'):+.2f}\n{wtl(g)}" for g in groups[1:]]])
        table(["Work", "Sec", "Tabu $", *[SHORT[s] for s in ALTS]], data, [.5, .45, .85, .85, .85, .85, .85, .85, .85], size=9)

    heading("Appendix E Search mechanics by fleet")
    p("Reference phase means over 160 cases per row (four workloads, ten seeds, four budgets). Best at is the share of the phase elapsed when the best schedule appeared; late is the share of cases whose best appeared in the final 10%. [1]")
    table(["Configuration", "Fleet", "Thousand moves / s", "Steps / s", "Moves / step", "Best at", "Late", "Saving"],
          [[LABEL[s], f, f"{avg(by(s, fleet=f), 'move_rate'):.1f}", f"{avg(by(s, fleet=f), 'steps_rate'):,.1f}", f"{avg(by(s, fleet=f), 'moves_per_step'):,.0f}",
            f"{avg(by(s, fleet=f), 'best_fraction') * 100:.0f}%", f"{avg(by(s, fleet=f), 'late_best') * 100:.0f}%",
            "Control" if s == "TABU" else signed(avg(by(s, fleet=f), "saving"))] for s in SOLVERS for f in FLEETS], [1.5, .6, 1.0, .85, .9, .7, .6, .75], size=9)
    sub("Savings by workload and fleet")
    table(["Configuration", "Workload", "5 techs", "10 techs", "20 techs", "50 techs"],
          [[LABEL[s], WORK_LABEL[w], *[signed(avg(by(s, fleet=f, workload=w), "saving")) for f in FLEETS]] for s in ALTS for w in WORKLOADS], [1.6, 1.3, 1.0, 1.0, 1.0, 1.0], size=9)

    heading("Appendix F Definitions and evidence")
    table(["Term", "Meaning in this report"], [
        ("Case", "One configuration run at one fleet, workload, seed and budget."),
        ("Matched pair", "Same fixture and starting schedule under the same seed and budget, comparing a configuration with Tabu."),
        ("Mean saving", "Average paired difference in accepted modeled dollars. Positive favors the configuration."),
        ("Reference and ceiling", "Cost phase result and its 2% fairness allowance, rounded down to cents."),
        ("Step", "Adoption of one move as the new current schedule."),
        ("Best at", "Time to the best reference schedule divided by reference phase time."),
        ("TIME_LIMIT", "The phase ended at its allocated time. It means neither proven optimal nor infeasible."),
        ("Contention ratio", "Median parallel over sequential move evaluation rate for identical cases.")], [1.85, 5.05])
    p(f"[1] Main archive: experiments/runs/{RUN.name}. Measured revision {REVISION}. Raw JSONL, manifest, configuration, cases and completion receipts are the primary evidence.")
    p(f"[2] Saved analysis {ANALYSIS.name} under that archive. analysis-provenance.json supplies the input hashes. Contention reports: {', '.join(f'{n} at once in {CONTENTION[n].parents[1].name}' for n in (4, 6, 10))}, each compared with {SEQUENTIAL.name}.")
    p("[3] Measured source: SolverBenchmark.java and SolverBenchmarkData.java under scheduler-service/src/test/java/dev/waterflex/scheduler/optimizer, and SolverEngine.java under calculation-engine/src/main/java/dev/waterflex/scheduler/optimizer. These define fixtures, configurations, phases and metrics.")
    p("[4] Timefold Solver 2.6.0 local search: acceptors, foragers and move selectors, as configured in the recorded solver XML of every case.")
    p("[5] Earlier study: docs/daily-budget-331df1f1. Its values are not pooled with this run; it ran ten cases at once with different workloads.")
    p(f"[6] Sequential confirmation: experiments/runs/{CONFIRM.name}, analysis {CONFIRM_ANALYSIS.name}, one case at a time.")
    p("Reproduce with docs/daily-tabu-variants-5d15129b/requirements.txt, then run report_data.py, build_report.py and verify_report.py there. The builder reads the archives, regenerates all ten charts and chart_data.json, and writes this editable Word file. EVIDENCE.md records methods.")

    settings = doc.settings.element
    for el in settings.findall(qn("w:updateFields")): settings.remove(el)
    update = OxmlElement("w:updateFields"); update.set(qn("w:val"), "true"); settings.append(update)
    stream = BytesIO(); doc.save(stream)
    with ZipFile(REFERENCE) as ref, ZipFile(stream) as built, ZipFile(OUTPUT, "w", ZIP_DEFLATED) as output:
        preserve = set(PRESERVE) | {n for n in ref.namelist() if "footer" in n}
        for name in sorted(built.namelist()):
            content = ref.read(name) if name in preserve else built.read(name)
            info = ZipInfo(name, (2026, 10, 11, 0, 0, 0)); info.compress_type = ZIP_DEFLATED
            output.writestr(info, content)
    print(OUTPUT)


if __name__ == "__main__": build()
