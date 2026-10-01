"""Check delivered tables, plot data and template parts against validated raw cases.

Run normally, without Python's -O flag. Visual pagination still requires rendering.
"""
import json
from hashlib import sha256
from math import isclose
from statistics import mean
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from build_report import HERE, OUTPUT, REFERENCE, REFERENCE_SHA, PRESERVE
from report_data import load

SOLVERS = ("TABU", "KOPT", "LATE_ACCEPTANCE", "SUBLIST")
LABELS = ("Tabu", "K opt", "Late acceptance", "Sublist")
FLEETS = (5, 10, 20, 50)
BUDGETS = (60, 90, 120, 240)
WORKLOADS = ("CLUSTERED", "DISPERSED")


def verify():
    config, rows = load()
    raw = {r["key"]: r["raw"] for r in rows}
    assert len(raw) == 1280

    def group(s, f=None, w=None, b=None, seed=None):
        selected = [(k, v) for k, v in raw.items() if k[4] == s
                    and (f is None or k[0] == f) and (w is None or k[1] == w)
                    and (b is None or k[3] == b * 1000) and (seed is None or k[2] == seed)]
        assert selected
        return selected

    def cost(v, stage="after"):
        return v[stage]["costCents"] / 100

    def savings(g):
        return [(raw[k[:-1] + ("TABU",)]["after"]["costCents"] - v["after"]["costCents"]) / 100 for k, v in g]

    def variance(v):
        return v["after"]["fairness"]["variance"]

    def counts(values):
        return "/".join(str(sum(test(x) for x in values)) for test in (lambda x:x>0, lambda x:x==0, lambda x:x<0))

    def dollars(value):
        return f"${value:,.2f}" if value >= 0 else f"-${-value:,.2f}"

    doc = Document(OUTPUT)
    assert len(doc.tables) == 24 and len(doc.inline_shapes) == 6 and len(doc.sections) == 1
    assert len([p for p in doc.paragraphs if p.style.name == "Heading 1"]) == 15
    text = "\n".join(p.text for p in doc.paragraphs)
    captions = [p for p in doc.paragraphs if p.style.name == "Caption"]
    assert len(captions) == 6
    for i, p in enumerate(captions, 1):
        assert p.text.startswith(f"Figure {i}.")
        assert all(str(r.font.color.rgb) == "000000" and r.italic for r in p.runs)
    assert "794 fairness improvements and 486 cost reductions" in text
    assert "Ten search seeds" in text and "K opt and sublist overlap exactly" in text

    def check_table(index, expected):
        actual = [[c.text for c in r.cells] for r in doc.tables[index].rows[1:]]
        assert actual == [[str(c) for c in r] for r in expected], (index, actual, expected)

    check_table(0, [[label, dollars(mean(cost(v) for _,v in group(s))),
                    "Control" if s == "TABU" else dollars(mean(savings(group(s)))),
                    "Control" if s == "TABU" else str(sum(x>0 for x in savings(group(s))))]
                   for s,label in zip(SOLVERS,LABELS)])
    check_table(3, [[f, group("TABU",f)[0][1]["appointments"],group("TABU",f)[0][1]["appointments"]*30] for f in FLEETS])
    check_table(6, [[b,*[f"{mean(savings(group(s,b=b))):+.2f}" for s in SOLVERS[1:]]] for b in BUDGETS])
    own=[]
    for s,label in zip(SOLVERS,LABELS):
        values=[(raw[k[:3]+(60000,s)]["after"]["costCents"]-v["after"]["costCents"])/100 for k,v in group(s,b=240)]
        assert len(values)==80
        own.append([label,dollars(mean(values)),counts(values)])
    check_table(7,own)
    check_table(8, [[label,counts(savings(group(s))),f"{min(savings(group(s))):+.2f}",f"{max(savings(group(s))):+.2f}"] for s,label in zip(SOLVERS[1:],LABELS[1:])])
    check_table(9, [[label,*[f"{cost(group(s,5,'CLUSTERED',60,17)[0][1],stage):.2f}" for stage in ('before','reference','after')],f"{variance(group(s,5,'CLUSTERED',60,17)[0][1]):.5f}"] for s,label in zip(SOLVERS,LABELS)])
    check_table(10, [[label,f"{mean(cost(v) for _,v in group(s,50,'DISPERSED')):,.2f}",f"{mean(variance(v) for _,v in group(s,50,'DISPERSED')):.5f}","Control" if s=='TABU' else sum(x<0 for x in savings(group(s,50,'DISPERSED')))] for s,label in zip(SOLVERS,LABELS)])
    elapsed=[]
    for b in BUDGETS:
        values=[v['elapsedMs']/1000 for k,v in raw.items() if k[3]==b*1000]
        elapsed.append([b,len(values),f"{min(values):.3f}",f"{mean(values):.3f}",f"{max(values):.3f}"])
    check_table(11,elapsed)
    check_table(12, [[label,f"{mean(sum(w['paidMinutes'] for w in v['after']['fairness']['workloads']) for _,v in group(s)):.2f}",*[f"{mean(v[field] for _,v in group(s)):.2f}" for field in ('paidWaitingAfter','changedAssignments','retimedAppointments')]] for s,label in zip(SOLVERS,LABELS)])
    check_table(13, [[label,f"{mean(cost(v,'reference') for _,v in group(s)):,.2f}",f"{mean(cost(v) for _,v in group(s)):,.2f}",f"{mean((v['after']['costCents']-v['reference']['costCents'])/100 for _,v in group(s)):.2f}"] for s,label in zip(SOLVERS,LABELS)])
    for i,f in enumerate(FLEETS):
        check_table(15+i*2, [[w[0],b,f"{mean(v['after']['costCents'] for _,v in group('TABU',f,w,b))/100:,.2f}",*[f"{mean(savings(group(s,f,w,b))):+.2f}\n{counts(savings(group(s,f,w,b)))}" for s in SOLVERS[1:]]] for w in WORKLOADS for b in BUDGETS])
        check_table(16+i*2, [[w[0],b,*[f"{mean(variance(v) for _,v in group(s,f,w,b)):.5f}" for s in SOLVERS]] for w in WORKLOADS for b in BUDGETS])

    data=json.loads((HERE/'generated/chart_data.json').read_text())
    def close(actual, expected):
        assert len(actual)==len(expected)
        assert all(isclose(a,e,rel_tol=1e-10,abs_tol=1e-10) for a,e in zip(actual,expected)), (actual,expected)
    groups=[(f,w) for f in FLEETS for w in WORKLOADS]
    assert data['cost_groups']['groups']==[list(g) for g in groups]
    assert data['cost_groups']['solvers']==list(SOLVERS[1:]) and data['cost_groups']['n_per_cell']==40
    for actual,(f,w) in zip(data['cost_groups']['values'],groups):
        close(actual,[mean(savings(group(s,f,w))) for s in SOLVERS[1:]])
    for d in data['budgets']:
        assert d['n_per_point']==40 and d['seconds']==list(BUDGETS)
        close(d['savings'],[mean(savings(group(d['solver'],w=d['workload'],b=b))) for b in BUDGETS])
    for d in data['variation']:
        assert d['n_per_seed']==32 and d['seeds']==sorted(config['daily']['seeds'])
        close(d['seed_means'],[mean(savings(group(d['solver'],seed=seed))) for seed in d['seeds']])
        close(d['case_savings'],sorted(savings(group(d['solver']))))
    for d in data['fairness']:
        assert d['n_per_point']==40 and d['groups']==[list(g) for g in groups]
        close(d['savings'],[mean(savings(group(d['solver'],f,w))) for f,w in groups])
        close(d['variance_reduction'],[mean(variance(raw[k[:-1]+('TABU',)])-variance(v) for k,v in group(d['solver'],f,w)) for f,w in groups])
    for d in data['throughput']:
        assert d['n_per_point']==80 and d['fleets']==list(FLEETS)
        phase=d['phase']+'Phase'
        close(d['thousand_moves_per_second'],[mean(v[phase]['moveEvaluations']/v[phase]['solveMs'] for _,v in group(d['solver'],f)) for f in FLEETS])
    for d in data['worked']:
        workloads=group(d['solver'],5,'CLUSTERED',60,17)[0][1]['after']['fairness']['workloads']
        assert d['technicians']==[w['technicianId'] for w in workloads]
        close(d['utilization_percent'],[w['paidMinutes']/w['regularCapacityMinutes']*100 for w in workloads])

    assert sha256(REFERENCE.read_bytes()).hexdigest()==REFERENCE_SHA
    with ZipFile(REFERENCE) as ref, ZipFile(OUTPUT) as out:
        for name in (*PRESERVE,*[n for n in ref.namelist() if 'footer' in n]):
            assert ref.read(name)==out.read(name),name
        for name in out.namelist():
            if name.endswith('.xml'):
                assert '\u2014' not in out.read(name).decode('utf-8'),name
        embedded={sha256(out.read(n)).hexdigest() for n in out.namelist() if n.startswith('word/media/')}
        assert embedded=={sha256(p.read_bytes()).hexdigest() for p in (HERE/'generated').glob('*.png')}
    assert doc.sections[0]._sectPr.xml==Document(REFERENCE).sections[0]._sectPr.xml
    assert all(shape._inline.docPr.get('descr') for shape in doc.inline_shapes)
    assert all(row._tr.get_or_add_trPr().find(qn('w:cantSplit')) is not None for t in doc.tables for row in t.rows)
    print('PASS: 1,280 raw cases; all numerical result tables; six chart datasets and embedded images; 24 tables; template parts and structure.')


if __name__ == '__main__':
    verify()
