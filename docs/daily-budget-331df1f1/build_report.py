"""Build the standalone Word account of daily budget run 331df1f1.

Requires python-docx 1.2.0. Run from any directory with PYTHONPATH pointing to
the isolated dependency folder, then run report_data.py for a fresh audit.
"""
from collections import defaultdict
from pathlib import Path
from statistics import mean

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from report_data import load

HERE = Path(__file__).resolve().parent
OUTPUT = HERE / "Daily_Budget_331df1f1_Report.docx"


def money(cents):
    return f"${cents / 100:,.2f}"


def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        line = OxmlElement(f"w:{edge}")
        line.set(qn("w:val"), "single")
        line.set(qn("w:sz"), "4")
        line.set(qn("w:color"), "D9D9D9")
        borders.append(line)
    t._tbl.tblPr.append(borders)
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for values in rows:
        t.add_row()
    for ri, (row, values) in enumerate(zip(t.rows, [headers] + rows)):
        for ci, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = str(value)
            shade(cell, "DDEBF5" if ri == 0 else ("F6FAFD" if ri % 2 == 0 else "FFFFFF"))
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(9)
                    if ri == 0:
                        run.bold = True
        if ri == 0:
            trpr = row._tr.get_or_add_trPr()
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            trpr.append(repeat)
    doc.add_paragraph()


def paragraph(doc, text, style=None):
    return doc.add_paragraph(text, style=style)


def build():
    config, rows = load()
    by_key = {(r["fleet"], r["workload"], r["seed"], r["budget_ms"], r["solver"]): r for r in rows}
    grouped = defaultdict(list)
    for r in rows:
        grouped[r["solver"]].append(r)
    paired = {}
    for solver in ("KOPT", "LATE_ACCEPTANCE", "SUBLIST"):
        paired[solver] = [(r, by_key[(r["fleet"], r["workload"], r["seed"], r["budget_ms"], "TABU")])
                          for r in grouped[solver]]

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.75)
    sec.left_margin = sec.right_margin = Inches(.83)
    styles = doc.styles
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        styles[name].font.name = "Calibri"
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(7)
    styles["Title"].font.size = Pt(21)
    styles["Title"].font.bold = True
    styles["Title"].paragraph_format.space_after = Pt(10)
    title_ppr = styles["Title"]._element.pPr
    if title_ppr is not None:
        for border in title_ppr.findall(qn("w:pBdr")):
            title_ppr.remove(border)
    styles["Heading 1"].font.size = Pt(14)
    styles["Heading 1"].paragraph_format.space_before = Pt(13)
    styles["Heading 1"].paragraph_format.space_after = Pt(5)
    styles["Heading 2"].font.size = Pt(11)
    styles["Heading 2"].paragraph_format.space_before = Pt(9)
    styles["Heading 2"].paragraph_format.space_after = Pt(4)
    footer = sec.footer.paragraphs[0]
    footer.text = "WaterFlex  |  Run 331df1f1  |  Methods in companion EVIDENCE.md"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.size = Pt(8)

    doc.add_paragraph("Daily Budget Experiment 331df1f1", "Title")
    paragraph(doc, "Modeled cost, fairness, and search throughput across 1,280 synthetic daily cases")
    paragraph(doc, "Decision for review: keep Tabu as the control while investigating alternatives by workload. All three alternatives reduce accepted modeled cost on average across their 320 matched cases, yet every one loses to Tabu in the 50 technician dispersed group. This archive supports a targeted follow-up, not a blanket solver switch or a production savings estimate.")

    doc.add_heading("Results at a glance", 1)
    paragraph(doc, "Each comparison holds the fixture, fleet, workload, search seed, and time budget fixed. Savings mean Tabu accepted cost less alternative accepted cost. A positive number favors the alternative. The accepted schedule is the outcome relevant to a dispatcher; the lower reference cost is a search milestone before the fairness phase.")
    summary_rows = []
    for solver, label in (("TABU", "Tabu control"), ("KOPT", "K opt"), ("LATE_ACCEPTANCE", "Late acceptance"), ("SUBLIST", "Sublist")):
        a = grouped[solver]
        savings = "Control" if solver == "TABU" else money(mean(int(c["accepted_cost"]) - int(r["accepted_cost"]) for r, c in paired[solver]))
        wins = "Control" if solver == "TABU" else f"{sum(int(r['accepted_cost']) < int(c['accepted_cost']) for r,c in paired[solver])} / 320"
        summary_rows.append((label, money(mean(int(r["accepted_cost"]) for r in a)), savings, wins,
                             f"{mean(float(r['fairness']) for r in a):.4f}"))
    table(doc, ("Solver", "Mean accepted cost", "Mean saving vs Tabu", "Cost wins", "Mean fairness variance"),
          summary_rows, (1.15, 1.36, 1.45, .85, 1.96))
    paragraph(doc, "Fairness variance is lower when paid workload is more evenly spread relative to technician capacity. These overall means mix different fleet sizes; compare fairness within a matched case or group. Tabu has the lowest overall mean variance, so average cost gains do not also mean average fairness gains.")

    doc.add_heading("Where the cost changes", 1)
    paragraph(doc, "The largest consistent gains occur in dispersed fleets of 10 and 20. The 50 technician dispersed group reverses the result in all ten seeds at all four budgets. The table gives mean accepted cost difference per matched case, in dollars. Each cell contains 40 pairings.")
    group_rows = []
    for fleet in (5, 10, 20, 50):
        for workload in ("CLUSTERED", "DISPERSED"):
            vals = []
            for solver in ("KOPT", "LATE_ACCEPTANCE", "SUBLIST"):
                pairs = [(r,c) for r,c in paired[solver] if int(r["fleet"]) == fleet and r["workload"] == workload]
                vals.append(money(mean(int(c["accepted_cost"]) - int(r["accepted_cost"]) for r,c in pairs)))
            group_rows.append((str(fleet), workload.title(), *vals))
    table(doc, ("Techs", "Workload", "K opt", "Late acceptance", "Sublist"), group_rows,
          (.7, 1.4, 1.55, 1.7, 1.55))
    paragraph(doc, "A negative value means the alternative costs more. In the 50 technician dispersed group, mean added accepted cost is $52.10 for K opt, $46.85 for late acceptance, and $47.40 for sublist. Tabu's mean accepted cost there is $4,476.81. The corresponding alternatives range from $4,523.67 to $4,528.91. This exception outweighs many smaller favorable cases in the pooled average.")

    doc.add_heading("Fairness and the accepted schedule", 1)
    paragraph(doc, "The daily search first discovers a low cost reference, then spends the remaining budget on fairness under an acceptance ceiling. A reference is therefore not the schedule finally accepted. Across all 1,280 cases, accepted cost averages $35.23 above its own reference. Reporting reference cost as if it were the accepted outcome would overstate the realized modeled cost gain.")
    paragraph(doc, "For a concrete archive case, the five technician clustered fixture at seed 17 and 60 seconds has a Tabu reference of $385.54 and accepted cost of $393.13. K opt accepts $389.44 and late acceptance $387.99. Their fairness variances are 0.0141 and 0.0141, respectively, versus Tabu's 0.0096. Both alternatives cost less in that case but distribute capacity adjusted paid work less evenly. The exact five technician case illustrates the tradeoff; it is not a fleet wide prediction.")
    paragraph(doc, "In the 50 technician dispersed group the cost loser can have better fairness. Mean accepted variance is 0.0225 for Tabu, 0.0132 for K opt, 0.0177 for late acceptance, and 0.0135 for sublist. A choice in this group needs an explicit willingness to pay roughly $47 to $52 more modeled cost per case for that balance. Policy still evaluates each actual proposal, rather than substituting group averages.")

    doc.add_heading("Time and search effort", 1)
    paragraph(doc, "The four wall clock budgets are 60, 90, 120, and 240 seconds. Every result records time limit termination for both search phases; elapsed times run from 60.008 to 240.053 seconds. The archive ran ten cases in parallel, with each JVM restricted to two active processors. This makes elapsed time a result for this contended setup, not a stand alone latency promise.")
    budget_rows = []
    for budget in (60000, 90000, 120000, 240000):
        values = []
        for solver in ("KOPT", "LATE_ACCEPTANCE", "SUBLIST"):
            pairs = [(r,c) for r,c in paired[solver] if int(r["budget_ms"]) == budget]
            values.append(money(mean(int(c["accepted_cost"]) - int(r["accepted_cost"]) for r,c in pairs)))
        budget_rows.append((str(budget//1000), *values))
    table(doc, ("Budget seconds", "K opt", "Late acceptance", "Sublist"), budget_rows,
          (1.55, 1.7, 1.85, 1.8))
    paragraph(doc, "Each budget row pools 80 matched cases per alternative. Average savings grow with the budget in this archive, but individual cases vary: K opt ranges from a $96.91 loss to a $51.14 gain against Tabu. The reference phase evaluates roughly 117,000 to 122,000 moves per second on average across solvers; the fairness phase roughly 80,600 to 82,900. Throughput is lower in the larger fleets and cannot by itself establish better schedules or scale to production hardware.")

    doc.add_heading("What this run supports", 1)
    paragraph(doc, "The 1,280 cases are four solvers by four fleet sizes by two workload shapes by ten search seeds by four budgets. They use eight fixed synthetic geography fixtures. The seeds change search paths on those fixtures; they do not sample ten customer territories. All 1,280 cases are paired, have accepted cost, reference cost, fairness, and elapsed time, and report zero violations. The archive records no failed cases. Every case reached both phase time limits. There is no own 15 second improvement measure in this archive, so no claim relies on it.")
    paragraph(doc, "A reasonable next experiment isolates the 50 technician dispersed failure mode and tests realistic territory, skill, appointment, and routing mixes. A pilot should measure admission, CPU contention, and decision latency before any default change. Nothing here measures payroll savings, customer outcomes, or production capacity.")

    assert "\u2014" not in "\n".join(p.text for p in doc.paragraphs)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
