"""Build "The WaterFlex scheduler: a briefing" for the President of WaterFlex from the validated sources. Install requirements.txt first.

The briefing tells the whole project at a technical level for a reader strong in finance and economics: what the scheduler does and how it scores
plans, how the booking and daily methods were tested and chosen (with the statistics behind each choice), what the performance work bought, and what
comes next as decisions. Every chart carries two captions, how to read it and what it means for WaterFlex. Every number is computed from the evidence
at build time, or read from the shipped VM configuration, except the quoted results EVIDENCE.md lists.
"""
import json
import re
import sys
from copy import deepcopy
from datetime import datetime
from hashlib import sha256
from pathlib import Path
from statistics import mean

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
sys.path.append(str(HERE.parent / "scheduler-explained-2026-10-14"))

from charts import by_label, by_run, make_charts, saving, wtl  # noqa: E402
from report_data import load  # noqa: E402

REFERENCE = HERE.parent / "WaterFlex_Scheduler_Algorithm_Comparison_Explained.docx"
REFERENCE_SHA = "5366f3017d6f18c15de295ba5fb9c1a014b6222911040f1c5e044845e57f832d"
OUTPUT = HERE / "WaterFlex_Scheduler_Briefing_2026-10-22.docx"
CONFIRM = "POST /api/v1/booking/holds/{holdId}/confirm 200"
DASHES = re.compile("[‒–—―]")


def money(v, cents=True):
    s = f"${abs(v):,.2f}" if cents else f"${abs(v):,.0f}"
    return "-" + s if v < 0 else s


def pct(v):
    return f"{v:.0f}%"


def build():
    assert sha256(REFERENCE.read_bytes()).hexdigest() == REFERENCE_SHA
    data = load()
    rec = make_charts(data, HERE / "generated")
    perf, runs, extras = data["perf"], data["perf"]["runs"], data["perf"]["extras"]
    per_minute = perf["per_minute"]
    F, M, L, C = data["daily"]["first"]["rows"], data["daily"]["follow"]["rows"], data["daily"]["follow"]["long"], data["daily"]["follow"]["confirm"]
    b24, f29 = data["booking24"], data["field29"]
    hd, ex = rec["headline"], rec["example_day"]

    doc = Document(REFERENCE)
    table_properties = deepcopy(doc.tables[0]._tbl.tblPr)
    for child in list(doc.element.body):
        if child.tag != qn("w:sectPr"):
            doc.element.body.remove(child)
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype.endswith(("/image", "/hyperlink")):
            del doc.part.rels[rid]
    doc.core_properties.title = "The WaterFlex scheduler: a briefing"
    doc.core_properties.subject = "What the scheduler does, how its booking and routing methods were tested, and what the performance work bought"
    doc.core_properties.author = doc.core_properties.last_modified_by = "WaterFlex"
    doc.core_properties.created = doc.core_properties.modified = datetime(2026, 10, 22)
    figures = []

    def clean(text):
        assert not DASHES.search(text), text[:60]
        return text

    def p(text, style=None, size=None):
        para = doc.add_paragraph(clean(text), style)
        para.paragraph_format.widow_control = True
        if size:
            for run in para.runs:
                run.font.size = Pt(size)
        return para

    def lead(head, text, size=None):
        para = doc.add_paragraph()
        para.paragraph_format.widow_control = True
        a = para.add_run(clean(head) + " "); a.bold = True
        b = para.add_run(clean(text))
        if size:
            a.font.size = b.font.size = Pt(size)
        return para

    def part(text):
        h = p(text, "Heading 1"); h.paragraph_format.keep_with_next = True; h.paragraph_format.page_break_before = True

    sections = []

    def heading(text):
        """A numbered section heading; the number is the section's place in the document, whatever the source text says."""
        sections.append(re.sub(r"^\d+ ", "", text))
        h = p(f"{len(sections)} {sections[-1]}", "Heading 2"); h.paragraph_format.keep_with_next = True

    def shade(cell, fill):
        props = cell._tc.get_or_add_tcPr()
        s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:fill"), fill); props.append(s)
        margins = OxmlElement("w:tcMar")
        for side in ("top", "left", "bottom", "right"):
            n = OxmlElement("w:" + side); n.set(qn("w:w"), "90"); n.set(qn("w:type"), "dxa"); margins.append(n)
        props.append(margins)

    def table(headers, rows, widths, size=9):
        assert abs(sum(widths) - 6.9) < .01, widths
        t = doc.add_table(rows=1, cols=len(headers)); t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t._tbl.replace(t._tbl.tblPr, deepcopy(table_properties))
        for col, w in zip(t.columns, widths):
            col.width = Inches(w)
        for i, values in enumerate([headers, *rows]):
            row = t.rows[0] if i == 0 else t.add_row()
            trpr = row._tr.get_or_add_trPr(); trpr.append(OxmlElement("w:cantSplit"))
            if i == 0:
                trpr.append(OxmlElement("w:tblHeader"))
            for j, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
                cell.width = Inches(width); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell.text = clean(str(value))
                shade(cell, "DCE6EF" if i == 0 else ("FFFFFF" if i % 2 else "F4F6F8"))
                para = cell.paragraphs[0]; para.paragraph_format.space_after = Pt(2); para.paragraph_format.space_before = Pt(2); para.paragraph_format.line_spacing = 1.0
                if j > 0 and len(headers) > 2:
                    para.alignment = 1
                if len(rows) <= 12 and i < len(rows):
                    para.paragraph_format.keep_with_next = True
                for run in para.runs:
                    run.font.name = "Calibri"; run.font.size = Pt(size); run.bold = i == 0
        spacer = p(""); spacer.paragraph_format.space_after = Pt(0); spacer.paragraph_format.line_spacing = Pt(3)

    def box(title, paragraphs):
        """A shaded "How it works" box: one cell, a bold title and plain paragraphs."""
        t = doc.add_table(rows=1, cols=1); t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t._tbl.replace(t._tbl.tblPr, deepcopy(table_properties))
        t.columns[0].width = Inches(6.9)
        row = t.rows[0]; row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        cell = row.cells[0]; cell.width = Inches(6.9); shade(cell, "F4F6F8")
        first = cell.paragraphs[0]; r = first.add_run(clean(title)); r.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = RGBColor(0x28, 0x7A, 0x4C)
        first.paragraph_format.keep_with_next = True
        for text in paragraphs:
            para = cell.add_paragraph(); run = para.add_run(clean(text)); run.font.size = Pt(10)
            para.paragraph_format.space_after = Pt(3)
        spacer = p(""); spacer.paragraph_format.space_after = Pt(0); spacer.paragraph_format.line_spacing = Pt(4)

    # Charts with open space print a little narrower so two fit on a page; labels stay above 5 pt.
    narrow = {"two_phase", "daily_heatmap", "wtl", "convergence", "out_of_sample", "engine", "holds", "time_split", "harness_vs_stack", "fairness_trade", "diminishing"}

    def figure(name, title, read, means):
        assert name not in figures, name
        figures.append(name)
        para = p(""); para.alignment = 1
        shape = para.add_run().add_picture(str(HERE / "generated" / f"{name}.png"), width=Inches(5.6 if name in narrow else 6.85))
        shape._inline.docPr.set("descr", clean(f"Figure {len(figures)}. {title} {read}"))
        para.paragraph_format.keep_with_next = True
        cap = p(f"Figure {len(figures)}. {title}", "Caption"); cap.paragraph_format.keep_with_next = True; cap.paragraph_format.space_after = Pt(2)
        for run in cap.runs:
            run.font.color.rgb = RGBColor(0, 0, 0); run.bold = True; run.italic = False
        r = lead("How to read this chart.", read, 9.5); r.paragraph_format.keep_with_next = True; r.paragraph_format.space_after = Pt(2)
        m = lead("What it means for WaterFlex.", means, 9.5); m.paragraph_format.space_after = Pt(9)

    # ---- The values the narrative uses, with the facts it states checked first --------------------------------------------
    b5 = by_label(runs["baseline"], "booking-cold-c5")
    vm_b, jc_b = by_run(runs["vm"], "base-spread2", "booking"), by_run(runs["routing"], "j-ch", "booking")
    final = by_run(runs["priority"], "j2-slot", "booking")
    ab_before, ab_after = extras["abComparison"]
    tk_first, tk_follow = [r for r in F if r["solver"] == "TABU_KOPT"], [r for r in M if r["solver"] == "TABU_KOPT"]
    s3_follow = [r for r in M if r["solver"] == "TABU_KOPT_SIZE_3"]
    twins = {r["key"][:-1]: r for r in tk_follow}
    direct = [twins[r["key"][:-1]]["accepted_cost"] - r["accepted_cost"] for r in s3_follow]
    w_first, w_follow = wtl(tk_first), wtl(tk_follow)
    late50 = mean(saving(F, s, fleet=50) for s in ("KOPT", "SUBLIST", "LATE_ACCEPTANCE"))
    held = {v: f29["booking"][("held", v)] for v in ("INSERTION", "BOUNDED", "SHARED")}
    lat = b24["latency"]
    fr = data["config"]["frontier"]
    closes_2000 = sum(next(e["gap"] for e in fr[f] if e["cap_ms"] == 250) - next(e["gap"] for e in fr[f] if e["cap_ms"] == 2000) for f in fr)
    closes_3000 = sum(next(e["gap"] for e in fr[f] if e["cap_ms"] == 250) - next(e["gap"] for e in fr[f] if e["cap_ms"] == 3000) for f in fr)
    lay = data["config"]["layouts"]
    best = {v: max(x["perMinute"] for x in lay.values() if x["vcpu"] == v) for v in (2, 3, 4)}
    fm15, fm60 = data["fairness"]["summary"]["15000"], data["fairness"]["summary"]["60000"]
    nb = data["numbering"]["summary"]["b"]
    fc15, fc60 = fm15["c"]["cost"], fm60["c"]["fairness"]
    assert fc15["medianDifferenceCents"] > 0 and fc60["lessFairBeyondNoise"] == 1
    fix = data["numbering"]["fix"]["baseline"]
    prof = extras["profiling"]["summary"]
    stack_b, shipped = data["stack"]["b"], data["stack"]["shipped"]
    s3 = data["quoted"]["study3"]
    dim = rec["diminishing"]
    year = rec["illustrative_year"]
    allowance = (ex["after"]["costCents"] - ex["reference"]["costCents"]) / ex["reference"]["costCents"] * 100
    confirm_before, confirm_after = jc_b["ops"]["confirm"]["ms"]["p50"], final["ops"]["confirm"]["ms"]["p50"]
    split = rec["time_split"]
    oos_all = rec["out_of_sample"]["all"]
    st_first, st_follow = oos_all["first"], oos_all["follow"]
    hist = rec["paired_hist"]["stats"]
    noise = {fl: hist[fl]["sd_cost"] / hist[fl]["sd_diff"] for fl in hist}
    conv = rec["convergence"]["best_fraction"]
    stair = rec["booking_staircase"]

    assert st_first["lo"] > 0 and st_follow["lo"] > 0 and st_first["sign_p"] < 1e-4 and st_follow["sign_p"] < 1e-4
    assert all(s["lo"] > 0 for s in rec["out_of_sample"]["follow"]) and all(s["lo"] > 0 for s in rec["out_of_sample"]["first"] if s)
    assert min(noise.values()) > 10
    assert all(conv[s][-1] > 0.9 for s in ("KOPT", "LATE_ACCEPTANCE")) and all(conv[s][-1] < 0.75 for s in ("TABU", "TABU_KOPT"))
    assert stair[-1]["p99"] < 1000 and stair[0]["p95"] > 5500
    assert abs(hist["50"]["median"] - hist["50"]["mean"]) < 1 and sum(r["saving"] < 0 for r in tk_first) / len(tk_first) < 0.15
    assert st_follow["days"] == 120 and len({r["key"][:-2] for r in tk_follow}) == 120

    assert all(saving(F, "TABU_KOPT", fleet=f) > 0 for f in (5, 10, 20, 50)) and all(saving(M, "TABU_KOPT", fleet=f) > 0 for f in (5, 10, 20, 30, 50))
    assert all(saving(F, s, fleet=50) < 0 for s in ("KOPT", "SUBLIST", "LATE_ACCEPTANCE"))
    assert held["BOUNDED"]["served"] > held["INSERTION"]["served"] > held["SHARED"]["served"]
    assert all(e["after_lost"] == 0 for e in rec["holds"])
    assert ex["after"]["costCents"] < ex["before"]["costCents"] and ex["after"]["variance"] < ex["before"]["variance"] < ex["reference"]["variance"] and allowance <= 2.0
    assert all(b > a for a, b in zip(dim["extra"]["50"][1:], dim["extra"]["50"][2:])) and all(v[-1] >= 0 for v in dim["extra"].values())
    assert fm15["b"]["fairness"]["lessFair"] == 0 and fm60["b"]["fairness"]["lessFair"] == 0 and not data["fairness"]["decision"]["c"]["promoted"]
    assert best[4] > best[3] > best[2] and lay["2x1"]["perMinute"] > lay["1x2"]["perMinute"]
    assert all(v["REFERENCE"] > 1.1 for v in shipped.values()) and all(v["FAIRNESS"] > 2 for v in stack_b.values())
    assert 0 < closes_2000 < closes_3000
    gap = {f: {e["cap_ms"]: e["gap"] for e in fr[f]} for f in fr}
    assert all(g[0] - g[250] < 0.8 and g[1000] - g[3000] > g[0] - g[1000] for g in gap.values())
    assert sum(direct) > 0

    # ---- Front matter -----------------------------------------------------------------------------------------------------
    p("The WaterFlex scheduler: a briefing", "Title")
    p("A technical briefing: what it does, how its booking and routing methods were tested and chosen, and what the performance work bought", "Subtitle")
    p("October 22 2026  |  Prepared for the President of WaterFlex  |  Covers the work from the September demonstration to today")
    p("This briefing explains the scheduler project at a technical level, with the business consequence kept in view. It has three aims: to say what the scheduler does for WaterFlex and how it decides, to show how the competing booking and routing methods were tested and why the chosen ones won, and to show what the performance work bought. It goes inside the machinery: the scoring rule the searches optimize, the algorithms and their settings, the statistics behind each decision, and the way the server admits and times its work.")
    p("Every chart has two short captions. The first says how to read it: what an axis, a bar or a dot stands for. The second says what it means for WaterFlex. Shaded \"How it works\" boxes carry the formulas and settings; a reader who wants only the argument can skip them without losing the thread.")
    p("Three cautions apply throughout, and they are repeated where they matter. First, the dollar figures are modeled operating cost: technicians' paid time at their hourly rates plus mileage, computed by the scheduler. They are not billed revenue or payroll. Second, the daily-plan studies run on synthetic days built to resemble WaterFlex work, not on recorded customer days; the booking tests use random Omaha addresses on the real Omaha road map. Third, each result is read against its own study's control. Results from different studies are never added together, because each study used its own test days and its own baseline.")

    part("Summary")
    figure("headline", "Five results in one page",
           "Each card is one result. The large number is the before and after, or the gain. The line beneath says what was measured and against what. The green card is the result that is worth money every day; the others are operating results.",
           "The scheduler books customers faster and loses none of their chosen windows, offers more appointments when the calendar is full, and plans each day for less modeled cost than the method it replaced.")
    lead("What the scheduler is.", "A separate service, on its own server, that WaterFlex Software calls for two jobs: finding arrival windows while a customer is on the phone, and rearranging a booked day overnight so it costs less without breaking any promise. It keeps no copy of WaterFlex's records.")
    lead("What was tested.", f"Three booking methods and eight daily-planning methods were compared on matched synthetic days. The comparisons ran to {len(F) + len(M) + len(L) + len(C) + len(data['daily']['first']['confirm']):,} daily-plan cases and {b24['requests']:,} September booking requests. Two configuration studies tuned the booking search, and a series of load and engine studies measured speed.")
    lead("What was chosen.", f"Bounded rearrangement for booking. It served the most customers in the September test. Tabu + K opt for the daily plan: it saved {money(hd['tk_saving'])} a day on average over the old method and was the only method that saved money at every fleet size. It held up on 40 new test fixtures it had never seen.")
    lead("How sure we are.", f"The daily choice is not a close call. Tabu + K opt's mean saving has a 95 percent interval of {money(st_first['lo'])} to {money(st_first['hi'])} a day in the study that chose it and {money(st_follow['lo'])} to {money(st_follow['hi'])} on the new days. It came out ahead on {st_first['day_wins']} of {st_first['days']} synthetic days and behind on {st_first['day_losses']}, and on the new days {st_follow['day_wins']} of {st_follow['days']} and {st_follow['day_losses']}. Each split would arise by chance less than once in ten thousand times.")
    lead("What the performance work bought.", f"Booking capacity at five desks went from {hd['vm_before']} to {hd['vm_after']:,} searches a minute on the rehearsal stack. Lost customer selections went from {pct(hd['lost_before'])} to none. The engine now tries {hd['speed_lo']:.2f} to {hd['speed_hi']:.2f} times as many schedule changes per second as it did a week ago, measured through the public interface.")
    lead("What remains.", "Measure on the production server across the real network, replay real days, finish the integration on the WaterFlex Software side, and encrypt the traffic. Part VII lists these as decisions.")

    # ---- Part I -------------------------------------------------------------------------------------------------------------
    part("Part I  What the scheduler is, in business terms")
    heading("1 Two jobs")
    p("The scheduler does two different jobs with different economics. Booking is a high-volume, low-latency job: many desks search at once, and every search must answer in seconds while a customer waits. The daily plan is a low-volume, compute-heavy job: one metro day at a time, given minutes rather than seconds, rearranging work that is already booked to make it cheaper.")
    figure("two_jobs", "The scheduler's two jobs",
           "Read each row from left to right. The pale boxes are WaterFlex Software and its users; the blue boxes are the scheduler's work; the green box is the result.",
           "Booking protects revenue: a customer gets an answer in under a second and is never double booked. The daily plan protects margin: it lowers the cost of work already sold without moving any promised window.")

    heading("2 What cost means, shown on one day")
    p("Every comparison in this briefing is about cost, so it is worth seeing what cost is. The scheduler prices a day as each technician's paid minutes at the regular hourly rate, any overtime minutes at the overtime rate, and the miles driven at the mileage rate, rounded once to the cent. Paid time covers service, driving and any paid waiting, including the drive home at the end of the day. Overtime is never added by either job. Separately, it measures fairness: how evenly the paid work is spread across technicians.")
    box("How it works: the two numbers every plan is judged on", [
        "Modeled cost. For a whole metro day, cost = regular paid minutes x regular rate / 60 + overtime minutes x overtime rate / 60 + miles x mileage rate. The sum is formed over the whole fleet and rounded once, half up, to the cent, so no rounding error can accumulate technician by technician. Each rate has at most four decimal places, hourly rates are at most $500 and mileage at most $25 a mile; the engine relies on those limits to do the sum exactly in 64-bit whole numbers.",
        "Fairness. Each technician's utilization u is paid minutes divided by available regular minutes c. The fleet's utilization u-bar is total paid minutes over total available minutes. Unfairness is the capacity-weighted variance: the sum of c x (u - u-bar) squared, divided by the sum of c. A technician with a longer day counts for more, and an idle technician who is qualified for the day's work still counts, so leaving someone idle is visible. Zero means every technician is equally loaded.",
    ])
    figure("example_day", "One synthetic day, before and after",
           "Each group of three bars is one technician. Grey is the day as WaterFlex sent it, blue is the plan after the cost search, and green is the plan finally proposed. The dashed line is a full regular day. The cards on the right give each plan's modeled cost and its variance (lower is more even).",
           f"On this day the cost search found {money((ex['before']['costCents'] - ex['reference']['costCents']) / 100)} of savings but loaded four technicians and idled two. The fairness search then spent {money((ex['after']['costCents'] - ex['reference']['costCents']) / 100)} of it, {allowance:.1f} percent, inside the 2 percent allowance the client sets, to spread the work evenly. The plan proposed is still {money((ex['before']['costCents'] - ex['after']['costCents']) / 100)} cheaper than the day as sent.")
    box("How it works: the fairness allowance", [
        "The daily plan runs two searches in sequence. The first, the cost search, finds the cheapest arrangement it can, with overtime minimized first. The second, the fairness search, keeps that overtime and may spend up to a set share of the cost the first one reached, 2 percent by default and chosen by each client: its ceiling is the cost search's result x 1.02, rounded down to the cent. Note that the ceiling is measured from the cheaper plan, not from the day as sent.",
        "In economic terms the allowance is a budget for equity: a client can buy a more even workload at a known, capped price. Booking uses no allowance. It ranks offers by added cost alone, and fairness only breaks exact ties.",
    ])

    heading("3 How a search ranks two plans")
    p("Searching means comparing millions of candidate plans, so the comparison rule is the heart of the system. The scheduler does not add costs and penalties into one number with weights. It uses a lexicographic score: seven levels compared in a fixed order, the way a dictionary orders words. Two plans are compared on the first level; only if they tie there does the second level count, and so on. No saving on a lower level can make up for a worse value on a higher one, so no amount of money can buy a late arrival.")
    figure("score_ladder", "How the daily search ranks two plans",
           "Read from the top. Each band is one level of the score, with what it counts and its unit. Blue bands are hard levels, promises and limits that must all be zero for a plan to be usable. Pale bands are soft levels, what the search minimizes once the hard levels are clear. This is a schematic and carries no data.",
           "The ordering is a policy written into the arithmetic: promises and legal limits come before money, and money comes before nothing. A weighted sum would need someone to choose how many dollars a late arrival is worth; this design never asks that question.")
    p("Two details make the levels exact rather than approximate. Every level is a whole number: seconds, minutes, cents or counts, and the spread of work is carried in units of one billionth. And the search's fast number is not the final word. When a search ends, its plan is re-scored from scratch by a separate evaluator using 34-digit decimal arithmetic, and only then judged by the acceptance rule. That rule accepts a plan only if it is independently feasible, adds no overtime, stays inside the cost ceiling, and improves overtime, fairness or cost. A cheaper plan does not displace a fairer one that is already inside the allowance. Every outcome carries a reason code, so a planner can see why a proposal was accepted or refused.")

    heading("4 Where it sits")
    figure("architecture", "Where the scheduler sits",
           "Boxes are services and arrows are calls. WaterFlex Software on the left sends the facts of a call to one address. Inside the dashed boundary, the scheduler's own server sends booking calls to one pool and daily plans to another.",
           "The scheduler is an add-on, not a second system of record. WaterFlex Software keeps every customer, job and technician; the scheduler keeps only what it needs to keep promises safe. Booking and overnight planning run on separate pools, so a long overnight plan never slows a customer on the phone.")
    p("The server is a virtual machine of 8 virtual CPUs and 32 GB of memory, reached only through the published interface (/api/v1). Many WaterFlex clients can share it: each client is a separate tenant with its own keys, and the database itself refuses to show one tenant's rows to another. The scheduler is stateless in the sense that matters: every call carries the facts it needs (technicians, jobs, rates, the client's policy settings), so a replica can be added or replaced without moving data, and the scheduler never holds policy settings of its own.")
    p("A front proxy sends daily-plan calls to the daily pool and everything else to the booking pool, and spreads each pool's calls over its replicas. The road map service answers drive times from the real Omaha road network; a separate address service turns addresses into map points. Part VI shows how the 8 virtual CPUs are divided among these parts.")

    heading("5 Why nothing is double booked")
    figure("booking_life", "A booking's life, and why nothing is double booked",
           "Follow the arrows from the search to the booked appointment. The diamond is the one check that matters; the dashed line is what happens if it fails.",
           "Two desks can never sell the same slot. When two desks collide, the cost is one more search measured in fractions of a second, not an angry customer.")
    p("The check rests on a timestamp WaterFlex Software already keeps for each technician and day. The scheduler compares it exactly and never guesses. A hold reserves a placeholder in a common arrangement of the day that every active hold and confirmed appointment must fit together, so two holds can never both rely on the same minutes. A confirm takes locks on the days it touches in a fixed order, which rules out two confirms each waiting on the other. It then re-validates the whole day independently and writes the assignment, route order and times in one transaction. A repeated confirm returns the existing appointment rather than booking twice, so a desk that times out can simply resend.")
    p("The scheduler also remembers the last 32 timestamps each day has moved past. That lets it tell an old snapshot arriving late from a genuine edit made in WaterFlex Software. Part V shows the measurement that made this necessary.")

    # ---- Part II ------------------------------------------------------------------------------------------------------------
    part("Part II  How we know: the evidence method")
    heading("5 Controlled experiments, decided in advance")
    p("Every choice in this briefing was made the same way, and the method will be familiar from controlled experiments. A study runs today's method and the candidate on the same synthetic days, with the same random seeds and the same seconds of computer time. Each case is then compared with its own twin, so a hard day is compared with the same hard day. The decision rule is written down before the run. For a speed change, the rule is a median gain of at least 1.05 times on the part of the work it targets, and no slowdown elsewhere. The result therefore cannot move the goalposts.")
    figure("paired_study", "How one study decides",
           "Read from top to bottom. Both sides get identical inputs, each case is paired with its twin, and the median of the pairs is held against the bar. A change that passes is run again through the public interface on a copy of the production setup before it ships.",
           "A promoted change has passed twice: once in a controlled experiment, and once in conditions close to production. This is the same discipline as pre-registration and out-of-sample testing, applied to software.")
    p("Two further habits matter. When a study chooses a winner, a follow-up tests that winner on new days it has never seen, the equivalent of out-of-sample validation. And every number in a report is recomputed from the raw files by a separate checking program before the report is built.")
    p(f"Pairing is what makes small differences visible. At 50 technicians the modeled cost of a synthetic day under the old method varies from day to day with a standard deviation of {money(hist['50']['sd_cost'], False)}. The paired difference between the two methods on the same day varies with a standard deviation of only {money(hist['50']['sd_diff'])}. Comparing each day with its twin removes the day-to-day variation, {noise['50']:.0f} times over at that size and at least {min(noise.values()):.0f} times at every size, which is why a saving of a few dollars on days costing hundreds can be detected at all.")
    box("How it works: the statistics in this briefing", [
        "Medians for speed. Speed is reported as the median of paired ratios (after over before, on the same day and seed). Ratios are skewed, since a few easy days can run several times faster, so the median is the honest summary; the mean of ratios would overstate.",
        "Intervals that respect the design. A 95 percent interval for a mean saving is found by the bootstrap: the study is resampled 10,000 times and the mean recomputed each time. Each synthetic day was run at several time budgets, and those runs are not independent, so the resampling draws whole days with all their runs. Resampling single runs would make the intervals look narrower than the evidence supports.",
        "Sign tests on days. A sign test asks how likely a split of wins and losses at least this lopsided would be if the two methods were equally good. It is counted on days (each day's mean saving), not on runs, for the same reason.",
        "Noise floors. Some searches stop on a clock, so two identical runs can differ. Where that matters the noise is measured directly, by running the same case twice, and a difference inside it counts as a tie.",
        "Bars written in advance. A speed change is promoted only if its median ratio reaches 1.05 on the phase it targets and it is no slower on the other. A quality change must also be no dearer and never less fair beyond noise. Writing the bar first is the software equivalent of pre-registration.",
    ])
    heading("6 Five weeks of studies")
    figure("timeline", "Five weeks of studies",
           "Each row is one study, placed on the dates it ran, with its size at the end of the bar. Colors mark the kind of study: booking, daily-plan quality, load and capacity, or engine speed.",
           "The choices rest on many thousands of matched cases, not on a handful of demonstrations. Two sizes are quoted from reports whose raw files are kept elsewhere, and they say so.")
    heading("7 What the evidence cannot say")
    p("The daily studies use synthetic days with a synthetic travel matrix, built to resemble WaterFlex's work in four shapes (clustered, dispersed, tight windows, mixed skills), each generated from a fixed seed so any case can be rebuilt exactly. The booking and load studies use random addresses on the real Omaha road map. Load numbers come from a rehearsal copy of the production setup running on one workstation, not from the production server across the network. The ratios between methods are what transfer; the absolute rates belong to the test machine. None of it is recorded customer traffic, and Part VII names the steps that would close that gap.")

    # ---- Part III -----------------------------------------------------------------------------------------------------------
    part("Part III  The booking methods tested")
    heading("8 Insertion, bounded rearrangement, or a larger search")
    p("A booking search fits one new job into days that are already booked. The simplest method, insertion, tries the job in every gap without moving anyone else. Bounded rearrangement also allows small, limited reshuffles nearby, which can open a better window. Larger searches allow more reshuffling still. The trade is familiar: more search can find cheaper slots and serve more customers, at the price of a longer wait.")
    figure("booking_methods", "Two ways to fit a new job into a booked day",
           "Each row is a technician's route. The left panel places the new visit into a gap without moving anything; the right panel lets nearby visits shift within limits. This is a schematic and carries no data.",
           "Bounded rearrangement is the middle ground: it finds better slots than insertion without the open-ended search time of the larger methods.")
    figure("booking_choice", "Choosing the booking method",
           "Left: each circle is a method, placed by how many of 480 test requests it served (up) and its typical worst-case wait (right, on a log scale where each step is about three times longer). Right: what bounded rearrangement saved per test case over insertion, in paid minutes (blue) and modeled dollars (green), by kind of day.",
           f"Bounded rearrangement served the most customers ({held['BOUNDED']['served']} of 480, against {held['INSERTION']['served']} for insertion). On clustered days it saved {money(b24['by_workload']['CLUSTERED']['cost'])} a case. The larger shared search was slower and served fewer ({held['SHARED']['served']}), so it was set aside, along with two other large searches screened in the same study.")
    p(f"The price of bounded rearrangement was time. One request at a time, its worst-case wait (the 95th percentile, the wait only one search in twenty exceeds) rose from {lat[('INSERTION', 1)]['p95'] / 1000:.2f} s to {lat[('BOUNDED', 1)]['p95'] / 1000:.2f} s in September. The later work made both faster, and Part V shows how. The decision was to keep bounded rearrangement, with insertion as the guaranteed first answer.")
    box("How it works: bounded rearrangement as a beam search", [
        "Insertion tries the new job in every gap of every eligible technician on every date in the horizon, moving nobody. Each placement is timed exactly: arrivals, waiting and the drive home are recomputed so the route keeps every promised window and its paid and overtime limits.",
        "Bounded rearrangement then explores a limited tree of changes on the same date: moving one existing visit to another route, swapping two visits, or reversing a stretch of a route. It considers the six most promising routes, always including the least loaded eligible technician, looks at most two changes deep, and keeps only the best eight partial arrangements at each depth (a beam search of width eight). It stops after 500 arrangements per window. The bounds are what keep the search inside a phone call.",
        "Candidates that add any overtime are discarded. The rest are ranked strictly by the cost they add to the day. Up to four distinct regular windows are offered, each two hours long and held for ten minutes; several technicians able to serve one window count as one choice.",
    ])
    figure("deadline_budget", "How a booking search spends its five seconds",
           "Time runs left to right across the client's deadline, at most five seconds. Each bar is a stage of the search. The insertion pass always runs; the optional polish (dashed) runs only when there are plenty of windows and nobody is waiting; the required refinement runs only when insertion found too few windows. The last second is reserved. Bar widths are schematic except the marked 2, 4 and 5 second points.",
           "The design guarantees an answer inside the deadline whatever the load. A search that runs out of time says so and returns what it has, rather than claiming the calendar is full.")
    p("The deadline covers everything the server does, not just calculation: waiting for a free slot, fetching drive times and talking to the database all come out of the same five seconds. Exploration stops at four seconds so the last second can independently re-check the offers and store the holds. A search that is incomplete reports itself as incomplete, never as no capacity, so WaterFlex Software can tell a busy server from a full calendar. When other desks are waiting, the polish is skipped so nobody queues behind it, and the host is told whether the polish ran.")
    heading("9 How long to polish")
    figure("refinement_frontier", "How long to let the booking search polish",
           "Each line is one fleet size on a fully booked schedule. Moving right along a line, the polish is allowed 0, 0.25, 1, 2 and 3 seconds; each point shows the median search time (across) and the extra modeled cost per booking compared with the best search a five-second deadline allows (up). The grey band is measurement noise: two identical searches differ by up to about $0.80.",
           f"This is a cost-of-time frontier. Two seconds of polish recovers most of what three seconds would ({money(closes_2000)} of {money(closes_3000)} summed over the five fleets) for about a second less per search. Two seconds is now the default; before this study it was a quarter second, which on full schedules bought almost nothing.")
    p("Why is there noise at all? The polish stops on a clock, not on a count of steps, so how far it gets depends on what else the machine is doing at that moment. The study ran identical searches twice to measure that spread directly, and treats any difference inside it as a tie. The shape of the frontier is not the textbook concave curve. On fully booked schedules a quarter second of polish bought nothing measurable, and at every fleet size more of the gap closed after the first second than within it. That is why the old quarter-second setting was a false economy.")
    p("The same study tested the threshold that decides when the polish is required, and found it never fired: the first pass always found enough windows. It also set the time reserved for saving the offers safely at the end of a search. With a half or a quarter second reserved, every long search was refused at the point of publishing, so the reserve may not be set below 0.8 seconds and stays at 1 second by default.")
    heading("10 When the calendar is full")
    figure("horizon", "Booked-out weeks: a horizon that moves",
           "Bars compare the fixed horizon (grey, always the next ten weekdays) with the dynamic horizon (green, which looks past fully booked days, up to fifteen weekdays). Left: the share of searches that offered nothing. Right: windows offered per search, out of four. Three technicians with the first twelve weekdays booked out.",
           f"In a busy week the fixed horizon turned away half of its callers ({pct(hd['empty_fixed'])}) with no window at all. The dynamic horizon cut that to {pct(hd['empty_dynamic'])}, and only once it reached its fifteen-day cap. A caller offered nothing is a sale at risk; this change keeps the calendar selling.")

    # ---- Part IV ------------------------------------------------------------------------------------------------------------
    part("Part IV  The daily-plan methods tested")
    heading("11 The contenders")
    p("Every method tested is a local search, run by the open-source Timefold solver. It starts from the day as sent, tries a small change, keeps it or not, and repeats until the time runs out. Each change is a move, and each round of choosing one is a step. The methods differ in two things: the moves they may try (the neighborhood) and the rule for accepting a step that makes the plan worse, which is how a search climbs out of a dead end.")
    table(["Method", "Moves tried, by share of draws", "Acceptance rule", "Per step"], [
        ["Tabu (old default)", "Relocate one visit 50%, swap two visits 50%", "Tabu: parts of the plan changed in the last 7 steps may not change again", "Best of up to 1,000 acceptable moves"],
        ["Tabu + K opt (chosen)", "Relocate 45%, swap 45%, move a run of 2 to 4 visits (maybe reversed) 5%, K opt 5%", "Tabu, memory of 7", "Best of up to 1,000 acceptable moves"],
        ["Tabu + K opt size 3", "As Tabu + K opt", "Tabu, memory of 3", "Best of up to 1,000 acceptable moves"],
        ["Late acceptance", "Relocate 50%, swap 50%", "Accept if no worse than the plan 400 steps ago", "First acceptable move"],
        ["K opt, sublist moves", "As Tabu + K opt, or without K opt", "Late acceptance, 400 steps", "First acceptable move"],
    ], [1.35, 2.15, 2.1, 1.3])
    p("Tabu search is a disciplined hill climber: at each step it looks at up to a thousand acceptable moves and takes the best, even if the best makes the plan slightly worse, while a short memory forbids undoing recent changes so it cannot circle. Late acceptance is looser: it takes the first move that is no worse than the plan it held 400 steps earlier, so it wanders more freely and covers ground faster, but each step is less considered. K opt is a larger move: it cuts two or three links in a route and reconnects the pieces another way, which can untangle a route that single relocations cannot. Tabu + K opt combines Tabu's discipline with K opt's reach. Every move is scored incrementally: the engine recalculates only the routes a move touches, plus the fleet-wide sums for cost and fairness, which is why it can try many thousands of moves a second.")
    figure("two_phase", "A daily plan's search: two phases, then a check",
           "The top row is the pipeline: the day as sent, the cost search, the fairness search and the acceptance check. The bottom row shows three kinds of move a search tries, with the moved visits highlighted. This is a schematic and carries no data.",
           "The methods differ only in the moves and the acceptance rule, so a fair comparison gives each the same days and the same seconds and compares the cost they reach.")
    heading("12 Which method saves money")
    figure("daily_heatmap", "Which daily search saves money, by fleet size",
           "Each cell is a method's average saving per day against Tabu, the old default, at one fleet size. Blue cells save money and brown cells cost more; darker means larger. The right column averages all fleet sizes. The chosen method is outlined.",
           f"Tabu + K opt is the only method that saves money at every fleet size, {money(hd['tk_saving'])} a day on average. The late acceptance family was better for mid-size fleets but lost about {money(-late50)} a day at 50 technicians. A blanket switch to it would have taxed WaterFlex's largest customers to subsidize the middle.")
    figure("wtl", "Wins, ties and losses against Tabu, case by case",
           "Each bar counts matched days: to the right, days where the method beat Tabu; to the left, days where it lost. The numbers give wins, ties and losses, and the right column the mean saving. The top panel is the study that chose Tabu + K opt; the bottom panel is the follow-up on new days.",
           f"An average can hide a few big wins among many losses. Here it does not: Tabu + K opt won {w_first[0]} of {len(tk_first)} matched cases and lost {w_first[2]}.")
    figure("paired_distribution", "Every matched day, not just the average",
           "Each panel is one fleet size. Each bar counts matched cases whose saving over Tabu fell in a $2.50 band: blue to the right of the black zero line saved money, brown to the left cost more. The green line is the mean and the green band its 95 percent interval.",
           f"The distributions sit to the right of zero and are not driven by a few outliers: the median saving is close to the mean ({money(hist['50']['median'])} against {money(hist['50']['mean'])} at 50 technicians). The losses are real but small and rare. At 5 technicians there is little to save at all, which is what one expects when a day has few ways to be rearranged.")
    figure("convergence", "When each search found its best plan",
           "Each line is a method. For each fleet size the point shows, on average, how far through the cost search the method was when it found the plan it finally returned. A point near 100 percent means the search was still improving when its time ran out.",
           f"At 50 technicians the late acceptance family found its best plan in the last few percent of its time ({conv['LATE_ACCEPTANCE'][-1] * 100:.0f} percent), so it was still improving when the clock stopped. Tabu and Tabu + K opt had settled by about {conv['TABU_KOPT'][-1] * 100:.0f} percent. That is consistent with why late acceptance lost at 50 technicians: on a large day its looser steps cover ground too slowly for the time given.")
    figure("out_of_sample", "The choice held on days it had never seen",
           f"Each row is a fleet size in one study: blue for the study that chose Tabu + K opt, green for the follow-up on 40 new test fixtures, each run from three new seeds ({st_follow['days']} synthetic days in all). The square is the mean saving per day over Tabu and the whiskers its 95 percent interval, from resampling whole days. On the right, days won and lost, and the sign test.",
           f"No interval touches zero, in or out of sample, at any fleet size. On the new days Tabu + K opt came out ahead on {st_follow['day_wins']} of {st_follow['days']} days. A choice that holds on fresh data, with intervals this far from zero, is not an artifact of the data that chose it.")
    p(f"The follow-up also found a stronger candidate: Tabu + K opt with a shorter memory (size 3). It beat Tabu + K opt by {money(mean(direct))} a day on the new days and won at every fleet size from 10 technicians up, but trailed slightly at 5. Part VII proposes offering it as a client choice before considering it as a default.")
    heading("13 What more search time buys")
    figure("diminishing", "What more search time buys",
           "Each line is one fleet size. It shows how much more Tabu + K opt saves per day when given 90, 120 or 240 seconds instead of 60. The list on the right gives each fleet's saving at 60 seconds over the day as sent, the base the extra is added to.",
           f"Returns diminish quickly. A 50-technician day gains {money(dim['extra']['50'][1])} from the first extra 30 seconds and only {money(dim['extra']['50'][3] - dim['extra']['50'][2])} from the last 120. Small fleets gain almost nothing. Compute time is cheap, but not free: the daily pool's cores are shared by every client's overnight work, so the client limit of 240 seconds is the right ceiling.")
    p("The base figures on the right deserve a caution. They are savings over the synthetic day as sent, which was built by booking jobs one at a time, so they show the scale of what a daily plan can recover rather than a forecast for a real WaterFlex day. The comparison between methods, which is what the decision rested on, does not depend on that base.")
    box("Illustrative only: what the method choice could be worth in a year", [
        f"This box scales one measured number by a stated assumption. It is not a forecast. The measured number is Tabu + K opt's average saving over Tabu per metro day, by fleet size, from the study that chose it ({year['source']}). The assumption is the number of working days a year. Annual value = saving per day x working days, for each metro of that size.",
        "The table below shows three assumptions for working days. Substitute your own; the saving per day is the only measured input. It is modeled operating cost on synthetic days, so a replay of real days (Part VII) is the step that would turn it into an estimate.",
    ])
    days = year["days"]
    table(["Technicians in the metro", "Measured saving per day", *[f"{d} working days" for d in days]],
          [[fl, money(v), *[money(v * d, cents=False) for d in days]] for fl, v in year["by_fleet"].items()], [1.5, 1.5, 1.3, 1.3, 1.3])
    heading("14 Fairness as a budgeted trade-off")
    p("How evenly work is spread is a policy choice, not an engineering fact, so it was tested with care. The fairness study compared the existing way of measuring evenness with two faster formulas and four other measures. The rule written in advance was: promote only if it is faster, never less fair beyond noise, and no dearer. The chosen formula, b, measures the same spread with whole-number arithmetic that a computer evaluates far faster.")
    figure("fairness_trade", "A fairer measure at the same cost",
           "Each dot is one synthetic day. Across: how much the day's modeled cost changed with formula b. Up: the spread of work with formula b as a share of the spread before (below 1.00 is more even). The shaded area is fairer than before. The card summarizes the study.",
           f"Formula b makes the fairness search {nb['fairnessPhaseSpeedup']['median']:.1f} times as productive. Because it can try more rearrangements in the same seconds, {nb['fairness']['fairer']} of 36 days ended more even and none less even, at a median cost change of {money(nb['cost']['medianDifferenceCents'] / 100)}. Formula c was as fast but failed the cost and fairness rules, so it was not shipped.")
    box("How it works: why formula b is faster but measures the same thing", [
        "The previous measure computed the variance in 34-digit decimal arithmetic on every move. That is exact, but decimal arithmetic runs in software, and a profile showed the variance alone took 57 percent of the fairness search's time.",
        "Formula b keeps the definition and changes the arithmetic. The paid-minute and capacity sums are kept as exact whole numbers, the variance is computed from them in the processor's native floating point, and the result is rounded once to whole units of one billionth. A move never accumulates rounding error, because each evaluation starts again from the exact sums.",
        "Is a billionth fine enough? Moving a single paid minute between two technicians with 480-minute days, in a fleet of 50, changes the variance by at least 170 of those units, so any real change is far above the resolution. And acceptance still re-reads the 34-digit variance, so the faster number can change which plan the search returns, never whether a plan is accepted.",
        f"Formula c computed the same variance from running sums updated move by move. It was as fast, but at 15 seconds its median cost was {fc15['medianDifferenceCents']:.0f} cents a day above the baseline, against a bar of no dearer, and at 60 seconds {fc60['lessFairBeyondNoise']} day ended less fair beyond noise. It failed both rules and was not shipped.",
    ])
    p("One alternative measure, load balance, looked fairer still on small and mid-size fleets in one study. Changing the fairness policy is a business decision, so it is listed in Part VII as a head-to-head test rather than adopted.")

    # ---- Part V -------------------------------------------------------------------------------------------------------------
    part("Part V  The performance work, and why speed is money")
    heading("15 Why speed is money")
    p("The daily plan has a fixed time budget, so a faster engine does not finish sooner; it tries more rearrangements in the same seconds and so finds a cheaper plan. On the booking side, speed is capacity: the same server answers more desks, which defers the next hardware purchase. Both are productivity gains on a fixed input.")
    figure("engine", "More moves per second buys cheaper days",
           "Each dot is one paired synthetic day, placed by how many rearrangements per second (in thousands) the engine tried before a change (across) and after it (up). The dotted line is no change; the dashed line is twice as many. Light dots are 20-technician days and dark dots 50-technician days.",
           f"One engine change, road lookup tables, roughly doubled the work per second ({prof['costPhaseSpeedup']['median']:.2f} times). In the same seconds, {prof['afterCostCentsWinTieLoss']['afterWins']} of 24 days ended cheaper, with no change to the search method itself. Speed turned directly into money.")
    heading("16 Booking capacity, one change at a time")
    figure("booking_staircase", "Booking capacity, one change at a time",
           "Bars are booking searches answered per minute by five desks after each change, left to right in the order they were made. Below, each search's typical wait (hollow dot) and worst-case wait, the 95th percentile (filled dot), against the 5.5 second target. The dotted line divides two test setups, so compare bars only within each half.",
           f"In the second half, on the copy of the production setup, capacity rose from {hd['vm_before']} to {hd['vm_after']:,} searches a minute with the worst-case wait well under a second. In the first half the worst-case wait fell from over six seconds, outside the target, to about one.")
    p("Each step answered the question the previous step raised. Lost holds came first, then a status flag that had misreported complete searches as incomplete. Next came the optional polish, which was holding a slot while other desks waited. Then routing, and finally the order in which waiting work is served.")
    p("The median hides the tail, and the tail is what a caller remembers. The table gives the full latency profile at five desks for each stage: the median, the 95th and 99th percentiles, and the single slowest request. In risk terms the 95th percentile is the booking equivalent of value at risk: the wait that only one search in twenty exceeds.")
    table(["Stage (five desks)", "Search p50", "p95", "p99", "Slowest", "Confirm p50", "Confirm p99"],
          [[" ".join(e["name"]) + (" *" if e["half"] == 2 else ""), f"{e['p50']:,} ms", f"{e['p95']:,} ms", f"{e['p99']:,} ms", f"{e['max']:,} ms",
            f"{e['confirm']['p50']:,} ms", f"{e['confirm']['p99']:,} ms"] for e in stair], [1.6, 0.85, 0.85, 0.85, 0.95, 0.9, 0.9])
    p("Stages marked * ran on the rehearsal stack beside five overnight planners; the others ran the scheduler alone on one machine. Compare within each group only. The slowest request is a single observation and moves from run to run; the percentiles are the stable measures.", size=9.5)
    figure("admission", "How the booking pool admits work",
           "The dashed box is one booking replica. Requests wait in a line on the left and run in slots on the right. Confirms (green) are taken before any waiting search and have a slot of their own. Below, what happens when the line is full, and the daily pool, which runs separately. This is a schematic and carries no data.",
           "Admission control is what keeps the tail short under load. A server that accepted every request at once would slow every request down; this one runs as many searches as it has cores, keeps a short line, and refuses the excess at once so a desk can retry, instead of letting every caller wait.")
    p("The slots are set to the replica's cores because each search's calculation runs on one thread: a third search on two cores would compete for the same cores and make all three slower. The line holds at most sixteen requests, and a request that finds it full is refused immediately with a busy answer and a suggested retry time. Confirms jump the line because they finish a sale the customer has already made and take only milliseconds; before that change a confirm could wait behind several five-second searches. Overnight daily plans run on a separate pool, so they never take a slot a caller needs.")
    figure("holds", "Lost holds, before and after the fix",
           "For one, five and ten desks, the brown bar is the share of customers' chosen windows that were lost when the desk confirmed, in the first load test; the green bar is the same share after the fix.",
           f"At five desks more than a third of chosen windows ({pct(hd['lost_before'])}) were being lost at confirm, a customer told one thing and given another. The cause was a race the design had not anticipated: an old snapshot arriving late looked like a fresh edit. The fix recognizes old snapshots and asks the desk to search again. No hold has been lost in any run since.")
    figure("time_split", "Where a booking search's time goes",
           "Each bar is the time the server spent on the median booking search at five desks, split into waiting for a free slot (grey), road routing (blue), calculation (amber) and the database (purple). Shorter is better.",
           f"Most of a search's time was never calculation; it was waiting and routing. Once routing was fixed, the wait shrank too, because each search held its slot for less time. Confirms now go ahead of waiting searches, so a customer's confirmation takes {confirm_after} ms at the median instead of {confirm_before} ms.")
    heading("17 The bottleneck lesson")
    figure("bottleneck", "The bottleneck lesson",
           "Three separate measurements. Left: a faster road graph answered a batch of routing questions in a fraction of the time. Middle: in the September booking test the same faster graph served the same number of customers. Right: in October, after the earlier fixes, it multiplied the searches answered in a minute.",
           f"An improvement only pays where it relieves the binding constraint. In September booking was held up by other things, so a ten times faster road graph bought nothing. In October routing had become the constraint, and the same idea multiplied capacity ({rec['bottleneck']['october']['before']} to {rec['bottleneck']['october']['after']:,} searches in a minute).")
    figure("ch_mechanism", "How a faster road graph answers with less work",
           "Junctions sit along the road from left to right, raised by how important they are. The solid green lines are a query climbing from the start (S) and the target (T) until they meet; the dashed lines are shortcuts added once, when the map is prepared. The cards give that one-time cost. This is a schematic and carries no data.",
           "The answer is the same to the meter; only the work changes. The price is a short preparation step each time the map is updated.")
    p("A contraction hierarchy is prepared once per map version. Every junction is ranked by importance, and junctions are then set aside from least to most important; whenever setting one aside would break a shortest path between its neighbors, a shortcut edge is added that stands for the path through it. At query time the search moves only upward in importance, from the start and from the destination at once, and stops where the two meet. Because the preparation adds exactly the shortcuts needed, the route found is the true shortest one, not an approximation. A booking search needs a table of drive times between every pair of stops it considers, so the routing service also answers the legs of one request in parallel on its cores. That was the step before the faster graph in the staircase.")
    heading("18 Engine productivity: every change, shipped or not")
    figure("engine_steps", "Engine speed: every change, shipped or not",
           "Each bar is one engine change, measured in its own study against the code just before it. Bar length is the median speed ratio on the part of the work the change targets, on a log scale where 2x and 4x are equally spaced. The dashed line is the 1.05x bar. Green changes shipped; grey ones did not, with the reason.",
           "The ratios are not multiplied together, because each was measured against a different baseline. Three of nine changes did not ship: two were fast but failed the cost or fairness rules, and one was slower. The discipline rejects speed that costs money.")
    p(f"Three of the bars are quoted from a study whose raw files are kept on a separate branch: long score and money ({s3['lm_over_main'][0]:.3f} times), the W and A changes on top ({s3['wa_over_lm'][0]:.3f} times) and the route-walk cascade ({s3['cr11_c_over_lm']:.2f} times). The W change on its own measured {s3['lmw_over_lm'][0]:.3f} times, just under the bar. It shipped on the owner's decision, for a consistent gain that is larger on bigger fleets. That is recorded here as a judgment call, not a pass.")
    p("How were these changes found? Not by guessing. Each study recorded where the engine's processor time went, sampled thousands of times a second while it solved. Before the numbering fix, a fifth of the cost search's time went to looking up a road-matrix entry by text label. After formula b, the fairness search's time split roughly into the scheduler's own code (30 percent), the solver engine (28), decimal score arithmetic (16), collections (16) and decimal money (6). That profile named the next target: the decimal score and money, which became the long score and money change. Each change goes after the largest remaining slice.")
    box("How it works: what each engine change did", [
        "Road lookup tables. Drive times were looked up by a text key built from the two endpoints on every leg of every move. Each endpoint now gets a number once, and the times sit in a table indexed by those numbers, with minutes precomputed. The same answers, without building a string per leg.",
        "Time as whole numbers. Route timing used general date and duration objects. It now counts nanoseconds from the start of the shift in plain 64-bit integers, exact for any time the interface accepts, and converts back to dates only on the way out.",
        "Road matrix numbering. The endpoint numbers were still found through a lookup on each leg. Each visit and route now carries its number, tagged with the version of the facts it came from, so a lookup is a single comparison.",
        "Long score and money. Each score level was an arbitrary-precision decimal. Because every level is now a whole number within known bounds (the rate limits in Part I guarantee the cost fits), each is held in a 64-bit integer. Every plan compares exactly as before; only the arithmetic is cheaper.",
        "W and A. Each visit keeps its time window and duration in nanoseconds, built once, and its matrix number in one field. W measured 1.047x, A no measurable change; both shipped by the owner's decision.",
        "Rejected: the whole-number score (S). It moved the score to 64-bit integers and the variance to units of a billionth in one step. It was faster in the fairness search but made 3 days dearer beyond noise, so it failed its bar. Its profile showed the decimal cost lay in the variance formula, not the score type. That led to formula b, and the long score shipped later on its own.",
        "Rejected: the route-walk cascade. It let the solver maintain arrival times itself, updating them from the changed visit onward, instead of the scheduler re-walking the whole route. It measured about 0.63x: slower, not faster.",
        "The gate for all of them. A scoring exactness test replays fixed solves before and after each change and requires the same score, the same number of moves, the same arrangement and the same independent evaluation. Speed is never bought with a different answer.",
    ])
    heading("19 Confirmed through the public interface")
    figure("harness_vs_stack", "The workstation result held through the public API",
           "Each row is a change and a fleet size. The amber dot is the speed ratio measured in the controlled workstation study; the green dot is the same change measured through the public interface on the copy of the production setup. Rows with only a green dot had no workstation figure for that fleet size.",
           f"The gains survive the trip from the laboratory to the full service, a little smaller as expected. The fairness formula gives {stack_b[10]['FAIRNESS']:.2f} to {stack_b[50]['FAIRNESS']:.2f} times, and this week's engine changes {shipped[50]['REFERENCE']:.2f} to {shipped[10]['REFERENCE']:.2f} times. The workstation study of those engine changes measured {s3['wa_over_main'][0]:.2f} times overall (quoted). Every plan stayed within its cost ceiling.")

    # ---- Part VI ------------------------------------------------------------------------------------------------------------
    part("Part VI  Capacity and the cost to run")
    heading("20 What another core buys")
    figure("capacity", "What another core buys at the booking desk",
           "Bars are booking searches answered per minute by ten clients at once, for seven ways of arranging the booking pool on 2, 3 and 4 virtual CPUs. A label such as 1 x 4 means one copy of the service with four working slots. Green is the arrangement chosen.",
           f"The marginal product of a core is real but falling: the best layout answered {best[2]:,.0f}, {best[3]:,.0f} and {best[4]:,.0f} searches a minute on 2, 3 and 4 CPUs. At 3 and 4 CPUs one copy with a slot per CPU was best. At 2 CPUs two single-slot copies won, a reversal the study measured but did not explain.")
    p("Under ten clients at once the optional polish was skipped on almost every search, so capacity is set by the first pass. The production server's 8 CPUs are split between booking, overnight planning, routing and the database, and the final split should be set from a measurement on that server.")
    compose = (HERE.parent.parent / "deploy/vm/docker-compose.yml").read_text(encoding="utf-8")
    limits = {m[0]: (float(m[1]), m[3]) for m in re.findall(r'cpus: "\$\{(\w+)_CPUS:-([\d.]+)\}", memory: "\$\{(\w+)_MEMORY:-(\w+)\}"', compose)}
    parts = [("SCHEDULER_BOOKING", "Booking pool, per replica", "Searches and confirms; one replica by default"),
             ("SCHEDULER_DAILY", "Daily pool, per replica", "Overnight plans, two at once; one replica by default"),
             ("ROUTING", "Road map service", "Drive times on the Omaha road graph"),
             ("NOMINATIM", "Address service", "Turns addresses into map points"),
             ("SCHEDULER_DB", "Database", "Only what keeps promises safe: holds and versions"),
             ("SCHEDULER_PROXY", "Front proxy", "The one address; sends each call to its pool")]
    assert all(k in limits for k, *_ in parts), sorted(limits)
    cpu_total = sum(limits[k][0] for k, *_ in parts)
    p(f"The table gives the shipped limits. They are ceilings, not reservations: they add up to {cpu_total:g} CPUs on an 8-CPU machine, a deliberate overcommitment, because the parts peak at different times. Booking peaks in business hours and the daily pool overnight. A limit stops one part from starving the others when they do collide.")
    table(["Part", "CPU limit", "Memory limit", "Role"], [[name, f"{limits[k][0]:g}", limits[k][1].upper().replace("G", " GB").replace("M", " MB"), role] for k, name, role in parts],
          [1.9, 0.9, 1.1, 3.0])
    p("The capacity grid above suggests the main lever: booking throughput rose with every core given to the booking pool, and a second booking replica or a larger one is a configuration change, not new code. Spreading one overnight plan over two cores is a separate question, taken up in Part VII.")
    heading("21 What still has to be measured")
    p("Every capacity number here comes from a rehearsal copy that applies the production server's CPU limits on a workstation. The latency target is defined at the WaterFlex desk, across the network, and that has not yet been measured. Until it is, the absolute numbers are the rehearsal machine's, and the ratios between configurations are what can be relied on.")

    # ---- Part VII -----------------------------------------------------------------------------------------------------------
    part("Part VII  What comes next, as decisions")
    figure("roadmap", "What comes next, as decisions",
           "Three lanes: measurements that would firm up the numbers, decisions that follow from them, and the integration steps before production traffic. Arrows suggest an order. The two green steps are the ones the others wait on. This is a schematic and carries no data.",
           "The highest-value next step is the measurement on the production server. Production traffic waits on the integration lane.")
    items = [
        ("Measure on the production server.", "Repeat the booking load test across the real network from the WaterFlex side. It decides whether the latency target holds where it is defined, and sizes the booking pool. Cost: a day of engineering and the server time."),
        ("Replay real days.", "Run recorded WaterFlex days, with real road times, through the daily plan. This turns the modeled savings into an estimate for WaterFlex's own work. No new code is needed; the interface accepts a recorded day as it is."),
        ("Fairness head-to-head.", "Test the load balance measure against the shipped formula on 5 to 50 technicians. A change of fairness policy is a business decision, so the study informs it rather than makes it."),
        ("Offer the stronger daily method as a choice.", f"Tabu + K opt size 3 beat the default by {money(mean(direct))} a day on new days. Offering it as an opt-in costs nothing; making it the default should wait for the real-day replay."),
        ("Multithreaded solving.", "Planned work with a gate. Today each daily plan uses one core, and several plans run side by side. Spreading one plan over two cores would let it try more rearrangements per second, but it halves how many plans run at once. It goes to a client only if a paired study shows cheaper plans, every score stays the same, and the overnight queue stays acceptable."),
        ("Confirm the rate limits.", "The engine's faster money arithmetic accepts rates up to four decimal places, $500 an hour and $25 a mile. WaterFlex Software should confirm that no stored rate exceeds them before deployment."),
        ("Encrypt the traffic.", "Encryption between WaterFlex Software and the scheduler, and inside the scheduler, is not yet in place. It is a precondition of production traffic."),
        ("Side by side, then cutover.", "Run WaterFlex Software's integration beside the current path until both agree, then switch."),
    ]
    for head, text in items:
        lead(head, text)

    # ---- Appendices ---------------------------------------------------------------------------------------------------------
    part("Appendix A  Glossary")
    table(["Term", "Meaning"], [
        ["Booking search", "One request for arrival windows for a new job, answered while the customer waits."],
        ["Daily plan", "A proposal that rearranges one booked metro day to cost less, keeping every promised window and adding no overtime."],
        ["Modeled cost", "Paid minutes at hourly rates plus mileage, computed by the scheduler. Not payroll or revenue."],
        ["Synthetic day", "A test day generated to resemble WaterFlex work, so methods can be compared on identical inputs."],
        ["Paired comparison", "Two methods run on the same day, seed and seconds; each result is compared with its twin."],
        ["Median", "The middle value: half the results are above it and half below. It resists a few extreme cases."],
        ["Typical and worst-case wait", "The median (p50) and the 95th percentile (p95) of the wait at the desk; one search in twenty is slower than the p95."],
        ["Speed ratio", "Rearrangements tried per second after a change, divided by before, on the same day."],
        ["Promotion bar", "The decision rule written before a study runs, for example a 1.05x median gain."],
        ["Fairness variance", "How unevenly paid work is spread across technicians; lower is more even."],
        ["Fairness allowance", "The share of the cost saving a client lets the fairness search spend, 2 percent by default."],
        ["Hold", "A window offered to a customer and reserved for ten minutes while they choose."],
        ["Rehearsal stack", "A copy of the production setup, with its CPU limits, run on a workstation."],
        ["Tabu, K opt, late acceptance", "Local search methods that differ in which changes they try and when they accept a worse step."],
        ["Contraction hierarchy", "A road graph prepared with shortcuts so a route is found with far less work, with identical answers."],
        ["Lexicographic score", "A score of several levels compared in order; a lower level counts only when every higher level ties."],
        ["Hard and soft levels", "Hard levels are promises and limits that must be zero; soft levels are what the search minimizes."],
        ["Move, step, neighborhood", "A move is one small change to a plan; a step is choosing and applying one; the neighborhood is the set of moves a method may try."],
        ["Tabu memory", "The number of recent steps whose changes may not be undone, 7 by default."],
        ["Beam search", "A search that keeps only the best few partial answers at each depth; booking keeps eight, two changes deep."],
        ["Bootstrap interval", "A range for a mean found by resampling the data many times; here whole days are resampled."],
        ["Sign test", "The chance of a win and loss split at least this lopsided if two methods were equally good."],
        ["p50, p95, p99", "The waits that half, 19 in 20 and 99 in 100 requests stay under."],
        ["Replica, slot, admission", "A copy of the service; one search it can run at once; the rule for which waiting request runs next or is refused."],
        ["Scoring exactness test", "A test that replays fixed solves and requires the same score, moves and plan after an engine change."],
    ], [1.9, 5.0])
    part("Appendix B  The studies behind this briefing")
    rows = [[s["start"] if s["start"] == s["end"] else f"{s['start']} to {s['end']}", s["name"], s["size"]] for s in rec["timeline"]]
    decisions = {"Booking methods": "Bounded rearrangement kept", "Field validation": "Keep Tabu, study fleet size and budget", "Search time budgets": "No blanket switch",
                 "Daily search variants": "Tabu + K opt chosen", "Follow-up on new days": "Choice held; size 3 found", "Load and admission": "Six fixes, one at a time",
                 "Engine profile": "Road lookup tables shipped", "Booking horizon": "Dynamic horizon", "Booking configuration": "Two-second polish, layout",
                 "Time and score types": "Whole-number time shipped", "Fairness measures": "Formula b shipped", "Road matrix numbering": "Shipped",
                 "Long score and money": "Shipped, with W and A", "Through the public API": "Confirmed on the stack"}
    table(["Dates", "Study", "Size", "Decision"], [[*r, decisions[r[1]]] for r in rows], [1.6, 1.8, 1.6, 1.9])
    p("Every study has its own folder of evidence and a method note in the repository. This briefing's own note, EVIDENCE.md beside the document, lists each source and how it was checked.", size=9.5)

    assert len(figures) == len(set(figures)) == 32, len(figures)
    names = set((HERE / "generated").glob("*.png"))
    assert {HERE / "generated" / f"{n}.png" for n in figures} <= names
    (HERE / "generated" / "figures.txt").write_text("\n".join(figures) + "\n", encoding="utf-8")
    doc.save(OUTPUT)
    print("wrote", OUTPUT.name, len(figures), "figures")


if __name__ == "__main__":
    build()
