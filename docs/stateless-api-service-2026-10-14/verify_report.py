"""Check the delivered tables and template parts against the source documents.

Parses the contract, the design document, the guide, the runbook and the conformance kit again on its own, without
report_data.py's loader, and compares every document-derived table cell for cell. Run normally, without Python's -O
flag. Visual pagination still requires rendering.
"""
import json
import re
from hashlib import sha256
from zipfile import ZipFile

import yaml
from docx import Document
from docx.oxml.ns import qn
from build_report import HERE, OUTPUT, REFERENCE, REFERENCE_SHA, PRESERVE
from report_data import PULL_REQUESTS, SPEC, DESIGN, GUIDE, RUNBOOK, CASES


def rows_between(text, start, end):
    segment = text[text.index(start):text.index(end, text.index(start))]
    lines = [l for l in segment.splitlines() if l.startswith("| ")][2:]
    return [[c.strip() for c in l.strip()[1:-1].split("|")] for l in lines]


def plain(text):
    return re.sub(r"\*\*(.+?)\*\*", r"\1", text.replace("`", "")).replace("[docs/integration-guide.md](integration-guide.md)", "docs/integration-guide.md")


def verify():
    spec = yaml.safe_load(SPEC.read_text(encoding="utf-8"))
    design, guide, runbook = (p.read_text(encoding="utf-8") for p in (DESIGN, GUIDE, RUNBOOK))
    cases = re.findall(r'name: "([^"]+)"', CASES.read_text(encoding="utf-8"))
    doc = Document(OUTPUT)
    assert len(doc.tables) == 27 and len(doc.inline_shapes) == 5 and len(doc.sections) == 1
    assert len([p for p in doc.paragraphs if p.style.name == "Heading 1"]) == 18
    captions = [p for p in doc.paragraphs if p.style.name == "Caption"]
    assert len(captions) == 5
    for i, p in enumerate(captions, 1):
        assert p.text.startswith(f"Figure {i}.") and all(str(r.font.color.rgb) == "000000" and r.italic for r in p.runs)

    def cells(index):
        return [[c.text for c in r.cells] for r in doc.tables[index].rows[1:]]

    def check(index, expected):
        actual = cells(index)
        assert actual == [[str(c) for c in r] for r in expected], (index, actual, expected)

    operations = [(m.upper(), path, op["operationId"], op["summary"]) for path, item in spec["paths"].items() for m, op in item.items() if m in ("get", "post", "delete", "put")]
    check(4, [[o[2], f"{o[0]} {o[1].replace('/api/v1', '')}", o[3]] for o in operations])
    check(2, [[plain(a), plain(b)] for a, b in rows_between(guide, "## 1. Topology and ownership", "Every calculating call carries a **snapshot**")])
    check(11, [[plain(c) for c in r] for r in rows_between(guide, "## 7. Errors, idempotency and retries", "**Timeouts to set at the caller")])
    errors = spec["components"]["schemas"]["Problem"]["properties"]["error"]["enum"]
    assert {r[1] for r in cells(11) if r[1] in errors} == set(errors), "every error code of the spec appears in the guide's table"
    check(14, [[plain(a), plain(b)] for a, b in rows_between(design, "## Thin store", "Every key starts with")])
    check(15, [[plain(a), plain(b)] for a, b in rows_between(design, "### What is stored", "### Tenant isolation")])
    check(16, [[plain(c) for c in r] for r in rows_between(design, "## Route cache", "How a lookup works")])
    check(17, [[plain(c) for c in r] for r in rows_between(runbook, "## Settings", "## Install")])
    check(20, [cases[i:i + 3] + [""] * (3 - len(cases[i:i + 3])) for i in range(0, len(cases), 3)])
    check(22, [[plain(c) for c in r] for r in rows_between(design, "### Contract changes", "### Latency target")])
    check(23, [[f"#{n}", d, ph, t] for n, d, ph, t in PULL_REQUESTS])
    schemas = sorted(spec["components"]["schemas"])
    check(24, [schemas[i:i + 4] + [""] * (4 - len(schemas[i:i + 4])) for i in range(0, len(schemas), 4)])
    check(25, [[plain(a), plain(b)] for a, b in rows_between(design, "### Endpoints", "### Contract changes")])
    phases = []
    for _, _, ph, _ in PULL_REQUESTS:
        if ph not in phases: phases.append(ph)
    phase_rows = cells(19)
    assert [r[0] for r in phase_rows] == phases
    for r, ph in zip(phase_rows, phases):
        numbers = [n for n, _, p_, _ in PULL_REQUESTS if p_ == ph]
        assert r[1] == ", ".join(f"#{n}" for n in numbers), ph
        assert r[2] == f"{min(d for _, d, p_, _ in PULL_REQUESTS if p_ == ph)[5:]} to {max(d for _, d, p_, _ in PULL_REQUESTS if p_ == ph)[5:]}"
    history = re.findall(r"^- (\d+\.\d+\.\d+) \((\d{4}-\d{2}-\d{2})\)", spec["info"]["description"].split("Version history:")[1], re.M)
    assert history[0][0] == spec["info"]["version"]
    text = "\n".join(p.text for p in doc.paragraphs)
    assert f"version {spec['info']['version']}" in text and f"{len(operations)} operations" in text and f"{len(schemas)} object schemas" in text and f"{len(errors)} error codes" in text
    assert f"{len(cases)} cases" in text and f"{len(PULL_REQUESTS)} pull requests" in text and f"1.0.0 on {history[-1][1]}" in text

    data = json.loads((HERE / "generated/chart_data.json").read_text())
    assert data["tiers"] == [r[0] for r in rows_between(design, "## Route cache", "How a lookup works")]
    assert data["history"]["phases"] == phases and data["history"]["versions"] == [list(v) for v in history]
    for ph in phases:
        assert data["history"]["counts"][ph] == [sum(1 for _, d, p_, _ in PULL_REQUESTS if d == day and p_ == ph) for day in data["history"]["days"]]

    assert sha256(REFERENCE.read_bytes()).hexdigest() == REFERENCE_SHA
    with ZipFile(REFERENCE) as reference, ZipFile(OUTPUT) as built:
        for name in (*PRESERVE, *[n for n in reference.namelist() if "footer" in n]):
            assert reference.read(name) == built.read(name), name
        for name in built.namelist():
            if name.endswith(".xml"):
                xml = built.read(name).decode("utf-8")
                assert "—" not in xml and "–" not in xml, name
        embedded = {sha256(built.read(n)).hexdigest() for n in built.namelist() if n.startswith("word/media/")}
        assert embedded == {sha256(p.read_bytes()).hexdigest() for p in (HERE / "generated").glob("*.png")}
    assert doc.sections[0]._sectPr.xml == Document(REFERENCE).sections[0]._sectPr.xml
    assert all(shape._inline.docPr.get("descr") for shape in doc.inline_shapes)
    assert all(row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None for t in doc.tables for row in t.rows)
    print(f"PASS: {len(operations)} operations, {len(schemas)} schemas, {len(errors)} error codes, {len(history)} spec versions, {len(cases)} conformance cases, {len(PULL_REQUESTS)} pull requests; 12 document-derived tables; five figures; template parts.")


if __name__ == "__main__":
    verify()
