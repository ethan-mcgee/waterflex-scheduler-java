"""Build the stateless API service report from the contract, the design document, the guide, the runbook and the
conformance kit, in the selected explainer's format.

Install requirements.txt first. No network or services.
"""
from copy import deepcopy
from datetime import datetime
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from charts import make_charts
from report_data import load, plain

HERE = Path(__file__).resolve().parent
REFERENCE = HERE.parent / "WaterFlex_Scheduler_Algorithm_Comparison_Explained.docx"
REFERENCE_SHA = "5366f3017d6f18c15de295ba5fb9c1a014b6222911040f1c5e044845e57f832d"
OUTPUT = HERE / "Stateless_API_Service_2026-10-14_Report.docx"
PRESERVE = ("word/styles.xml", "word/stylesWithEffects.xml", "word/numbering.xml", "word/theme/theme1.xml", "word/fontTable.xml")


def build():
    assert sha256(REFERENCE.read_bytes()).hexdigest() == REFERENCE_SHA
    data = load()
    assets = HERE / "generated"
    make_charts(data, assets)
    doc = Document(REFERENCE)
    table_properties = deepcopy(doc.tables[0]._tbl.tblPr)
    for child in list(doc.element.body):
        if child.tag != qn("w:sectPr"): doc.element.body.remove(child)
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype.endswith(("/image", "/hyperlink")): del doc.part.rels[rid]
    doc.core_properties.title = "WaterFlex stateless API service explained"
    doc.core_properties.subject = f"What the scheduler's public API is, how WaterFlex Software uses it, and how it was built; contract {data['version']}"
    doc.core_properties.author = doc.core_properties.last_modified_by = "WaterFlex"
    doc.core_properties.created = doc.core_properties.modified = datetime(2026, 10, 14)
    figures = []

    def p(text, style=None):
        assert "—" not in text and "–" not in text, text
        para = doc.add_paragraph(text, style); para.paragraph_format.widow_control = True
        return para

    def heading(text):
        h = p(text, "Heading 1"); h.paragraph_format.page_break_before = True

    def sub(text): p(text, "Heading 2")

    def table(headers, rows, widths, size=10):
        assert abs(sum(widths) - 6.9) < .01
        t = doc.add_table(rows=1, cols=len(headers)); t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t._tbl.replace(t._tbl.tblPr, deepcopy(table_properties))
        for col, w in zip(t.columns, widths): col.width = Inches(w)
        for i, values in enumerate([headers, *rows]):
            assert len(values) == len(headers), values
            row = t.rows[0] if i == 0 else t.add_row()
            trpr = row._tr.get_or_add_trPr(); trpr.append(OxmlElement("w:cantSplit"))
            if i == 0: trpr.append(OxmlElement("w:tblHeader"))
            for j, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
                text = str(value); assert "—" not in text and "–" not in text, text
                cell.width = Inches(width); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell.text = text
                props = cell._tc.get_or_add_tcPr()
                shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "DCE6EF" if i == 0 else ("FFFFFF" if i % 2 else "F4F6F8")); props.append(shade)
                margins = OxmlElement("w:tcMar")
                for side in ("top", "left", "bottom", "right"):
                    n = OxmlElement("w:" + side); n.set(qn("w:w"), "65"); n.set(qn("w:type"), "dxa"); margins.append(n)
                props.append(margins)
                para = cell.paragraphs[0]; para.paragraph_format.space_after = Pt(2); para.paragraph_format.space_before = Pt(2); para.paragraph_format.line_spacing = 1.0
                if len(headers) > 2 and j > 0 and len(text) < 24: para.alignment = 1
                if len(rows) <= 12 and i < len(rows): para.paragraph_format.keep_with_next = True  # Short tables stay on one page.
                for run in para.runs: run.font.name = "Calibri"; run.font.size = Pt(size); run.bold = i == 0
        spacer = p(""); spacer.paragraph_format.space_after = Pt(0); spacer.paragraph_format.line_spacing = Pt(3)

    def figure(name, caption):
        figures.append(name)
        caption = f"Figure {len(figures)}. {caption}"
        para = p(""); shape = para.add_run().add_picture(str(assets / f"{name}.png"), width=Inches(6.85))
        shape._inline.docPr.set("descr", caption); para.paragraph_format.keep_with_next = True
        cap = p(caption, "Caption"); cap.paragraph_format.space_after = Pt(7)
        for run in cap.runs:
            run.font.color.rgb = RGBColor(0, 0, 0); run.bold = False; run.italic = True

    ops, errors, history, cases, prs = data["operations"], data["errors"], data["history"], data["cases"], data["pull_requests"]
    phases = []
    for _, _, ph, _ in prs:
        if ph not in phases: phases.append(ph)
    first_day, last_day = min(d for _, d, _, _ in prs), max(d for _, d, _, _ in prs)

    p("WaterFlex stateless API service explained", "Title")
    p("What the scheduler's public API is, how WaterFlex Software uses it, and how it was built", "Subtitle")
    p(f"October 14 2026  |  Contract version {data['version']}, pull requests #{prs[0][0]} to #{prs[-1][0]}")
    p(f"The WaterFlex scheduler is now a service that WaterFlex Software calls over HTTP and nothing else touches. It keeps no customer, job or technician records; every call carries a snapshot of the facts it needs, and it answers with time windows to offer a customer or whole timed routes for a day, which WaterFlex Software writes into its own records with a compare-and-set. The contract has {len(ops)} operations, {len(data['schemas'])} object schemas and {len(errors)} error codes, reached version 1.0.0 on 2026-10-13 and {data['version']} with five additive changes the next day, and is proven against a running scheduler by a conformance kit of {len(cases)} cases. This report explains what the service is, how a host uses it safely under concurrency, and the {len(prs)} pull requests between {first_day} and {last_day} that built it. [1, 2, 3]")
    sub("What the service offers")
    table(["Facet", "What it is"], [
        ("Surface", f"{len(ops)} operations under /api/v1 plus an open /health, all served, all idempotent on a host requestId"),
        ("Inputs", "A snapshot per call: rates, policy, technicians, technician-days with lastModified, committed appointments, coordinates"),
        ("Outputs", "Booking offers held for 10 minutes; daily and repair proposals; commit and confirm receipts listing every assignment to write; timed and drawn routes"),
        ("State kept", "Holds, offer sets, proposals, receipts and request claims, keyed by host IDs, purged after the deployment's retention period; a route cache"),
        ("Concurrency", "Host lastModified per technician-day compared for exact equality; row locks, versions and row-level security in the scheduler's own Postgres"),
        ("Clients", "One tenant per WaterFlex client, chosen by the bearer token, isolated by tenant_id and database policies; fair turns between clients"),
        ("Deployment", "Its own VM: a proxy, booking and daily replica pools of one jar, a routing service per metro, optional Nominatim, its own database")], [1.1, 5.8])
    p("The design decision behind all of this was taken on 2026-10-07: a thin scheduler store. The scheduler owns only what it must to keep customer promises safe when several desks book at once, and WaterFlex Software stays the record. The sections below follow the integration guide's order, then explain what is underneath, and end with how the service was built and what still stands between it and production traffic.")
    sub("How to read this report")
    table(["Sections", "What you will learn"], [("1 to 3", "What the service is and is not, where it runs, and the contract's operations and conventions."),
                                                ("4 to 7", "The snapshot and its timestamp rules, booking, daily and repair proposals, and every error with its remedy."),
                                                ("8 to 11", "Underneath: many clients and replicas, the thin store, routing and geocoding, and operating the service."),
                                                ("12 to 14", "How it was built, how an integration is proven, and what remains before production traffic."),
                                                ("Appendices", "Every contract change, every pull request, every schema, definitions and sources.")], [1.45, 5.45])

    heading("1 What the service is, and what it is not")
    p("The scheduler is an add-on to WaterFlex Software. WaterFlex Software owns master data (customers, addresses, jobs, appointments, technicians, qualifications, shifts, time off, depots) and the business record of every appointment. The scheduler owns only what it must to keep customer promises safe under concurrency: booking holds and the arrangement they depend on, the offer sets and proposals it issued, and the receipts of what was committed. Calculation and routing are stateless: complete data in, results out. [1]")
    table(["WaterFlex Software owns", "The scheduler owns"], [[plain(a), plain(b)] for a, b in data["ownership"][1]], [3.45, 3.45])
    p("Three consequences follow. The service has no screens: WaterFlex Software renders every map and schedule, and the scheduler returns data only. The service reads no policy of its own: every rule a client chooses (its fairness budget, its daily search method and times, its booking deadline and window threshold) arrives with every call, and a missing value is refused, never assumed. And no client library ships: the product stays Java, scaled on its own VM, and a host integrates from any language against the OpenAPI contract.")
    sub("What it is not")
    p("It is not a database of schedules, so it cannot drift from the host's record; the host's timestamp per technician-day decides whether what the scheduler computed still applies. It is not a batch system, so overnight runs belong to the host, which queues a proposal per metro-day at the time each client chooses. And it is not reachable by anything but the API: no shared database, no callbacks, no client code from this repository.")

    heading("2 Where it runs")
    figure("topology", "The scheduler VM and its host. WaterFlex Software calls one address; behind it a proxy sends booking calls to the booking pool and proposals to the daily pool, each replica embeds the calculation engine, and the VM holds a routing service per metro, an optional Nominatim and the scheduler's own database.")
    p("The scheduler runs on its own VM with a partitioned share of compute, on the same private network as the rest of WaterFlex Software, and is reached only through https://<scheduler-vm>/api/v1. On the VM, scheduler-service may run as several replicas of the same jar behind one load balancer. Booking replicas answer booking searches, confirms and route calls; daily replicas take daily and repair proposals. Every replica answers every route, holds no state of its own and shares the one thin store, so which replica answers a call never changes the answer. [1, 2]")
    table(["Component", "Role", "Talks to"], [
        ("scheduler-proxy (nginx, :8000)", "The one address a host calls; sends proposals and request status calls to the daily pool and everything else to the booking pool", "Both pools"),
        ("Booking replicas", "Booking offers, select, release, confirm, route evaluation and geometry; embedded calculation engine", "Routing, Nominatim, the database"),
        ("Daily replicas", "Daily and repair proposals, synchronous or queued; several calculations at once", "Routing, Nominatim, the database"),
        ("routing-service per metro", "Turns coordinates into directed road legs from that metro's graph; sees coordinates only", "Its map volume"),
        ("Nominatim (optional)", "Locates an address-only location as an exact house, in NOMINATIM mode", "Its own database"),
        ("scheduler-db (Postgres 16)", "The thin store, tenants and tokens, and the route cache; created and migrated by the scheduler itself", "Nothing else"),
        ("Management port (:9000)", "Health with components and Prometheus metrics; private network only", "Operators")], [1.7, 3.5, 1.7])
    p("The calculation engine stays embedded in each replica by default (SCHEDULER_CALCULATION_MODE=EMBEDDED). Moving it to a remote solver pool is a configuration change already wired in, kept as a scaling lever rather than a default. Each served metro has its own routing service loaded with that metro's road graph; routing sees only coordinates, never client data, so every client in a metro shares it.")

    heading("3 The contract")
    p(f"The machine-readable contract is docs/api/openapi-v1.yaml, OpenAPI 3.1, version {data['version']}. Every operation is marked implemented; a test validates every example in the spec against its schema and fails if the implemented operations differ from the routes the scheduler serves, and the Java types mirror every object schema, so the spec, the code and the examples cannot drift apart unnoticed. [2]")
    table(["Operation", "Method and path", "What it does"], [[o["id"], f"{o['method']} {o['path'].replace('/api/v1', '')}", o["summary"]] for o in ops], [1.45, 2.35, 3.1], size=9)
    sub("Conventions every call follows")
    table(["Convention", "Rule"], [
        ("Tenant", "Derived from the bearer token; a request never names a tenant, and a tenantId in a body is a 400."),
        ("requestId", "Every write carries a host UUID that is its idempotency key across all of the tenant's endpoints: same body replays the stored answer, different body is a 400."),
        ("IDs", "WaterFlex Software's own opaque strings, 1 to 128 characters."),
        ("Money and rates", "Canonical decimal strings such as \"20\" or \"0.67\", never JSON numbers."),
        ("Times", "Instants are RFC 3339 with an offset; service dates are calendar dates in the metro's time zone."),
        ("Strictness", "Unknown fields, duplicate keys, missing required values and dangling references are rejected with 400; nothing is assumed or filled in."),
        ("Overtime", "Never assigned by any operation; a repair that would need it is reported infeasible."),
        ("Retention", "Requests, proposals, offers and receipts are kept for SCHEDULER_RETENTION_DAYS, then purged; an older requestId runs as a new request.")], [1.3, 5.6])
    p(f"Within version 1, changes are additive only (a new operation, optional request field, response field, error code or enum value) and raise the minor version; each is listed in the spec's version history and in the design document's change table before it is deployed. A breaking change is published as /api/v2 beside v1. The contract reached 1.0.0 on {history[-1][1]} and is at {history[0][0]} after {len(history) - 1} additive versions (section 12 and Appendix A).")

    heading("4 The snapshot and the lastModified rules")
    p("Every calculating call carries a Snapshot for one metro: its rates and policy, the technicians it covers, one TechnicianDay per technician and date, and every committed appointment on those technician-days. The scheduler computes a canonical hash of it, the input revision, and every proposal and offer records which revision it was computed from. [1, 3]")
    table(["Part", "What it holds"], [
        ("metroId, timeZone", "The metro whose routing service answers, and the zone of its service dates and 6 a.m. cutoff."),
        ("rates", "regularHourly, overtimeHourly, mileagePerMile, travelBufferPct, travelBufferMinutes: the client's own cost model."),
        ("policy", "fairnessBudget, dailySearch (TABU_KOPT or TABU), costSearchSeconds, fairnessSearchSeconds, bookingDeadlineMs, regularWindowThreshold, optional bookingRefinementMs."),
        ("technicians", "ID and qualifications (the services the technician may perform)."),
        ("technicianDays", "Per technician and date: lastModified, shift, absences (required, [] when none), start and end locations, maxPaidMinutes."),
        ("appointments", "Per committed appointment: ID, technician, date, service, duration, arrival window, location, sequence and plannedStart (required)."),
        ("locations", "Coordinates, used exactly as sent; or an address, located only in NOMINATIM mode and only as an exact house.")], [1.4, 5.5])
    p("On each technician-day, sequence and plannedStart must put the appointments in the same order; a contradiction is a 422, never reordered. A location that cannot be located is never guessed: in a proposal its technician-day is left exactly as it is and listed in skippedTechnicianDays, and in a booking search an unlocatable job is a 422.")
    sub("The concurrency token")
    p("The scheduler cannot lock rows it does not own, so commits use optimistic concurrency on both sides, and the token is the lastModified timestamp WaterFlex Software already keeps per technician and day. Every snapshot includes it; a proposal or offer records the values it was computed from; a commit or confirm requires the submitted values to equal the recorded ones; and the host then writes the business record with a compare-and-set on the same values. It is compared for exact equality at full precision and never ordered, so clock skew between servers cannot cause a wrong decision; at worst an unrelated change forces a retry. [1]")
    table(["Rule for the host", "Why"], [
        ("It changes on every change that affects the day: an appointment created, edited, cancelled or deleted; a shift, availability or absence change; a start or end location change; a qualification change.", "A missed update is the one way an integration can overwrite a dispatcher's change."),
        ("A move touches both sides: moving an appointment from technician A on Tuesday to technician B on Wednesday updates both technician-days.", "A proposal for either day would otherwise apply to facts that moved."),
        ("A shift template change updates every future day it affects.", "Those days' routes were computed against the old shift."),
        ("It is sent exactly as stored: not rounded, reformatted or regenerated between the snapshot and the commit.", "Equality is exact; a reformatted value is a different value.")], [4.2, 2.7])

    heading("5 Booking")
    p("Booking finds arrival windows for a new job, holds them for 10 minutes, and turns the customer's choice into an appointment the host writes. A search covers exactly the dates the host asks for (1 to 21), returns at most the client's offerLimit windows (1 to 4) of the client's windowHours (1 to 8) each, and answers within the client's bookingDeadlineMs (at most 5 seconds). An offer may move other customers' appointments to make room; those moves come back in the confirm receipt. [1, 3]")
    figure("booking", "The booking sequence and its refusals. Each 409 and a failed receipt write lead to the same place: release the set (or let the hold expire) and search again from current records with a new requestId.")
    table(["Step", "Call", "Answer"], [
        ("1 Search", "POST /booking/offers with the job, horizon, offerLimit, windowHours and the snapshot of every technician-day in the horizon", "201 OfferSet: offers with windows, expiresAt, searchComplete, refinement, incompleteDates, skippedTechnicianDays"),
        ("2 Select", "POST /booking/offers/{offerId}/select", "200: that hold kept with its original expiry, the others in the set released"),
        ("3 Confirm", "POST /booking/holds/{holdId}/confirm with a fresh snapshot of the hold's metro and date only", "200 receipt: every appointment on each technician-day whose route changes, the new job under its job ID, and the timestamps to compare"),
        ("4 Write", "In one host transaction: compare each listed lastModified, write every listed assignment, move the timestamps", "On any difference write nothing, release, search again"),
        ("Release", "POST /booking/offers/{offerId}/release", "200 even when nothing was still held; unreleased holds expire after 10 minutes")], [.85, 3.2, 2.85], size=9)
    sub("How a search decides what it can still promise")
    p("Each metro-day has one stored common arrangement: every technician's route with the host's appointments and every active hold, including moves pending offers depend on. Each search and confirm compares the host's snapshot with it, technician-day by technician-day, and nothing is guessed: a hold either still fits the current facts exactly as stored, or it is lost and reported.")
    table(["Rule", "Technician-day in the snapshot", "What the scheduler does"], [
        ("0", "Carries a lastModified the scheduler has already moved past (it remembers the last 32), or the unchanged timestamp with exactly the route from before a confirm receipt the host has not written yet", "Refused with 409 STALE before anything is reconciled, so nothing is lost; changed lists the newest timestamp seen"),
        ("1", "lastModified unchanged", "As stored"),
        ("2", "Changed, and its appointments in order are exactly what the scheduler expects (for example after the host wrote a receipt)", "Accepted; the new timestamp is recorded"),
        ("3", "Changed in any other way (a dispatcher added, removed or moved an appointment)", "Replaces the stored route when that route has no hold and no pending move; otherwise every hold on that date is lost and the date restarts from the host's routes"),
        ("4", "The whole arrangement, re-evaluated with the snapshot's current windows, durations, shifts and absences", "If it no longer fits, every hold on that date is lost the same way")], [.45, 3.1, 3.35], size=9)
    p("Rule 0 is what keeps several booking desks safe. A request built from older records can arrive after a newer receipt was already used, and without the rule the older routes looked like a dispatcher's edit and lost every hold on the date. With it the late request is told to rebuild from current records, and the other desks' holds stand. The search itself runs two passes within the deadline: an insertion pass that places the job on each day and finds the regular windows, and a refinement that improves cost, optional when the first pass already found more distinct windows than the client's threshold, and shed while other customers wait for a slot.")

    heading("6 Daily, repair and queued proposals")
    figure("daily", "Daily and repair proposals: synchronous or queued, then the commit with the host's current timestamps. The lower left lists the refusals and the cancel.")
    p("POST /daily/proposals sends the snapshot of one metro day and gets back a DailyProposal: proposed routes with planned times, unresolved demand, the policy decision and a reason. The scheduler parses strictly, claims the requestId, applies the 6 a.m. cutoff in the metro's zone (a frozen day is a 422), routes the day through its metro's service under its own 30 s allowance, and then runs the calculation for the client's cost and fairness seconds plus a 5 s reserve inside the replica's admission limits. Only an IMPROVED decision can be committed; NO_IMPROVEMENT and REJECTED_BY_POLICY return the current routes. [1, 3]")
    table(["Decision", "Meaning", "Reason codes that may begin the reason"], [
        ("IMPROVED", "The calculation's plan was accepted; commit it", "COST_REDUCTION, FAIRNESS_IMPROVEMENT, OVERTIME_REDUCTION, COMPLETE_ZERO_OVERTIME_DEMAND"),
        ("REJECTED_BY_POLICY", "A valid plan was found and policy rejected it; current routes returned", "NO_POLICY_IMPROVEMENT, COST_CEILING_EXCEEDED, FAIRNESS_REGRESSION_WITHIN_ALLOWANCE, OVERTIME_PROHIBITED, OVERTIME_TARGET_MISMATCH"),
        ("NO_IMPROVEMENT", "No valid plan, or nothing to judge; current routes returned", "NO_VALID_CANDIDATE, INCOMPLETE_OR_INFEASIBLE_DEMAND, or a sentence only")], [1.4, 2.4, 3.1])
    p("The commit sends the host's current lastModified for every technician-day the proposal covers, skipped ones included. Equal timestamps mean the facts are unchanged, the cutoff is checked again, and a receipt lists every assignment on the proposal's routes; the host writes them with a compare-and-set. A proposal is committed at most once: its row is locked for the commit and the receipt table is unique per proposal. A repair proposal is the daily path with an absence added to the absent technician's day, in a mode where IMPROVED means every appointment is placed without overtime; anything else means a dispatcher decides.")
    sub("Queued proposals and cancellation")
    p("With the header Prefer: respond-async a proposal is queued and answered 202 at once with a RequestStatus; the host polls GET /requests/{requestId} no sooner than Retry-After and reads QUEUED, RUNNING, then COMPLETED with the status and body the call itself would have answered. Any replica's worker takes queued requests, the waiting request of the client whose last one started longest ago first, so one client's thirty overnight days never hold back another's single day. A client may have 50 requests queued or running; a request still waiting after 15 minutes completes with 429 without running. Overnight optimization is the host's to schedule with this queue; the scheduler has no batch of its own.")
    p("DELETE /requests/{requestId} cancels a proposal, booking search or confirm still running under that requestId, best effort, from any replica: a replica that does not run the call records the cancel on its claim, and the replica running it stops the call within about a second. The cancelled call answers 409 CANCELLED, stores nothing and releases its requestId, so retrying the same request later runs it again.")

    heading("7 Errors, idempotency and retries")
    p("Every error is a JSON Problem with an error code and a message, except 401. A host branches on the code, never on the message, which is for people and may change. The table is the integration guide's, which the spec's enum and this report's verifier keep in step. [3]")
    table(["Status", "Code", "Meaning", "What to do"], [[plain(s), plain(c), plain(m), plain(w)] for s, c, m, w in data["problems"][1]], [.45, 1.6, 2.6, 2.25], size=9)
    table(["Caller timeout for", "Set it to"], [
        ("Booking offers and confirm", "The client's bookingDeadlineMs plus about 3 seconds"),
        ("Daily and repair proposals, synchronous", "costSearchSeconds + fairnessSearchSeconds plus about 15 seconds"),
        ("Queued proposals", "A second for the 202; give up polling after the queue's 15 minutes plus the search time"),
        ("Everything else", "10 to 30 seconds"),
        ("After any client-side timeout", "Repeat the same requestId: the stored answer comes back, or a 429 while the call still runs")], [2.4, 4.5])
    p("A request that ended in 429, 503 or 409 CANCELLED released its requestId, so repeating it runs it again; a 400, 404, 409 of another kind or 422 is stored and replayed, so repeating it unchanged returns the same answer. That split is what lets a host retry blindly after a timeout without ever double-booking.")

    heading("8 Many clients and many replicas")
    p("One Postgres database and one set of tables serve every client. Each WaterFlex client is a tenant with its own API tokens (wfs_ plus 32 random bytes, stored only as SHA-256 digests, issued and revoked with the scheduler's own command line); the server maps the token to tenant_id, and a request cannot choose a tenant. Every client-owned table has a tenant_id column that starts every key and index, so two clients can both have technician 42. [1]")
    table(["Guard", "How it works"], [
        ("Token to tenant", "The bearer token decides the tenant; a revoked token or a disabled tenant gets 401."),
        ("Keys", "Every primary key, unique constraint and index starts with tenant_id."),
        ("Row-level security", "Each transaction runs as the scheduler_tenant database role with app.tenant_id set; policies restrict every client-owned table to that tenant, an unset tenant sees nothing, and that role cannot read any other table. A missed filter in application code fails closed."),
        ("Least privilege", "The runtime login is a member of scheduler_app, never the schema owner, so a runtime bug can never alter the schema; migrations run as the owner with --command=migrate."),
        ("Shared road legs", "Route legs are not tenant data and are shared across clients in a metro; the cache holds coordinates at about 1 m, is reachable only by the scheduler, and is purged within 30 days.")], [1.3, 5.6])
    sub("Several replicas on one database")
    p("Correctness lives in the database. Holds are protected by row locks and unique constraints on (tenant, technician, date) and by the timestamp comparison inside one transaction, so any replica can serve any request and the load balancer needs no sticky sessions. The two scheduled tasks, the weekly route cache cleanup and the daily retention purge, are idempotent on every replica. There is no hold expiry sweeper: expiry is a time, applied when a day is next used. What stays per replica by design is admission (the search slots, with queued bookings started by client: the client with the fewest running, then the one whose last booking started longest ago), the memory tier of the route cache, and the in-memory list of running calls, which the cancel flag in the database reaches from any replica.")
    p("A replica's role (SCHEDULER_ROLE booking, daily or all) decides only which background work runs there: a booking replica takes no queued daily work, so a queued day never takes a slot its bookings need, and a daily replica may run several calculations at once. The pools scale apart, each with its own CPU and memory. Measured on the VM rehearsal stack, two pools at the same total capacity nearly doubled the days optimized per minute and cut a booking's tail, at the cost of confirms waiting behind searches; the performance report covers it.")

    heading("9 The thin store and retention")
    p("The public API has its own tables, created by the scheduler's own migration and owned by nothing else; the portal's database holds none of them. Persisted proposal and booking state is validated again on every read, and a stored answer is validated again before it is replayed. [1]")
    table(["Table", "Holds"], [[plain(t), plain(h)] for t, h in data["tables"][1]], [1.6, 5.3], size=9)
    table(["Data", "Typical lifetime"], [[plain(d), plain(l)] for d, l in data["stored"][1]], [2.8, 4.1])
    p("The retention period is the deployment's SCHEDULER_RETENTION_DAYS (1 to 3650), required with no default: the scheduler does not start without it. A daily purge deletes rows older than the period in batches of 5,000 through a database function, children before parents; a row that a newer row still references stays until both have aged, and the function refuses a cutoff less than a day old. A host keeps its own copy of any receipt it needs longer.")

    heading("10 Routing and geocoding")
    p("Coordinates in the snapshot are authoritative: when a location arrives with latitude and longitude the scheduler routes from them as given and never geocodes again. Nominatim is the fallback only for a location that arrives with an address and no coordinates, only in NOMINATIM mode, and only as an exact house: ZIP, house number and street must agree, a mailing city is accepted when the ZIP agrees, and a street-level or ambiguous result is not located. Every drawn stop says where its coordinates came from (HOST or NOMINATIM), and each stored hold notes its job's source, so a bad route can be traced to its input. [1]")
    figure("tiers", "The four route cache tiers in lookup order, from the design document. A pair of coordinates is answered by the first tier that holds it; a leg computed by the routing service is written back to the database and memory tiers in a separate short transaction.")
    table(["Tier", "Where", "Key", "Lifetime and size", "Shared by"], [[plain(c) for c in row] for row in data["tiers"][1]], [1.0, 1.5, 1.4, 1.9, 1.1], size=9)
    p("None of the tiers is a source of truth: every entry is recomputable from the road graph, and every entry is bound to a routing identity, a hash of the graph and map version that the routing service reports and checks on every call. A map update, or a switch between the flexible and contraction hierarchy graphs, changes the identity, clears the memory tiers and leaves the database tier unused for the old identity, so a changed road network can never be mixed with old legs. Each map version is prepared as a whole (the merged OpenStreetMap extract, both road graphs, tiles, and TIGER address ranges for Nominatim), validated with preview services, and activated by switching symlinks.")

    heading("11 Operating the service")
    p("The runbook (docs/operations-runbook.md) is the operator's manual: install, start, stop and upgrade order, health, logs, metrics, retention, tokens, maps, backup and restore, scaling and failure modes. Its settings table is reproduced here from the document. [4]")
    table(["Setting", "Required", "Meaning"], [[plain(s), plain(r), plain(m)] for s, r, m in data["settings"][1]], [1.9, .75, 4.25], size=9)
    table(["Signal", "Where", "What it tells an operator"], [
        ("/health", "Port 8000, open", "ok or degraded with each component (schema, database, routing per metro with its identity, calculation, geocoding) and the replica's role; 503 when degraded"),
        ("Request log line", "Each replica's log", "One line per API call: operation, status, elapsed time, and how it was spent (slot wait, routing, calculation, database, routing pairs, cache hits, legs fetched, CPU), with tenant and requestId on every line of the call"),
        ("Prometheus metrics", "Management port 9000, private", "Request timers by operation and status, admission rejections, routing legs and cache hits by tier, solve seconds, invalid refinements, retention rows deleted"),
        ("Routing /health", "Each routing service", "Readiness, map version, prepared graph (flexible or CH) and the routing identity")], [1.3, 1.5, 4.1])
    p("Upgrades deploy a new solver pool beside the old one when the calculation protocol or the Timefold artifact changes, then switch the scheduler; a remote result must come from the same artifact and model versions as the caller, deliberately. A map refresh starts every route cache cold, which the first searches on a new map pay for.")

    heading("12 How it was built")
    figure("history", "Pull requests merged per day in the stateless API work, by phase, with the contract versions published on each day. The design was proposed on 2026-10-07; the service was serving every operation by 2026-10-09, stable at 1.0.0 on 2026-10-13, and measured under load from then on.")
    p(f"The work ran in phases that the design document names. S1 hardened what the stateless service would inherit (one overnight run per metro-day across replicas, cancellation routed to the solving replica, a capped route cache, a token on the routing service). S2 added tenants, tokens, the draft OpenAPI contract and strict Java types. S3 to S5 built the daily path from a request instead of the database, the thin store, and the commit. Booking and repair followed. S6 turned the portal into the first client, feature by feature behind a setting, until P7 removed the scheduler's own endpoints and database path and P8 dropped their tables. Then the contract gaps were closed, the integration package built (operations hardening, the scheduler's own database, the conformance kit, the guide and contract 1.0.0), and the service measured and improved under load. [1, 5]")
    table(["Phase", "Pull requests", "Dates", "What it delivered"], [
        [ph, ", ".join(f"#{n}" for n, _, p_, _ in prs if p_ == ph), f"{min(d for _, d, p_, _ in prs if p_ == ph)[5:]} to {max(d for _, d, p_, _ in prs if p_ == ph)[5:]}",
         {"Design": "The stateless API and thin store proposal, and the deployment, tenancy and route cache decisions",
          "S1 hardening": "Replica-safe overnight runs, routed cancellation, a capped route cache, routing authentication",
          "S2 contract": "Tenants and tokens, the draft OpenAPI 3.1 contract, strict Java types that mirror it",
          "S3 daily": "The daily solver input from a request, per-metro routing, address rules, TIGER, daily proposals served",
          "S4 thin store": "The api_* tables for requests and proposals",
          "S5 commit": "The daily commit with STALE and NOT_COMMITTABLE",
          "Booking": "Facts from a request, offers, select, release, confirm, Nominatim fallback",
          "Repair": "Repair proposals without overtime",
          "S6 portal": "The portal as a switchable multi-client API client: snapshots, booking, dispatch, repair, overnight, local tokens",
          "P7 retirement": "Every portal path moved to the API; the scheduler's own endpoints, database path and benchmarks deleted",
          "P8 cleanup": "Unused tables and code dropped; the design document brought up to date",
          "Client policy": "Window length, daily search method and times chosen per client",
          "Contract gaps": "Refinement fallback and reason codes, per-client policy, host cancellation, coordinate sources",
          "Integration": "Readiness, logs, metrics, retention, least privilege, the scheduler's own database, the conformance kit, the guide and contract 1.0.0",
          "Performance": "The load tool and the measured fixes: holds kept, early refusal, searchComplete, shedding, queued proposals, replica pools"}[ph]] for ph in phases], [1.0, 1.65, 1.0, 3.25], size=9)
    p("Two proofs carried the migration. Equivalence was proven where it matters: an integration test showed that a request describing a seeded day produced exactly the daily and booking facts the database path produced, before that path was deleted. And the portal moved one feature at a time behind a setting with smoke tests of both paths, so each step could be compared with the one before it, until the setting itself was removed.")

    heading("13 Proving an integration")
    p(f"A host proves its integration with the conformance kit (infra/host-conformance), a standalone script that runs {len(cases)} named cases against any scheduler over the network with a token for a dedicated tenant and writes a pass/fail report. The cases are the guide's sequences and refusals, with synthetic but routable snapshots inside a given area; a CI job runs the kit against the scheduler and a fixture router so the kit cannot drift from the service. [6]")
    rows = [cases[i:i + 3] + [""] * (3 - len(cases[i:i + 3])) for i in range(0, len(cases), 3)]
    table(["Cases", "", ""], rows, [2.3, 2.3, 2.3], size=9)
    p("The same folder holds the load tool that plays WaterFlex Software at scale (booking desks and daily planners against a stack), whose measurements are the subject of the performance report. Before go-live the guide's checklist asks for TLS on every connection, one tenant token per client with a rotation procedure, every served metro in ROUTING_METRO_URLS, a tested lastModified that moves on every change including both sides of a move, receipt writes that are a compare-and-set, a new requestId per request and the same one per retry, callers that branch on error codes with the timeouts of section 7, an overnight schedule per client and metro, receipts kept on the host's side beyond the retention period, and the conformance kit passing against the production scheduler.")

    heading("14 What remains")
    table(["Item", "State", "Who decides"], [
        ("TLS", "Not set up between WaterFlex Software and the scheduler, nor from the scheduler to routing and Nominatim (the remote solver already requires https off loopback). Deferred by owner decision on 2026-10-13 and recorded as a precondition of production traffic in the guide, the runbook and the design document. Until then the API binds to a private interface only.", "Owner, before production traffic"),
        ("The VM itself", "The deploy/vm compose is the VM's shape and has been rehearsed locally; the latency target is defined at the WaterFlex caller across the network and has not been measured there yet.", "Owner and WaterFlex Software"),
        ("Google Maps terms", "Whether WaterFlex Software's agreement allows Google-geocoded coordinates to be passed to a non-Google routing engine and cached for up to 30 days; if not, the host sends addresses only.", "WaterFlex Software"),
        ("Dedicated databases", "Whether any client contractually requires one; every row already carries tenant_id, so a client can be moved without an API change.", "Owner"),
        ("Remote mode", "Wired in and covered by CI, not performance-tested; a short embedded-versus-remote study is planned.", "Owner"),
        ("Remote solve cancellation", "Tracked in memory per solver replica; a load balancer must hash on the Solve-Request-Id header so a cancel reaches the solving replica, as the design document shows for nginx.", "Operator")], [1.3, 4.3, 1.3], size=9)
    p("Everything else WaterFlex Software needs is on its own side and documented in the guide: build snapshots with a lastModified that follows the rules, write receipts with a compare-and-set, hold one tenant token per client, route each served metro, and run its own overnight schedule against the daily proposal endpoint. The portal in this repository is the worked example of each of these.")
    sub("How this report was checked")
    p("Every table that describes the contract or the service is read from its source when the report is built: the operations, error codes, schemas and version history from docs/api/openapi-v1.yaml; the endpoint, contract change, thin store, lifetime and route cache tables from docs/stateless-api-design.md; the ownership and error tables from docs/integration-guide.md; the settings from docs/operations-runbook.md; and the case names from the conformance kit. verify_report.py parses the same files again on its own and checks every such table cell for cell, along with the template parts. The pull request list is the one typed input and names its source. [1 to 6]")

    heading("Appendix A Every contract change")
    p("The design document's change table, newest first: what changed, the pull request or spec version it came with, and what WaterFlex Software does about it. [1]")
    table(["Change", "Since", "What WaterFlex Software does"], [[plain(c), plain(s), plain(w)] for c, s, w in data["changes"][1]], [2.5, .5, 3.9], size=9)

    heading("Appendix B Every pull request")
    p("The merged pull requests of the stateless API work, in order, as listed in the repository on 2026-10-14. Pull requests about other work in the same period (the Timefold audit, benchmark studies, the dispatch board, CI) are not listed. [5]")
    table(["PR", "Merged", "Phase", "Title"], [[f"#{n}", d, ph, t] for n, d, ph, t in prs], [.45, .8, 1.1, 4.55], size=9)

    heading("Appendix C Schemas and endpoints")
    p(f"The {len(data['schemas'])} object schemas of the contract, and the design document's endpoint table. [1, 2]")
    schemas = data["schemas"]
    rows = [schemas[i:i + 4] + [""] * (4 - len(schemas[i:i + 4])) for i in range(0, len(schemas), 4)]
    table(["Schemas", "", "", ""], rows, [1.725] * 4, size=9)
    table(["Endpoint", "Purpose"], [[plain(e), plain(purpose)] for e, purpose in data["endpoints"][1]], [2.0, 4.9], size=9)

    heading("Appendix D Definitions and evidence")
    table(["Term", "Meaning in this report"], [
        ("Snapshot", "The facts one call needs: rates, policy, technicians, technician-days with lastModified, committed appointments."),
        ("Input revision", "A canonical hash of a snapshot; every proposal and offer records the one it was computed from."),
        ("Technician-day", "One technician on one service date; the unit the lastModified token and every receipt are keyed by."),
        ("Hold", "A slot kept for 10 minutes for one offer; a set holds one per offer until one is selected."),
        ("Common arrangement", "The scheduler's stored routes for a metro-day with every active hold and pending move."),
        ("Receipt", "The assignments a commit or confirm asks the host to write, with the timestamps to compare."),
        ("Claim", "The row that holds a requestId while its call runs; released by 429, 503 and CANCELLED, kept with the stored answer otherwise."),
        ("Routing identity", "A hash of the road graph and map version, reported by the routing service and checked on every call."),
        ("Tenant", "One WaterFlex client as the scheduler sees it: chosen by the token, isolated by tenant_id and row-level security.")], [1.6, 5.3])
    p("[1] docs/stateless-api-design.md: the design, its decisions (marked Decided), the contract change table, the thin store, the route cache, the portal phases, remaining gaps and open questions.")
    p(f"[2] docs/api/openapi-v1.yaml, version {data['version']}: the machine-readable contract with its version history; web/scripts/openapi-v1.test.ts validates its examples and the served routes; PublicApiContractTest keeps the Java types in step.")
    p("[3] docs/integration-guide.md: the host's guide, with every sequence, payload, error and the go-live checklist.")
    p("[4] docs/operations-runbook.md: the operator's manual, with the settings table reproduced in section 11.")
    p("[5] The repository's merged pull requests, listed on 2026-10-14, as recorded in report_data.py.")
    p("[6] infra/host-conformance: the conformance kit (src/cases.ts names the cases) and the load tool.")
    p("Reproduce with docs/stateless-api-service-2026-10-14/requirements.txt, then run report_data.py, build_report.py and verify_report.py there. The builder reads the documents, regenerates the five figures and chart_data.json, and writes this editable Word file. EVIDENCE.md records methods.")

    settings = doc.settings.element
    for el in settings.findall(qn("w:updateFields")): settings.remove(el)
    update = OxmlElement("w:updateFields"); update.set(qn("w:val"), "true"); settings.append(update)
    stream = BytesIO(); doc.save(stream)
    with ZipFile(REFERENCE) as ref, ZipFile(stream) as built, ZipFile(OUTPUT, "w", ZIP_DEFLATED) as output:
        preserve = set(PRESERVE) | {n for n in ref.namelist() if "footer" in n}
        for name in sorted(built.namelist()):
            content = ref.read(name) if name in preserve else built.read(name)
            info = ZipInfo(name, (2026, 10, 14, 0, 0, 0)); info.compress_type = ZIP_DEFLATED
            output.writestr(info, content)
    print(OUTPUT)


if __name__ == "__main__": build()
