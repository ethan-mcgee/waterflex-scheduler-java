"""Build the daily search follow-up report from validated evidence, connected to the first Tabu variant study.

Install requirements.txt first. No network, services or benchmark execution.
"""
from copy import deepcopy
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from statistics import mean, median
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from charts import (SOLVERS, ALTS, LONG_SOLVERS, LABEL, FLEETS, EARLIER_FLEETS, BUDGETS, LONG_FLEETS, LONG_BUDGETS, WORKLOADS, WORK_LABEL,
                    select, pair, avg, versus_tabu, make_charts)
from report_data import load, earlier, MAIN, MAIN_ANALYSIS, LONG, LONG_ANALYSIS, CONFIRM, CONFIRM_ANALYSIS, REVISION

HERE = Path(__file__).resolve().parent
REFERENCE = HERE.parent / "WaterFlex_Scheduler_Algorithm_Comparison_Explained.docx"
REFERENCE_SHA = "5366f3017d6f18c15de295ba5fb9c1a014b6222911040f1c5e044845e57f832d"
OUTPUT = HERE / "Daily_Search_Followup_6454fb05_Report.docx"
PRESERVE = ("word/styles.xml", "word/stylesWithEffects.xml", "word/numbering.xml", "word/theme/theme1.xml", "word/fontTable.xml")
SHORT = {"TABU": "Tabu", "TABU_KOPT": "Tabu+K", "TABU_KOPT_SIZE_3": "Tabu+K 3", "TABU_SIZE_3": "Size 3", "KOPT": "K opt", "LATE_ACCEPTANCE": "Late acc."}
FIRST = "the first study"


def dollars(value):
    return f"${value:,.2f}" if value >= 0 else f"-${-value:,.2f}"


def signed(value):
    return f"{value:+.2f}"


def counts(values):
    return "/".join(str(sum(test(v) for v in values)) for test in (lambda x: x > 0, lambda x: x == 0, lambda x: x < 0))


def wtl(rows, field="saving"):
    return counts([r[field] for r in rows])


def gain(rows, solver, short_ms, long_ms, **kw):
    """Same solver, fixture and seed: the shorter budget's accepted cost minus the longer budget's, in integer cents."""
    short = {r["key"][:4]: r for r in select(rows, solver=solver, budget_ms=short_ms, **kw)}
    return [short[r["key"][:4]]["raw"]["after"]["costCents"] - r["raw"]["after"]["costCents"] for r in select(rows, solver=solver, budget_ms=long_ms, **kw)]


def exact_dollars(cents, divisor=1):
    """Mean of integer cents divided exactly, rounded half up, so a true half cent always prints the same way."""
    value = (Decimal(sum(cents)) / Decimal(len(cents) * divisor * 100)).quantize(Decimal("0.01"), ROUND_HALF_UP)
    return f"${value:,}" if value >= 0 else f"-${-value:,}"


def build():
    assert sha256(REFERENCE.read_bytes()).hexdigest() == REFERENCE_SHA
    data = load()
    main_config, long_config = data["main"][0], data["long"][0]
    M, L, C = pair(data["main"][1], "TABU"), pair(data["long"][1], "TABU_KOPT"), pair(data["confirm"][1], "TABU")
    E = pair(earlier(), "TABU")
    assets = HERE / "generated"
    charts = make_charts(M, L, C, E, assets)
    doc = Document(REFERENCE)
    table_properties = deepcopy(doc.tables[0]._tbl.tblPr)
    for child in list(doc.element.body):
        if child.tag != qn("w:sectPr"): doc.element.body.remove(child)
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype.endswith(("/image", "/hyperlink")): del doc.part.rels[rid]
    doc.core_properties.title = "WaterFlex daily search follow-up explained"
    doc.core_properties.subject = "Out of sample test of the Tabu variant findings, run 6454fb05 with long budget and sequential arms"
    doc.core_properties.author = doc.core_properties.last_modified_by = "WaterFlex"
    doc.core_properties.created = doc.core_properties.modified = datetime(2026, 10, 13)
    figures = []

    def p(text, style=None):
        assert chr(0x2014) not in text and chr(0x2013) not in text  # No em or en dashes.
        para = doc.add_paragraph(text, style); para.paragraph_format.widow_control = True
        return para

    def heading(text):
        h = p(text, "Heading 1"); h.paragraph_format.page_break_before = True

    def sub(text): p(text, "Heading 2")

    def table(headers, rows, widths, size=10):
        assert abs(sum(widths) - 6.9) < .01, widths
        t = doc.add_table(rows=1, cols=len(headers)); t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t._tbl.replace(t._tbl.tblPr, deepcopy(table_properties))
        for col, w in zip(t.columns, widths): col.width = Inches(w)
        for i, values in enumerate([headers, *rows]):
            assert len(values) == len(headers), values
            row = t.rows[0] if i == 0 else t.add_row()
            trpr = row._tr.get_or_add_trPr(); trpr.append(OxmlElement("w:cantSplit"))
            if i == 0: trpr.append(OxmlElement("w:tblHeader"))
            for j, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
                cell.width = Inches(width); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell.text = str(value)
                props = cell._tc.get_or_add_tcPr()
                shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "DCE6EF" if i == 0 else ("FFFFFF" if i % 2 else "F4F6F8")); props.append(shade)
                margins = OxmlElement("w:tcMar")
                for side in ("top", "left", "bottom", "right"):
                    n = OxmlElement("w:" + side); n.set(qn("w:w"), "65"); n.set(qn("w:type"), "dxa"); margins.append(n)
                props.append(margins)
                para = cell.paragraphs[0]; para.paragraph_format.space_after = Pt(2); para.paragraph_format.space_before = Pt(2); para.paragraph_format.line_spacing = 1.0
                if len(headers) > 2 and j > 0: para.alignment = 1
                if len(rows) <= 12 and i < len(rows): para.paragraph_format.keep_with_next = True  # Short tables stay on one page.
                for run in para.runs: run.font.name = "Calibri"; run.font.size = Pt(size); run.bold = i == 0
        spacer = p(""); spacer.paragraph_format.space_after = Pt(0); spacer.paragraph_format.line_spacing = Pt(3)

    def figure(name, caption):
        figures.append(name)
        caption = f"Figure {len(figures)}. {caption}"
        para = p(""); shape = para.add_run().add_picture(str(assets / f"{name}.png"), width=Inches(6.85))
        shape._inline.docPr.set("descr", caption); para.paragraph_format.keep_with_next = True
        cap = p(caption, "Caption"); cap.paragraph_format.space_after = Pt(7)
        for run in cap.runs:
            run.font.color.rgb = RGBColor(0, 0, 0); run.bold = False; run.italic = True

    def by(solver, rows=M, **kw):
        return select(rows, solver=solver, **kw)

    def save(solver, rows=M, **kw):
        return avg(by(solver, rows, **kw), "saving")

    # Facts every claim below depends on, checked before any prose uses them.
    n_alt = len(by("TABU_KOPT"))
    assert n_alt == 240 and len(M) == 1200 and len(L) == 192 and len(C) == 32
    assert all(save("TABU_KOPT", fleet=f) > 0 for f in FLEETS) and all(save("TABU_KOPT", workload=w) > 0 for w in WORKLOADS)
    assert all(save("TABU_KOPT_SIZE_3", fleet=f) > save("TABU_KOPT", fleet=f) for f in FLEETS if f >= 10)
    assert save("TABU_KOPT_SIZE_3", fleet=5) < save("TABU_KOPT", fleet=5)
    best_overall = max(ALTS, key=lambda s: save(s))
    assert best_overall == "TABU_KOPT_SIZE_3" == data["decision"]["candidate"]
    tk_twins = {r["key"][:-1]: r for r in by("TABU_KOPT")}

    def over_tk(**kw):
        return [tk_twins[r["key"][:-1]]["accepted_cost"] - r["accepted_cost"] for r in by("TABU_KOPT_SIZE_3", **kw)]

    kept = [r for r in M if r["kept_start"]]
    assert {r["fleet"] for r in kept} <= {5, 10} and all(r["raw"]["acceptanceReason"] == "FAIRNESS_REGRESSION_WITHIN_ALLOWANCE" for r in kept)
    rule_old = lambda f: "KOPT" if f <= 20 else "TABU_KOPT"
    old_rule = mean(avg(select(E, solver=rule_old(f), fleet=f), "saving") for f in EARLIER_FLEETS)
    new_rule_four = mean(save(rule_old(f), fleet=f) for f in EARLIER_FLEETS)
    rules = [("Tabu + K opt everywhere", lambda f: "TABU_KOPT"), ("Tabu + K opt size 3 everywhere", lambda f: "TABU_KOPT_SIZE_3"),
             (f"K opt to 20, then Tabu + K opt ({FIRST})", rule_old), ("K opt to 20, then Tabu + K opt size 3", lambda f: "KOPT" if f <= 20 else "TABU_KOPT_SIZE_3")]
    rule_value = {name: mean(save(choice(f), fleet=f) for f in FLEETS) for name, choice in rules}
    hindsight = {f: max(ALTS, key=lambda s: save(s, fleet=f)) for f in FLEETS}
    hindsight_value = mean(save(hindsight[f], fleet=f) for f in FLEETS)
    first_tk = avg(select(E, solver="TABU_KOPT"), "saving")
    curve = {(c["source"], c["solver"], c["seconds"]): c["saving"] for c in charts["time_curve"]}
    seq = [r["saving"] for r in by("TABU_KOPT_SIZE_3", C)]
    par_twins = {r["key"][:-1]: r for r in by("TABU_KOPT_SIZE_3")}
    par = [par_twins[r["key"][:-1]]["saving"] for r in by("TABU_KOPT_SIZE_3", C)]
    agree = sum((a > 0) == (b > 0) and (a < 0) == (b < 0) for a, b in zip(seq, par))
    move_ratio = {s: median(next(m for m in M if m["key"] == r["key"])["move_rate"] / r["move_rate"] for r in by(s, C)) for s in ("TABU", "TABU_KOPT_SIZE_3")}

    p("WaterFlex daily search follow-up explained", "Title")
    p("Testing the Tabu variant findings on new schedules, longer searches and one combined method", "Subtitle")
    p("October 13 2026  |  Runs 6454fb05 and a317757e with sequential confirmation 877fd9f7")
    tk, tk3 = by("TABU_KOPT"), by("TABU_KOPT_SIZE_3")
    direct = over_tk()
    p(f"This follow-up tested the conclusions of the Tabu variant study (run 5d15129b) on {len({r['raw']['datasetFingerprint'] for r in M})} new synthetic fixtures that study never saw, "
      f"with the whole time budget spent on the cost search. Tabu + K opt, now the default daily search for existing clients, again beat Tabu at every fleet size: "
      f"{sum(r['saving'] > 0 for r in tk)} wins, {sum(r['saving'] == 0 for r in tk)} ties and {sum(r['saving'] < 0 for r in tk)} losses over {n_alt} matched cases, "
      f"saving {dollars(avg(tk, 'saving'))} per case. The new combination, Tabu + K opt size 3, did better still. It saved {dollars(avg(tk3, 'saving'))} per case "
      f"and beat Tabu + K opt itself in {sum(v > 0 for v in direct)} of {len(direct)} cases. Rerun one case at a time it beat Tabu in {sum(v > 0 for v in seq)} of {len(seq)}. [1, 3]")
    sub("What the comparison tells us")
    table(["Configuration", "Mean accepted cost", "Mean saving vs Tabu", "Wins / ties / losses"],
          [[LABEL[s], dollars(avg(by(s), "accepted_cost")), "Control" if s == "TABU" else dollars(save(s)), "Control" if s == "TABU" else wtl(by(s))]
           for s in SOLVERS], [1.9, 1.6, 1.7, 1.7])
    p(f"It also answers the four questions {FIRST} left open. Tabu size 3 and the K opt moves combine from 10 technicians up. "
      f"The fleet size rule {FIRST} picked in hindsight held up on new schedules ({dollars(new_rule_four)} per case against {dollars(old_rule)}), but Tabu + K opt size 3 "
      f"matches it without any rule. The late acceptance family is slow at 50 technicians, not broken: it reaches Tabu + K opt only after 480 to 600 seconds, more than "
      f"the 240 a client can allow. Every cheaper method spreads work less evenly than Tabu when no fairness search follows. Costs are modeled on synthetic schedules, "
      f"not payroll savings.")
    sub("How to read this report")
    table(["Sections", "What you will learn"], [("1 to 4", "Why the follow-up was run, what was tested, what changed from the first study and the configurations."),
                                                ("5", "Whether Tabu + K opt holds up on new schedules."),
                                                ("6 and 7", "Why tabu size 3 and K opt combine, and whether a fleet size rule is needed."),
                                                ("8", "What longer searches buy, on one time axis across both studies."),
                                                ("9 and 10", "Workload balance without a fairness phase, and case to case variation."),
                                                ("11", "The one at a time confirmation."),
                                                ("12 and appendices", "Decisions supported by the evidence, every fleet group, mechanics and sources.")], [1.45, 5.45])

    heading("1 Why this follow-up was run")
    p(f"The Tabu variant study ran 4,480 cases on 16 fixtures and found that Tabu combined with the K opt moves saved {dollars(first_tk)} per case against Tabu. "
      "It was the only configuration ahead at every fleet size, workload and budget. On October 12 2026 Tabu + K opt became a per client choice on the portal's "
      "Solver settings page and the default for existing clients. [5, 8]")
    p("That study chose its winner on the same sixteen fixtures it measured, and its section 14 left four questions open. This follow-up was designed to answer them "
      "on schedules that played no part in choosing anything:")
    long_la_600 = save("LATE_ACCEPTANCE", L, fleet=50, budget_ms=600000)
    table(["Open question", "How this study tests it", "Short answer", "Section"], [
        ("Does Tabu + K opt hold up?", "40 new fixtures, new search seeds, cost search only.", f"Yes. {dollars(avg(tk, 'saving'))} per case, ahead at every fleet size.", "5"),
        ("Combine both ideas?", "A new configuration, Tabu + K opt with tabu size 3.", f"Yes, from 10 technicians up: {dollars(mean(direct))} per case more than Tabu + K opt.", "6"),
        ("Choose by fleet size?", "Score the hindsight rule on new fixtures, and add 30 technicians.", f"The rule held up, but one method matches it.", "7"),
        ("Large fleet budgets", "240, 480 and 600 second searches at 30 and 50 technicians.", f"Late acceptance draws level with Tabu + K opt at 600 s ({signed(long_la_600)}).", "8"),
        ("Balance tradeoff", "Record workload variance of every cost search result.", "Cheaper methods are less even; a fairness search closes most of the gap.", "9")], [1.5, 2.15, 2.55, .7])

    heading("2 What was tested")
    p("One case is one configuration improving one fixed daily schedule for one search seed and one time budget. In a matched comparison every configuration receives "
      "the same starting schedule. The experiment rearranges appointments that are already booked; it does not admit new booking requests. [1, 6]")
    arms = [("Main", main_config, M), ("Long budget", long_config, L), ("Confirmation", data["confirm"][0], C)]
    table(["Arm", "Configurations", "Fleets", "Fixture seeds", "Search seeds", "Budgets (s)", "Cases", "At once"],
          [[name, len(cfg["daily"]["solvers"]), ", ".join(map(str, cfg["daily"]["fleets"])), ", ".join(map(str, cfg["daily"]["fixture_seeds"])),
            ", ".join(map(str, cfg["daily"]["seeds"])), ", ".join(map(str, cfg["daily"]["budgets_seconds"])), f"{len(rows):,}", cfg["daily"]["parallel_cases"]]
           for name, cfg, rows in arms], [1.05, 1.05, .9, .75, .85, .95, .65, .7])
    p(f"The main arm is 5 configurations x 5 fleets x 4 workloads x 2 fixtures x 3 seeds x 2 budgets = {len(M):,} cases, so each configuration has {n_alt}. "
      "A fleet, workload and budget group holds six cases per configuration: two fixtures times three seeds. The long budget arm compares four configurations on 30 and 50 "
      "technician days at much longer budgets. The confirmation reruns the main arm's best configuration beside Tabu one case at a time. All "
      f"{len(M) + len(L) + len(C):,} results are present and paired, report zero violations and zero overtime, and every search ended at its time limit.")
    sub("Cost search only")
    p("In the first study every case ran two searches: a cost search with two thirds of the budget, then a fairness search that could spend up to 2% more to spread work "
      "more evenly. This study removes the fairness search and gives the cost search the whole budget. That isolates the question the study asks: how good is each "
      "search method at finding cheap schedules?")
    p(f"The policy check still runs on every result. It keeps the cheaper schedule unless the starting schedule was already within 2% of it and the cheaper schedule "
      f"spreads work less evenly; then it keeps the start. That happened in {len(kept)} of {len(M):,} main cases, all on 5 and 10 technician days, "
      f"{sum(r['solver'] == 'TABU' for r in kept)} of them for Tabu. Those cases count with the starting cost, so they appear as ties or small differences.")
    sub("New fixtures")
    p("A fixture is the synthetic day being improved: its technicians, visits, skills, arrival windows and travel times. The benchmark now takes a fixture seed. "
      "Seed 0 reproduces the first study's sixteen fixtures byte for byte, which a fingerprint test enforces. Other seeds keep each workload's structure but change "
      "the geography: which technicians start with six visits instead of two, and the travel times within the workload's ranges. The main arm uses seeds 1 and 2, "
      "the long budget arm seed 3, and no fixture is shared between them. The search seeds 211, 223 and 227 had not been used before. [6]")
    table(["Workload", "How the fixture is built", "What makes it hard"], [
        ("Clustered", "Four travel groups: short trips within a group, longer ones between groups. 30 minute visits.", "Few big wins; gains come from keeping each route inside one group."),
        ("Dispersed", "Irregular one way travel of 3 to 10 minutes per leg with no group structure. 30 minute visits.", "Many routes of similar cost; reordering pays, and the search space is rough."),
        ("Mixed skill", "Shorter legs of 1.5 to 5 minutes. Only every third technician can do the scarce service, which half of their visits need.", "Many moves are illegal, so feasible improvements are sparse."),
        ("Tight window", "The same short legs, 45 minute visits and arrival windows starting every 90 instead of 120 minutes.", "Long jobs packed close together leave little slack for reordering.")], [1.15, 3.1, 2.65])
    first_improvement = {f: (1 - avg(select(E, solver="TABU", fleet=f), "reference_cost") / avg(select(E, solver="TABU", fleet=f), "start_cost")) * 100 for f in EARLIER_FLEETS}
    table(["Technicians", "Appointments", "Mean starting cost $", "Mean Tabu result $", "Tabu improvement", "First study, cost phase"],
          [[f, by("TABU", fleet=f)[0]["raw"]["appointments"], f"{avg(by('TABU', fleet=f), 'start_cost'):,.2f}", f"{avg(by('TABU', fleet=f), 'accepted_cost'):,.2f}",
            f"{(1 - avg(by('TABU', fleet=f), 'accepted_cost') / avg(by('TABU', fleet=f), 'start_cost')) * 100:.1f}%",
            f"{first_improvement[f]:.1f}%" if f in first_improvement else "Not tested"] for f in FLEETS], [1.0, 1.1, 1.3, 1.3, 1.1, 1.1])
    p("The new fixtures leave Tabu about as much room to improve as the old ones did, so the two studies are working on comparable problems. As before, five "
      "technician days are already close to as good as these searches can make them, and the configurations separate on larger fleets.")

    heading("3 What changed from the first study")
    table(["Dimension", "First study (5d15129b)", "This study"], [
        ("Fixtures", "16, fixture seed 0", "40 in the main arm (seeds 1 and 2), 8 in the long arm (seed 3)"),
        ("Search seeds", "Ten seeds", "211, 223 and 227, never used before"),
        ("Fleet sizes", "5, 10, 20 and 50", "Adds 30, where the first study's fleet rule switches"),
        ("Search phases", "Cost search for two thirds of the budget, then a fairness search within 2%", "Cost search for the whole budget"),
        ("Cost search time", "40, 60, 80 and 160 seconds (budgets of 60 to 240)", "60 and 120 seconds; 240, 480 and 600 in the long arm"),
        ("Configurations", "Seven, including Tabu size 15 and sublist", "Drops size 15 and sublist; adds Tabu + K opt size 3"),
        ("Cases at once", "Six; confirmation one at a time", "Six; confirmation one at a time"),
        ("Measured revision", "99be00b", REVISION[:7])], [1.55, 2.6, 2.75])
    p("Because of the phase change, budgets do not line up by their label. A 60 second case here searches for cost for 60 seconds; a 60 second case in the first study "
      "searched for cost for only 40. The fair comparison is by seconds of cost search. The first study's 90 second budget gave its cost search 60 seconds, the same as "
      "this study's 60 second cases, and sections 5 and 8 compare on that basis.")
    p("The first study's headline savings were measured after its fairness search. That search cost every configuration about the same, so its savings before and after "
      "fairness were nearly equal (its section 4). The headline numbers of the two studies are therefore comparable in size, though not identical in definition.")

    heading("4 The configurations")
    p("Every configuration uses the same Timefold engine, score rules, starting schedules and budgets. They differ in which kinds of moves they try, the acceptor and "
      "the forager. The first report's sections 2 and 3 explain these ideas in detail. [5, 7]")
    table(["Configuration", "Moves tried", "Acceptor", "Accepted per step", "Arms"], [
        ("Tabu (control)", "Change, swap", "Entity tabu 7", "1,000", "All three"),
        ("Tabu + K opt", "Change, swap, sublist, 2 to 3 opt", "Entity tabu 7", "1,000", "Main, long budget"),
        ("Tabu + K opt size 3", "Change, swap, sublist, 2 to 3 opt", "Entity tabu 3", "1,000", "Main, confirmation"),
        ("Tabu size 3", "Change, swap", "Entity tabu 3", "1,000", "Main"),
        ("K opt", "Change, swap, sublist, 2 to 3 opt", "Late acceptance 400", "1", "Main, long budget"),
        ("Late acceptance", "Change, swap", "Late acceptance 400", "1", "Long budget")], [1.6, 2.1, 1.35, .95, .9])
    table(["Term", "Meaning in this report"], [
        ("Step", "The moment the search adopts one candidate move and changes its current schedule."),
        ("Moves per step", "How many candidates the search examined before adopting one. Tabu examines until it has 1,000 allowed ones, or 10,000 in all."),
        ("Tabu size", "How many recently moved items the search may not touch again. A longer list forbids more candidates, so each step costs more examinations."),
        ("Late acceptance", "Adopts the first candidate no worse than the schedule of 400 steps ago. Many cheap steps, each less informed."),
        ("Best found at", "The share of the search that had elapsed when its final best schedule first appeared. Near 100% means still improving at the deadline.")], [1.55, 5.35])
    sub("Why Tabu + K opt size 3 was built")
    p("The first study found two separate improvements on Tabu. Tabu size 3 forbids fewer items, so each step finds its 1,000 allowed candidates sooner and the search "
      "takes more steps. Tabu + K opt adds sublist and K opt moves, which restructure routes, and Tabu's forager adopts them whenever they are the best candidate. "
      "One improvement changes how often the search steps, the other what each step can do. If they do not interfere, combining them should keep both gains.")

    heading("5 Does Tabu + K opt hold up on new schedules?")
    p("Each cell averages 48 paired differences: four workloads, two fixtures, three seeds and two budgets, holding fixture and starting schedule fixed. Read the sign "
      "before the size. Blue cells mean the configuration accepted cheaper schedules than Tabu; tan cells mean Tabu was cheaper. [1]")
    figure("cost_fleet", "Mean accepted cost saving vs Tabu in dollars per case on the new fixtures. Each fleet cell has 48 pairs; the All fleets column has 240. "
                         "Tabu + K opt size 3 is bold.")
    p(f"Tabu + K opt is positive at every fleet size and on every workload, and largest at 50 technicians ({signed(save('TABU_KOPT', fleet=50))}). That is the same "
      f"shape {FIRST} found. Tabu + K opt size 3 sits above it from 10 technicians up. K opt is the strongest method at 10 and 20 technicians "
      f"and the weakest at 50 ({signed(save('KOPT', fleet=50))}), again as before. Tabu size 3 helps everywhere except 5 technician days.")
    table(["Wins / ties / losses", "5 techs", "10 techs", "20 techs", "30 techs", "50 techs"],
          [[LABEL[s], *[wtl(by(s, fleet=f)) for f in FLEETS]] for s in ALTS], [1.9, 1.0, 1.0, 1.0, 1.0, 1.0])
    p("Each entry counts 48 matched cases in exact cents. On 5 technician days most comparisons are ties, because every method finds the same schedule or the policy "
      "check keeps the start. From 10 technicians up, the two Tabu + K opt variants lose very few cases.")
    sub("The same answer at the same search time")
    figure("replication", "Saving vs Tabu after 60 seconds of cost search. Hatched: first study fixtures (its 90 second budget, 40 pairs per bar). "
                          "Solid: new fixtures (60 second cases, 24 pairs per bar). Cost search savings, before any fairness search.")
    rep = {d["solver"]: d for d in charts["replication"]}
    p(f"Holding the cost search time fixed at 60 seconds, each method's pattern across fleet sizes repeats on the new schedules. Tabu + K opt saves "
      f"{signed(rep['TABU_KOPT']['earlier'][3])} and {signed(rep['TABU_KOPT']['new'][3])} at 50 technicians on the old and new fixtures. Tabu size 3 is slightly worse "
      f"than Tabu on 5 technician days in both. K opt wins at 10 and 20 technicians in both and loses at 50 in both, by "
      f"{dollars(-rep['KOPT']['earlier'][3])} and {dollars(-rep['KOPT']['new'][3])}. The sizes move a few dollars from fixture to fixture, but no sign flips.")
    p(f"So the first study's conclusions were not an accident of its sixteen fixtures. Pooled over its whole matrix, Tabu + K opt saved {dollars(first_tk)} per case "
      f"there and {dollars(avg(tk, 'saving'))} here.")

    heading("6 Combining tabu size 3 with K opt")
    p("Tabu + K opt size 3 changes one setting of Tabu + K opt: the tabu list holds 3 recently moved items instead of 7. Its effect runs through the step economics "
      "the first study described. [1, 5]")
    figure("steps", "Cost search means per configuration and fleet, 48 cases per point, on log scales. Left: steps per second. Right: candidate moves examined per step.")
    table(["Fleet", "Tabu + K opt steps / s", "Size 3 steps / s", "Tabu + K opt moves / step", "Size 3 moves / step", "Size 3 vs Tabu + K opt", "W/T/L"],
          [[f, f"{avg(by('TABU_KOPT', fleet=f), 'steps_rate'):.1f}", f"{avg(by('TABU_KOPT_SIZE_3', fleet=f), 'steps_rate'):.1f}",
            f"{avg(by('TABU_KOPT', fleet=f), 'moves_per_step'):,.0f}", f"{avg(by('TABU_KOPT_SIZE_3', fleet=f), 'moves_per_step'):,.0f}",
            signed(mean(over_tk(fleet=f))), counts(over_tk(fleet=f))] for f in FLEETS], [.6, 1.05, .95, 1.15, 1.0, 1.15, 1.0])
    at20 = lambda s, field: avg(by(s, fleet=20), field)
    p(f"At 20 technicians Tabu + K opt examines about {at20('TABU_KOPT', 'moves_per_step'):,.0f} candidates per step and takes {at20('TABU_KOPT', 'steps_rate'):.1f} "
      f"steps per second. With a shorter tabu list fewer candidates are forbidden, so the size 3 version finds its 1,000 allowed ones after about "
      f"{at20('TABU_KOPT_SIZE_3', 'moves_per_step'):,.0f} and steps {at20('TABU_KOPT_SIZE_3', 'steps_rate') / at20('TABU_KOPT', 'steps_rate'):.1f} times as often. "
      f"Both examine about the same number of moves per second ({at20('TABU_KOPT', 'move_rate'):.0f} and {at20('TABU_KOPT_SIZE_3', 'move_rate'):.0f} thousand), "
      "so the extra steps are free. Every one of them is still the best of up to 1,000 candidates, including the structural moves.")
    figure("combination", "Left: mean saving vs Tabu by fleet, 48 pairs per point, with the sum of the two single improvements dashed. Right: Tabu + K opt size 3 against "
                          "Tabu + K opt on the same 48 cases per fleet, counted in exact cents.")
    share = mean(save("TABU_KOPT_SIZE_3", fleet=f) / (save("TABU_KOPT", fleet=f) + save("TABU_SIZE_3", fleet=f)) for f in FLEETS if f >= 10)
    assert .5 < share < 1
    p(f"The two improvements mostly add up. From 10 technicians up the combination keeps on average {share * 100:.0f}% of the sum of the two separate gains, and it "
      f"beats Tabu + K opt in {sum(v > 0 for f in FLEETS if f >= 10 for v in over_tk(fleet=f))} of {sum(len(over_tk(fleet=f)) for f in FLEETS if f >= 10)} cases there. "
      "A likely reason they do not add up completely is that they overlap: both help the search keep moving where Tabu would otherwise stall. That is an "
      "interpretation; the archive records the outcome, not the cause.")
    five = over_tk(fleet=5)
    p(f"On 5 technician days the combination never beats Tabu + K opt: {counts(five)} wins, ties and losses, {signed(mean(five))} per case. Its best schedule appears "
      f"after {avg(by('TABU_KOPT_SIZE_3', fleet=5), 'best_fraction') * 100:.0f}% of the search on average. With only 18 visits, a list of 3 lets the search return to "
      f"arrangements it just left, the same weakness {FIRST} saw in Tabu size 3. That reading is an interpretation consistent with the early best time, not something "
      "the archive records directly. The loss is small in dollars because five technician days have little room to improve.")

    heading("7 Choosing a method by fleet size")
    p(f"{FIRST.capitalize()} found that K opt beat everything at 10 and 20 technicians and lost badly at 50, so it suggested a rule: K opt up to 20 technicians, Tabu + K opt "
      "above. This study scores that rule on new fixtures and adds 30 technician days to locate the switch point. [1, 5]")
    figure("fleet_budget", "Mean saving vs Tabu by fleet at each cost search time, 24 pairs per point.")
    p(f"K opt's advantage depends on both fleet size and time. At 30 technicians it saves {signed(save('KOPT', fleet=30, budget_ms=60000))} with 60 seconds and "
      f"{signed(save('KOPT', fleet=30, budget_ms=120000))} with 120; at 50 technicians it loses {dollars(-save('KOPT', fleet=50, budget_ms=60000))} and "
      f"{dollars(-save('KOPT', fleet=50, budget_ms=120000))}. The Tabu family barely changes between the two times. Section 8 shows why: K opt is still improving when "
      "its time runs out, and Tabu is not.")
    figure("convergence", "Left: mean time of the best schedule as a percentage of the search. Right: percentage of cases whose best appeared in the final 10%. 48 cases per point.")
    p(f"At 50 technicians K opt found its best schedule after {avg(by('KOPT', fleet=50), 'best_fraction') * 100:.0f}% of the search on average, and "
      f"{avg(by('KOPT', fleet=50), 'late_best') * 100:.0f}% of its cases were still improving in the final 10%. The Tabu family's best arrived about three quarters of "
      f"the way through. K opt's late acceptance search takes about {avg(by('KOPT', fleet=50), 'steps_rate'):,.0f} cheap steps per second at 50 technicians. "
      "That is plenty on a small day and not enough on a 200 visit day.")
    table(["Rule", "5", "10", "20", "30", "50", "Mean"],
          [[name, *[signed(save(choice(f), fleet=f)) for f in FLEETS], dollars(rule_value[name])] for name, choice in rules], [2.6, .7, .7, .7, .7, .7, .8])
    p(f"Each fleet column is a mean over 48 cases, and the last column weights the five fleets equally. {FIRST.capitalize()} reported {dollars(old_rule)} for its rule, "
      f"chosen in hindsight on fleets 5, 10, 20 and 50. On the new fixtures the same rule scores {dollars(new_rule_four)} on those four fleets, almost exactly the same: "
      f"it held up out of sample. With 30 technicians included it scores {dollars(rule_value[rules[2][0]])}.")
    p(f"But a single method now matches it. Tabu + K opt size 3 everywhere scores {dollars(rule_value[rules[1][0]])} without depending on fleet size or on K opt "
      f"having enough time. The best rule in hindsight on this data, K opt up to 20 technicians and Tabu + K opt size 3 above, scores {dollars(hindsight_value)}. "
      "That figure is selected on the same data it is measured on, so it is optimistic, just as the first study's rule was before this test.")
    assert [hindsight[f] for f in FLEETS] == ["KOPT", "KOPT", "KOPT", "TABU_KOPT_SIZE_3", "TABU_KOPT_SIZE_3"]
    assert abs(hindsight_value - rule_value[rules[3][0]]) < 1e-9

    heading("8 What longer searches buy")
    p("The long budget arm asks how far the late acceptance family improves when time is not the constraint. It runs Tabu + K opt (the control), Tabu, K opt and late "
      "acceptance on 30 and 50 technician days with 240, 480 and 600 seconds of cost search. Each point has only 8 pairs (four workloads, one fixture, two seeds), so "
      "read the direction more than the exact size. [2]")
    figure("long_budget", "Mean saving vs Tabu + K opt by cost search time, 8 pairs per point. Above zero is cheaper than Tabu + K opt.")
    table(["Configuration", "Fleet", "240 s", "480 s", "600 s"],
          [[LABEL[s], f, *[f"{save(s, L, fleet=f, budget_ms=b * 1000):+.2f} ({wtl(by(s, L, fleet=f, budget_ms=b * 1000))})" for b in LONG_BUDGETS]]
           for s in LONG_SOLVERS[1:] for f in LONG_FLEETS], [1.75, .8, 1.45, 1.45, 1.45])
    tabu_long = by("TABU", L)
    assert all(r["saving"] < 0 for r in tabu_long)
    p(f"Tabu loses every one of its {len(tabu_long)} cases to Tabu + K opt, by about $10 to $12 per case at every budget. Its gap does not close with time: the K opt moves "
      f"find schedules Tabu does not reach at any of these budgets. K opt is ahead of Tabu + K opt on 30 technician days from 240 seconds. On 50 technician days it "
      f"is behind at 240 seconds and ahead from 480 ({signed(save('KOPT', L, fleet=50, budget_ms=480000))} and {signed(save('KOPT', L, fleet=50, budget_ms=600000))}). "
      f"Late acceptance on 50 technician days moves from {signed(save('LATE_ACCEPTANCE', L, fleet=50, budget_ms=240000))} to {signed(long_la_600)}, level with "
      "Tabu + K opt. On 30 technician days it stays a few dollars behind at every budget.")
    table(["Configuration", "Fleet", "240 to 480 s", "Per minute", "480 to 600 s", "Per minute"],
          [[LABEL[s], f, exact_dollars(g1 := gain(L, s, 240000, 480000, fleet=f)), exact_dollars(g1, 4), exact_dollars(g2 := gain(L, s, 480000, 600000, fleet=f)),
            exact_dollars(g2, 2)] for s in LONG_SOLVERS for f in LONG_FLEETS], [1.75, .7, 1.15, 1.0, 1.25, 1.05])
    p("This table compares each configuration with itself, holding fixture and seed fixed, over 8 pairs. Tabu + K opt gains about $1 to $3 from the extra minutes; "
      "its search is mostly done by 240 seconds. On 50 technician days K opt and late acceptance gain about $10 between 240 and 480 seconds, but their gain per minute "
      "falls in the last two minutes. They are still improving at 600 seconds, more slowly.")
    sub("One time axis across both studies")
    figure("time_curve", "Mean saving vs Tabu at 50 technicians by seconds of cost search. Hollow points: first study (cost search of its 60, 90, 120 and 240 second budgets, "
                         "40 pairs each, fixture seed 0). Filled: this study's main arm (60 and 120 s, 24 pairs) and long arm (240 to 600 s, 8 pairs, fixture seed 3). "
                         "Fixtures differ between sources.")
    assert curve[("main", "KOPT", 120)] < 0 < curve[("long", "KOPT", 240)]
    assert curve[("5d15129b", "LATE_ACCEPTANCE", 160)] < 0 < curve[("long", "LATE_ACCEPTANCE", 240)]
    p(f"Measured by seconds of cost search, the two studies trace one curve. K opt at 50 technicians is {dollars(-curve[('5d15129b', 'KOPT', 40)])} behind Tabu after 40 "
      f"seconds, {dollars(-curve[('main', 'KOPT', 120)])} behind after 120, {dollars(curve[('long', 'KOPT', 240)])} ahead after 240 and "
      f"{dollars(curve[('long', 'KOPT', 600)])} ahead after 600. Late acceptance follows almost the same path. Tabu + K opt stays between "
      f"{dollars(min(v for (src, s, t), v in curve.items() if s == 'TABU_KOPT'))} and {dollars(max(v for (src, s, t), v in curve.items() if s == 'TABU_KOPT'))} "
      "ahead of Tabu at every time, because its gain comes from better steps, not more of them.")
    p("This answers the first study's budget question. On a 50 technician day the late acceptance family passes Tabu somewhere between 120 and 240 seconds of cost "
      "search, and reaches Tabu + K opt only at 480 to 600 seconds. A client's daily search may spend at most 240 seconds in total. Within that limit the Tabu + K opt "
      "methods remain the better choice for large fleets.")

    heading("9 Balance without a fairness phase")
    p("Fairness variance measures how unevenly paid time falls across technicians, relative to each one's capacity. Lower is more even. For teaching arithmetic: two "
      "technicians with equal 480 minute days and 240 and 480 paid minutes have utilizations of 0.50 and 1.00, and a variance of 0.0625. These illustrative values are "
      "not experiment observations. Because this study has no fairness search, every number here is the balance of the cheapest schedule each method found. [1, 6]")
    figure("fairness", "Each point averages 48 cases for one configuration and fleet; labels give the fleet. Right is cheaper than Tabu; up is more balanced than Tabu. "
                       "Variance differences are multiplied by 1,000.")
    table(["Configuration", "5 techs", "10 techs", "20 techs", "30 techs", "50 techs", "All"],
          [[LABEL[s], *[f"{avg(by(s, fleet=f), 'fairness'):.4f}" for f in FLEETS], f"{avg(by(s), 'fairness'):.4f}"] for s in SOLVERS], [1.9, .8, .85, .85, .85, .85, .8])
    assert all(avg(by(s), "fairness") > avg(by("TABU"), "fairness") for s in ALTS)
    p("Tabu produces the most even workloads, and every cheaper method is less even from 10 technicians up. The order roughly follows the savings: Tabu size 3 gives "
      "up the least balance, and Tabu + K opt size 3 and K opt the most. A cheaper schedule usually packs visits onto fewer, denser routes. This is the "
      "same tradeoff the first study saw, now measured without a fairness search to soften it.")
    sub("What the fairness search would recover")
    first_var = lambda s, f, stage: mean(r["raw"][stage]["fairness"]["variance"] for r in select(E, solver=s, fleet=f))
    table(["Fleet", "Starting schedule", "Tabu, cost search", "Tabu + K opt, cost search", "Tabu, after fairness", "Tabu + K opt, after fairness"],
          [[f, f"{first_var('TABU', f, 'before'):.4f}", f"{first_var('TABU', f, 'reference'):.4f}", f"{first_var('TABU_KOPT', f, 'reference'):.4f}",
            f"{first_var('TABU', f, 'after'):.4f}", f"{first_var('TABU_KOPT', f, 'after'):.4f}"] for f in EARLIER_FLEETS], [.7, 1.1, 1.15, 1.35, 1.2, 1.4])
    spend50 = mean((r["raw"]["after"]["costCents"] - r["raw"]["reference"]["costCents"]) / 100 for r in select(E, solver="TABU", fleet=50))
    p(f"This table uses the first study's archive, where each case recorded the variance before and after its fairness search (160 cases per cell). On 50 technician "
      f"days the cost search left Tabu at {first_var('TABU', 50, 'reference'):.4f} and Tabu + K opt at {first_var('TABU_KOPT', 50, 'reference'):.4f}, similar to this "
      f"study. The fairness search then brought them to {first_var('TABU', 50, 'after'):.4f} and {first_var('TABU_KOPT', 50, 'after'):.4f}, for about "
      f"{dollars(spend50)} per day within the 2% allowance. It removed most of the imbalance and most of the gap between the two methods.")
    p("So the variance differences in this section are an upper bound. A client whose fairness search time is above zero gets a much more even schedule from any of "
      "these methods. A client who sets it to zero, which the Solver settings page allows, gets the balance shown here.")

    heading("10 What individual cases, fixtures and seeds show")
    p(f"A win counts a lower accepted cost in exact cents, regardless of size. The table covers all {n_alt} matched cases per configuration, including ties and losses. [1]")
    table(["Configuration", "Wins / ties / losses", "Worst saving $", "Median saving $", "Best saving $"],
          [[LABEL[s], wtl(by(s)), signed(min(r["saving"] for r in by(s))), signed(median(r["saving"] for r in by(s))), signed(max(r["saving"] for r in by(s)))]
           for s in ALTS], [1.9, 1.5, 1.15, 1.2, 1.15])
    figure("variation", "Left: mean saving per fixture seed and search seed, 40 cases per point. Right: cumulative distribution of all 240 case savings per configuration. "
                        "These describe observed variation, not confidence intervals.")
    combos = [(x, s) for x in (1, 2) for s in (211, 223, 227)]
    combo_mean = lambda s, x, seed: avg(by(s, fixture_seed=x, seed=seed), "saving")
    assert all(combo_mean(s, x, seed) > 0 for s in ALTS for x, seed in combos)
    assert all(combo_mean("TABU_KOPT_SIZE_3", x, seed) > combo_mean("TABU_KOPT", x, seed) for x, seed in combos)
    p("On the left, every configuration saves money on average for every fixture and seed combination, and Tabu + K opt size 3 is above Tabu + K opt in all six. "
      "The ranking does not depend on one lucky fixture or seed.")
    p("On the right, pick a dollar value and read up to a curve: its height is the share of cases saving that much or less. K opt has a long left tail of 50 "
      "technician losses. Tabu + K opt size 3's curve sits furthest right through most of its range: many wins of $10 to $25 and few losses.")

    heading("11 Checking the result one case at a time")
    full = sum(r["receipt"]["full_load_fraction"] >= .95 for r in M)
    p(f"The main and long arms ran six cases at once, each on its own physical core, as in the first study. {full:,} of {len(M):,} main cases ran with every slot busy "
      "for at least 95% of their time; the rest are the run's tail. The first study measured that six at once costs every solver about 15% of its search rate, "
      "which keeps comparisons between solvers fair. The 2026-10-10 contention check was reused rather than repeated. [4, 5]")
    p(f"After the main arm, the driver ranked the configurations by mean accepted cost saving and picked {LABEL[data['decision']['candidate']]}. It reran it beside Tabu "
      "one case at a time on fixture seed 1, seeds 211 and 223, 20 and 50 technicians, every workload and 120 seconds: 32 cases with no contention at all. "
      "The first study's confirmation of Tabu + K opt, run the same way, found $5.54 per case and 50 wins in 64. [3, 5]")
    figure("confirmation", "Each point is one of 16 matched cases: Tabu + K opt size 3 saving vs Tabu with six at once (horizontal) and one at a time (vertical). "
                           "The diagonal is perfect agreement.")
    cases = by("TABU_KOPT_SIZE_3", C)
    groups = [("All", lambda r: True)] + [(f"{f} technicians", lambda r, f=f: r["fleet"] == f) for f in (20, 50)] + [(WORK_LABEL[w], lambda r, w=w: r["workload"] == w) for w in WORKLOADS]
    table(["Group", "Pairs", "One at a time $", "Six at once $", "W/T/L one at a time"],
          [[name, sum(map(test, cases)), signed(mean(v for r, v in zip(cases, seq) if test(r))), signed(mean(v for r, v in zip(cases, par) if test(r))),
            counts([v for r, v in zip(cases, seq) if test(r)])] for name, test in groups], [1.7, .8, 1.5, 1.4, 1.5])
    p(f"Run one at a time, Tabu + K opt size 3 saves {dollars(mean(seq))} per case against {dollars(mean(par))} on the same cases six at once, and wins {sum(v > 0 for v in seq)} "
      f"of {len(seq)}. The sign matches in {agree} of {len(seq)}. Running alone, Tabu evaluated moves {1 / move_ratio['TABU']:.2f} times as fast and the size 3 "
      f"version {1 / move_ratio['TABU_KOPT_SIZE_3']:.2f} times, close to the 15% slowdown the first study measured. Both methods got the same boost and their "
      "difference barely moved. These cases also appeared in the main arm, so the check rules out a contention artifact; it is not an independent sample of new schedules.")

    heading("12 What should happen next")
    behind = [(f, w) for f in FLEETS if f >= 10 for w in WORKLOADS if save("TABU_KOPT_SIZE_3", fleet=f, workload=w) <= save("TABU_KOPT", fleet=f, workload=w)]
    assert behind == [(10, "CLUSTERED")], behind
    p("The evidence supports four conclusions. Tabu + K opt, the current default, holds its advantage over Tabu on new schedules. Tabu + K opt size 3 is better still "
      "at every fleet size from 10 technicians up and for every fixture and seed tested, and the advantage survived running alone. Of the 16 fleet and workload groups "
      "from 10 technicians up it trails Tabu + K opt only on clustered 10 technician days, by cents. A fleet size rule is not needed. Longer searches would favor K opt "
      "on large fleets, but only beyond the 240 seconds a client can allow.")
    table(["Decision", "Evidence", "Suggested next step"], [
        ("Offer Tabu + K opt size 3?", f"{dollars(mean(direct))} per case more than Tabu + K opt; ahead from 10 technicians; slightly worse at 5.",
         "Add it as a third daily search choice, then replay real customer days before making it the default."),
        ("Build a fleet size rule?", "The first study's rule held up, and one method matches it.", "No selector for now."),
        ("Longer daily searches?", "K opt passes Tabu + K opt at 50 technicians only from about 480 seconds.", "Keep the 240 second limit; revisit only for overnight runs."),
        ("Balance", "Cheaper methods are less even without a fairness search; with one, most of the gap closes.",
         "Measure Tabu + K opt size 3 with the fairness search on before adoption."),
        ("Real schedules", "Every fixture is synthetic, with a synthetic travel matrix.", "Replay recorded days with measured road travel through both Tabu + K opt variants.")], [1.55, 2.75, 2.6])
    sub("How the calculations were checked")
    p("Checks cover the archived configuration and case hashes, every raw result hash in each saved analysis, unique and complete case coverage, fixture identity and "
      "starting schedules between matched pairs, no fixture shared between the main and long arms, zero overtime and violations, the policy decision of every case, "
      "and an independent recomputation of fairness variance from paid minutes. The first study's archive is loaded through its own validated loader. "
      "verify_report.py recomputes every table and chart dataset from raw cases. [1, 2, 3, 5]")

    heading("Appendix A Main arm by fleet")
    p("Each row holds six matched cases per configuration (two fixtures, three seeds) at the stated cost search time. Work gives the workload initial. Savings are "
      "Tabu accepted cost minus the configuration's accepted cost in dollars per case, with wins, ties and losses below in exact cents; each entry totals six. [1]")
    for f in FLEETS:
        sub(f"{f} technicians")
        rows = []
        for w in WORKLOADS:
            for b in BUDGETS:
                groups = [by(s, fleet=f, workload=w, budget_ms=b * 1000) for s in SOLVERS]
                # Integer cents first, so an exact half cent rounds the same way as in verify_report.py.
                rows.append([w[0], b, f"{mean(r['raw']['after']['costCents'] for r in groups[0]) / 100:,.2f}", *[f"{avg(g, 'saving'):+.2f}\n{wtl(g)}" for g in groups[1:]]])
        table(["Work", "Sec", "Tabu $", *[SHORT[s] for s in ALTS]], rows, [.6, .6, 1.1, 1.15, 1.15, 1.15, 1.15], size=9)

    heading("Appendix B Search mechanics by fleet")
    p("Cost search means over 48 cases per row (four workloads, two fixtures, three seeds, two times). Best at is the share of the search elapsed when the best schedule "
      "appeared; late is the share of cases whose best appeared in the final 10%. [1]")
    table(["Configuration", "Fleet", "Thousand moves / s", "Steps / s", "Moves / step", "Best at", "Late", "Saving"],
          [[LABEL[s], f, f"{avg(by(s, fleet=f), 'move_rate'):.1f}", f"{avg(by(s, fleet=f), 'steps_rate'):,.1f}", f"{avg(by(s, fleet=f), 'moves_per_step'):,.0f}",
            f"{avg(by(s, fleet=f), 'best_fraction') * 100:.0f}%", f"{avg(by(s, fleet=f), 'late_best') * 100:.0f}%",
            "Control" if s == "TABU" else signed(save(s, fleet=f))] for s in SOLVERS for f in FLEETS], [1.6, .55, 1.0, .85, .9, .7, .6, .7], size=9)
    sub("Savings by workload and fleet")
    table(["Configuration", "Workload", "5 techs", "10 techs", "20 techs", "30 techs", "50 techs"],
          [[LABEL[s], WORK_LABEL[w], *[signed(save(s, fleet=f, workload=w)) for f in FLEETS]] for s in ALTS for w in WORKLOADS],
          [1.6, 1.05, .85, .85, .85, .85, .85], size=9)

    heading("Appendix C Long budget arm by workload")
    p("Each entry holds two matched cases (one fixture, two seeds). Savings are Tabu + K opt accepted cost minus the configuration's accepted cost in dollars per case, "
      "with wins, ties and losses below. The Tabu + K opt column gives the control's mean accepted cost. [2]")
    rows = []
    for f in LONG_FLEETS:
        for w in WORKLOADS:
            for b in LONG_BUDGETS:
                groups = [by(s, L, fleet=f, workload=w, budget_ms=b * 1000) for s in LONG_SOLVERS]
                rows.append([f, w[0], b, f"{mean(r['raw']['after']['costCents'] for r in groups[0]) / 100:,.2f}", *[f"{avg(g, 'saving'):+.2f}\n{wtl(g)}" for g in groups[1:]]])
    table(["Fleet", "Work", "Sec", "Tabu + K opt $", *[SHORT[s] for s in LONG_SOLVERS[1:]]], rows, [.7, .7, .7, 1.3, 1.15, 1.15, 1.2], size=9)

    heading("Appendix D Definitions and evidence")
    table(["Term", "Meaning in this report"], [
        ("Case", "One configuration run at one fleet, workload, fixture, seed and budget."),
        ("Fixture seed", "Selects the synthetic day. Seed 0 is the first study's fixture; seeds 1 to 3 are new."),
        ("Matched pair", "Same fixture and starting schedule under the same seed and budget, comparing a configuration with the arm's control."),
        ("Mean saving", "Average paired difference in accepted modeled dollars. Positive favors the configuration."),
        ("Cost only", "The cost search receives the whole budget; no fairness search runs."),
        ("Kept start", "The policy kept the starting schedule because it was within 2% of the cheaper one and more even."),
        ("Best at", "Time to the best schedule divided by search time."),
        ("TIME_LIMIT", "The search ended at its allocated time. It means neither proven optimal nor infeasible.")], [1.85, 5.05])
    p(f"[1] Main arm: experiments/runs/{MAIN.name}, analysis {MAIN_ANALYSIS.name}. Measured revision {REVISION}. Raw JSONL, manifest, configuration, cases and "
      "completion receipts are the primary evidence.")
    p(f"[2] Long budget arm: experiments/runs/{LONG.name}, analysis {LONG_ANALYSIS.name}.")
    p(f"[3] Sequential confirmation: experiments/runs/{CONFIRM.name}, analysis {CONFIRM_ANALYSIS.name}, one case at a time.")
    p("[4] Driver status .scratch/bench-2026-10-12/status.json (local only): stage order, the confirmation candidate and its ranking. Each analysis-provenance.json "
      "supplies its input hashes.")
    p("[5] First study: docs/daily-tabu-variants-5d15129b (report and EVIDENCE.md) and experiments/runs/20261010T040557Z-daily-tabu-variants-p6-5d15129b. Its values are "
      "compared here but never pooled with this study's.")
    p("[6] Measured source: SolverBenchmark.java and SolverBenchmarkData.java under scheduler-service/src/test/java/dev/waterflex/scheduler/optimizer, SolverEngine.java "
      "and SchedulingPolicy.java under calculation-engine/src/main/java/dev/waterflex/scheduler/optimizer. These define fixtures, configurations, phases and the policy check.")
    p("[7] Timefold Solver 2.6.0 local search: acceptors, foragers and move selectors, as configured in the recorded solver XML of every case.")
    p("[8] Pull requests #133 (per client daily search choice and search times) and #134 (fixture seeds, cost only mode and Tabu + K opt size 3).")
    p("Reproduce with docs/daily-search-followup-6454fb05/requirements.txt, then run report_data.py, build_report.py and verify_report.py there. The builder reads the "
      f"archives, regenerates all {len(figures)} charts and chart_data.json, and writes this editable Word file. EVIDENCE.md records methods.")

    settings = doc.settings.element
    for el in settings.findall(qn("w:updateFields")): settings.remove(el)
    update = OxmlElement("w:updateFields"); update.set(qn("w:val"), "true"); settings.append(update)
    stream = BytesIO(); doc.save(stream)
    with ZipFile(REFERENCE) as ref, ZipFile(stream) as built, ZipFile(OUTPUT, "w", ZIP_DEFLATED) as output:
        preserve = set(PRESERVE) | {n for n in ref.namelist() if "footer" in n}
        for name in sorted(built.namelist()):
            content = ref.read(name) if name in preserve else built.read(name)
            info = ZipInfo(name, (2026, 10, 13, 0, 0, 0)); info.compress_type = ZIP_DEFLATED
            output.writestr(info, content)
    print(OUTPUT)


if __name__ == "__main__": build()
