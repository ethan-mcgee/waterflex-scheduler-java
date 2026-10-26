"""Build the road matrix numbering study report from the validated evidence and the reference explainer's template.

Install requirements.txt first. No network, services or benchmark execution. The two diagrams in diagrams/ are
rendered to generated/ separately, and stacks.json is written once from the raw recordings (see EVIDENCE.md); the
charts are regenerated here.
"""
import json
from copy import deepcopy
from datetime import datetime
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from statistics import median
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from charts import LABEL, SHORT, case_label, make_charts, short_frame
from report_data import ARMS, FIX, FLEETS, MEASURES, REPS, beyond_noise, case_mean, cases, index, load, ratios, reference_mean

HERE = Path(__file__).resolve().parent
REFERENCE = HERE.parent / "WaterFlex_Scheduler_Algorithm_Comparison_Explained.docx"
REFERENCE_SHA = "5366f3017d6f18c15de295ba5fb9c1a014b6222911040f1c5e044845e57f832d"
OUTPUT = HERE / "Road_Matrix_Numbering_2026-10-20_Report.docx"
EARLIER = HERE.parent / "evidence" / "solver-variance-2026-10-20.json"
PRESERVE = ("word/styles.xml", "word/stylesWithEffects.xml", "word/numbering.xml", "word/theme/theme1.xml", "word/fontTable.xml")
DIAGRAMS = ("numbering-lookup", "numbering-arms")
MEASURE_NAME = {"variance": "Variance", "maximumUtilization": "Highest utilization", "utilizationSpread": "Utilization spread",
                "absoluteDeviation": "Absolute deviation", "loadBalance": "Load balance"}
GROUPS = ("decimal: variance", "variance (non-decimal)", "decimal: score impacts", "decimal: money", "collections", "engine", "scheduler code",
          "score engine", "time", "other")
SITES = ("identity probe (RoadMatrix.number)", "segment options map (RouteTimingSearch.segments)", "timing search's own work (RouteTimingSearch.segments)",
         "BigDecimal score impacts and other decimal work", "BigDecimal money (Monetary)", "Timefold's own maps")
RATE = {"referencePhase": "referenceMovesPerSecond", "fairnessPhase": "fairnessMovesPerSecond"}
DASHES = (chr(0x2014), chr(0x2013))  # Em and en dash, by code point so this file holds neither.


def speed(entry):
    return f"{entry['median']:.2f}x ({entry['min']:.2f} to {entry['max']:.2f})"


def cents(value):
    """Signed cents, to the quarter cent a median of two-repetition means can reach."""
    return f"{value:+,.2f}".rstrip("0").rstrip(".")


def plain(value):
    """An unsigned amount of cents for prose that says 'less' or 'cheaper' itself."""
    assert value >= 0
    return f"{value:,.2f}".rstrip("0").rstrip(".")


def milli(value):
    """A variance as thousandths, the unit every variance in the report is shown in."""
    text = f"{Decimal(value) * 1000:.2f}"
    return "0.00" if text == "-0.00" else text


def build():
    assert sha256(REFERENCE.read_bytes()).hexdigest() == REFERENCE_SHA
    document = load()
    assets = HERE / "generated"
    for name in DIAGRAMS:
        assert (assets / f"{name}.png").exists(), f"render diagrams/{name}.html first (EVIDENCE.md)"
    stacks = json.loads((HERE / "stacks.json").read_text(encoding="utf-8"))
    earlier = json.loads(EARLIER.read_text(encoding="utf-8-sig"))["budgets"]["15000"]
    make_charts(document, assets)
    doc = Document(REFERENCE)
    table_properties = deepcopy(doc.tables[0]._tbl.tblPr)
    for child in list(doc.element.body):
        if child.tag != qn("w:sectPr"): doc.element.body.remove(child)
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype.endswith(("/image", "/hyperlink")): del doc.part.rels[rid]
    doc.core_properties.title = "WaterFlex scheduler road matrix numbering explained"
    doc.core_properties.subject = "Numbering visits once in the scoring path, then options b, c and load balance measured again on the faster engine, 2026-10-20"
    doc.core_properties.author = doc.core_properties.last_modified_by = "WaterFlex"
    doc.core_properties.created = doc.core_properties.modified = datetime(2026, 10, 20)
    figures = []

    def p(text, style=None):
        assert not any(dash in text for dash in DASHES)
        para = doc.add_paragraph(text, style); para.paragraph_format.widow_control = True
        return para

    def heading(text):
        h = p(text, "Heading 1"); h.paragraph_format.page_break_before = True

    def sub(text): p(text, "Heading 2")

    def table(headers, data, widths, size=10):
        assert abs(sum(widths) - 6.9) < .01
        t = doc.add_table(rows=1, cols=len(headers)); t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t._tbl.replace(t._tbl.tblPr, deepcopy(table_properties))
        for col, w in zip(t.columns, widths): col.width = Inches(w)
        for i, values in enumerate([headers, *data]):
            assert len(values) == len(headers)
            row = t.rows[0] if i == 0 else t.add_row()
            trpr = row._tr.get_or_add_trPr(); trpr.append(OxmlElement("w:cantSplit"))
            if i == 0: trpr.append(OxmlElement("w:tblHeader"))
            for j, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
                text = str(value)
                assert not any(dash in text for dash in DASHES)
                cell.width = Inches(width); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell.text = text
                props = cell._tc.get_or_add_tcPr()
                shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "DCE6EF" if i == 0 else ("FFFFFF" if i % 2 else "F4F6F8")); props.append(shade)
                margins = OxmlElement("w:tcMar")
                for side in ("top", "left", "bottom", "right"):
                    n = OxmlElement("w:" + side); n.set(qn("w:w"), "65"); n.set(qn("w:type"), "dxa"); margins.append(n)
                props.append(margins)
                para = cell.paragraphs[0]; para.paragraph_format.space_after = Pt(2); para.paragraph_format.space_before = Pt(2); para.paragraph_format.line_spacing = 1.0
                if len(headers) > 2 and j > 0: para.alignment = 1
                if len(data) <= 12 and i < len(data): para.paragraph_format.keep_with_next = True  # Short tables stay on one page.
                for run in para.runs: run.font.name = "Calibri"; run.font.size = Pt(size); run.bold = i == 0
        spacer = p(""); spacer.paragraph_format.space_after = Pt(0); spacer.paragraph_format.line_spacing = Pt(3)

    def figure(name, caption, width=6.85):
        figures.append(name)
        caption = f"Figure {len(figures)}. {caption}"
        para = p(""); shape = para.add_run().add_picture(str(assets / f"{name}.png"), width=Inches(width))
        shape._inline.docPr.set("descr", caption); para.paragraph_format.keep_with_next = True; para.alignment = 1
        cap = p(caption, "Caption"); cap.paragraph_format.space_after = Pt(7)
        for run in cap.runs:
            run.font.color.rgb = RGBColor(0, 0, 0); run.bold = False; run.italic = True

    # The numbers the narrative refers to by name.
    paired = document["paired"]
    by = index(paired["runs"])
    all_cases = cases(by)
    fix, summary, comparison, profiles = paired["fix"], paired["summary"], paired["comparison"], document["profiles"]
    fp, fb = fix["baseline"], fix["b"]
    sb, sc = summary["b"], summary["c"]
    lb = comparison["load-balance"]
    manifest = document["manifest"]
    total_runs = len(paired["runs"])

    def rate(arm, phase, fleet=None):
        return median(by[(arm, c, r)][RATE[phase]] for c in all_cases if fleet is None or c[0] == fleet for r in REPS)

    def fleet_ratio(arm, against, fleet, field="variance"):
        values = [float(case_mean(by, arm, c, field) / case_mean(by, against, c, field)) for c in all_cases if c[0] == fleet and case_mean(by, against, c, field) > 0]
        return median(values)

    def ratio(arm, name):
        return comparison[arm]["measures"][name]["medianRatioToBaseline"]

    def cost_diff(arm, against):
        return median(float(case_mean(by, arm, c, "costCents") - case_mean(by, against, c, "costCents")) for c in all_cases)

    def ref_cost_diff(arm, against, fleet=None):
        return median(float(case_mean(by, arm, c, "referenceCostCents") - case_mean(by, against, c, "referenceCostCents")) for c in all_cases if fleet is None or c[0] == fleet)

    def variance_change(arm, against):
        return median(float(case_mean(by, arm, c, "variance") / case_mean(by, against, c, "variance")) for c in all_cases if case_mean(by, against, c, "variance") > 0)

    def ref_variance_change(arm, against):
        return median(float(reference_mean(by, arm, c, "variance") / reference_mean(by, against, c, "variance")) for c in all_cases if reference_mean(by, against, c, "variance") > 0)

    noise = {(arm, old, field): beyond_noise(by, arm, old, field) for arm, old in (("baseline", "old"), ("b", "old-b"), ("load-balance", "b")) for field in ("variance", "costCents")}
    lb_lower = sum(1 for c in all_cases if case_mean(by, "load-balance", c, "variance") < case_mean(by, "b", c, "variance"))
    lb_higher = sum(1 for c in all_cases if case_mean(by, "load-balance", c, "variance") > case_mean(by, "b", c, "variance"))
    lb_same = len(all_cases) - lb_lower - lb_higher
    c_fail = [x for x in paired["pairs"] if x["arm"] == "c" and Decimal(x["meanVarianceDifference"]) > Decimal(x["baselineVarianceSpread"])]
    assert len(c_fail) == 1
    c_fail = c_fail[0]
    site = lambda name, s: stacks["profiles"][name]["sitesPercent"].get(s, 0.0)
    b_same_10 = sum(1 for c in all_cases if c[0] == 10 and case_mean(by, "b", c, "variance") == case_mean(by, "baseline", c, "variance"))
    earlier_b = earlier["summary"]["b"]["fairnessPhaseSpeedup"]["median"]
    earlier_lb, earlier_bv = earlier["comparison"]["load-balance"]["measures"]["variance"]["medianRatioToBaseline"], earlier["comparison"]["b"]["measures"]["variance"]["medianRatioToBaseline"]

    p("WaterFlex scheduler road matrix numbering explained", "Title")
    p("Numbering visits once in the scoring path, then options b, c and load balance measured again on the faster engine", "Subtitle")
    p(f"October 20 2026  |  Road matrix numbering study of pull request #187, started {manifest['startedAt'][11:16]} UTC on 2026-10-20")
    p(f"The fairness measure study this morning found that, once the fairness variance is cheap, both phases of the daily search spend most of their time in the same place: about a fifth of all samples looking up each visit's number in the road matrix by its identity text, and about a tenth building a small map on every route evaluation. Commit ed97d8b removes both without changing any scoring rule. This study measured that fix against the code before it, and measured options b, c and Timefold's load balance again on top of it, in {total_runs} solves over 36 fleets of 10, 30 and 50 technicians at the 15 s production budget. The fix makes production's cost phase {fp['referencePhaseSpeedup']['median']:.2f} times faster and option b's fairness phase {fb['fairnessPhaseSpeedup']['median']:.2f} times faster, and the extra moves end on cheaper plans: a median {plain(-cost_diff('baseline', 'old'))} cents less per day for production. Option b again passes every bar, now at {sb['fairnessPhaseSpeedup']['median']:.1f} times the fixed production's fairness-phase rate. Option c fails the fairness bar in one case. Load balance, on these fleet sizes, ends fairer than b on every measure, which this morning's study did not show and which needs its own test before it can change any decision.")
    sub("What the measurements tell us")
    table(["Arm", "What it changes", "Cost-phase speed", "Fairness-phase speed", "Variance", "Cost", "Verdict"], [
        ["The fix, production", "Numbers visits and routes once", f"{fp['referencePhaseSpeedup']['median']:.2f}x", f"{fp['fairnessPhaseSpeedup']['median']:.2f}x",
         f"{variance_change('baseline', 'old'):.2f}", f"{cents(cost_diff('baseline', 'old'))} c", "Not slower in either phase; promote"],
        ["The fix, option b", "The same, under b's formula", f"{fb['referencePhaseSpeedup']['median']:.2f}x", f"{fb['fairnessPhaseSpeedup']['median']:.2f}x",
         f"{variance_change('b', 'old-b'):.2f}", f"{cents(cost_diff('b', 'old-b'))} c", "The gain carries over to b"],
        ["b", "Variance in binary64", f"{sb['referencePhaseSpeedup']['median']:.2f}x", f"{sb['fairnessPhaseSpeedup']['median']:.2f}x", f"{ratio('b', 'variance'):.2f}",
         f"{cents(sb['cost']['medianDifferenceCents'])} c", "Passes every bar; promote"],
        ["c", "Variance from running sums", f"{sc['referencePhaseSpeedup']['median']:.2f}x", f"{sc['fairnessPhaseSpeedup']['median']:.2f}x", f"{ratio('c', 'variance'):.2f}",
         f"{cents(sc['cost']['medianDifferenceCents'])} c", "Fails the fairness bar in one case"],
        ["load-balance", "Timefold load balance", f"{median(ratios(by, 'load-balance', 'baseline', 'referenceMovesPerSecond')):.2f}x", f"{lb['fairnessPhaseSpeedup']['median']:.2f}x",
         f"{ratio('load-balance', 'variance'):.2f}", f"{cents(lb['cost']['medianDifferenceCents'])} c", "Fairer than b here; test head to head"]],
        [1.05, 1.35, .7, .75, .65, .65, 1.75], size=9)
    p("The two fix rows compare the fixed tree with the same variant built from the commit before the fix, case by case and repetition by repetition; their variance and cost columns are against that old variant. The other rows compare with the fixed production. Each speed is the median over 72 paired solves of moves per second divided by the comparison's; each variance is the median over the cases of the final variance divided by the comparison's (lower is fairer); each cost is the median over the 36 cases of the final cost minus the comparison's, in cents. Every solve ended with no constraint violation and inside its cost ceiling.")
    p("All of it comes from one evidence file, docs/evidence/solver-matrix-2026-10-20.json, and, for the table of callers in section 2, from the run's own flight recordings. None of it is customer data; the fleets are synthetic.")
    sub("How to read this report")
    table(["Sections", "What you will learn"], [
        ("1", "What was measured: the two trees, the six arms, the 36 fleets and the bars."),
        ("2", "What the profile found, and how the fix removes it."),
        ("3 and 4", "How much faster each phase became, by fleet size, and what the extra moves bought."),
        ("5", "Options b and c against the bars on the faster engine."),
        ("6", "Load balance measured again, and why it now leads b."),
        ("7", "Where the time goes now."),
        ("8 and 9", "Conclusions, and what should be tested or implemented next."),
        ("Appendices", "Every case for the fix, for b and c and for load balance, every profile, definitions and sources.")], [1.45, 5.45])

    heading("1 What was measured")
    p("The daily search runs in two phases. The cost phase looks for the cheapest plan that breaks no rule; its result is the reference plan. The fairness phase starts from the reference and looks for a plan that spreads the work more evenly, at no more than the reference cost plus the client's fairness allowance (two percent here). Every candidate move in either phase is scored by re-walking the routes it changed: for each leg, the scorer finds the leg's travel time in the road matrix, a table indexed by a number for each technician and visit. [1, 2]")
    p("Two code trees were built from their commits and run side by side. Tree o is the commit before the fix, the code this morning's study ran. Tree n is commit ed97d8b, the fix. Production and option b ran on both trees, so the fix gets its own paired number for each; options b, c and load balance ran on tree n and were compared with tree n's production, as this morning's study compared them with tree o's.")
    figure("numbering-arms", "The six arms and what each is compared with. Every comparison pairs the same case and repetition, and the order of the trees and of the arms within a tree rotated from case to case.", width=6.2)
    table(["Arm", "Tree, variant", "What it is", "Judged by"], [
        ("old", "o, TABU_KOPT", "Production before the fix", "The reference for the fix"),
        ("old-b", "o, TABU_KOPT_BINARY_VARIANCE", "Option b before the fix", "The reference for b's own gain from the fix"),
        ("baseline", "n, TABU_KOPT", "Production with the fix", "The fix: at least as fast as old in both phases"),
        ("b", "n, TABU_KOPT_BINARY_VARIANCE", "Production's variance in binary64, carried as rounded units of 1e-9", "Four bars, against this baseline"),
        ("c", "n, TABU_KOPT_INCREMENTAL_VARIANCE", "Production's variance from three exact running sums", "Four bars, against this baseline"),
        ("load-balance", "n, TABU_KOPT_LOAD_BALANCE", "Timefold's loadBalance over utilization, unweighted", "Every measure, no bars")], [.95, 2.15, 2.3, 1.5], size=9)
    table(["Dimension", "Values in this study"], [
        ("Fleets", f"{', '.join(manifest['sizes'])} technicians; workloads {', '.join(w.title().replace('_', ' ') for w in manifest['workloads'])}; seeds {', '.join(manifest['seeds'])}: 36 cases on fixture seed {manifest['fixtureSeed']}"),
        ("Budget", "15 s, the production default: 10 s cost phase, 5 s fairness phase"),
        ("Repetitions", "Two of every case; one JVM per case and tree, the tree order alternating and the arm order rotating"),
        ("Search", "TABU_KOPT at NO_ASSERT, the production move selectors; fairness allowance 2 percent"),
        ("Machine", f"One {manifest['processors']}-processor workstation; each solve limited to 2 processors, 1.5 GiB heap and the serial collector ({' '.join(manifest['jvm'])})"),
        ("Trees", f"o: {manifest['revisions']['o'][:7]} (before the fix); n: {manifest['revisions']['n'][:7]} (the fix)"),
        ("Profiles", "A flight recording of every arm's two phases on seed 17, one per fleet and workload: twelve recordings per arm and phase")], [1.3, 5.6])
    sub("The bars")
    p("The fix changes no scoring rule; ScoringExactnessTest shows that a step-limited solve with the fix ends on exactly the same score, moves, arrangement and independent evaluation as before. So it is judged as the timing change of the representation study was: promoted if it is not slower in either phase. Options b and c are held to the four bars the owner set for this morning's study, against the fixed production.")
    table(["Bar", "Test"], [
        ("Fix", "The median paired moves per second of the fixed production is at least the old production's, in both phases, with zero violations."),
        ("Speed (b, c)", "The median of the 72 paired fairness-phase speedups is at least 1.25x."),
        ("Fairness (b, c)", "No case is less fair than the baseline by more than that case's own spread between the baseline's two repetitions."),
        ("Cost (b, c)", "The median per-case cost difference is not above zero."),
        ("Ceiling (b, c)", "Every accepted result is inside its cost ceiling, with zero violations.")], [1.3, 5.6])

    heading("2 What the profile found")
    p("A flight recording samples where the CPU is many times a second and keeps the stack of method calls at each sample. The evidence of this morning's study kept only the method at the top of each stack, which put an immutable map's lookup (ImmutableCollections$MapN.probe) at about a fifth of all samples but could not say who called it. Reading the recordings again with their full stacks answered that, and this run's recordings of tree o show the same thing. [3]")
    table(["Site", "Old production, cost phase", "Fixed production, cost phase", "b before the fix, fairness phase", "b after the fix, fairness phase"],
          [[name, *[f"{site(profile, name):.1f}%" for profile in ("old cost", "baseline cost", "old-b fairness", "b fairness")]] for name in SITES],
          [2.3, 1.15, 1.15, 1.15, 1.15], size=9)
    p(f"Two sites held close to a third of each phase profiled. The identity probe ({site('old cost', SITES[0]):.1f} percent of production's cost phase) is the road matrix turning a visit's identity text into its number, once for every leg of every route the move touched. The map's keys were cut out of the matrix's 'from>to' keys, so they are never the same text object as the visit's own identity, and every lookup compared the whole identity character by character. The segment options map ({site('old cost', SITES[1]):.1f} percent) is a LinkedHashMap from end position to timing that the route timing search built and threw away on every call, even when only one entry was ever read.")
    figure("numbering-lookup", "How a leg finds its cell in the road matrix, before and after the fix. The kept number carries the fact revision it came from, so a visit reused under another revision looks itself up again instead of reading a stale cell.")
    p("The fix keeps each visit's and route's matrix number on the visit and route themselves, together with the fact revision that gave it; a lookup is then one reference comparison and a field read. A revision that reuses the same visits (repairing a day after an absence, or booking's per-route plans) finds a different revision on the kept number and looks it up again; a regression test shows this, and fails if the revision check is removed. The options map became an array indexed by end position, walked in the same ascending order the map was, and a frozen route keeps its regular capacity and eligibility for its revision. Both former sites are gone from the recordings of tree n. The timing search's own work now holds the share that remains: it is the same loop with the lookups inlined into it, a larger share of a faster phase.")

    heading("3 What the fix did to speed")
    table(["Comparison, phase", "All fleets", "10 technicians", "30 technicians", "50 technicians"],
          [[f"{'Production' if arm == 'baseline' else 'Option b'}, {'cost' if phase == 'referencePhase' else 'fairness'} phase", speed(fix[arm][phase + "Speedup"]),
            *[f"{fix[arm]['byFleet'][phase][str(f)]['median']:.2f}x" for f in FLEETS]]
           for arm in FIX for phase in RATE], [1.9, 1.6, 1.1, 1.15, 1.15])
    figure("fix_speed", "Moves per second on the fixed tree over the same variant before the fix, every pair of every case, by fleet and phase. The black bar is the median; the medians are under each fleet, production first.")
    p(f"Production's cost phase evaluates {fp['referencePhaseSpeedup']['median']:.2f} times as many moves per second, every fleet size gains, and the gain is largest on small fleets: {fp['byFleet']['referencePhase']['10']['median']:.2f}x at 10 technicians against {fp['byFleet']['referencePhase']['50']['median']:.2f}x at 50. The work the fix removed is per leg re-walked, which depends on route length, not on the fleet. What grows with the fleet is the fleet-wide part of every move, Timefold's own bookkeeping and the fleet-level cost and fairness sums, so the saved work is a smaller share of a 50-technician move. That reading is an interpretation of the profile, not a separate measurement.")
    p(f"Production's fairness phase gains much less, {fp['fairnessPhaseSpeedup']['median']:.2f}x, because it is still dominated by the DECIMAL128 variance: {profiles['baseline fairness']['samplesPercent']['decimal: variance']:.0f} percent of its samples are decimal variance arithmetic, which the fix does not touch. Under option b, where the variance is cheap, the fairness phase gains {fb['fairnessPhaseSpeedup']['median']:.2f}x, close to the cost phase's gain. The fix and b add up: each removes a different cost.")
    figure("rates", "Median moves per second by arm, phase and fleet size. Production's fairness phase (dark grey, right) stays low on both trees because of its decimal variance; every arm with a cheap variance runs its fairness phase near its cost phase's rate.")
    table(["Arm", "Cost phase, 10 / 30 / 50", "Fairness phase, 10 / 30 / 50"],
          [[LABEL[arm], " / ".join(f"{rate(arm, 'referencePhase', f) / 1000:,.0f}k" for f in FLEETS), " / ".join(f"{rate(arm, 'fairnessPhase', f) / 1000:,.0f}k" for f in FLEETS)] for arm in ARMS],
          [2.3, 2.3, 2.3])
    p("The absolute rates are the harness's, limited to two processors on a workstation; the paired ratios are what transfers to the service. On the fixed tree, option b's fairness phase evaluates about as many moves per second as the cost phase, which was this morning's finding for the old tree as well.")

    heading("4 What the extra moves bought")
    p("A faster phase does not finish sooner: it evaluates more moves in the same seconds, and the search can end on a different plan. Because the fix changes no rule, any difference in the final plans is the extra search, and this section reads what it found.")
    table(["Comparison", "Reference cost, median", "Final cost, median", "Cheaper / dearer beyond noise", "Final variance, median ratio", "Less / more fair beyond noise"],
          [[f"{'Production' if arm == 'baseline' else 'Option b'}, fixed against old", f"{cents(ref_cost_diff(arm, old))} c", f"{cents(cost_diff(arm, old))} c",
            f"{noise[(arm, old, 'costCents')][1]} / {noise[(arm, old, 'costCents')][0]}", f"{variance_change(arm, old):.3f}",
            f"{noise[(arm, old, 'variance')][0]} / {noise[(arm, old, 'variance')][1]}"] for arm, old in FIX.items()],
          [1.75, 1.0, .95, 1.1, 1.0, 1.1], size=9)
    figure("fix_outcome", "Left: each case's final cost on the fixed tree minus the old one, by fleet (circles 10, squares 30, triangles 50 technicians). Right: the same cost difference against the final variance ratio; points left of the vertical line are cheaper, points above the horizontal line less fair.")
    p(f"The cost phase found cheaper reference plans: a median {plain(-ref_cost_diff('baseline', 'old'))} cents less per day for production and {plain(-ref_cost_diff('baseline', 'old', 30))} less at 30 technicians, and cheaper beyond the old production's own spread between repetitions in {noise[('baseline', 'old', 'costCents')][1]} of 36 cases with none dearer. The fairness phase then starts from that cheaper plan, under a ceiling two percent above it, so its ceiling is lower too. Production's final variance is close to the old one at the median ({variance_change('baseline', 'old'):.3f}) but more often a little higher than lower: {noise[('baseline', 'old', 'variance')][0]} cases less fair beyond the old spread and {noise[('baseline', 'old', 'variance')][1]} fairer. The reference plans were as fair as before (median ratio {ref_variance_change('baseline', 'old'):.3f}), so this is the fairness phase working inside a tighter ceiling, not the cost phase making plans less even. Option b shows the same pattern: {plain(-cost_diff('b', 'old-b'))} cents cheaper at the median, {noise[('b', 'old-b', 'costCents')][1]} cases cheaper beyond noise, and {noise[('b', 'old-b', 'variance')][0]} less fair and {noise[('b', 'old-b', 'variance')][1]} fairer beyond noise.")
    p("This is the policy working as written: cost comes first, and fairness may spend up to two percent above whatever the cost phase found. A cheaper cost phase moves the whole day down in cost; the allowance moves with it. Whether to spend some of the saving on fairness is a policy question, the size of the allowance, and not something the fix should decide.")
    figure("acceptance", "How each arm's 72 results were accepted against the host's plan: as a fairness improvement (solid) or as a cost reduction at no loss of fairness (hatched).")
    p(f"The acceptance reasons agree. The fixed production is accepted as a cost reduction more often than the old one ({comparison['baseline']['acceptance'].get('COST_REDUCTION', 0)} of 72 runs against {comparison['old']['acceptance'].get('COST_REDUCTION', 0)}), because its plans are cheaper and its fairness phase, still slow, cannot always beat the host's plan on variance inside the lower ceiling. Every arm with a cheap variance improves fairness in about 65 of 72 runs.")

    heading("5 Options b and c on the faster engine")
    table(["Option", "Fairness-phase speed", "Cost-phase speed", "Fairer / same / less fair / beyond noise", "Median cost difference (cheaper / ties / dearer)", "Inside ceiling", "Passes"],
          [[arm, speed(s["fairnessPhaseSpeedup"]), f"{s['referencePhaseSpeedup']['median']:.3f}x",
            f"{s['fairness']['fairer']} / {s['fairness']['same']} / {s['fairness']['lessFair']} / {s['fairness']['lessFairBeyondNoise']}",
            f"{cents(s['cost']['medianDifferenceCents'])} c ({s['cost']['cheaper']} / {s['cost']['ties']} / {s['cost']['dearer']})",
            "Yes" if s["bars"]["ceiling"] else "No", "Yes" if s["passes"] else "No"] for arm, s in (("b", sb), ("c", sc))],
          [.7, 1.4, .85, 1.3, 1.45, .6, .6], size=9)
    figure("speed", "Fairness-phase moves per second of b, c and load balance over the fixed production in the same case and repetition, by fleet. The dashed line is the 1.25x bar.")
    p(f"Option b's fairness phase now runs {sb['fairnessPhaseSpeedup']['median']:.2f} times the fixed production's, against {earlier_b:.2f} times the old production's this morning; both the fix and b speed up their own share of the work, and the ratio between them grows. The gain still grows with the fleet, from {sb['fairnessPhaseSpeedupByFleet']['10']['median']:.2f}x at 10 technicians to {sb['fairnessPhaseSpeedupByFleet']['50']['median']:.2f}x at 50, because production's variance is recomputed over the whole fleet on every move. Option c is as fast in the fairness phase ({sc['fairnessPhaseSpeedup']['median']:.2f}x) and, as this morning, slower in the cost phase ({sc['referencePhaseSpeedup']['median']:.3f}x), where it maintains running sums that only the fairness phase reads.")
    figure("pairs", "Final variance of each case relative to the fixed production's (the mean of two repetitions on each side). The grey bar reaches production's value plus its own spread between repetitions; a marker above its bar fails the fairness bar.")
    p(f"Option b was fairer than the fixed production in {sb['fairness']['fairer']} of 36 cases and the same in {sb['fairness']['same']}, and never less fair; its median variance is {sb['fairness']['medianVarianceRatio']:.2f} of production's. The cases where it ties are mostly the small fleets: at 10 technicians b and production end on the same variance in {b_same_10} of 12 cases, because a 10-technician day has few even arrangements and both searches find them. At 30 and 50 technicians b's median variance is {fleet_ratio('b', 'baseline', 30):.2f} and {fleet_ratio('b', 'baseline', 50):.2f} of production's.")
    figure("cost", "Final cost of each case minus the fixed production's, in dollars. The black bar is the median, which the cost bar holds at or below zero.")
    p(f"Both options meet the cost bar at the median ({cents(sb['cost']['medianDifferenceCents'])} cents for b, {cents(sc['cost']['medianDifferenceCents'])} for c). Their single cases scatter a few dollars either way: each spends almost all of its allowance, so the final plans sit close to the same ceiling and the cents are which plan the search ended on.")
    sub("The case where option c failed")
    figure("c_case", f"Final variance in each repetition of {c_fail['technicians']} {c_fail['workload'].title()} {c_fail['seed']}, the one case where option c ended less fair than the fixed production beyond its spread. Pale bars are repetition 1, solid repetition 2.", width=5.6)
    table(["Arm", "Repetition 1", "Repetition 2", "Mean"],
          [[SHORT[arm], *[milli(by[(arm, (c_fail['technicians'], c_fail['workload'], c_fail['seed']), r)]["measures"]["variance"]) for r in REPS],
            milli(case_mean(by, arm, (c_fail['technicians'], c_fail['workload'], c_fail['seed']), "variance"))] for arm in ("old", "baseline", "b", "c", "load-balance")],
          [2.4, 1.5, 1.5, 1.5])
    p(f"The fixed production's two repetitions ended at {milli(c_fail['baselineVariance'][0])} and {milli(c_fail['baselineVariance'][1])} thousandths, a spread of {milli(c_fail['baselineVarianceSpread'])}. Option c's first repetition was fairer than both; its second ended at {milli(c_fail['armVariance'][1])}, which takes c's mean {milli(c_fail['meanVarianceDifference'])} above production's, beyond the spread. Option b ended at {milli(by[('b', (c_fail['technicians'], c_fail['workload'], c_fail['seed']), 1)]['measures']['variance'])} in both repetitions. As this morning, one unlucky repetition fails c, and c has no advantage over b to set against it.")

    heading("6 Load balance measured again")
    p(f"This morning, on fleets of 20 and 50 technicians, load balance was within 0.01 of b or behind it on every measure: at 15 s its median variance was {earlier_lb:.2f} of production's against b's {earlier_bv:.2f}. On the fleets of this study it is ahead of b on all five measures, its own and production's alike.")
    table(["Arm", *[MEASURE_NAME[m] for m in MEASURES], "Median cost difference", "Fairness-phase speed"],
          [[LABEL[arm], *[f"{ratio(arm, m):.2f} ({comparison[arm]['measures'][m]['lowerThanBaseline']}/{comparison[arm]['measures'][m]['higherThanBaseline']})" for m in MEASURES],
            f"{cents(comparison[arm]['cost']['medianDifferenceCents'])} c", f"{comparison[arm]['fairnessPhaseSpeedup']['median']:.2f}x"] for arm in ("old", "old-b", "b", "c", "load-balance")],
          [1.45, .75, .75, .75, .75, .75, .9, .8], size=9)
    p("Each measure cell is the median over the 36 cases of the arm's value divided by the fixed production's, then in brackets the cases lower and higher; lower is fairer on every measure. The old and old-b rows read the trees before the fix on the same scale; they are fairer than the fixed production on the variance because their ceilings sat higher, as section 4 explains.")
    figure("heatmap", "The table above as one picture: each arm's median ratio to the fixed production on each measure. Teal is fairer, red less fair; bold is the measure the arm searched on.", width=5.8)
    figure("lb_vs_b", "Each case's final variance under load balance against option b's, by fleet. Points below the dashed line are cases where load balance ended fairer than b.", width=4.6)
    table(["Fleet", "Load balance lower / higher than b", "Beyond b's own spread, lower / higher", "b's variance over production", "Load balance's variance over production"],
          [[f"{f} technicians", f"{sum(1 for c in all_cases if c[0] == f and case_mean(by, 'load-balance', c, 'variance') < case_mean(by, 'b', c, 'variance'))} / {sum(1 for c in all_cases if c[0] == f and case_mean(by, 'load-balance', c, 'variance') > case_mean(by, 'b', c, 'variance'))}",
            " / ".join(str(v) for v in reversed(beyond_noise({k: v for k, v in by.items() if k[1][0] == f}, "load-balance", "b", "variance"))),
            f"{fleet_ratio('b', 'baseline', f):.2f}", f"{fleet_ratio('load-balance', 'baseline', f):.2f}"] for f in FLEETS],
          [1.2, 1.5, 1.6, 1.2, 1.4])
    p(f"Load balance ended with a lower variance than b in {lb_lower} cases, higher in {lb_higher} and the same in {lb_same}; beyond b's own spread between repetitions, {noise[('load-balance', 'b', 'variance')][1]} lower and {noise[('load-balance', 'b', 'variance')][0]} higher. Its lead is largest where this morning's study had no fleets: at 10 technicians, where b mostly ties production and load balance does not ({fleet_ratio('load-balance', 'baseline', 10):.2f} against b's {fleet_ratio('b', 'baseline', 10):.2f}), and at 30 ({fleet_ratio('load-balance', 'baseline', 30):.2f} against {fleet_ratio('b', 'baseline', 30):.2f}). At 50 technicians, the one size both studies share, the two are close ({fleet_ratio('load-balance', 'baseline', 50):.2f} against {fleet_ratio('b', 'baseline', 50):.2f}), as they were this morning.")
    p(f"It also costs a little more. Load balance's median cost is {cents(median(float(case_mean(by, 'load-balance', c, 'costCents') - case_mean(by, 'b', c, 'costCents')) for c in all_cases))} cents above b's, and dearer than b beyond b's spread in {noise[('load-balance', 'b', 'costCents')][0]} cases against {noise[('load-balance', 'b', 'costCents')][1]} cheaper; its cost phase runs at {median(ratios(by, 'load-balance', 'baseline', 'referenceMovesPerSecond')):.2f}x production's, because Timefold maintains its load balance collector on every move of both phases. A plausible reading, not tested here, is that an unweighted measure that counts every technician equally finds even days on small fleets that a capacity-weighted variance scores as ties. Either way, two studies now disagree on fleet sizes they did not share, and a policy change would need a test designed for it: load balance and b head to head on 10, 20, 30 and 50 technicians, at 15 s and 60 s, with the same bars b was held to.")

    heading("7 Where the time goes now")
    figure("profile", "Share of CPU samples in each phase by kind of work, before and after the fix, from twelve flight recordings per arm and phase. The collections share falls from about half to about a sixth; production's fairness phase is still its decimal variance (red).")
    table(["Kind of work", "Old, cost", "Fixed, cost", "Old, fairness", "Fixed, fairness", "old-b, fairness", "b, fairness"],
          [[group, *[f"{profiles[name]['samplesPercent'].get(group, 0):.1f}%" for name in ("old cost", "baseline cost", "old fairness", "baseline fairness", "old-b fairness", "b fairness")]]
           for group in GROUPS if any(profiles[name]["samplesPercent"].get(group, 0) > 0 for name in ("old cost", "baseline cost", "old fairness", "baseline fairness", "old-b fairness", "b fairness"))],
          [1.6, .85, .85, .9, .9, .9, .9], size=9)
    figure("leaves", "The ten leaf frames holding the most of option b's fairness-phase samples, before and after the fix. A leaf frame is the method running when a sample was taken.")
    leaves_b = profiles["b fairness"]["leafFramesPercent"]
    p(f"After the fix the leaders are spread out, and no single collection call stands out. The route timing search itself is the largest leaf ({leaves_b['dev.waterflex.scheduler.optimizer.RouteTimingSearch.segments']:.1f} percent), now with the number lookups inlined into it. BigDecimal holds about {profiles['b fairness']['samplesPercent']['decimal: score impacts'] + profiles['b fairness']['samplesPercent']['decimal: money']:.0f} percent: score impacts in BendableBigDecimalScore ({profiles['b fairness']['samplesPercent']['decimal: score impacts']:.0f} percent) and money ({profiles['b fairness']['samplesPercent']['decimal: money']:.0f} percent, Monetary.cents for each route). Timefold's engine and score bookkeeping hold about {profiles['b fairness']['samplesPercent']['engine'] + profiles['b fairness']['samplesPercent']['score engine']:.0f} percent, including the maps it keeps for itself ({site('b fairness', SITES[5]):.1f} percent).")
    table(["Next candidate", "Share now", "What it would take"], [
        ("BigDecimal score impacts and money", f"{profiles['b fairness']['samplesPercent']['decimal: score impacts'] + profiles['b fairness']['samplesPercent']['decimal: money']:.0f}%",
         "The long score of the representation study (S) gained 1.04x when BigDecimal was a smaller share; with it now about a fifth, S is worth measuring again on this tree. Money could also be computed per route in integer cents with the same rounding."),
        ("The route timing search", f"{leaves_b['dev.waterflex.scheduler.optimizer.RouteTimingSearch.segments']:.0f}%",
         "Re-walking only the changed part of a route needs arrival shadow variables (CR11): a larger change, gated by the same exactness test."),
        ("Timefold's engine and maps", f"{profiles['b fairness']['samplesPercent']['engine'] + profiles['b fairness']['samplesPercent']['score engine']:.0f}%",
         "Not ours to change. Nearby selection would cut the moves evaluated rather than their cost.")], [1.9, .8, 4.2], size=9)

    heading("8 Conclusions")
    table(["Question", "Answer from the evidence"], [
        ("Was the identity lookup what section 6 of the fairness report suspected?", f"Yes. With full stacks, the immutable map probe is all RoadMatrix.number, {site('old cost', SITES[0]):.0f} percent of production's cost phase, and the per-call options map in the timing search is another {site('old cost', SITES[1]):.0f} percent."),
        ("Did removing them make the search faster?", f"Yes. Production's cost phase evaluates {fp['referencePhaseSpeedup']['median']:.2f}x as many moves per second and option b's fairness phase {fb['fairnessPhaseSpeedup']['median']:.2f}x, with no change of rule. Production's fairness phase gains only {fp['fairnessPhaseSpeedup']['median']:.2f}x because it is still its decimal variance."),
        ("Does the gain depend on the fleet?", f"It is largest on small fleets ({fp['byFleet']['referencePhase']['10']['median']:.2f}x at 10 technicians, {fp['byFleet']['referencePhase']['50']['median']:.2f}x at 50), because the work removed is per leg while the fleet-wide part of a move grows with the fleet."),
        ("What did the extra moves buy?", f"Cheaper days: a median {plain(-cost_diff('baseline', 'old'))} cents less for production, cheaper beyond noise in {noise[('baseline', 'old', 'costCents')][1]} of 36 cases and dearer in none. Final fairness is about the same at the median and a little lower in some cases, because the allowance is measured from a cheaper plan."),
        ("Is option b still the right formula?", f"Yes. On the faster engine it passes every bar, is {sb['fairnessPhaseSpeedup']['median']:.1f}x faster than production's fairness phase, and was never less fair in any case."),
        ("And option c?", "No. It failed the fairness bar in one case, again, and is slower than b in the cost phase."),
        ("Should the policy change to load balance?", "Not on this evidence alone. It led b on every measure on these fleets, mostly on 10 and 30 technicians, which this morning's study did not cover, and it costs a little more. It needs a head-to-head test before the owner decides.")], [2.4, 4.5])
    p("Two cautions carry over from the fairness report. The fleets are synthetic and at most 50 technicians, and each solve was limited to two processors on a workstation, so the absolute rates are the harness's, not the VM's; the paired ratios are what transfers. And fairer is fairer by the measure's definition: whether planners see the difference is a question for the pilot.")

    heading("9 What should happen next")
    table(["Next step", "What to do", "What it decides"], [
        ("Keep the fix", "ed97d8b is in this pull request. Merge it with option b's formula; it changes no rule and passes ScoringExactnessTest unchanged.", "Production speed in both phases."),
        ("Ship option b", "Replace DayConstraintProvider.fairness with b's binary64 units, keeping SchedulingPolicy.measure (DECIMAL128) for acceptance and reporting. Run the exactness test and the FULL_ASSERT campaign.", "The production fairness formula."),
        ("Load balance head to head", "b against load balance on 10, 20, 30 and 50 technicians, at 15 s and 60 s, on fresh seeds, with b's four bars applied to load balance against b. Read the busiest technician and the cost as well as the variance.", "Whether the policy should change."),
        ("Remove c", "c failed a bar in both studies. Remove it from main with file edits once b ships, keeping it on the study branch.", "A smaller solver surface."),
        ("Revisit the allowance", "The fix made days cheaper and left the allowance relative to the cheaper plan. If the owner wants some of the saving spent on fairness, that is the allowance's size, a client setting, not a code change.", "How the saving is used."),
        ("Measure the long score again", "Repeat the representation study's S arm on the fixed tree with b, now that BigDecimal is about a fifth of both phases.", "The next speedup."),
        ("Daily load with the fix and b", "Repeat the VM rehearsal stack's daily load with both, to see the gain at the service.", "Whether the gain shows at the API.")], [1.55, 3.7, 1.65], size=9)
    sub("How the calculations were checked")
    p("Every table and chart value is computed from the evidence file at build time; nothing is typed in, except the table of callers in section 2, which stacks.py writes from the run's own recordings and stores with the SHA-256 of each one. report_data.py checks the evidence before use: the run is the full study on the two pinned commits; every arm, case and repetition is present exactly once; no run has a violation or ends over its ceiling. It then recomputes from the individual runs every summary the decisions rest on (the fix's speedups overall and by fleet, the cost differences, the fairness bar counts, every comparison ratio and acceptance count) and requires them to equal the recorded summary. verify_report.py recomputes the tables a second time from the raw JSON without the builder's helpers, and checks the figure captions, the template parts and the absence of dashes. [4]")

    heading("Appendix A Every case for the fix")
    p("Each row is one case: the reference plan's and the final plan's cost on the fixed tree minus the old one, in cents, and the final variance of both, in thousandths, all as means of two repetitions; first for production, then for option b. [4]")
    table(["Techs", "Workload", "Seed", "Production: reference / final cost", "Variance old / fixed", "b: reference / final cost", "Variance old-b / b"],
          [[c[0], c[1].replace("_", " ").title(), c[2],
            *[x for arm, old in FIX.items() for x in (f"{cents(float(case_mean(by, arm, c, 'referenceCostCents') - case_mean(by, old, c, 'referenceCostCents')))} / {cents(float(case_mean(by, arm, c, 'costCents') - case_mean(by, old, c, 'costCents')))}",
                                                    f"{milli(case_mean(by, old, c, 'variance'))} / {milli(case_mean(by, arm, c, 'variance'))}")]] for c in all_cases],
          [.45, .95, .45, 1.45, 1.15, 1.3, 1.15], size=9)

    heading("Appendix B Every paired case for b and c")
    p("Each row is one case: the fixed production's final variance in its two repetitions, its spread, the option's two final variances, all in thousandths, and the mean cost difference in cents. The fairness bar fails where the option's mean exceeds production's by more than the spread. [4]")
    table(["Arm", "Techs", "Workload", "Seed", "Baseline variance", "Spread", "Option variance", "Mean difference", "Cost difference"],
          [[x["arm"], x["technicians"], x["workload"].replace("_", " ").title(), x["seed"], " / ".join(milli(v) for v in x["baselineVariance"]), milli(x["baselineVarianceSpread"]),
            " / ".join(milli(v) for v in x["armVariance"]), milli(x["meanVarianceDifference"]), cents(x["meanCostDifferenceCents"])]
           for x in sorted(paired["pairs"], key=lambda x: (x["arm"], x["technicians"], x["workload"], x["seed"]))],
          [.4, .45, 1.0, .45, 1.15, .65, 1.15, .8, .85], size=9)

    heading("Appendix C Every case for load balance")
    p("Each row is one case: the final variance of the fixed production, b and load balance, in thousandths, and load balance's highest utilization and cost against b's, as means of two repetitions. [4]")
    table(["Techs", "Workload", "Seed", "Variance: production / b / load balance", "Highest utilization: b / load balance", "Cost, load balance minus b"],
          [[c[0], c[1].replace("_", " ").title(), c[2], " / ".join(milli(case_mean(by, arm, c, "variance")) for arm in ("baseline", "b", "load-balance")),
            " / ".join(f"{float(case_mean(by, arm, c, 'maximumUtilization')):.3f}" for arm in ("b", "load-balance")),
            f"{cents(float(case_mean(by, 'load-balance', c, 'costCents') - case_mean(by, 'b', c, 'costCents')))} c"] for c in all_cases],
          [.5, 1.0, .5, 2.0, 1.65, 1.25], size=9)

    heading("Appendix D Every profile")
    p("Each column is one arm and phase on seed 17: the share of its CPU samples by kind of work, over twelve recordings (one per fleet and workload). A blank is a kind with no samples. [3]")
    names = [f"{arm} {phase}" for arm in ARMS for phase in ("cost", "fairness")]
    for chunk in (names[:6], names[6:]):
        table(["Kind of work", *[n.replace(" cost", ", cost").replace(" fairness", ", fair.") for n in chunk]],
              [[group, *[f"{profiles[n]['samplesPercent'][group]:.1f}" if group in profiles[n]["samplesPercent"] else "" for n in chunk]] for group in GROUPS] +
              [["Samples", *[f"{sum(profiles[n]['samples'].values()):,}" for n in chunk]]], [1.8, *([5.1 / len(chunk)] * len(chunk))], size=9)

    heading("Appendix E Definitions and sources")
    table(["Term", "Meaning in this report"], [
        ("Tree", "One commit, frozen and built once: o before the fix, n with it."),
        ("Arm", "One tree and solver variant, run in every case."),
        ("Case", "One fleet size, workload and seed; each was solved twice by every arm."),
        ("Reference plan", "The plan the cost phase ends with; the fairness phase starts from it."),
        ("Allowance and ceiling", "The fairness phase may spend up to 2 percent above the reference plan's cost; the ceiling is that cost times 1.02, rounded down to a cent."),
        ("Fact revision", "The immutable facts one plan is scored against: technicians, visits, roads and rates. The fix keeps a visit's number together with the revision that gave it."),
        ("Moves per second", "Candidate moves a phase evaluated per second; more moves in the same budget is a more thorough search, not a faster answer."),
        ("Speedup", "An arm's moves per second divided by its comparison's in the same case and repetition; the report gives the median of these ratios."),
        ("Beyond noise", "A case's difference is beyond noise when it exceeds the comparison arm's own difference between its two repetitions in that case."),
        ("Leaf frame and site", "A leaf frame is the method at the top of a sampled stack. A site in section 2 also names the scheduler method that called it, from the full stack."),
        ("Thousandths", "Variances in this report are shown multiplied by 1000.")], [1.6, 5.3])
    p("[1] docs/scheduler-policy.md: the fairness phase, the allowance, the acceptance rule and the numeric contract.")
    p("[2] The fix, commit ed97d8b: PlanVisit and TechRoute keep their road matrix numbers (roadNumber) and a route its scoring facts per revision; RouteTimeline and RouteTimingSearch read them; the regression test is RouteTimelineTest.roadNumbersFollowTheRevisionTheyAreScoredAgainst; the gate is ScoringExactnessTest.")
    p("[3] Profiles: the profiles block of the evidence, attributed by infra/representation_study.py's frame groups, and stacks.json in this folder, written by stacks.py from the run's recordings in experiments/runs/matrix-study-20261020T155528/profiles, which are kept outside git.")
    p("[4] Evidence: docs/evidence/solver-matrix-2026-10-20.json, written by infra/matrix_study.py from the run infra/run-matrix-study.ps1 made (manifest, every run, the fix's pairs, b and c against the bars, every arm on every measure, the decision and the profiles). This morning's figures for comparison come from docs/evidence/solver-variance-2026-10-20.json. docs/solver-benchmarks.md, 'Road matrix numbering (PR 187)', is the narrative kept with the code.")
    p("Reproduce with docs/road-matrix-numbering-2026-10-20/requirements.txt: render the two diagrams as EVIDENCE.md says, then run report_data.py, build_report.py and verify_report.py there. The builder regenerates the charts and chart_data.json and writes this editable Word file.")

    settings = doc.settings.element
    for el in settings.findall(qn("w:updateFields")): settings.remove(el)
    update = OxmlElement("w:updateFields"); update.set(qn("w:val"), "true"); settings.append(update)
    stream = BytesIO(); doc.save(stream)
    with ZipFile(REFERENCE) as ref, ZipFile(stream) as built, ZipFile(OUTPUT, "w", ZIP_DEFLATED) as output:
        preserve = set(PRESERVE) | {n for n in ref.namelist() if "footer" in n}
        for name in sorted(built.namelist()):
            content = ref.read(name) if name in preserve else built.read(name)
            info = ZipInfo(name, (2026, 10, 20, 0, 0, 0)); info.compress_type = ZIP_DEFLATED
            output.writestr(info, content)
    print(OUTPUT, len(doc.tables), "tables", len(figures), "figures")


if __name__ == "__main__": build()
