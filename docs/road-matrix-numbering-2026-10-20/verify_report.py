"""Check the delivered tables, chart data and template parts against the raw evidence file.

Recomputes from docs/evidence/solver-matrix-2026-10-20.json without the builder's or loader's helpers. Run normally,
without Python's -O flag. Visual pagination still requires rendering.
"""
import json
from decimal import Decimal, localcontext
from hashlib import sha256
from statistics import median
from zipfile import ZipFile

from docx import Document
from build_report import DASHES, HERE, OUTPUT, PRESERVE, REFERENCE, REFERENCE_SHA

EVIDENCE = HERE.parent / "evidence" / "solver-matrix-2026-10-20.json"
MEASURES = ("variance", "maximumUtilization", "utilizationSpread", "absoluteDeviation", "loadBalance")
NAMES = {"old": "old (before the fix)", "old-b": "old-b (b before the fix)", "b": "b (binary64)", "c": "c (running sums)", "load-balance": "load-balance"}
RATES = {"referencePhase": "referenceMovesPerSecond", "fairnessPhase": "fairnessMovesPerSecond"}
FIX = (("baseline", "old"), ("b", "old-b"))
FLEETS = (10, 30, 50)


def signed(value):
    return f"{value:+,.2f}".rstrip("0").rstrip(".")


def thousandths(value):
    text = f"{Decimal(value) * 1000:.2f}"
    return "0.00" if text == "-0.00" else text


def verify():
    raw = json.loads(EVIDENCE.read_text(encoding="utf-8-sig"))
    paired = raw["paired"]
    doc = Document(OUTPUT)
    assert len(doc.tables) == 23 and len(doc.inline_shapes) == 14 and len(doc.sections) == 1, (len(doc.tables), len(doc.inline_shapes))
    assert len([p for p in doc.paragraphs if p.style.name == "Heading 1"]) == 14
    captions = [p for p in doc.paragraphs if p.style.name == "Caption"]
    assert len(captions) == 14
    for i, p in enumerate(captions, 1):
        assert p.text.startswith(f"Figure {i}.") and all(str(r.font.color.rgb) == "000000" and r.italic for r in p.runs)
    for shape in doc.inline_shapes:
        assert shape._inline.docPr.get("descr", "").startswith("Figure ")

    def rows(index):
        return [[c.text for c in r.cells] for r in doc.tables[index].rows[1:]]

    def check(index, expected):
        actual = rows(index)
        assert actual == [[str(c) for c in r] for r in expected], (index, actual[:2], expected[:2])

    runs = {(r["arm"], r["technicians"], r["workload"], r["seed"], r["rep"]): r for r in paired["runs"]}
    keys = sorted({(r["technicians"], r["workload"], r["seed"]) for r in paired["runs"]})
    assert len(keys) == 36 and len(runs) == 432 == len(paired["runs"])

    def value(arm, case, rep, field):
        run = runs[(arm, *case, rep)]
        return Decimal(run["measures"][field]) if field in MEASURES else Decimal(run[field])

    def mean(arm, case, field):
        return (value(arm, case, 1, field) + value(arm, case, 2, field)) / 2

    def pairs(arm, against, rate, fleet=None):
        return [runs[(arm, *c, rep)][rate] / runs[(against, *c, rep)][rate] for c in keys if fleet is None or c[0] == fleet for rep in (1, 2)]

    def noise(arm, against, field):
        worse = better = 0
        for c in keys:
            own = [value(against, c, rep, field) for rep in (1, 2)]
            gap = mean(arm, c, field) - mean(against, c, field)
            worse += gap > max(own) - min(own); better += gap < -(max(own) - min(own))
        return worse, better

    # Table 6: the fix's speedups, overall and by fleet.
    expected = []
    for arm, old in FIX:
        for phase, rate in RATES.items():
            every = pairs(arm, old, rate)
            expected.append([f"{'Production' if arm == 'baseline' else 'Option b'}, {'cost' if phase == 'referencePhase' else 'fairness'} phase",
                             f"{round(median(every), 3):.2f}x ({round(min(every), 3):.2f} to {round(max(every), 3):.2f})",
                             *[f"{round(median(pairs(arm, old, rate, f)), 3):.2f}x" for f in FLEETS]])
    check(6, expected)
    assert paired["decision"]["fix"]["promoted"] and all(median(pairs("baseline", "old", rate)) >= 1 for rate in RATES.values())

    # Table 8: what the fix's extra moves bought.
    expected = []
    for arm, old in FIX:
        ref = median(float(sum(Decimal(runs[(arm, *c, rep)]["referenceCostCents"]) - Decimal(runs[(old, *c, rep)]["referenceCostCents"]) for rep in (1, 2)) / 2) for c in keys)
        cost = median(float(mean(arm, c, "costCents") - mean(old, c, "costCents")) for c in keys)
        variance = median(float(mean(arm, c, "variance") / mean(old, c, "variance")) for c in keys if mean(old, c, "variance") > 0)
        cw, cb = noise(arm, old, "costCents"); vw, vb = noise(arm, old, "variance")
        expected.append([f"{'Production' if arm == 'baseline' else 'Option b'}, fixed against old", f"{signed(ref)} c", f"{signed(cost)} c", f"{cb} / {cw}", f"{variance:.3f}", f"{vw} / {vb}"])
    check(8, expected)

    # Table 9: b and c against the bars, recomputed from the runs.
    expected = []
    for arm in ("b", "c"):
        speed = pairs(arm, "baseline", "fairnessMovesPerSecond"); reference = pairs(arm, "baseline", "referenceMovesPerSecond")
        fairer = same = less = beyond = 0
        for c in keys:
            base = [value("baseline", c, rep, "variance") for rep in (1, 2)]
            gap = mean(arm, c, "variance") - sum(base) / 2
            fairer += gap < 0; same += gap == 0; less += gap > 0; beyond += gap > max(base) - min(base)
        diffs = [float(mean(arm, c, "costCents") - mean("baseline", c, "costCents")) for c in keys]
        within = all(runs[(arm, *c, rep)]["costCents"] <= runs[(arm, *c, rep)]["costCeilingCents"] for c in keys for rep in (1, 2))
        passes = median(speed) >= 1.25 and beyond == 0 and median(diffs) <= 0 and within
        assert passes == paired["summary"][arm]["passes"] == paired["decision"][arm]["promoted"]
        expected.append([arm, f"{round(median(speed), 3):.2f}x ({round(min(speed), 3):.2f} to {round(max(speed), 3):.2f})", f"{round(median(reference), 3):.3f}x",
                         f"{fairer} / {same} / {less} / {beyond}",
                         f"{signed(median(diffs))} c ({sum(d < 0 for d in diffs)} / {sum(d == 0 for d in diffs)} / {sum(d > 0 for d in diffs)})", "Yes" if within else "No", "Yes" if passes else "No"])
    check(9, expected)
    assert paired["decision"]["b"]["promoted"] and not paired["decision"]["c"]["promoted"]

    # Table 11: every arm on every measure against the fixed production.
    expected = []
    for arm in ("old", "old-b", "b", "c", "load-balance"):
        cells = []
        for name in MEASURES:
            both = [(mean(arm, c, name), mean("baseline", c, name)) for c in keys]
            cells.append(f"{round(median([float(a / b) for a, b in both if b > 0]), 3):.2f} ({sum(a < b for a, b in both)}/{sum(a > b for a, b in both)})")
        diffs = [float(mean(arm, c, "costCents") - mean("baseline", c, "costCents")) for c in keys]
        expected.append([NAMES[arm], *cells, f"{signed(median(diffs))} c", f"{round(median(pairs(arm, 'baseline', 'fairnessMovesPerSecond')), 3):.2f}x"])
    check(11, expected)

    # Table 12: load balance against b, by fleet.
    expected = []
    for f in FLEETS:
        sized = [c for c in keys if c[0] == f]
        lower = sum(mean("load-balance", c, "variance") < mean("b", c, "variance") for c in sized)
        higher = sum(mean("load-balance", c, "variance") > mean("b", c, "variance") for c in sized)
        out_lower = out_higher = 0
        for c in sized:
            own = [value("b", c, rep, "variance") for rep in (1, 2)]
            gap = mean("load-balance", c, "variance") - mean("b", c, "variance")
            out_higher += gap > max(own) - min(own); out_lower += gap < -(max(own) - min(own))
        over = lambda arm: median(float(mean(arm, c, "variance") / mean("baseline", c, "variance")) for c in sized if mean("baseline", c, "variance") > 0)
        expected.append([f"{f} technicians", f"{lower} / {higher}", f"{out_lower} / {out_higher}", f"{over('b'):.2f}", f"{over('load-balance'):.2f}"])
    check(12, expected)

    # Tables 5 and 13: the call-site and profile shares.
    stacks = json.loads((HERE / "stacks.json").read_text(encoding="utf-8"))
    for row in rows(5):
        assert row[1:] == [f"{stacks['profiles'][n]['sitesPercent'].get(row[0], 0.0):.1f}%" for n in ("old cost", "baseline cost", "old-b fairness", "b fairness")], row
    for name, profile in stacks["profiles"].items():
        assert len(profile["recordings"]) == 12 and profile["samples"] > 0, name
    for row in rows(13):
        assert row[1:] == [f"{raw['profiles'][n]['samplesPercent'].get(row[0], 0):.1f}%" for n in ("old cost", "baseline cost", "old fairness", "baseline fairness", "old-b fairness", "b fairness")], row

    # Table 17: every case for the fix.
    expected = []
    for c in keys:
        cells = []
        for arm, old in FIX:
            ref = float(sum(Decimal(runs[(arm, *c, rep)]["referenceCostCents"]) - Decimal(runs[(old, *c, rep)]["referenceCostCents"]) for rep in (1, 2)) / 2)
            cells += [f"{signed(ref)} / {signed(float(mean(arm, c, 'costCents') - mean(old, c, 'costCents')))}", f"{thousandths(mean(old, c, 'variance'))} / {thousandths(mean(arm, c, 'variance'))}"]
        expected.append([c[0], c[1].replace("_", " ").title(), c[2], *cells])
    check(17, expected)

    # Table 18: every paired case for b and c.
    expected = []
    for arm in ("b", "c"):
        for c in keys:
            base = [value("baseline", c, rep, "variance") for rep in (1, 2)]
            own = [value(arm, c, rep, "variance") for rep in (1, 2)]
            expected.append([arm, c[0], c[1].replace("_", " ").title(), c[2], " / ".join(thousandths(v) for v in base), thousandths(max(base) - min(base)),
                             " / ".join(thousandths(v) for v in own), thousandths(sum(own) / 2 - sum(base) / 2),
                             signed(float(mean(arm, c, "costCents") - mean("baseline", c, "costCents")))])
    check(18, expected)

    # Table 19: every case for load balance.
    expected = [[c[0], c[1].replace("_", " ").title(), c[2], " / ".join(thousandths(mean(arm, c, "variance")) for arm in ("baseline", "b", "load-balance")),
                 " / ".join(f"{float(mean(arm, c, 'maximumUtilization')):.3f}" for arm in ("b", "load-balance")),
                 f"{signed(float(mean('load-balance', c, 'costCents') - mean('b', c, 'costCents')))} c"] for c in keys]
    check(19, expected)

    # Chart data agrees with the evidence.
    chart = json.loads((HERE / "generated" / "chart_data.json").read_text(encoding="utf-8"))
    for arm, old in FIX:
        for rate in RATES.values():
            for f in FLEETS:
                assert [round(v, 4) for v in pairs(arm, old, rate, f)] == chart["fixSpeed"][arm][rate][str(f)], (arm, rate, f)
    for arm in ("b", "c", "load-balance"):
        for f in FLEETS:
            assert [round(v, 4) for v in pairs(arm, "baseline", "fairnessMovesPerSecond", f)] == chart["speed"][arm][str(f)]
        assert chart["cost"][arm] == [float(mean(arm, c, "costCents") - mean("baseline", c, "costCents")) for c in keys]
    for arm, ratios in chart["heatmap"].items():
        assert ratios == {m: paired["comparison"][arm]["measures"][m]["medianRatioToBaseline"] for m in MEASURES}
    for row in chart["profile"]:
        assert row["shares"] == raw["profiles"][f"{row['arm']} {row['phase']}"]["samplesPercent"]
    for arm, counts in chart["acceptance"].items():
        got = [r["acceptanceReason"] for r in paired["runs"] if r["arm"] == arm]
        assert counts == {reason: got.count(reason) for reason in set(got)}, arm

    # Template parts preserved byte for byte, and no em or en dash anywhere in the document XML.
    assert sha256(REFERENCE.read_bytes()).hexdigest() == REFERENCE_SHA
    with ZipFile(REFERENCE) as ref, ZipFile(OUTPUT) as out:
        for name in set(PRESERVE) | {n for n in ref.namelist() if "footer" in n}:
            assert ref.read(name) == out.read(name), name
        for name in out.namelist():
            if name.endswith(".xml"):
                text = out.read(name).decode("utf-8")
                assert not any(dash in text for dash in DASHES), name
        body = out.read("word/document.xml").decode("utf-8")
        assert body.count("<w:cantSplit/>") == sum(len(t.rows) for t in doc.tables)
    print("verified", OUTPUT.name)


if __name__ == "__main__":
    # Exact decimal arithmetic: the default 28 digits would round two equal 34-digit variances apart when averaged.
    with localcontext() as context:
        context.prec = 80
        verify()
