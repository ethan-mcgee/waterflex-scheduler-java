"""Check the delivered tables, chart data and template parts against the raw evidence file.

Recomputes from docs/evidence/solver-load-balance-2026-10-23.json without the builder's or loader's helpers, finding each
table by its header row. Run normally, without Python's -O flag. Visual pagination still requires rendering.
"""
import json
import sys
from decimal import Decimal
from hashlib import sha256
from statistics import median
from zipfile import ZipFile

from docx import Document
from build_report import HERE, OUTPUT, PRESERVE, REFERENCE, REFERENCE_SHA

EVIDENCE = HERE.parent / "evidence" / "solver-load-balance-2026-10-23.json"
ARMS = ("b", "load-balance")
CODE = {"FAIRNESS_IMPROVEMENT": "F", "COST_REDUCTION": "C", "OVERTIME_REDUCTION": "O"}


def signed(value):
    return f"{value:+,.1f}" if value != int(value) else f"{int(value):+,}"


def thousandths(value):
    text = f"{Decimal(value) * 1000:.2f}"
    return "0.00" if text == "-0.00" else text


def title(workload):
    return workload.replace("_", " ").title()


def verify(evidence=EVIDENCE, output=OUTPUT, charts=HERE / "generated"):
    raw = json.loads(evidence.read_text(encoding="utf-8-sig"))
    doc = Document(output)
    captions = [p for p in doc.paragraphs if p.style.name == "Caption"]
    assert len(captions) == len(doc.inline_shapes) == 11 and len(doc.sections) == 1, (len(captions), len(doc.inline_shapes))
    for i, p in enumerate(captions, 1):
        assert p.text.startswith(f"Figure {i}.") and all(str(r.font.color.rgb) == "000000" and r.italic for r in p.runs)
    for shape in doc.inline_shapes:
        assert shape._inline.docPr.get("descr", "").startswith("Figure ")

    def tables(first):
        found = [t for t in doc.tables if t.rows[0].cells[0].text == first]
        assert found, first
        return found

    def rows(t):
        return [[c.text for c in r.cells] for r in t.rows[1:]]

    keys = sorted(raw["budgets"], key=int)
    runs = {key: {(r["arm"], r["technicians"], r["workload"], r["fixtureSeed"], r["seed"], r["rep"]): r for r in raw["budgets"][key]["runs"]} for key in keys}
    cases = sorted({(r["technicians"], r["workload"], r["fixtureSeed"], r["seed"]) for r in raw["budgets"][keys[0]]["runs"]})
    reps = sorted({r["rep"] for r in raw["budgets"][keys[0]]["runs"]})

    def mean(key, arm, case, field):
        values = [Decimal(runs[key][(arm, *case, rep)]["measures"][field]) if field in raw["budgets"][key]["measures"] else Decimal(runs[key][(arm, *case, rep)][field])
                  for rep in reps]
        return sum(values) / len(values)

    def median_ratio(key, field, keep=lambda c: True):
        values = [float(mean(key, "load-balance", c, field) / mean(key, "b", c, field)) for c in cases if keep(c) and mean(key, "b", c, field) > 0]
        return median(values) if values else None

    # The appendices: every case at every budget.
    appendices = tables("Techs")
    assert len(appendices) == len(keys)
    for t, key in zip(appendices, keys):
        expected = []
        for c in cases:
            b, lb = mean(key, "b", c, "variance"), mean(key, "load-balance", c, "variance")
            expected.append([str(c[0]), title(c[1]), str(c[2]), thousandths(b), thousandths(lb), f"{float(lb / b):.2f}" if b > 0 else "n/a",
                             f"{float(mean(key, 'b', c, 'maximumUtilization')):.3f}", f"{float(mean(key, 'load-balance', c, 'maximumUtilization')):.3f}",
                             signed(float(mean(key, "load-balance", c, "costCents") - mean(key, "b", c, "costCents"))),
                             " / ".join("".join(CODE.get(runs[key][(arm, *c, r)]["acceptanceReason"], "R") for r in reps) for arm in ARMS)])
        assert rows(t) == expected, (key, rows(t)[:2], expected[:2])

    # The workload table of section 2.
    workloads = sorted({c[1] for c in cases})
    fmt = lambda v: "n/a" if v is None else f"{v:.2f}"
    expected = [[title(w), *[fmt(median_ratio(key, "variance", lambda c, w=w: c[1] == w)) for key in keys],
                 *[fmt(median_ratio(key, "loadBalance", lambda c, w=w: c[1] == w)) for key in keys]] for w in workloads]
    assert rows(tables("Workload")[0]) == expected

    # The speed table of section 6: medians of paired moves-per-second ratios by fleet.
    field = {"referencePhase": "referenceMovesPerSecond", "fairnessPhase": "fairnessMovesPerSecond"}
    fleets = sorted({c[0] for c in cases})
    speed = lambda key, phase, keep: median([runs[key][("load-balance", *c, r)][field[phase]] / runs[key][("b", *c, r)][field[phase]] for c in cases if keep(c) for r in reps])
    expected = [[f"{f} technicians", *[f"{speed(key, phase, lambda c, f=f: c[0] == f):.2f}x" for key in keys for phase in field]] for f in fleets]
    expected.append(["All", *[f"{raw['budgets'][key]['speed'][phase]['all']['median']:.2f}x" for key in keys for phase in field]])
    assert rows(tables("Fleet")[0]) == expected
    for key in keys:
        for phase in field:
            assert abs(round(speed(key, phase, lambda c: True), 3) - raw["budgets"][key]["speed"][phase]["all"]["median"]) <= .001

    # The rule table of section 3, and the decision.
    rule_rows = rows(tables("Part, budget")[0])
    assert len(rule_rows) == 4 * len(keys)
    for i, key in enumerate(keys):
        budget = raw["budgets"][key]
        variance = [(mean(key, "load-balance", c, "variance"), mean(key, "b", c, "variance")) for c in cases]
        beyond = sum(a > b * Decimal("1.05") for a, b in variance)
        relative = median([(mean(key, "load-balance", c, "costCents") - mean(key, "b", c, "costCents")) / mean(key, "b", c, "costCents") for c in cases])
        holds = {"variance": budget["measures"]["variance"]["medianRatioInterval"][1] < 1 and beyond <= len(cases) // 10,
                 "maximum": round(median_ratio(key, "maximumUtilization"), 3) <= 1, "cost": relative <= Decimal("0.001"),
                 "ceiling": all(r["costCents"] <= r["costCeilingCents"] for r in budget["runs"] if r["acceptanceReason"] in CODE)
                            and all(r == 0 for r in budget["violations"].values())}
        assert holds == budget["rule"], (key, holds, budget["rule"])
        assert [row[3] for row in rule_rows[4 * i:4 * i + 4]] == ["Yes" if holds[p] else "No" for p in ("variance", "maximum", "cost", "ceiling")]
    assert raw["decision"]["switch"] == all(all(raw["budgets"][key]["rule"].values()) for key in keys)

    # Chart data agrees with the evidence.
    chart = json.loads((charts / "chart_data.json").read_text(encoding="utf-8"))
    for key in keys:
        assert sorted(p["ratio"] for p in chart["ratios"][key]) == sorted(float(mean(key, "load-balance", c, "variance") / mean(key, "b", c, "variance"))
                                                                          for c in cases if mean(key, "b", c, "variance") > 0)
        for phase in field:
            for f in fleets:
                assert chart["speed"][key][phase][str(f)] == [round(runs[key][("load-balance", *c, r)][field[phase]] / runs[key][("b", *c, r)][field[phase]], 4)
                                                              for c in cases if c[0] == f for r in reps]

    # Template parts preserved byte for byte, and no em or en dash anywhere in the document XML.
    assert sha256(REFERENCE.read_bytes()).hexdigest() == REFERENCE_SHA
    with ZipFile(REFERENCE) as ref, ZipFile(output) as out:
        for name in set(PRESERVE) | {n for n in ref.namelist() if "footer" in n}:
            assert ref.read(name) == out.read(name), name
        for name in out.namelist():
            if name.endswith(".xml"):
                text = out.read(name).decode("utf-8")
                assert "—" not in text and "–" not in text, name
        body = out.read("word/document.xml").decode("utf-8")
        assert body.count("<w:cantSplit/>") == sum(len(t.rows) for t in doc.tables)
    print("verified", output.name)


if __name__ == "__main__":
    verify()
