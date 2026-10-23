"""Build the load balance head-to-head report from the validated evidence and the reference explainer's template.

Install requirements.txt first. No network, services or benchmark execution. The two diagrams in diagrams/ are rendered
to generated/ separately (see EVIDENCE.md); the charts are regenerated here.
"""
from copy import deepcopy
from datetime import datetime
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from statistics import median
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from charts import LABEL, make_charts, seconds, workload_name
from report_data import (ARMS, COST_ALLOWANCE, LESS_FAIR_ALLOWED, MEASURES, PHASES, PROFILE_GROUPS, TOLERANCE, allowance_used, case_mean, cases,
                         cost_differences, earlier_ratio, index, load, perfectly_even, ratios, reps, same_plan, speedups)

HERE = Path(__file__).resolve().parent
REFERENCE = HERE.parent / "WaterFlex_Scheduler_Algorithm_Comparison_Explained.docx"
REFERENCE_SHA = "5366f3017d6f18c15de295ba5fb9c1a014b6222911040f1c5e044845e57f832d"
OUTPUT = HERE / "Load_Balance_Head_To_Head_2026-10-23_Report.docx"
PRESERVE = ("word/styles.xml", "word/stylesWithEffects.xml", "word/numbering.xml", "word/theme/theme1.xml", "word/fontTable.xml")
DIAGRAMS = ("study-design", "weighting")
MEASURE_NAME = {"variance": "Variance", "maximumUtilization": "Highest utilization", "utilizationSpread": "Utilization spread",
                "absoluteDeviation": "Absolute deviation", "loadBalance": "Load balance"}
PHASE_NAME = {"referencePhase": "cost phase", "fairnessPhase": "fairness phase"}


def cents(value):
    return f"{value:+,.1f}" if value != int(value) else f"{int(value):+,}"


def milli(value):
    """A variance as thousandths, the unit every variance in the report is shown in."""
    text = f"{Decimal(value) * 1000:.2f}"
    return "0.00" if text == "-0.00" else text


def pct(value, places=0):
    return f"{value * 100:.{places}f}%"


def span(entry):
    """A median ratio with its 95 percent interval."""
    if entry["medianRatioToB"] is None:
        return "n/a"
    low, high = entry["medianRatioInterval"] or (entry["medianRatioToB"], entry["medianRatioToB"])
    return f"{entry['medianRatioToB']:.2f} ({low:.2f} to {high:.2f})"


def frame_name(frame):
    """Class and method of a leaf frame; a lambda's generated class (a hex address) is named by the class that declared it."""
    parts = frame.split(".")
    if parts[-2].startswith("0x"):
        return parts[-3].replace("$$Lambda", "") + " lambda." + parts[-1]
    return ".".join(parts[-2:])


def yes(flag):
    return "Yes" if flag else "No"


def build(evidence=None, full=True, output=OUTPUT, assets=HERE / "generated"):
    """The report from the full evidence; {@code evidence}, {@code full} and the paths are for a dry run on the smoke evidence."""
    assert sha256(REFERENCE.read_bytes()).hexdigest() == REFERENCE_SHA
    document = load(evidence, full) if evidence else load()
    for name in DIAGRAMS:
        assert (assets / f"{name}.png").exists(), f"render diagrams/{name}.html first (EVIDENCE.md)"
    make_charts(document, assets)
    doc = Document(REFERENCE)
    table_properties = deepcopy(doc.tables[0]._tbl.tblPr)
    for child in list(doc.element.body):
        if child.tag != qn("w:sectPr"): doc.element.body.remove(child)
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype.endswith(("/image", "/hyperlink")): del doc.part.rels[rid]
    doc.core_properties.title = "WaterFlex scheduler: load balance against the variance"
    doc.core_properties.subject = "Timefold's load balance against production's capacity-weighted variance as the daily fairness policy, 2026-10-23"
    doc.core_properties.author = doc.core_properties.last_modified_by = "WaterFlex"
    doc.core_properties.created = doc.core_properties.modified = datetime(2026, 10, 23)
    figures = []

    def p(text, style=None):
        assert "—" not in text and "–" not in text
        para = doc.add_paragraph(text, style); para.paragraph_format.widow_control = True
        return para

    def heading(text):
        h = p(text, "Heading 1"); h.paragraph_format.page_break_before = True

    def sub(text): p(text, "Heading 2")

    def table(headers, data, widths, size=10):
        assert abs(sum(widths) - 6.9) < .01
        t = doc.add_table(rows=1, cols=len(headers)); t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t._tbl.replace(t._tbl.tblPr, deepcopy(table_properties))
        for col, w in zip(t.columns, widths): col.width = Inches(w)
        for i, values in enumerate([headers, *data]):
            assert len(values) == len(headers)
            row = t.rows[0] if i == 0 else t.add_row()
            trpr = row._tr.get_or_add_trPr(); trpr.append(OxmlElement("w:cantSplit"))
            if i == 0: trpr.append(OxmlElement("w:tblHeader"))
            for j, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
                text = str(value)
                assert "—" not in text and "–" not in text
                cell.width = Inches(width); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell.text = text
                props = cell._tc.get_or_add_tcPr()
                shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "DCE6EF" if i == 0 else ("FFFFFF" if i % 2 else "F4F6F8")); props.append(shade)
                margins = OxmlElement("w:tcMar")
                for side in ("top", "left", "bottom", "right"):
                    n = OxmlElement("w:" + side); n.set(qn("w:w"), "65"); n.set(qn("w:type"), "dxa"); margins.append(n)
                props.append(margins)
                para = cell.paragraphs[0]; para.paragraph_format.space_after = Pt(2); para.paragraph_format.space_before = Pt(2); para.paragraph_format.line_spacing = 1.0
                if len(headers) > 2 and j > 0: para.alignment = 1
                if len(data) <= 12 and i < len(data): para.paragraph_format.keep_with_next = True  # Short tables stay on one page.
                for run in para.runs: run.font.name = "Calibri"; run.font.size = Pt(size); run.bold = i == 0
        spacer = p(""); spacer.paragraph_format.space_after = Pt(0); spacer.paragraph_format.line_spacing = Pt(3)

    def figure(name, caption, width=6.85):
        figures.append(name)
        caption = f"Figure {len(figures)}. {caption}"
        para = p(""); shape = para.add_run().add_picture(str(assets / f"{name}.png"), width=Inches(width))
        shape._inline.docPr.set("descr", caption); para.paragraph_format.keep_with_next = True; para.alignment = 1
        cap = p(caption, "Caption"); cap.paragraph_format.space_after = Pt(7)
        for run in cap.runs:
            run.font.color.rgb = RGBColor(0, 0, 0); run.bold = False; run.italic = True

    # The numbers the narrative refers to by name.
    manifest = document["manifest"]
    budgets = document["budgets"]
    keys = sorted(budgets, key=int)
    short_key, long_key = keys[0], keys[-1]
    by = {key: index(budgets[key]) for key in keys}
    all_cases = cases(budgets[short_key])
    repetitions = reps(budgets[short_key])
    fleets = sorted({c[0] for c in all_cases})
    workloads = sorted({c[1] for c in all_cases})
    profiles = document.get("profiles", {})
    switch = document["decision"]["switch"]
    total_runs = sum(len(budgets[key]["runs"]) for key in keys)
    n = len(all_cases)
    allowed = int(LESS_FAIR_ALLOWED * n)

    def m(key, name):
        return budgets[key]["measures"][name]

    def rate(key, arm, field):
        return median([r[field] for r in budgets[key]["runs"] if r["arm"] == arm])

    def keep_workload(w):
        return lambda c: c[1] == w

    def keep_fleet(f):
        return lambda c: c[0] == f

    def ratio_median(key, name, keep):
        values = ratios(by[key], all_cases, repetitions, name, keep)
        return median(values) if values else None

    def fmt(value):
        return "n/a" if value is None else f"{value:.2f}"

    failed = {key: [rule for rule, ok in budgets[key]["rule"].items() if not ok] for key in keys}
    same_final = {key: sum(1 for c in all_cases if same_plan(by[key], c, repetitions, "final")) for key in keys}
    same_reference = {key: sum(1 for c in all_cases if same_plan(by[key], c, repetitions, "reference")) for key in keys}
    # Reference cost, load balance minus b, per case: what the slower cost phase leaves before any fairness is spent.
    reference_gap = {key: [sum(by[key][("load-balance", c, r)]["referenceCostCents"] - by[key][("b", c, r)]["referenceCostCents"] for r in repetitions) / len(repetitions)
                           for c in all_cases] for key in keys}
    mixed = "MIXED_SHIFT"
    mixed_cases = [c for c in all_cases if c[1] == mixed]

    def lean(key, field, keep=lambda c: True):
        """Cases where load balance's mean is below and above b's on a field."""
        pairs = [(case_mean(by[key], "load-balance", c, field, repetitions), case_mean(by[key], "b", c, field, repetitions)) for c in all_cases if keep(c)]
        return sum(a < b for a, b in pairs), sum(a > b for a, b in pairs)

    def fleet_mean_cost(key, f):
        values = cost_differences(by[key], all_cases, repetitions, keep_fleet(f))
        return sum(values) / len(values)
    rule_name = {"variance": "variance", "maximum": "highest utilization", "cost": "cost", "ceiling": "ceiling"}

    p("WaterFlex scheduler: load balance against the variance", "Title")
    p("Timefold's load balance tested head to head against production's capacity-weighted variance as the daily fairness policy", "Subtitle")
    p(f"October 23 2026  |  Load balance head-to-head of PR J, run {manifest['startedAt'][:10]} on revision {manifest['revision'][:7]}")
    verdict = ("It meets every part of the rule at both budgets, so the evidence supports switching the policy; the switch itself is the owner's decision."
               if switch else
               "It does not meet the rule: " + "; ".join(f"at {seconds(key)} it fails on {', '.join(rule_name[r] for r in failed[key])}" for key in keys if failed[key])
               + ". The evidence does not support switching the policy.")
    p(f"The daily search ends with a fairness phase that may spend up to two percent more cost to even out how busy each technician is. Production measures that with the capacity-weighted variance of each technician's utilization (option b, shipped in PR H). The road matrix study found Timefold's load balance fairer than b on all five fairness measures at 10 to 50 technicians and 15 s, a little dearer, and slower in its cost phase. This study tested the two head to head in {total_runs} solves: {n} different days at each of two budgets, {seconds(short_key)} (the production default) and {seconds(long_key)}, including a new workload in which technicians work 4, 6 and 9 hour days, the one place where the two measures count technicians differently. On production's own measure, load balance's final variance is a median {m(short_key, 'variance')['medianRatioToB']:.2f} of b's at {seconds(short_key)} and {m(long_key, 'variance')['medianRatioToB']:.2f} at {seconds(long_key)}, at a median cost difference of {cents(budgets[short_key]['cost']['medianDifferenceCents'])} and {cents(budgets[long_key]['cost']['medianDifferenceCents'])} cents a day. Those medians are ties: the two arms ended on the same plan on {same_final[short_key]} and {same_final[long_key]} of the {n} days. On the days they differ, load balance was less fair than b on production's measure more often than it was fairer ({m(short_key, 'variance')['higher']} against {m(short_key, 'variance')['lower']} days at {seconds(short_key)}, {m(long_key, 'variance')['higher']} against {m(long_key, 'variance')['lower']} at {seconds(long_key)}), most of all on the mixed-shift days, and its slower cost phase left a dearer starting plan on about one day in six. {verdict}")
    sub("What the measurements tell us")
    table(["Reading", *[f"At {seconds(key)}" for key in keys]], [
        ["Variance, load balance over b (95% interval)", *[span(m(key, "variance")) for key in keys]],
        [f"Cases less fair by more than {pct(TOLERANCE)} (allowed: {allowed})", *[f"{m(key, 'variance')['higherBeyondTolerance']} of {n}" for key in keys]],
        ["Highest utilization, load balance over b", *[span(m(key, "maximumUtilization")) for key in keys]],
        ["Load balance unfairness, load balance over b", *[span(m(key, "loadBalance")) for key in keys]],
        ["Median cost difference (relative)", *[f"{cents(budgets[key]['cost']['medianDifferenceCents'])} c ({pct(budgets[key]['cost']['medianRelativeDifference'], 3)})" for key in keys]],
        ["Moves per second, cost / fairness phase", *[f"{budgets[key]['speed']['referencePhase']['all']['median']:.2f}x / {budgets[key]['speed']['fairnessPhase']['all']['median']:.2f}x" for key in keys]],
        ["Meets the rule", *[yes(budgets[key]["passes"]) for key in keys]]], [3.1, 1.9, 1.9])
    p(f"Each ratio is the median over the {n} cases of load balance's final value divided by b's in the same case; lower is fairer, and the bracket is the 95 percent bootstrap interval of that median. The cost is the median over cases of load balance's final cost minus b's, in cents, and as a share of b's cost. Speeds are medians of paired ratios of moves per second. Every one of the {total_runs} solves ended with no constraint violation and every accepted result inside its cost ceiling.")
    p(f"The report explains each result through what the solver recorded: the five fairness measures of every final plan, its cost and ceiling, why it was accepted, how far into each phase it found its best plan, its moves per second, and where the CPU went in flight recordings of both arms. All of it comes from one evidence file, docs/evidence/solver-load-balance-2026-10-23.json. None of it is customer data; the days are synthetic.")
    sub("How to read this report")
    table(["Sections", "What you will learn"], [
        ("1", "What was measured: the two arms, the 120 days, the five measures and the rule written before the run."),
        ("2", "Why the two measures can disagree, and the workload built to show it."),
        ("3", "Load balance against each part of the rule, at both budgets."),
        ("4", "Where load balance does better or worse: by fleet size and by workload."),
        ("5", "What it costs, how much allowance each arm spent and why each result was accepted."),
        ("6 and 7", "Search speed, convergence, and where the CPU goes under each measure."),
        ("8 and 9", "Conclusions, and what should happen next."),
        ("Appendices", "Every case at both budgets, every profile by fleet, definitions and sources.")], [1.45, 5.45])

    heading("1 What was measured")
    p("The daily search runs in two phases. The cost phase looks for the cheapest plan that breaks no rule; its result is the reference plan. The fairness phase starts from the reference and looks for a plan that spreads the work more evenly, at no more than the reference's cost plus the fairness allowance, two percent here. Its objective is the fairness measure, then cost. [1, 2]")
    p("Both arms run the same cost phase from the same day. They differ only in the fairness objective and in the measure their result is accepted on: production's rule, with the arm's own measure in place of the variance. Every final plan is then read on all five measures, so each arm is judged on the other's measure as well as its own.")
    figure("study-design", "How one case runs. Both arms share the day, the cost phase's budget and the cost ceiling; the fairness phase searches on the arm's own measure, the result is accepted on it, and the final plan is read on all five.", width=6.6)
    table(["Arm", "Variant", "Fairness objective", "Accepted on"], [
        ("b", "TABU_KOPT", "Production: the capacity-weighted variance of utilization, in binary64 units of 1e-9", "Variance"),
        ("load balance", "TABU_KOPT_LOAD_BALANCE", "Timefold's loadBalance collector over each technician's utilization in millionths, unweighted (docs 13.4.5.9, chapter 15)", "Load balance")], [1.0, 2.1, 2.8, 1.0], size=9)
    table(["Dimension", "Values in this study"], [
        ("Days", f"{', '.join(manifest['sizes'])} technicians; workloads {', '.join(workload_name(w) for w in manifest['workloads'])}; fixture seeds {', '.join(manifest['fixtureSeeds'])} (five drawn days of each shape), solver seed {', '.join(manifest['seeds'])}: {n} cases, each a different problem"),
        ("Budgets", f"{seconds(short_key)} (the production default: two thirds cost, one third fairness) and {seconds(long_key)}, to see whether the result holds when the fairness search has room"),
        ("Repetitions", f"{manifest['repetitions']}: the rule is read across cases, so time went to more distinct days rather than reruns of the same one"),
        ("Search", "TABU_KOPT at NO_ASSERT, the production move selectors; fairness allowance 2 percent"),
        ("Machine", f"One {manifest['processors']}-processor workstation; each case its own JVM, both arms in it with their order alternating, limited to 2 processors, 1.5 GiB heap and the serial collector"),
        ("Tree", f"One frozen build of {manifest['revision'][:7]}: main 52339c8 (the long score, long money, cached times and road numbers on visits) plus the Mixed Shift fixture, a test-only change"),
        ("Profiles", f"A flight recording of both arms' phases at {seconds(short_key)} on fixture seed {', '.join(manifest['profileFixtureSeeds'])}, one per phase per solve")], [1.3, 5.6])
    sub("The five fairness measures")
    p("Each measure counts the technicians production's variance counts: those qualified for the day's demand, with regular capacity. Lower is fairer on every one.")
    table(["Measure", "How to interpret it"], [
        ("Variance", "Production's measure: the capacity-weighted mean of the squared distance between each technician's utilization and the fleet's. Shown in thousandths."),
        ("Highest utilization", "The busiest technician's share of regular hours."),
        ("Utilization spread", "The busiest technician's utilization minus the least busy one's."),
        ("Absolute deviation", "The capacity-weighted mean distance of each utilization from the fleet's, without squaring."),
        ("Load balance", "Timefold's loadBalance unfairness over each utilization in millionths, every technician counted once.")], [1.55, 5.35])
    sub("The rule, written before the run")
    p("The owner set the rule before the run started. Load balance replaces b only if every part holds at both budgets. Speed and convergence are reported, not judged.")
    table(["Part", "Test"], [
        ("Variance", f"The 95 percent bootstrap interval of the median per-case variance ratio lies below 1, and at most {pct(LESS_FAIR_ALLOWED)} of cases ({allowed} of {n}) are less fair than b by more than {pct(TOLERANCE)}."),
        ("Highest utilization", "The median per-case ratio is at most 1: the busiest technician is no busier."),
        ("Cost", f"The median per-case cost difference is at most {pct(COST_ALLOWANCE, 1)} of b's cost."),
        ("Ceiling", "Every accepted result of both arms is inside its cost ceiling, with zero violations.")], [1.5, 5.4])
    p("The bootstrap interval answers whether a median ratio below 1 would survive a different set of days: it resamples the cases 10,000 times with a fixed seed and takes the middle 95 percent of the resampled medians. A single repetition cannot measure a case's own noise, so a fixed tolerance takes the place of the spread between repetitions the earlier studies used.")

    heading("2 Why the two measures can disagree")
    p("The variance and load balance both reward a plan where utilizations sit close together, but they count technicians differently. The variance weights each technician's squared gap by their regular minutes, so a 9-hour technician's gap counts 2.25 times a 4-hour technician's, and the fleet's mean is total paid minutes over total capacity. Load balance counts every technician once, whatever their hours. On a day where every technician works the same hours, the weights are all equal and the two differ only in how they combine the gaps. On a day with mixed shifts they can prefer different plans.")
    figure("weighting", "How each measure counts a technician. Bar length is how much one technician's distance from the fleet counts. Production's variance counts a 9-hour day 2.25 times a 4-hour day; load balance counts each once.", width=6.6)
    p("Every earlier benchmark fixture gave every technician a 540-minute day, so no earlier study could see this difference. This study adds a Mixed Shift workload: technicians on 4, 6 and 9 hour days by rank, with their loads scaled to the shift. The other five workloads keep nine-hour days.")
    table(["Workload", *[f"Variance ratio, {seconds(key)}" for key in keys], *[f"Load balance ratio, {seconds(key)}" for key in keys]],
          [[workload_name(w), *[fmt(ratio_median(key, "variance", keep_workload(w))) for key in keys], *[fmt(ratio_median(key, "loadBalance", keep_workload(w))) for key in keys]]
           for w in workloads], [1.7, 1.3, 1.3, 1.3, 1.3])
    p(f"Each cell is the median over the workload's {n // len(workloads)} cases of load balance's final value over b's, on production's measure and on load balance's own. Below 1 on the left means load balance's plans are fairer by production's definition; below 1 on the right, fairer by its own.")
    p(f"Mixed Shift is where the two measures part. On its {len(mixed_cases)} days load balance's plans were less fair than b's by production's measure on {lean(short_key, 'variance', lambda c: c[1] == mixed)[1]} days at {seconds(short_key)} and {lean(long_key, 'variance', lambda c: c[1] == mixed)[1]} at {seconds(long_key)}, with median ratios of {fmt(ratio_median(short_key, 'variance', keep_workload(mixed)))} and {fmt(ratio_median(long_key, 'variance', keep_workload(mixed)))}, while on load balance's own measure they were slightly fairer ({fmt(ratio_median(short_key, 'loadBalance', keep_workload(mixed)))} and {fmt(ratio_median(long_key, 'loadBalance', keep_workload(mixed)))}). The likely reason is the weighting: load balance evens out utilization counting every technician once, so it accepts a larger gap on a 9-hour technician to close one on a 4-hour technician, and the variance, which counts the 9-hour gap more than twice as heavily, reads the result as less fair. On the five workloads with equal days the median ratio is 1.00 or close to it on both measures.")
    figure("cross", f"Each case at {seconds(long_key)}: load balance's plan against b's on load balance's own measure (across) and on production's variance (up), log scales. Lower left is fairer on both; upper left is fairer by load balance's definition but less fair by production's; the colors are the workloads.")

    heading("3 Load balance against the rule")
    table(["Part, budget", "Reading", "Test", "Holds"],
          [row for key in keys for row in (
              [f"Variance, {seconds(key)}", f"median {span(m(key, 'variance'))}; {m(key, 'variance')['higherBeyondTolerance']} of {n} less fair by more than {pct(TOLERANCE)}",
               f"interval below 1; at most {allowed} cases", yes(budgets[key]["rule"]["variance"])],
              [f"Highest utilization, {seconds(key)}", f"median {span(m(key, 'maximumUtilization'))}", "median at most 1", yes(budgets[key]["rule"]["maximum"])],
              [f"Cost, {seconds(key)}", f"median {cents(budgets[key]['cost']['medianDifferenceCents'])} c, {pct(budgets[key]['cost']['medianRelativeDifference'], 3)} of b's cost",
               f"at most {pct(COST_ALLOWANCE, 1)}", yes(budgets[key]["rule"]["cost"])],
              [f"Ceiling, {seconds(key)}", f"{sum(budgets[key]['violations'].values())} violations, {len(budgets[key]['outsideCeiling'])} accepted results over a ceiling", "none", yes(budgets[key]["rule"]["ceiling"])])],
          [1.6, 3.0, 1.5, .8], size=9)
    figure("ratios", "Each case's final variance under load balance over b's, sorted, on a log scale. The black line is the median, the grey band its 95 percent interval, the dashed line the 5 percent tolerance. Points are colored by workload. A very large ratio is a day where b came close to a perfectly even plan, so a small absolute gap divides into a large number; the appendices give the variances themselves.")
    for key in keys:
        reading = m(key, "variance")
        p(f"At {seconds(key)}, load balance's variance was lower than b's in {reading['lower']} of {n} cases and higher in {reading['higher']}, {reading['higherBeyondTolerance']} of them by more than {pct(TOLERANCE)}. Its median ratio is {span(reading)}.")
    figure("measures", "On each measure and budget, how many of the 120 days load balance's final plan was fairer than b's (left of the line) or less fair (right), split at the 5 percent tolerance; ties are centered on the line. Every median ratio and interval is exactly 1.00, so the lean is the reading.")
    for key in keys:
        p(f"At {seconds(key)}: " + "; ".join(f"{MEASURE_NAME[name].lower()} {span(m(key, name))}, lower in {m(key, name)['lower']} and higher in {m(key, name)['higher']} cases" for name in MEASURES) + ".")
    sub("How often the two arms reach the same plan")
    table(["Budget", "Same reference plan", "Same final plan", "b's final plan perfectly even", "Load balance's perfectly even"],
          [[seconds(key), *[sum(1 for c in all_cases if same_plan(by[key], c, repetitions, stage)) for stage in ("reference", "final")],
            *[sum(1 for c in all_cases if perfectly_even(by[key], c, repetitions, arm)) for arm in ARMS]] for key in keys],
          [1.0, 1.4, 1.3, 1.6, 1.6])
    p(f"A plan is the same when both arms ended on the same cost and the same value on all five measures; perfectly even means a variance of zero, every counted technician at the fleet's utilization. These counts are derived from the recorded runs, not recorded by the harness. A median ratio of exactly 1 follows from them: when a large share of days ends on the same plan under both measures, the middle case is one of them.")
    earlier = earlier_ratio()
    p(f"The road matrix study is often quoted as load balance at 0.58 of production's variance against b's 0.74. Both figures were medians against the old decimal production, not against each other, and two medians against a third arm do not give the ratio between the two. Read the way this study reads it, each case's load balance variance over b's, the road matrix study's median is {earlier['median']:.2f} over its {earlier['ratios']} cases with a nonzero variance ({', '.join(map(str, earlier['fleets']))} technicians, four workloads, one drawn day, two repetitions), with {earlier['same']} of {earlier['cases']} cases on the same plan. [4]")

    heading("4 By fleet size and workload")
    figure("grid", "The median variance ratio, load balance over b, in each workload and fleet size (five days per cell). Teal is fairer under load balance, red less fair.")
    for key in keys:
        sub(f"At {seconds(key)}")
        table(["Group", "Variance ratio", "Highest utilization ratio", "Median cost difference", "Cases"],
              [[f"{f} technicians", fmt(ratio_median(key, "variance", keep_fleet(f))), fmt(ratio_median(key, "maximumUtilization", keep_fleet(f))),
                f"{cents(median(cost_differences(by[key], all_cases, repetitions, keep_fleet(f))))} c", sum(1 for c in all_cases if c[0] == f)] for f in fleets] +
              [[workload_name(w), fmt(ratio_median(key, "variance", keep_workload(w))), fmt(ratio_median(key, "maximumUtilization", keep_workload(w))),
                f"{cents(median(cost_differences(by[key], all_cases, repetitions, keep_workload(w))))} c", sum(1 for c in all_cases if c[1] == w)] for w in workloads],
              [1.6, 1.9, 1.3, 1.3, .8], size=9)

    heading("5 What it costs and why results were accepted")
    figure("cost", "Each case's final cost under load balance minus b's, in dollars, by fleet size; pale at the short budget, solid at the long. The black bar is the median.")
    for key in keys:
        c = budgets[key]["cost"]
        p(f"At {seconds(key)} load balance ended cheaper than b in {c['cheaper']} cases, the same in {c['ties']} and dearer in {c['dearer']}; the median difference is {cents(c['medianDifferenceCents'])} cents and the mean {cents(c['meanDifferenceCents'])}.")
    sub("The starting plan")
    p(f"Both arms run the same cost objective in the cost phase, so any difference in the reference plan comes from how many moves each phase evaluated in its share of the budget. The reference plans were the same on {same_reference[short_key]} and {same_reference[long_key]} of the {n} days. Where they differ, load balance's was dearer on {sum(g > 0 for g in reference_gap[short_key])} days and cheaper on {sum(g < 0 for g in reference_gap[short_key])} at {seconds(short_key)}, and dearer on {sum(g > 0 for g in reference_gap[long_key])} and cheaper on {sum(g < 0 for g in reference_gap[long_key])} at {seconds(long_key)}. This is derived from the recorded reference costs. Because the cost ceiling is set from the reference plan, a dearer start carries through to a dearer final plan: at {fleets[-1]} technicians load balance's final plans cost a mean {cents(round(fleet_mean_cost(short_key, fleets[-1]), 1))} cents a day more than b's at {seconds(short_key)} and {cents(round(fleet_mean_cost(long_key, fleets[-1]), 1))} at {seconds(long_key)}, while the median day costs the same.")
    figure("allowance", "Left: the share of the 2 percent allowance each run spent above its own reference plan; the black bar is the median. Right: how each arm's results were accepted, fairness improvements darkest.")
    reasons = sorted({r for key in keys for arm in ARMS for r in budgets[key]["acceptance"][arm]})
    table(["Arm, budget", "Allowance spent (median)", *[r.replace("_", " ").capitalize() for r in reasons]],
          [[f"{LABEL[arm]}, {seconds(key)}", pct(median([u for u in (allowance_used(r) for r in budgets[key]["runs"] if r["arm"] == arm) if u is not None])),
            *[budgets[key]["acceptance"][arm].get(r, 0) for r in reasons]] for key in keys for arm in ARMS],
          [1.9, 1.5, *([3.5 / len(reasons)] * len(reasons))])

    heading("6 Search speed and convergence")
    p("The two arms differ in when their fairness work runs. Production's constraint filters on the fairness target before it computes the variance, so the variance is computed only in the fairness phase. Load balance's constraint groups the routes into Timefold's loadBalance collector first and filters on the target after, so the collector is updated on every move of both phases, even though its penalty only counts once the fairness phase has a target. [2] The table gives the median of paired ratios, load balance's moves per second over b's in the same case.")
    table(["Fleet", *[f"{PHASE_NAME[phase].capitalize()}, {seconds(key)}" for key in keys for phase in PHASES]],
          [[f"{f} technicians", *[f"{median(speedups(by[key], all_cases, repetitions, phase, keep_fleet(f))):.2f}x" for key in keys for phase in PHASES]] for f in fleets] +
          [["All", *[f"{budgets[key]['speed'][phase]['all']['median']:.2f}x" for key in keys for phase in PHASES]]],
          [1.3, *([5.6 / (2 * len(keys))] * (2 * len(keys)))])
    figure("speed", "Moves per second of load balance over b in the same case, in each phase, by fleet size and budget. The black bar is the median.")
    p(f"In absolute terms b's cost phase evaluated a median {rate(short_key, 'b', 'referenceMovesPerSecond'):,.0f} moves per second at {seconds(short_key)} and load balance's {rate(short_key, 'load-balance', 'referenceMovesPerSecond'):,.0f}; these are the harness's rates, not the service's. Fewer moves in the same seconds means a less thorough cost search, which shows in the reference plan's cost before any fairness is spent.")
    table(["Arm, phase", *[f"Median share of budget to best, {seconds(key)}" for key in keys], *[f"Runs past 90%, {seconds(key)}" for key in keys]],
          [[f"{LABEL[arm]}, {PHASE_NAME[phase]}", *[pct(budgets[key]["convergence"][arm][phase]["medianShareOfBudget"]) if budgets[key]["convergence"][arm][phase]["medianShareOfBudget"] is not None else "n/a" for key in keys],
            *[budgets[key]["convergence"][arm][phase]["overNinetyPercent"] for key in keys]] for arm in ARMS for phase in PHASES],
          [2.1, *([4.8 / (2 * len(keys))] * (2 * len(keys)))])
    figure("convergence", "How far into each phase's budget each arm found its best plan, the median over cases. A share near 100 percent means the search was still improving when time ran out.")

    heading("7 Where the time goes")
    if profiles:
        names = [f"{arm} {phase}" for phase in ("cost", "fairness") for arm in ARMS if f"{arm} {phase}" in profiles]
        figure("profile", f"Share of CPU samples in each phase by kind of work, from the flight recordings at {seconds(short_key)}. Score engine is Timefold's constraint streams, where the load balance collector runs.")
        table(["Kind of work", *[n.replace("load-balance", "load balance").replace(" cost", ", cost").replace(" fairness", ", fairness") for n in names]],
              [[group, *[f"{profiles[name]['all']['samplesPercent'].get(group, 0):.1f}%" for name in names]] for group in PROFILE_GROUPS
               if any(profiles[name]["all"]["samplesPercent"].get(group, 0) > 0 for name in names)],
              [1.9, *([5.0 / len(names)] * len(names))])
        leaves_lb = profiles["load-balance cost"]["all"]["leafFramesPercent"]
        leaves_b = profiles["b cost"]["all"]["leafFramesPercent"]
        table(["Leaf frame, load balance's cost phase", "Load balance", "b"],
              [[frame_name(frame), f"{share:.1f}%", f"{leaves_b[frame]:.1f}%" if frame in leaves_b else "below top 20"]
               for frame, share in list(leaves_lb.items())[:10]], [4.3, 1.3, 1.3], size=9)
        p("Each leaf frame is the method running when a sample was taken, as a share of the phase's samples; b's share is shown where the frame is among b's top twenty.")
        share = lambda name, group: profiles[name]["all"]["samplesPercent"].get(group, 0)
        p(f"The two cost phases do the same search, so their profiles differ only by the collector. Collections take {share('load-balance cost', 'collections'):.1f} percent of load balance's cost-phase samples against {share('b cost', 'collections'):.1f} percent of b's, and the score engine {share('load-balance cost', 'score engine'):.1f} against {share('b cost', 'score engine'):.1f}; the rest of the profile, led by PlanVisit.roadNumber and the route walk, is in the same order in both. Reading the extra collections time as the collector's own per-technician map is an interpretation: the recordings keep leaf frames, not the callers that would confirm it. In the fairness phase, load balance shows {share('load-balance fairness', 'decimal: score impacts'):.1f} percent in decimal score work that b does not have, most likely the collector's unfairness, a BigDecimal, converted to the long score on each move.")

    heading("8 Conclusions")
    rows = [
        ("Is load balance fairer than the variance by production's own measure?",
         "No. " + " ".join(f"At {seconds(key)} its median ratio is {span(m(key, 'variance'))}; it was fairer on {m(key, 'variance')['lower']} days and less fair on {m(key, 'variance')['higher']}, {m(key, 'variance')['higherBeyondTolerance']} of them by more than {pct(TOLERANCE)}." for key in keys)),
        ("Does it make the busiest day lighter?", "No. " + " ".join(f"At {seconds(key)} the median ratio is {span(m(key, 'maximumUtilization'))}, lighter on {m(key, 'maximumUtilization')['lower']} days and heavier on {m(key, 'maximumUtilization')['higher']}." for key in keys)),
        ("What does it cost?", " ".join(f"At {seconds(key)} a median {cents(budgets[key]['cost']['medianDifferenceCents'])} cents a day and a mean {cents(budgets[key]['cost']['meanDifferenceCents'])}." for key in keys) + f" The mean comes from {fleets[-1]}-technician days, where the slower cost phase leaves a dearer starting plan."),
        ("Does it slow the search?", "Yes. " + " ".join(f"At {seconds(key)} its cost phase runs at {budgets[key]['speed']['referencePhase']['all']['median']:.2f}x b's and its fairness phase at {budgets[key]['speed']['fairnessPhase']['all']['median']:.2f}x." for key in keys)),
        ("Why is the median ratio exactly 1?", f"The two arms ended on the same plan on {same_final[short_key]} and {same_final[long_key]} of {n} days; the middle day is one of those."),
        ("Where do the measures disagree?", f"On mixed shifts: load balance's plans are less fair by production's variance ({fmt(ratio_median(short_key, 'variance', keep_workload(mixed)))} and {fmt(ratio_median(long_key, 'variance', keep_workload(mixed)))}) and slightly fairer by its own ({fmt(ratio_median(short_key, 'loadBalance', keep_workload(mixed)))} and {fmt(ratio_median(long_key, 'loadBalance', keep_workload(mixed)))})."),
        ("Does the rule support switching?", ("Yes, at both budgets." if switch else "No. " + "; ".join(f"At {seconds(key)} it fails on {', '.join(rule_name[r] for r in failed[key])}" for key in keys if failed[key]) + "."))]
    table(["Question", "Answer from the evidence"], rows, [2.4, 4.5])
    p("Two cautions carry over. The days are synthetic and at most 50 technicians, and each solve was limited to two processors on a workstation, so the absolute rates are the harness's; the paired ratios are what transfers. And fairer is fairer by a measure's definition: whether planners see the difference is a question for the pilot.")

    heading("9 What should happen next")
    if switch:
        table(["Next step", "What to do", "What it decides"], [
            ("Owner decision", "Decide whether to make load balance the daily fairness policy on this evidence.", "The policy."),
            ("Stack verification", "If yes, verify on the VM rehearsal stack through the public API, with shipped routing and admission defaults, before calling it shipped.", "Whether the result holds at the service."),
            ("Acceptance and contract", "Move acceptance and the numeric contract in scheduler-policy.md to the load balance reading, and gate with ScoringExactnessTest and the FULL_ASSERT campaign.", "A consistent policy end to end.")], [1.45, 3.7, 1.75])
    else:
        table(["Next step", "What to do", "What it decides"], [
            ("Keep the variance", "Make no policy change. Record in scheduler-policy.md that load balance was measured head to head and why it was not chosen.", "The owner's policy decision, on this evidence."),
            ("Remove the study arm", "Remove TABU_KOPT_LOAD_BALANCE and FairnessMeasures.LoadBalance from main with file edits, keeping them on this branch.", "A smaller solver surface."),
            ("Keep Mixed Shift", "Keep the Mixed Shift fixture in the benchmark set so later fairness and speed studies see mixed capacities.", "Coverage of part-time technicians."),
            ("If load balance returns", "Filter on the fairness target before the loadBalance collector, as production's variance does, so the cost phase does not pay for it; and test a capacity-weighted load balance on Mixed Shift days.", "Whether the slower cost phase and the weighting are what held it back.")], [1.45, 3.7, 1.75])
    sub("How the calculations were checked")
    p("Every table and chart value is computed from the evidence file at build time; nothing is typed in. report_data.py checks the evidence before use: it is the full run on the pinned revision with the planned matrix; every arm, case and budget is present exactly once; no run has a violation or an accepted result over its ceiling. It then recomputes from the individual runs every reading the rule rests on, the bootstrap intervals included, and requires them to equal the recorded ones, and recomputes each part of the rule and the decision. verify_report.py recomputes the tables a second time from the raw JSON without the builder's helpers, and checks the figure captions, the template parts and the absence of dashes. [3]")

    for key, letter in zip(keys, "AB"):
        heading(f"Appendix {letter} Every case at {seconds(key)}")
        p("Each row is one day: both arms' final variance in thousandths, load balance's over b's, both arms' highest utilization, load balance's cost minus b's in cents, and how each result was accepted (F fairness improvement, C cost reduction, O overtime reduction, R rejected). [3]")
        code = {"FAIRNESS_IMPROVEMENT": "F", "COST_REDUCTION": "C", "OVERTIME_REDUCTION": "O"}
        table(["Techs", "Workload", "Day", "b variance", "Load bal. variance", "Ratio", "b highest", "Load bal. highest", "Cost difference", "Accepted b / LB"],
              [[c[0], workload_name(c[1]), c[2], milli(case_mean(by[key], "b", c, "variance", repetitions)), milli(case_mean(by[key], "load-balance", c, "variance", repetitions)),
                f"{float(case_mean(by[key], 'load-balance', c, 'variance', repetitions) / case_mean(by[key], 'b', c, 'variance', repetitions)):.2f}" if case_mean(by[key], "b", c, "variance", repetitions) > 0 else "n/a",
                f"{float(case_mean(by[key], 'b', c, 'maximumUtilization', repetitions)):.3f}", f"{float(case_mean(by[key], 'load-balance', c, 'maximumUtilization', repetitions)):.3f}",
                cents(float(case_mean(by[key], "load-balance", c, "costCents", repetitions) - case_mean(by[key], "b", c, "costCents", repetitions))),
                " / ".join("".join(code.get(by[key][(arm, c, r)]["acceptanceReason"], "R") for r in repetitions) for arm in ARMS)]
               for c in all_cases], [.45, 1.0, .45, .7, .75, .5, .7, .75, .8, .8], size=8.5)

    if profiles:
        heading("Appendix C Every profile by fleet")
        p("Each column is one arm and phase at one fleet size: the share of its CPU samples by kind of work. A blank is a kind with no samples. [3]")
        for arm in ARMS:
            sub(LABEL[arm].capitalize())
            columns = [(phase, f) for phase in ("cost", "fairness") for f in fleets if str(f) in profiles.get(f"{arm} {phase}", {}).get("byFleet", {})]
            table(["Kind of work", *[f"{phase}, {f}" for phase, f in columns]],
                  [[group, *[f"{profiles[f'{arm} {phase}']['byFleet'][str(f)]['samplesPercent'][group]:.1f}" if group in profiles[f"{arm} {phase}"]["byFleet"][str(f)]["samplesPercent"] else "" for phase, f in columns]]
                   for group in PROFILE_GROUPS] +
                  [["Recordings", *[profiles[f"{arm} {phase}"]["byFleet"][str(f)]["recordings"] for phase, f in columns]]],
                  [1.7, *([5.2 / len(columns)] * len(columns))], size=9)

    heading("Appendix D Definitions and sources")
    table(["Term", "Meaning in this report"], [
        ("Arm", "One fairness objective under test, run as its own solver variant in every case."),
        ("Case, day", "One fleet size, workload and fixture seed: a drawn day of that shape. Every case is a different problem."),
        ("Reference plan", "The plan the cost phase ends with; the fairness phase starts from it."),
        ("Allowance and ceiling", "The fairness phase may spend up to 2 percent above the reference plan's cost; the ceiling is that cost times 1.02, rounded down to a cent."),
        ("Utilization", "A technician's paid minutes divided by their regular capacity minutes for the day."),
        ("Ratio to b", "Load balance's final value on a measure divided by b's in the same case; the report gives the median over cases."),
        ("Bootstrap interval", "The middle 95 percent of the medians of 10,000 resamples of the cases, drawn with a fixed seed."),
        ("Moves per second", "Candidate moves a phase evaluated per second; more moves in the same budget is a more thorough search."),
        ("Share of budget to best", "The time a phase took to find its best plan over the phase's budget."),
        ("Leaf frame", "The method at the top of a sampled stack in a flight recording."),
        ("Thousandths", "Variances in this report are shown multiplied by 1000.")], [1.6, 5.3])
    p("[1] docs/scheduler-policy.md: the fairness phase, the allowance, the acceptance rule and the numeric contract.")
    p("[2] The constraint providers: calculation-engine/src/main/java/dev/waterflex/scheduler/optimizer, DayConstraintProvider (production) and FairnessMeasures.LoadBalance; Timefold Solver documentation 13.4.5.9 (the loadBalance collector) and chapter 15 (load balancing and fairness).")
    p("[3] Evidence: docs/evidence/solver-load-balance-2026-10-23.json, written by infra/load_balance_study.py from the run infra/run-load-balance-study.ps1 made (manifest, every run, both budgets' readings, the rule, the decision and the profiles). docs/solver-benchmarks.md, 'Load balance head to head', is the narrative kept with the code.")
    p("[4] The road matrix study: docs/evidence/solver-matrix-2026-10-20.json (its runs block), and docs/solver-benchmarks.md, 'Road matrix numbering (PR 187)'.")
    p("Reproduce with docs/load-balance-2026-10-23/requirements.txt: render the two diagrams as EVIDENCE.md says, then run report_data.py, build_report.py and verify_report.py there. The builder regenerates the charts and chart_data.json and writes this editable Word file.")

    settings = doc.settings.element
    for el in settings.findall(qn("w:updateFields")): settings.remove(el)
    update = OxmlElement("w:updateFields"); update.set(qn("w:val"), "true"); settings.append(update)
    stream = BytesIO(); doc.save(stream)
    with ZipFile(REFERENCE) as ref, ZipFile(stream) as built, ZipFile(output, "w", ZIP_DEFLATED) as archive:
        preserve = set(PRESERVE) | {n for n in ref.namelist() if "footer" in n}
        for name in sorted(built.namelist()):
            content = ref.read(name) if name in preserve else built.read(name)
            info = ZipInfo(name, (2026, 10, 23, 0, 0, 0)); info.compress_type = ZIP_DEFLATED
            archive.writestr(info, content)
    print(output)


if __name__ == "__main__": build()
