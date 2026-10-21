"""Check the delivered tables, chart data and template parts against the raw evidence file.

Recomputes from docs/evidence/api-load-2026-10-20-binary-variance.json without the builder's or loader's helpers. Run
normally, without Python's -O flag. Visual pagination still requires rendering.
"""
import json
import math
from decimal import Decimal
from hashlib import sha256
from statistics import median
from zipfile import ZipFile

from docx import Document
from build_report import HERE, OUTPUT, PRESERVE, REFERENCE, REFERENCE_SHA

EVIDENCE = HERE.parent / "evidence" / "api-load-2026-10-20-binary-variance.json"
FLEETS, SEEDS = (10, 20, 50), (1, 2, 3)


def nearest(values):
    ordered = sorted(values)
    return ordered[math.ceil(len(ordered) / 2) - 1]


def k(value):
    return f"{value / 1000:,.1f}k"


def verify():
    raw = json.loads(EVIDENCE.read_text(encoding="utf-8-sig"))
    doc = Document(OUTPUT)
    assert len(doc.tables) == 15 and len(doc.inline_shapes) == 9 and len(doc.sections) == 1, (len(doc.tables), len(doc.inline_shapes))
    assert len([p for p in doc.paragraphs if p.style.name == "Heading 1"]) == 11
    captions = [p for p in doc.paragraphs if p.style.name == "Caption"]
    assert len(captions) == 9
    for i, p in enumerate(captions, 1):
        assert p.text.startswith(f"Figure {i}.") and all(str(r.font.color.rgb) == "000000" and r.italic for r in p.runs)
    for shape in doc.inline_shapes:
        assert shape._inline.docPr.get("descr", "").startswith("Figure ")

    def rows(index):
        return [[c.text for c in r.cells] for r in doc.tables[index].rows[1:]]

    def check(index, expected):
        actual = rows(index)
        assert actual == [[str(c) for c in r] for r in expected], (index, actual[:2], expected[:2])

    runs = {(r["step"], r["image"], r["fleet"], r["seed"]): r["report"] for r in raw["runs"]}
    assert len(runs) == 20

    def receipts(step, image, fleet, seed):
        return runs[(step, image, fleet, seed)]["study"]["receipts"]

    def speed(step, image, fleet, seed, phase):
        return nearest([r["solves"][phase]["movesPerSecond"] for r in receipts(step, image, fleet, seed)])

    def ratio(step, fleet, seed, phase):
        return speed(step, "after", fleet, seed, phase) / speed(step, "before", fleet, seed, phase)

    def share(keys):
        return median([r["solves"][1]["movesPerSecond"] / r["solves"][0]["movesPerSecond"] for key in keys for r in receipts(*key)])

    def fairer(keys):
        return sum(Decimal(r["body"]["proposed"]["variance"]) < Decimal(r["body"]["reference"]["variance"]) for key in keys for r in receipts(*key))

    def call(key):
        return [o for o in runs[key]["operations"] if o["operation"] == "daily proposal"][0]["ms"]

    def daily_cpu(key):
        return [c for c in runs[key]["containers"] if c["name"].startswith("scheduler-vm-scheduler-daily")][0]["cpuPercent"]

    def keys(image, fleet):
        return [(1, image, fleet, seed) for seed in SEEDS]

    queued = {image: (2, image, 20, 4) for image in ("before", "after")}

    # Table 0: what the measurements tell us.
    expected = []
    for fleet in FLEETS:
        fair = [ratio(1, fleet, seed, 1) for seed in SEEDS]
        expected.append([f"{fleet} technicians", f"{median(fair):.2f}x ({min(fair):.2f} to {max(fair):.2f})", f"{median([ratio(1, fleet, s, 0) for s in SEEDS]):.2f}x",
                         f"{share(keys('before', fleet)):.2f} / {share(keys('after', fleet)):.2f}", f"{fairer(keys('before', fleet))} / {fairer(keys('after', fleet))} of 12",
                         f"{median([call(key)['p50'] for key in keys('before', fleet)]) / 1000:.1f} s / {median([call(key)['p50'] for key in keys('after', fleet)]) / 1000:.1f} s"])
    expected.append(["20, 2 planners queuing", f"{ratio(2, 20, 4, 1):.2f}x (one run)", f"{ratio(2, 20, 4, 0):.2f}x", f"{share([queued['before']]):.2f} / {share([queued['after']]):.2f}",
                     f"{fairer([queued['before']])} / {fairer([queued['after']])} of 12", f"{call(queued['before'])['p50'] / 1000:.1f} s / {call(queued['after'])['p50'] / 1000:.1f} s"])
    check(0, expected)

    # Table 4: the fairness phase as a share of the reference phase.
    check(4, [[f"{fleet} technicians", f"{share(keys('before', fleet)):.2f}", f"{1 / share(keys('before', fleet)):.1f} times", f"{share(keys('after', fleet)):.2f}",
               f"{1 / share(keys('after', fleet)):.2f} times"] for fleet in FLEETS])

    # Table 5: every run's phase speeds and ratios.
    expected = [[f"{fleet}, {seed}", k(speed(1, "before", fleet, seed, 0)), k(speed(1, "after", fleet, seed, 0)), f"{ratio(1, fleet, seed, 0):.2f}x",
                 k(speed(1, "before", fleet, seed, 1)), k(speed(1, "after", fleet, seed, 1)), f"{ratio(1, fleet, seed, 1):.2f}x"] for fleet in FLEETS for seed in SEEDS]
    expected.append(["20, 4, queued", k(speed(2, "before", 20, 4, 0)), k(speed(2, "after", 20, 4, 0)), f"{ratio(2, 20, 4, 0):.2f}x",
                     k(speed(2, "before", 20, 4, 1)), k(speed(2, "after", 20, 4, 1)), f"{ratio(2, 20, 4, 1):.2f}x"])
    check(5, expected)

    # Table 6: each pair's plans against their reference plans.
    def pair(image, fleet, seed, field):
        return [r["body"][side][field] for r in receipts(1, image, fleet, seed) for side in ("reference", "proposed")]

    expected = []
    for fleet in FLEETS:
        for seed in SEEDS:
            cells = [f"{fleet}, {seed}"]
            cells.append(" / ".join(f"{nearest([r['body']['reference']['costCents'] for r in receipts(1, image, fleet, seed)]):,}" for image in ("before", "after")))
            cells.append(" / ".join(f"{Decimal(sorted([r['body']['proposed']['variance'] for r in receipts(1, image, fleet, seed)], key=Decimal)[1]):.4f}" for image in ("before", "after")))
            diffs = [nearest([r["body"]["proposed"]["costCents"] - r["body"]["reference"]["costCents"] for r in receipts(1, image, fleet, seed)]) for image in ("before", "after")]
            cells.append(" / ".join(f"{d:+,}" if d else "0" for d in diffs))
            cells.append(f"{fairer([(1, 'before', fleet, seed)])} / {fairer([(1, 'after', fleet, seed)])} of 4")
            expected.append(cells)
    check(6, expected)

    # Table 7: the service by fleet.
    check(7, [[f"{fleet} technicians", " / ".join(f"{median([call(key)['p50'] for key in keys(image, fleet)]) / 1000:.1f}" for image in ("before", "after")),
               " / ".join(f"{max(call(key)['max'] for key in keys(image, fleet)) / 1000:.1f}" for image in ("before", "after")),
               " / ".join(f"{sum(runs[key]['server']['routingPairs'] for key in keys(image, fleet)):,}" for image in ("before", "after")),
               " / ".join(f"{median([daily_cpu(key)['mean'] for key in keys(image, fleet)]):.0f}%" for image in ("before", "after"))] for fleet in FLEETS])

    # Table 8: the queued step.
    for row in rows(8):
        if row[0].startswith("Queue wait"):
            assert row[1:] == [f"{nearest(runs[queued[i]]['daily']['queueWaitsMs'])} / {max(runs[queued[i]]['daily']['queueWaitsMs'])}" for i in ("before", "after")], row
        if row[0].startswith("Road pairs"):
            assert row[1:] == [f"{runs[queued[i]]['server']['routingPairs']:,}" for i in ("before", "after")], row

    # Appendix A: every proposal, in run order and log order.
    order = sorted(runs, key=lambda key: (key[0], key[2], key[3], ("before", "after").index(key[1])))
    expected = [[step, image, fleet, seed, i, f"{r['solves'][0]['movesPerSecond']:,}", f"{r['solves'][1]['movesPerSecond']:,}", f"{r['body']['reference']['costCents']:,}",
                 f"{r['body']['proposed']['costCents']:,}", f"{Decimal(r['body']['reference']['variance']):.4f}", f"{Decimal(r['body']['proposed']['variance']):.4f}"]
                for (step, image, fleet, seed) in order for i, r in enumerate(receipts(step, image, fleet, seed), 1)]
    assert len(expected) == 96
    check(12, expected)

    # Appendix B: every run, by start time.
    expected = [[key[0], key[1], key[2], key[3], runs[key]["startedAt"][11:19], runs[key]["daily"]["proposals"], f"{call(key)['p50'] / 1000:.1f}", f"{call(key)['max'] / 1000:.1f}",
                 f"{runs[key]['server']['routingPairs']:,}", f"{daily_cpu(key)['mean']:.0f}%", f"{runs[key]['daily']['decisions']['IMPROVED']} IMPROVED"]
                for key in sorted(runs, key=lambda key: runs[key]["startedAt"])]
    check(13, expected)

    # Chart data agrees with the evidence.
    chart = json.loads((HERE / "generated" / "chart_data.json").read_text(encoding="utf-8"))
    for p, phase in enumerate(("REFERENCE", "FAIRNESS")):
        for image in ("before", "after"):
            for fleet in FLEETS:
                assert chart["speed"][phase][image][str(fleet)] == [r["solves"][p]["movesPerSecond"] for key in keys(image, fleet) for r in receipts(*key)]
        for fleet in FLEETS:
            assert chart["ratio"][phase][str(fleet)] == [round(ratio(1, fleet, seed, p), 4) for seed in SEEDS]
        assert chart["ratioQueued"][phase] == round(ratio(2, 20, 4, p), 4)
    for image in ("before", "after"):
        for fleet in FLEETS:
            assert chart["gap"][image][str(fleet)] == [round(r["solves"][1]["movesPerSecond"] / r["solves"][0]["movesPerSecond"], 4) for key in keys(image, fleet) for r in receipts(*key)]
        assert chart["queued"][image]["waits"] == runs[queued[image]]["daily"]["queueWaitsMs"]

    # Template parts preserved byte for byte, and no em or en dash anywhere in the document XML.
    assert sha256(REFERENCE.read_bytes()).hexdigest() == REFERENCE_SHA
    with ZipFile(REFERENCE) as ref, ZipFile(OUTPUT) as out:
        for name in set(PRESERVE) | {n for n in ref.namelist() if "footer" in n}:
            assert ref.read(name) == out.read(name), name
        for name in out.namelist():
            if name.endswith(".xml"):
                text = out.read(name).decode("utf-8")
                assert "\u2014" not in text and "\u2013" not in text, name
        body = out.read("word/document.xml").decode("utf-8")
        assert body.count("<w:cantSplit/>") == sum(len(t.rows) for t in doc.tables)
    print("verified", OUTPUT.name)


if __name__ == "__main__":
    verify()
