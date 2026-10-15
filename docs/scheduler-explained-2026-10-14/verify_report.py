"""Verify the built report against its raw sources, without the loader's or builder's helpers.

Recomputes the chart datasets from the evidence files and experiment archives directly, checks the timeline's pull
request dates against the repository's merge history, and checks the preserved template parts, figure descriptions,
table row splitting, the absence of citation markers and Appendix C, and the absence of em and en dashes.
"""
import json
import re
import subprocess
from hashlib import sha256
from pathlib import Path
from statistics import mean, median
from zipfile import ZipFile

from docx import Document

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
EVIDENCE = ROOT / "docs/evidence"
OUTPUT = HERE / "WaterFlex_Scheduler_Explained_2026-10-14.docx"
REFERENCE = HERE.parent / "WaterFlex_Scheduler_Algorithm_Comparison_Explained.docx"
PRESERVE = ("word/styles.xml", "word/stylesWithEffects.xml", "word/numbering.xml", "word/theme/theme1.xml", "word/fontTable.xml")
FIRST = ROOT / "experiments/runs/20261010T040557Z-daily-tabu-variants-p6-5d15129b"
FOLLOW = ROOT / "experiments/runs/20261013T052707Z-daily-search-followup-6454fb05"
LONG = ROOT / "experiments/runs/20261013T103709Z-daily-search-long-budget-a317757e"
OFFERS, CONFIRM = "POST /api/v1/booking/offers 201", "POST /api/v1/booking/holds/{holdId}/confirm 200"
checked = []


def read(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def dollars(value):
    return f"-${-value:,.2f}" if value < 0 else f"${value:,.2f}"


def cases(run, keyed_by_fixture_seed):
    """Every raw result of a run, read from the files the analysis provenance names, as a dict of key, variant, cost, variance and best-found point."""
    analysis = sorted((run / "analysis").iterdir())[-1]
    out = []
    for item in read(analysis / "analysis-provenance.json")["inputs"]:
        path = run / item["name"].replace("\\", "/")
        assert sha256(path.read_bytes()).hexdigest() == item["sha256"], path
        raw = [r for r in map(json.loads, path.read_text().splitlines()) if r.get("type") == "result"][0]
        key = (raw["technicians"], raw["workload"], raw.get("fixtureSeed") if keyed_by_fixture_seed else None, raw["seed"], raw["budgetMs"])
        phase = raw["referencePhase"]
        out.append({"key": key, "variant": raw["variant"], "cost": raw["after"]["costCents"], "variance": raw["after"]["fairness"]["variance"],
                    "best": phase["timeToBestMs"] / phase["solveMs"]})
    return out


def savings(rows, control):
    """(key, variant, saving in dollars against the matched control) for every non-control case."""
    twins = {c["key"]: c["cost"] for c in rows if c["variant"] == control}
    return [(c["key"], c["variant"], (twins[c["key"]] - c["cost"]) / 100) for c in rows if c["variant"] != control]


def close(a, b):
    return abs(a - b) < 1e-9


def check(name, expected, actual):
    assert expected == actual, (name, expected, actual)
    checked.append(name)


def per_minute(report):
    op = [o for o in report["operations"] if o["operation"] == "booking offers"][0]
    return op["statuses"].get("201", 0) * 60 / report["settings"]["durationS"]


def op(report, name):
    return [o for o in report["operations"] if o["operation"] == name][0]


def main():
    doc = Document(OUTPUT)
    charts = read(HERE / "generated/chart_data.json")

    # Template parts, figures, rows and dashes.
    with ZipFile(REFERENCE) as ref, ZipFile(OUTPUT) as built:
        for name in set(PRESERVE) | {n for n in ref.namelist() if "footer" in n}:
            assert ref.read(name) == built.read(name), name
        xml = built.read("word/document.xml").decode("utf-8")
    assert "—" not in xml and "–" not in xml
    order = (HERE / "generated/figures.txt").read_text(encoding="utf-8").split()
    assert len(order) == len(set(order)) and {p.stem for p in (HERE / "generated").glob("*.png")} == set(order), "every figure placed once"
    descriptions = re.findall(r'<wp:docPr [^>]*descr="([^"]*)"', xml)
    assert len(descriptions) == len(order) and all(d.startswith(f"Figure {i}. ") for i, d in enumerate(descriptions, 1)), len(descriptions)
    text = "\n".join(para.text for para in doc.paragraphs) + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert not re.search(r"\[\d", text), "no citation markers"
    assert "Appendix C" not in text and "References" not in [para.text for para in doc.paragraphs if para.style.name.startswith("Heading")]
    assert xml.count("<w:tr>") + xml.count("<w:tr ") == xml.count("<w:cantSplit/>"), "every row must refuse to split"
    checked.append(f"template, {len(order)} figures, rows, dashes, no citations")

    # The timeline's pull request dates against the merge history.
    merged = {}
    for line in subprocess.run(["git", "log", "origin/main", "--merges", "--format=%ad %s", "--date=short"], cwd=ROOT, capture_output=True, text=True, check=True).stdout.splitlines():
        found = re.match(r"(\d{4}-\d{2}-\d{2}) Merge pull request #(\d+) ", line)
        if found: merged.setdefault(int(found.group(2)), found.group(1))
    dated = [e for e in charts["timeline"] if e["pr"] is not None]
    assert len(dated) == 12 and all(merged.get(e["pr"]) == e["date"] for e in dated), [(e["pr"], e["date"], merged.get(e["pr"])) for e in dated]
    checked.append("timeline merge dates")

    # The daily paired means by budget (Sept 29).
    summary = read(EVIDENCE / "scheduler-field-2026-09-29/summary.json")
    paired = {(d["budgetSeconds"], d["variant"]): d["meanSavingsCents"] / 100 for d in summary["dailyPaired"]}
    check("budgets29 chart", {"budgets": [15, 30, 60], "values": {v: [paired[(b, v)] for b in (15, 30, 60)] for v in ("CURRENT_UNCAPPED", "KOPT", "SUBLIST", "LATE_ACCEPTANCE", "RUIN_RECREATE")}},
          charts["budgets29"])

    # The two daily studies, from the raw result files.
    raw_first, raw_follow, raw_long = cases(FIRST, False), cases(FOLLOW, True), cases(LONG, True)
    first, follow = savings(raw_first, "TABU"), savings(raw_follow, "TABU")
    assert len(first) == 4480 * 6 // 7 and len(follow) == 1200 * 4 // 5 and len(raw_long) == 192
    for key, rows in (("first", first), ("follow", follow)):
        for entry in charts["wtl"][key]:
            own = [x for _, v, x in rows if v == entry["solver"]]
            assert (entry["wins"], entry["ties"], entry["losses"]) == (sum(x > 0 for x in own), sum(x == 0 for x in own), sum(x < 0 for x in own)) and close(entry["mean"], mean(own)), entry
    assert charts["wtl"]["chosen"] == {"first": "TABU_KOPT", "follow": "TABU_KOPT_SIZE_3"}
    checked.append("wtl chart")
    late = charts["late_mechanics"]
    for solver, values in late["best_fraction"].items():
        assert all(close(v, mean(c["best"] for c in raw_first if c["variant"] == solver and c["key"][0] == f)) for v, f in zip(values, late["fleets"])), solver
    for group, raw in zip(late["variance"], (raw_first, raw_follow)):
        assert all(close(e["value"], mean(c["variance"] for c in raw if c["variant"] == e["solver"])) for e in group["rows"]), group["label"]
    checked.append("late_mechanics chart")
    arm = savings(raw_long, "TABU_KOPT")
    for fleet, by_solver in charts["long_arm"]["fleets"].items():
        for solver, values in by_solver.items():
            assert all(close(v, mean(x for k, s, x in arm if s == solver and k[0] == int(fleet) and k[4] == b * 1000)) for v, b in zip(values, charts["long_arm"]["budgets"])), (fleet, solver)
    checked.append("long_arm chart")
    for key, rows, solvers, fleets in (("first", first, charts["daily_fleet"]["first"]["solvers"], (5, 10, 20, 50)), ("follow", follow, charts["daily_fleet"]["follow"]["solvers"], (5, 10, 20, 30, 50))):
        values = [[mean(x for k, v, x in rows if v == s and k[0] == f) for f in fleets] + [mean(x for k, v, x in rows if v == s)] for s in solvers]
        assert all(abs(a - b) < 1e-9 for ra, rb in zip(values, charts["daily_fleet"][key]["values"]) for a, b in zip(ra, rb)), key
    checked.append("daily_fleet chart")

    # The first load measurement (#151) and the fresh-seed comparison (#152).
    baseline = read(EVIDENCE / "api-load-2026-10-13.json")
    concurrency = read(EVIDENCE / "api-load-2026-10-13-booking-concurrency.json")
    base = {r["label"]: r for r in baseline["runs"]}
    fresh = {r["label"]: r for r in concurrency["coldRemeasurement"]["runs"]}
    check("holds chart", [{"desks": n, "before_lost": b["booking"]["holdUnavailable"], "before_selected": b["booking"]["selected"], "after_lost": a["booking"]["holdUnavailable"],
                           "after_selected": a["booking"]["selected"], "before_p95": op(b, "booking offers")["ms"]["p95"], "after_p95": op(a, "booking offers")["ms"]["p95"],
                           "before_per_minute": per_minute(b), "after_per_minute": per_minute(a), "stale_per_search": a["booking"]["stale"] / a["booking"]["searches"]}
                          for n, b, a in ((1, base["booking-cold-c1"], fresh["cold-recheck-c1 (seed 93)"]), (5, base["booking-cold-c5"], fresh["cold-recheck-c5 (seed 91)"]), (10, base["booking-cold-c10"], fresh["cold-recheck-c10 (seed 92)"]))],
          charts["holds"])

    # The routing runs (J0, J) and the chart datasets that use the VM runs.
    vm = {r["label"]: r for r in read(EVIDENCE / "api-load-2026-10-14-vm-baseline.json")["runs"]}
    routing = {r["label"]: r for r in read(EVIDENCE / "api-load-2026-10-14-routing.json")["runs"]}
    priority = {r["label"]: r for r in read(EVIDENCE / "api-load-2026-10-14-confirm-priority-and-jvm.json")["runs"]}
    remote = {r["label"]: r for r in read(EVIDENCE / "api-load-2026-10-14-embedded-vs-remote.json")["runs"]}
    cpu = lambda r, what: [c for c in r["booking"]["containers"] if what in c["name"]][0]["cpuPercent"]["mean"]
    per_leg = lambda r: r["serverBreakdown"][OFFERS]["routingMs"]["p50"] / r["serverBreakdown"][OFFERS]["fetched"]["p50"]
    j_runs = [(name, r) for name, r in (("J0", vm["base-spread2"]), ("J flex", routing["j-flex"]), ("J CH", routing["j-ch"]))]
    assert [e["run"] for e in charts["j_routing"]] == [n for n, _ in j_runs]
    for entry, (name, r) in zip(charts["j_routing"], j_runs):
        assert entry["searches"] == r["booking"]["booking"]["searches"] and close(entry["per_leg"], per_leg(r)) and close(entry["routing_cores"], cpu(r, "routing") / 100), entry
    checked.append("j_routing chart")
    attribution = [("J0 baseline", vm["base-spread2"]), ("J flexible", routing["j-flex"]), ("J CH", routing["j-ch"]), ("J2 shipped", priority["j2-slot"]), ("K remote", remote["k-remote5"])]
    expected = [{"stage": s, "operation": o, "elapsed_p50": r["serverBreakdown"][o]["elapsedMs"]["p50"], **{k: r["serverBreakdown"][o][k]["p50"] for k in ("queueMs", "routingMs", "calculationMs", "dbMs")}} for s, r in attribution for o in (OFFERS, CONFIRM)]
    check("attribution chart", expected, charts["attribution"])
    stages = [("#151", base["booking-cold-c5"]), ("#152", fresh["cold-recheck-c5 (seed 91)"]), ("#154", read(EVIDENCE / "api-load-2026-10-13-search-outcome.json")["runs"][0]),
              ("#155", [r for r in read(EVIDENCE / "api-load-2026-10-13-refinement-shedding.json")["runs"] if r["label"] == "after-c5 (seed 55)"][0]),
              ("J0", vm["base-spread2"]["booking"]), ("J flex", routing["j-flex"]["booking"]), ("J CH", routing["j-ch"]["booking"]), ("J2", priority["j2-slot"]["booking"])]
    assert stages[2][1]["label"] == "booking-cold-c5"
    check("progression chart", [{"stage": s, "label": r["label"], "searches_per_minute": per_minute(r), "p50": op(r, "booking offers")["ms"]["p50"], "p95": op(r, "booking offers")["ms"]["p95"]} for s, r in stages], charts["progression"])
    check("confirm chart", [{"stage": s, "run": run, "confirm_p50": op(r["booking"], "confirm")["ms"]["p50"], "confirm_p95": op(r["booking"], "confirm")["ms"]["p95"], "confirm_wait_p50": r["serverBreakdown"][CONFIRM]["queueMs"]["p50"],
                             "offers_p50": op(r["booking"], "booking offers")["ms"]["p50"], "offers_p95": op(r["booking"], "booking offers")["ms"]["p95"]}
                            for s, run, r in (("J CH", "j-ch", routing["j-ch"]), ("priority", "j2-confirm", priority["j2-confirm"]), ("priority,\nZGC / Parallel", "j2-gc", priority["j2-gc"]),
                                              ("priority +\nconfirm slot", "j2-slot", priority["j2-slot"]), ("slot, compact\nheaders", "j2-coh2", priority["j2-coh2"]))], charts["confirm"])
    check("remote chart", [{"run": run, "offers_p50": op(r["booking"], "booking offers")["ms"]["p50"], "offers_p95": op(r["booking"], "booking offers")["ms"]["p95"], "searches": r["booking"]["booking"]["searches"],
                            "answered_201": op(r["daily"], "daily proposal")["statuses"].get("201", 0), "answered_503": op(r["daily"], "daily proposal")["statuses"].get("503", 0), "daily_cpu_mean": cpu(r, "daily")}
                           for run, r in (("k-embedded", remote["k-embedded"]), ("k-remote", remote["k-remote"]), ("k-remote4", remote["k-remote4"]), ("k-remote5", remote["k-remote5"]))], charts["remote"])

    # The routing benchmarks and the paired solver study.
    for date, name in (("2026-09-24", "routing-ch-omaha-2026-09-24.json"), ("2026-10-14", "routing-ch-2026-10-14.json")):
        bench = read(EVIDENCE / name)
        assert bench["checks"]["exactDirectedParity"] and bench["routable"] == 132
        for engine in ("flexible", "CH-car-v1"):
            rows = [l for l in bench["latency"] if l["engine"] == engine]
            entry = [c for c in charts["ch"] if c["date"] == date and c["engine"] == engine][0]
            assert entry["p50"] == [l["p50Ms"] for l in rows] and entry["p95"] == [l["p95Ms"] for l in rows] and entry["concurrency"] == [l["concurrency"] for l in rows]
    checked.append("ch chart")
    prof = read(EVIDENCE / "solver-profiling-2026-10-14.json")
    pairs = prof["paired"]["pairs"]
    for entry in charts["paired_fleets"]:
        fp = [x for x in pairs if x["technicians"] == entry["fleet"]]
        assert len(fp) == 12 and entry["before_median"] == median(x["costPhaseMovesPerSecond"]["before"] for x in fp) and entry["after_median"] == median(x["costPhaseMovesPerSecond"]["after"] for x in fp)
        assert abs(entry["speedup"] - median(x["costPhaseMovesPerSecond"]["after"] / x["costPhaseMovesPerSecond"]["before"] for x in fp)) < 0.005, entry
        assert abs(entry["fairness_speedup"] - median(x["fairnessPhaseMovesPerSecond"]["after"] / x["fairnessPhaseMovesPerSecond"]["before"] for x in fp)) < 0.005, entry
        assert (entry["wins"], entry["ties"], entry["losses"]) == (sum(x["afterCostCents"]["after"] < x["afterCostCents"]["before"] for x in fp),
                                                                   sum(x["afterCostCents"]["after"] == x["afterCostCents"]["before"] for x in fp), sum(x["afterCostCents"]["after"] > x["afterCostCents"]["before"] for x in fp)), entry
    assert [e["fleet"] for e in charts["paired_fleets"]] == [20, 50]
    checked.append("paired_fleets chart")
    for side, raw in (("before", prof["profileBefore"]), ("after", prof["profileAfter"])):
        groups = charts["profile"][side]
        assert charts["profile"]["samples"][side] == raw["samples"] == sum(raw["recordings"].values())
        assert abs(sum(groups.values()) - sum(raw["leafFramesPercent"].values())) < 0.5, side
        assert groups["Immutable map probe"] == round(sum(v for k, v in raw["leafFramesPercent"].items() if k.startswith("java/util/ImmutableCollections$MapN")), 1), side
        assert groups["BigDecimal arithmetic"] >= round(sum(v for k, v in raw["leafFramesPercent"].items() if k.startswith("java/math/BigDecimal.")), 1), side
    checked.append("profile chart")
    assert len(charts["paired"]) == 24 and all(p["before"] == x["costPhaseMovesPerSecond"]["before"] and p["after"] == x["costPhaseMovesPerSecond"]["after"] for p, x in zip(charts["paired"], pairs))
    checked.append("paired chart")

    # The September booking study: service and latency.
    held = {b["variant"]: b for b in summary["booking"] if b["stage"] == "held"}
    assert [h["served"] for h in charts["booking_study"]["held"]] == [held[v]["served"] for v in ("INSERTION", "BOUNDED", "SHARED")]
    evidence = read(ROOT / "docs/booking-experiment-explainer/evidence.json")
    for entry in charts["booking_study"]["latency"]:
        values = sorted(a["elapsed_ms"] for r in evidence["cases"] if r["solver"] == entry["solver"] and r["concurrency"] == entry["concurrency"] for a in r["attempts"] if a["elapsed_ms"] is not None)
        assert values[max(0, -(-len(values) * 95 // 100) - 1)] == entry["p95"], entry  # the explainer's quantile: ceil(n * 0.95) - 1
    checked.append("booking_study chart")

    print("verified", len(checked), "items:", ", ".join(checked))


if __name__ == "__main__":
    main()
