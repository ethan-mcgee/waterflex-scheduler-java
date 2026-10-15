"""Build "WaterFlex scheduler explained" from the validated sources. Install requirements.txt first. No network or benchmark runs."""
import re
from copy import deepcopy
from datetime import datetime
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from statistics import mean, median
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from charts import SOLVER_LABEL, by_label, by_run, make_charts
from report_data import load, per_fleet, plain, wtl

HERE = Path(__file__).resolve().parent
REFERENCE = HERE.parent / "WaterFlex_Scheduler_Algorithm_Comparison_Explained.docx"
REFERENCE_SHA = "5366f3017d6f18c15de295ba5fb9c1a014b6222911040f1c5e044845e57f832d"
OUTPUT = HERE / "WaterFlex_Scheduler_Explained_2026-10-14.docx"
PRESERVE = ("word/styles.xml", "word/stylesWithEffects.xml", "word/numbering.xml", "word/theme/theme1.xml", "word/fontTable.xml")
OFFERS, CONFIRM = "POST /api/v1/booking/offers 201", "POST /api/v1/booking/holds/{holdId}/confirm 200"


def ms(value):
    return f"{value:,} ms"


def seconds(value_ms):
    return f"{value_ms / 1000:.1f} s"


def dollars(value):
    return f"-${-value:,.2f}" if value < 0 else f"${value:,.2f}"


def pct(part, whole):
    return f"{part / whole * 100:.0f}%"


def lat(op):
    return f"{op['ms']['p50']:,} / {op['ms']['p95']:,} ms"


def by(rows, solver, **kw):
    out = [r for r in rows if r["solver"] == solver and all(r[k] == v for k, v in kw.items())]
    assert out, (solver, kw)
    return out


def saving(rows, solver, **kw):
    return mean(r["saving"] for r in by(rows, solver, **kw))


def wtl_text(rows):
    w, t, l = wtl(rows)
    return f"{w} / {t} / {l}"


def build():
    assert sha256(REFERENCE.read_bytes()).hexdigest() == REFERENCE_SHA
    data = load()
    charts = make_charts(data, HERE / "generated")
    docs, perf, daily, b24, f29, r24 = data["docs"], data["perf"], data["daily"], data["booking24"], data["field29"], data["routing24"]
    runs, extras, benchmarks, per_minute, queue_wait = perf["runs"], perf["extras"], perf["benchmarks"], perf["per_minute"], perf["queue_wait"]
    doc = Document(REFERENCE)
    table_properties = deepcopy(doc.tables[0]._tbl.tblPr)
    for child in list(doc.element.body):
        if child.tag != qn("w:sectPr"): doc.element.body.remove(child)
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype.endswith(("/image", "/hyperlink")): del doc.part.rels[rid]
    doc.core_properties.title = "WaterFlex scheduler explained"
    doc.core_properties.subject = "The stateless scheduling API, the optimization work behind it, and what comes next"
    doc.core_properties.author = doc.core_properties.last_modified_by = "WaterFlex"
    doc.core_properties.created = doc.core_properties.modified = datetime(2026, 10, 14)
    figures = []

    def p(text, style=None):
        assert "—" not in text and "–" not in text, text[:60]
        para = doc.add_paragraph(text, style); para.paragraph_format.widow_control = True
        return para

    def heading(text, new_page=False):
        h = p(text, "Heading 1"); h.paragraph_format.keep_with_next = True; h.paragraph_format.page_break_before = new_page

    def sub(text): p(text, "Heading 2")

    def table(headers, rows, widths, size=10, keep=True):
        assert abs(sum(widths) - 6.9) < .01, widths
        t = doc.add_table(rows=1, cols=len(headers)); t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t._tbl.replace(t._tbl.tblPr, deepcopy(table_properties))
        for col, w in zip(t.columns, widths): col.width = Inches(w)
        for i, values in enumerate([headers, *rows]):
            assert len(values) == len(headers), values
            row = t.rows[0] if i == 0 else t.add_row()
            trpr = row._tr.get_or_add_trPr(); trpr.append(OxmlElement("w:cantSplit"))
            if i == 0: trpr.append(OxmlElement("w:tblHeader"))
            for j, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
                assert "—" not in str(value) and "–" not in str(value), value
                cell.width = Inches(width); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell.text = str(value)
                props = cell._tc.get_or_add_tcPr()
                shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "DCE6EF" if i == 0 else ("FFFFFF" if i % 2 else "F4F6F8")); props.append(shade)
                margins = OxmlElement("w:tcMar")
                for side in ("top", "left", "bottom", "right"):
                    n = OxmlElement("w:" + side); n.set(qn("w:w"), "65"); n.set(qn("w:type"), "dxa"); margins.append(n)
                props.append(margins)
                para = cell.paragraphs[0]; para.paragraph_format.space_after = Pt(2); para.paragraph_format.space_before = Pt(2); para.paragraph_format.line_spacing = 1.0
                if len(headers) > 2 and j > 0: para.alignment = 1
                if keep and len(rows) <= 12 and i < len(rows): para.paragraph_format.keep_with_next = True
                for run in para.runs: run.font.name = "Calibri"; run.font.size = Pt(size); run.bold = i == 0
        spacer = p(""); spacer.paragraph_format.space_after = Pt(0); spacer.paragraph_format.line_spacing = Pt(3)

    def figure(name, caption):
        figures.append(name)
        caption = f"Figure {len(figures)}. {caption}"
        para = p(""); shape = para.add_run().add_picture(str(HERE / "generated" / f"{name}.png"), width=Inches(6.85))
        shape._inline.docPr.set("descr", caption); para.paragraph_format.keep_with_next = True
        cap = p(caption, "Caption"); cap.paragraph_format.space_after = Pt(7)
        for run in cap.runs:
            run.font.color.rgb = RGBColor(0, 0, 0); run.bold = False; run.italic = True

    # The runs and rows the narrative refers to by name.
    b5, b1, b10 = (by_label(runs["baseline"], f"booking-cold-c{c}") for c in (5, 1, 10))
    b_daily = {c: by_label(runs["baseline"], f"daily-c{c}") for c in (1, 4)}
    fresh = {n: by_label(extras["coldRemeasurement"], f"cold-recheck-c{n} (seed {s})") for n, s in ((1, 93), (5, 91), (10, 92))}
    c10 = by_label(runs["concurrency"], "booking-cold-c10")
    a_daily = {n: by_label(runs["admission"], f"daily-c{n}") for n in (1, 4)}
    ab_before, ab_after = extras["abComparison"]
    outcome = extras["outcome"]
    fair = {(r["note"], r["desks"]): r for r in runs["queue"] if r["note"] in ("sync", "async")}
    roles = {run: (by_run(runs["roles"], run, "booking"), by_run(runs["roles"], run, "daily")) for run in ("all-1", "split-1", "split-2", "all-2")}
    vm_b, vm_d = by_run(runs["vm"], "base-spread2", "booking"), by_run(runs["vm"], "base-spread2", "daily")
    jf_b, jc_b, jc_d = by_run(runs["routing"], "j-flex", "booking"), by_run(runs["routing"], "j-ch", "booking"), by_run(runs["routing"], "j-ch", "daily")
    j2 = {run: by_run(runs["priority"], run, "booking") for run in ("j2-confirm", "j2-gc", "j2-slot", "j2-coh2")}
    final = j2["j2-slot"]
    k = {run: by_run(runs["remote"], run, "booking") for run in ("k-embedded", "k-remote", "k-remote4", "k-remote5")}
    k_d = {run: by_run(runs["remote"], run, "daily") for run in k}
    prof = extras["profiling"]; pairs, summary = prof["pairs"], prof["summary"]
    ch_new = {(l["engine"], l["concurrency"]): l for l in benchmarks["2026-10-14"]["latency"]}
    ch_old = {(l["engine"], l["concurrency"]): l for l in benchmarks["2026-09-24"]["latency"]}
    routing_cpu = lambda r: [x for x in r["containers"] if "routing" in x["name"]][0]["cpuPercent"]["mean"]
    daily_cpu = lambda r: [x for x in r["containers"] if "daily" in x["name"]][0]["cpuPercent"]["mean"]
    per_leg = lambda r: r["breakdown"][OFFERS]["routingMs"]["p50"] / r["breakdown"][OFFERS]["fetched"]["p50"]
    F, M, L, C = daily["first"]["rows"], daily["follow"]["rows"], daily["follow"]["long"], daily["follow"]["confirm"]
    FC = daily["first"]["confirm"]
    tk_first, tk_follow, s3_follow = by(F, "TABU_KOPT"), by(M, "TABU_KOPT"), by(M, "TABU_KOPT_SIZE_3")
    late_family = [r for r in F if r["solver"] in ("KOPT", "SUBLIST", "LATE_ACCEPTANCE")]
    tk_twins = {r["key"][:-1]: r for r in tk_follow}
    direct = [tk_twins[r["key"][:-1]]["accepted_cost"] - r["accepted_cost"] for r in s3_follow]
    held = {v: f29["booking"][("held", v)] for v in ("INSERTION", "BOUNDED", "SHARED")}
    omaha = r24["omaha"]
    total_runs = sum(len(v) for v in runs.values()) + len(extras["coldRemeasurement"]) + len(extras["abComparison"])
    daily_cases = len(F) + len(FC) + len(M) + len(L) + len(C)
    ops = docs["operations"]

    # Facts the narrative depends on, checked before any prose uses them.
    assert all(saving(F, "TABU_KOPT", fleet=f) > 0 for f in (5, 10, 20, 50)) and all(saving(F, "TABU_KOPT", workload=w) > 0 for w in ("CLUSTERED", "DISPERSED", "MIXED_SKILL", "TIGHT_WINDOW"))
    assert all(saving(M, "TABU_KOPT", fleet=f) > 0 for f in (5, 10, 20, 30, 50))
    assert all(saving(M, "TABU_KOPT_SIZE_3", fleet=f) > saving(M, "TABU_KOPT", fleet=f) for f in (10, 20, 30, 50)) and saving(M, "TABU_KOPT_SIZE_3", fleet=5) < saving(M, "TABU_KOPT", fleet=5)
    assert wtl(by(C, "TABU_KOPT_SIZE_3")) == (16, 0, 0)
    assert all(saving(F, s, fleet=50) < 0 for s in ("KOPT", "SUBLIST", "LATE_ACCEPTANCE"))
    assert held["BOUNDED"]["served"] > held["INSERTION"]["served"] > held["SHARED"]["served"]
    assert b24["paid_min"] > 0 and b24["cost"] > 0
    assert k["k-remote5"]["ops"]["booking offers"]["ms"]["p95"] > 1.05 * k["k-embedded"]["ops"]["booking offers"]["ms"]["p95"]
    assert summary["afterCostCentsWinTieLoss"]["afterWins"] == 21 and summary["violationsAfter"] == 0

    p("WaterFlex scheduler explained", "Title")
    p("The stateless scheduling API, the optimization work behind it, and what comes next", "Subtitle")
    p("October 14 2026  |  Covers the work from 2026-09-23 to 2026-10-14")
    p(f"Since the September demonstration the scheduler has become a service with no screens and no copy of anyone's data: WaterFlex Software sends a snapshot of a metro day or a booking request to /api/v1 and gets back offers, optimized routes or a receipt. Underneath, the same month produced a different daily search (Tabu + K opt, chosen over Tabu after {daily_cases:,} matched solver cases on synthetic days), a different road graph (a contraction hierarchy, after the time inside a booking search was attributed part by part), and a different way of admitting work (fair queues, replica pools and a confirm that never waits). This document explains the system as it stands, then tells each of those threads as it happened: the first iteration, what was wrong with it, the measurement that showed it, the change, and the numbers that justified keeping it. It closes with what the evidence says should happen next. [1, 2, 3]")
    sub("The headline numbers")
    table(["Thread", "First measured iteration", "Where it stands", "Section"], [
        ["Daily search", f"Tabu, retained as the default after the Sept 29 comparison; K opt trailed it by {dollars(-f29['daily'][(15, 'KOPT')])} per case at 15 s", f"Tabu + K opt, per client: {dollars(saving(F, 'TABU_KOPT'))} per case better than Tabu over {len(tk_first)} cases, ahead in every fleet, workload and budget; {dollars(saving(M, 'TABU_KOPT'))} on 40 unseen fixtures", "8"],
        ["Booking search", f"Bounded rearrangement served {held['BOUNDED']['served']} of 480 held-out requests against insertion's {held['INSERTION']['served']}, at a latency cost", f"Bounded, inside the client's deadline; optional refinement shed under load: {per_minute(ab_before):.0f} to {per_minute(ab_after):.0f} searches per minute at five desks", "9"],
        ["Routing", f"Flexible graph, two worker threads: {per_leg(vm_b):.2f} ms per fetched leg, routing {ms(vm_b['breakdown'][OFFERS]['routingMs']['p50'])} of a {ms(vm_b['breakdown'][OFFERS]['elapsedMs']['p50'])} search", f"CH graph versioned with the map: {ch_new[('flexible', 1)]['p50Ms']:.1f} to {ch_new[('CH-car-v1', 1)]['p50Ms']:.1f} ms per 16-pair request; {vm_b['booking']['searches']} to {jc_b['booking']['searches']} searches in 60 s", "10"],
        ["Admission", f"{pct(b5['booking']['holdUnavailable'], b5['booking']['selected'])} of selected holds lost at five desks; a queued day refused after {seconds(b_daily[4]['ops']['daily proposal']['ms']['p50'])}", f"No hold lost; fair queue; two pools; confirm {ms(jc_b['ops']['confirm']['ms']['p50'])} to {ms(final['ops']['confirm']['ms']['p50'])} p50", "11"],
        ["Inside the engine", "Road legs looked up by string key on every scored route", f"{summary['costPhaseSpeedup']['median']:.2f}x the moves per second in the same budget (median of 24 paired solves); embedded engine kept over a remote pool for booking", "12"]], [1.0, 2.5, 2.8, .6], size=9)
    p("Every figure in that table is recomputed from the evidence files in the repository when this document is built, and a verifier recomputes them again on its own. None of it is customer traffic: the solver cases are synthetic days with a synthetic travel matrix, and the load runs are random Omaha addresses against a rehearsal stack on one workstation. The sections say, each time, what a number can and cannot support.")
    sub("How to read this document")
    table(["Sections", "What you will learn"], [("1 to 6", "The system as it stands: what the service is, where it runs, how a request flows, the contract and its concurrency token, many clients and replicas, routing and geocoding."),
                                                ("7", "How the work was measured, and the discipline that keeps the numbers honest."),
                                                ("8 to 12", "The optimization threads in order: the daily search, the booking search, routing, admission under load, and the engine itself."),
                                                ("13 and 14", "What should happen next: performance, then solver quality."),
                                                ("Appendices", "Every study and evidence file, definitions, and sources.")], [1.45, 5.45])

    heading("1 What the scheduler is now, and what it is not")
    p("The scheduler is an add-on to WaterFlex Software. WaterFlex Software owns master data (customers, addresses, jobs, appointments, technicians, qualifications, shifts, time off, depots) and the business record of every appointment. The scheduler owns only what it must to keep customer promises safe when several desks book at once: the holds and the arrangement they depend on, the offer sets and proposals it issued, and the receipts of what was committed. Calculation and routing are stateless: complete data in, results out. That is the thin scheduler store, decided on 2026-10-07. [4, 5]")
    table(["WaterFlex Software owns", "The scheduler owns"], [[plain(a), plain(b)] for a, b in docs["ownership"][1]], [3.45, 3.45])
    sub("What changed since the demonstration")
    p("The scheduler shown in September read its own copy of the portal's data and ran its own jobs. The one described here keeps nothing it can recompute and runs nothing the host did not ask for. [4]")
    table(["Piece", "In September", "Now"], [list(row) for row in data["before_after"]], [1.1, 2.9, 2.9], size=9)
    p("Three consequences follow. The service has no screens, so every map and schedule is rendered by WaterFlex Software from data the scheduler returns. The service reads no policy of its own: every rule a client chooses (its fairness budget, its daily search method and times, its booking deadline and window threshold) arrives with every call, and a missing value is refused, never assumed. And no client library ships: the product stays Java on its own VM, and a host integrates from any language against the OpenAPI contract. The portal in this repository is now one such client, used for demonstrations and as the worked example of the integration; it is not part of the service and this document does not describe it.")

    heading("2 Where it runs")
    figure("topology", "The scheduler VM and its host. WaterFlex Software calls one address; behind it the proxy sends booking calls to the booking pool and proposals to the daily pool. Both pools embed the calculation engine and reach the same three shared pieces. The management port is private and not behind the proxy.")
    p("The scheduler runs on its own VM with a partitioned share of compute, on the same private network as the rest of WaterFlex Software, reached only through https://<scheduler-vm>/api/v1. On the VM, scheduler-service runs as replicas of one jar behind nginx. Booking replicas answer booking searches, confirms and route calls; daily replicas take daily and repair proposals. Every replica answers every route, holds no state of its own and shares the one thin store, so which replica answers a call never changes the answer. [4, 6]")
    table(["Component", "Role", "Talks to"], [
        ("scheduler-proxy (nginx, :8000)", "The one address a host calls; proposals and request status go to the daily pool, everything else to the booking pool", "Both pools"),
        ("Booking replicas", "Booking offers, select, release, confirm, route evaluation and geometry; embedded calculation engine", "Routing, Nominatim, the database"),
        ("Daily replicas", "Daily and repair proposals, synchronous or queued; several calculations at once", "Routing, Nominatim, the database"),
        ("routing-service per metro", "Turns coordinates into directed road legs from that metro's graph; sees coordinates only", "Its map volume"),
        ("Nominatim (optional)", "Locates an address-only location as an exact house, in NOMINATIM mode", "Its own database"),
        ("scheduler-db (Postgres 16)", "The thin store, tenants and tokens, and the route cache; created and migrated by the scheduler itself", "Nothing else"),
        ("solver-service (optional)", "A remote calculation pool a scheduler pool may be pointed at; measured in section 12", "Nothing else"),
        ("Management port (:9000)", "Health with components and Prometheus metrics; private network only", "Operators")], [1.7, 3.5, 1.7])
    p("The calculation engine stays embedded in each replica by default (SCHEDULER_CALCULATION_MODE=EMBEDDED). Moving it to the remote pool is a configuration change already wired in and measured (section 12); it is a lever for the daily pool, not a default. Each served metro has its own routing service loaded with that metro's road graph; routing sees only coordinates, never client data, so every client in a metro shares it.")

    heading("3 One request, end to end")
    p("Every calculating call carries a Snapshot of one metro: its rates and policy, the technicians it covers, one TechnicianDay per technician and date with the host's lastModified, and every committed appointment on those days. The scheduler computes a canonical hash of it (the input revision) and every proposal and offer records which revision it was computed from. What happens next is a short pipeline, and every stage of it can refuse. [4, 7]")
    sub("A daily proposal")
    table(["Stage", "What happens", "Refusal"], [
        ("Authenticate", "TenantAuthentication maps the bearer token to a tenant; a request never names one.", "401"),
        ("Parse", "The body is read strictly: unknown fields, duplicate keys, missing values and dangling references are 400; contradictions between sequence and plannedStart are 422.", "400, 422"),
        ("Claim", "The requestId is claimed in the thin store with a hash of the body: a replay returns the stored answer, a different body is 400, a running one is 429.", "400, 429"),
        ("Cutoff", "The 6 a.m. cutoff in the metro's zone; a frozen day is 422.", "422"),
        ("Admission check", "SearchAdmission estimates when the work ahead would finish; a request that could not start in time is refused at once with Retry-After.", "429 BUSY"),
        ("Route", "DailyPreparation fetches every directed leg between the day's points under its own 30 s allowance (section 6); an unlocatable technician-day is skipped and reported.", "503"),
        ("Calculate", "DailyOperation takes a slot and runs the cost search and the fairness search for the client's seconds plus a 5 s reserve, embedded or remote; a watchdog ends an overrunning phase.", "503, 409 CANCELLED"),
        ("Decide", "Policy compares the result with the day as sent: IMPROVED, REJECTED_BY_POLICY or NO_IMPROVEMENT with a reason code; a plan with any overtime cannot be returned.", ""),
        ("Store", "The proposal and the lastModified of each technician-day are stored and answered 201; a refusal that would repeat is stored too, one that would not (429, 503, cancel) releases the claim.", "")], [1.1, 4.9, .9], size=9)
    figure("daily", "Daily and repair proposals: synchronous or queued, then the commit with the host's current timestamps. The lower left lists the refusals and the cancel.")
    sub("A booking search")
    p(f"A booking search claims its requestId, bounds the whole call by the client's bookingDeadlineMs (at most 5 s), checks the cutoff and metro, builds the facts, takes an admission slot fairly across tenants, reconciles the stored arrangement of each date with the snapshot (section 4), routes the new job's legs, runs an insertion pass and then an optional refinement through the engine, validates the result independently, and publishes the offers as holds that last 10 minutes. If another booking changed the same day while it ran, it runs again, at most three times, while the deadline leaves time. The windows a customer sees are the insertion pass's; the refinement only improves cost. [7, 8]")
    figure("booking", "The booking sequence and its refusals. Each 409 and a failed receipt write lead to the same place: release the set (or let the hold expire) and search again from current records with a new requestId.")

    heading("4 The contract and the concurrency token")
    p(f"The machine-readable contract is docs/api/openapi-v1.yaml, OpenAPI 3.1, version {docs['version']}. Every operation is marked implemented; a test validates every example against its schema and fails if the implemented operations differ from the routes the scheduler serves, and the Java types mirror every object schema. Within version 1 a change is additive only and raises the minor version; the contract reached 1.0.0 on {docs['history'][-1][1]} and is at {docs['history'][0][0]} after {len(docs['history']) - 1} additive versions. [7]")
    table(["Operation", "Method and path", "What it does"], [[o["id"], f"{o['method']} {o['path'].replace('/api/v1', '')}", o["summary"]] for o in ops], [1.45, 2.35, 3.1], size=9, keep=False)
    sub("The timestamp the whole design rests on")
    p("The scheduler cannot lock rows it does not own, so every commit is optimistic on both sides, and the token is the lastModified timestamp WaterFlex Software already keeps per technician and day. Every snapshot includes it; a proposal or offer records the values it was computed from; a commit or confirm requires the submitted values to equal the recorded ones; and the host then writes the business record with a compare-and-set on the same values. It is compared for exact equality and never ordered, so clock skew between servers cannot cause a wrong decision; at worst an unrelated change forces a retry. The host must move it on every change that affects the day, on both sides of a move, and send it exactly as stored. [4, 8]")
    p("For booking, each metro-day has one stored common arrangement: every technician's route with the host's appointments and every active hold. Each search and confirm compares the host's snapshot with it, technician-day by technician-day, and nothing is guessed:")
    table(["Rule", "Technician-day in the snapshot", "What the scheduler does"], [
        ("0", "Carries a lastModified the scheduler has already moved past (it remembers the last 32), or the unchanged timestamp with exactly the route from before a confirm receipt the host has not written yet", "Refused with 409 STALE before anything is reconciled, so nothing is lost; the host rebuilds from current records"),
        ("1", "lastModified unchanged", "As stored"),
        ("2", "Changed, and its appointments in order are exactly what the scheduler expects (for example after the host wrote a receipt)", "Accepted; the new timestamp is recorded"),
        ("3", "Changed in any other way (a dispatcher added, removed or moved an appointment)", "Replaces the stored route when that route has no hold and no pending move; otherwise every hold on that date is lost and the date restarts from the host's routes"),
        ("4", "The whole arrangement, re-evaluated with the snapshot's current windows, durations, shifts and absences", "If it no longer fits, every hold on that date is lost the same way")], [.45, 3.1, 3.35], size=9)
    p("Rule 0 did not exist in the first version; section 11 shows what happened without it. The rest of the conventions are the same for every call: the tenant comes from the token; every write carries a host UUID that is its idempotency key across all of the tenant's endpoints; money and rates are decimal strings; instants are RFC 3339; and no operation ever assigns overtime, so a repair that would need it is reported infeasible rather than solved.")
    table(["Status", "Code", "Meaning", "What to do"], [[plain(s), plain(c), plain(m), plain(w)] for s, c, m, w in docs["problems"][1]], [.45, 1.6, 2.6, 2.25], size=9)

    heading("5 Many clients, many replicas, and the thin store")
    p("One Postgres database and one set of tables serve every client. Each WaterFlex client is a tenant with its own API tokens, stored only as SHA-256 digests and issued with the scheduler's own command line. Every client-owned table has a tenant_id column that starts every key and index, so two clients can both have technician 42. [4, 6]")
    table(["Guard", "How it works"], [
        ("Token to tenant", "The bearer token decides the tenant; a revoked token or a disabled tenant gets 401."),
        ("Row-level security", "Each transaction runs as the scheduler_tenant role with app.tenant_id set; policies restrict every client-owned table to that tenant and an unset tenant sees nothing, so a missed filter in application code fails closed."),
        ("Least privilege", "The runtime login is a member of scheduler_app, never the schema owner; migrations run as the owner with --command=migrate."),
        ("Fair turns", "Admission starts queued bookings by client (fewest running, then the one whose last booking started longest ago); queued proposals are taken in turn across clients."),
        ("Shared road legs", "Route legs are not tenant data and are shared across clients in a metro; the cache holds coordinates at about 1 m and is purged within 30 days.")], [1.3, 5.6])
    sub("Several replicas on one database")
    p("Correctness lives in the database: holds are protected by row locks and unique constraints on (tenant, technician, date) and by the timestamp comparison inside one transaction, so any replica can serve any request and the load balancer needs no sticky sessions. What stays per replica by design is admission (the search slots), the memory tier of the route cache, and the in-memory list of running calls, which the cancel flag in the database reaches from any replica within about a second. A replica's role (SCHEDULER_ROLE booking, daily or all) decides only which background work runs there: a booking replica takes no queued daily work, and a daily replica may run several calculations at once. Section 11 measures what the split bought.")
    sub("Queued proposals")
    p("With the header Prefer: respond-async a proposal is queued in a durable table and answered 202 at once; the host polls GET /requests/{requestId} and reads QUEUED, RUNNING, then COMPLETED with the status and body the call itself would have answered. Any replica's worker takes queued requests with SKIP LOCKED, the waiting request of the client whose last one started longest ago first, so one client's thirty overnight days never hold back another's single day. A client may have 50 requests queued or running; a request still waiting after 15 minutes completes with 429 without running. Overnight optimization is therefore the host's to schedule; the scheduler has no batch of its own.")
    sub("What is stored, and for how long")
    table(["Table", "Holds"], [[plain(t), plain(h)] for t, h in docs["tables"][1]], [1.6, 5.3], size=9)
    table(["Data", "Typical lifetime"], [[plain(d), plain(l)] for d, l in docs["stored"][1]], [2.8, 4.1])
    p("The retention period is the deployment's SCHEDULER_RETENTION_DAYS, required with no default: the scheduler does not start without it. A daily purge deletes rows older than the period in batches through a database function, children before parents. A host keeps its own copy of any receipt it needs longer.")

    heading("6 Routing and geocoding")
    p("Coordinates in the snapshot are authoritative: when a location arrives with latitude and longitude the scheduler routes from them as given and never geocodes again. Nominatim is the fallback only for a location that arrives with an address and no coordinates, only in NOMINATIM mode, and only as an exact house: ZIP, house number and street must agree, and a street-level or ambiguous result is not located. Every drawn stop says where its coordinates came from, so a bad route can be traced to its input. [4]")
    figure("tiers", "The four route cache tiers in lookup order, from the design document. A pair of coordinates is answered by the first tier that holds it; a leg computed by the routing service is written back to the database and memory tiers in a separate short transaction.")
    table(["Tier", "Where", "Key", "Lifetime and size", "Shared by"], [[plain(c) for c in row] for row in docs["tiers"][1]], [1.0, 1.5, 1.4, 1.9, 1.1], size=9)
    p("None of the tiers is a source of truth: every entry is recomputable from the road graph, and every entry is bound to a routing identity, a hash of the graph and map version that the routing service reports and checks on every call. A map update, or a switch between the flexible and contraction hierarchy graphs, changes the identity, clears the memory tiers and leaves the database tier unused for the old identity, so a changed road network can never be mixed with old legs. Each map version is prepared as a whole (the merged OpenStreetMap extract, both road graphs, tiles, and TIGER address ranges for Nominatim), validated with preview services, and activated by switching symlinks. Section 10 is the story of how the second graph got there.")

    heading("7 How the work was measured", new_page=True)
    p("Two kinds of evidence run through the rest of this document, and they answer different questions. Fixture studies ask which search finds cheaper schedules in the same seconds; load studies ask how the service behaves when many callers use it at once. Neither is customer traffic. [1, 2, 3]")
    table(["Kind", "What one case is", "How a comparison is made", "What it cannot say"], [
        ("Daily solver study", "One configuration improving one fixed synthetic day for one search seed and one time budget; a synthetic travel matrix, not roads", "Matched: every configuration gets the same starting schedule; a win is a lower accepted cost in exact cents; savings are in modeled dollars per case against the control", "Payroll savings, or how a real customer day behaves"),
        ("Booking study (September)", "Ten booking requests against one fixture day, for one search method, concurrency and cache state", "Paired on cases where both methods booked the same customers; service counted separately from latency", "What a production desk experiences"),
        ("Load study (October)", "A 60 or 90 s run of the load tool playing WaterFlex Software: desks searching, holding and confirming at random Omaha addresses; planners asking for optimized days", "Fresh seeds for every run (a seed run before leaves roads cached); every latency is the caller's, waits and retries included; one change per pull request", "Absolute capacity of the real VM; the rehearsal stack's limits are a shape")], [1.2, 2.0, 2.5, 1.2], size=9)
    p(f"The discipline is the same everywhere. Each study records its configuration, seeds, revision and raw results with hashes; a loader checks the shape of every file before a number is read (statuses add up to calls, matched pairs share a starting schedule, no constraint violation anywhere); the report is built from those files, and a verifier recomputes every table from the raw data without the builder's help. This document is built the same way from {len(perf['studies'])} load studies ({total_runs} runs), the solver profiling study, two routing benchmarks, the two daily search studies ({daily_cases:,} cases), the September booking study ({b24['requests']:,} requests) and the field validation summary. Appendix A lists each file.")
    table(["Measure", "How to read it"], [
        ("Saving per case", "The control's accepted cost minus the configuration's, in modeled dollars, for the same starting day; positive is cheaper than the control."),
        ("Wins / ties / losses", "Matched cases where the configuration's accepted cost was lower, equal or higher than the control's, in exact cents."),
        ("Searches per minute", "Booking searches answered 201 in a load run, scaled to a minute; refusals and resends do not count."),
        ("p50 / p95", "The latency that half, or 95 percent, of calls in a run stayed within, measured at the caller with every wait included. The booking design target is p95 within the client's 5 s deadline plus 500 ms."),
        ("Holds lost", "Confirms answered 409 HOLD_UNAVAILABLE out of the holds customers selected: a customer who chose a time and could not get it."),
        ("Moves per second", "Candidate moves the daily search evaluated per second of its phase; more moves in the same budget is a better search, not a faster answer.")], [1.55, 5.35])

    heading("8 The daily search: from Tabu to Tabu + K opt")
    p("A daily proposal rearranges a day that is already booked: appointments may move between technicians and change order, but every customer's promised window holds and no overtime is added. The search runs two phases, a cost search and then a fairness search that may spend up to a small allowance of the cost it found to spread work more evenly, and policy accepts the result only if it beats the day as sent. The question of September was which local search to run inside that budget. [9, 10]")
    sub("First iteration: Tabu, and a comparison that said to keep it")
    p(f"The demonstration's default was Tabu search with a 15 s budget: single-visit relocations and swaps, an entity tabu list of 7 that stops the search undoing a move it just made, and acceptance of a worse step when that is the best available, so the search can leave a local trap. The algorithm comparison of September 29 tested seven alternatives on {f29['daily_cases']} clustered cases over two held-out seeds and three budgets. At 15 s nothing beat Tabu except the uncapped version of the original search ({dollars(f29['daily'][(15, 'CURRENT_UNCAPPED')])} per case); the K opt family trailed it ({dollars(-f29['daily'][(15, 'KOPT')])} per case behind for K opt, {dollars(-f29['daily'][(15, 'LATE_ACCEPTANCE')])} for late acceptance), and a larger ruin-and-recreate neighborhood trailed most ({dollars(-f29['daily'][(15, 'RUIN_RECREATE')])}). At 60 s the picture reversed: K opt saved {dollars(f29['daily'][(60, 'KOPT')])} per case and late acceptance {dollars(f29['daily'][(60, 'LATE_ACCEPTANCE')])}. The report also noted that K opt's 15 s losses were concentrated in the 50-technician cases. The decision was to retain Tabu and study fleet size and budget together. [11]")
    table(["Budget", "Uncapped original", "K opt", "Sublist", "Late acceptance", "Ruin and recreate"],
          [[f"{b} s", *[dollars(f29["daily"][(b, v)]) for v in ("CURRENT_UNCAPPED", "KOPT", "SUBLIST", "LATE_ACCEPTANCE", "RUIN_RECREATE")]] for b in (15, 30, 60)], [.8, 1.3, 1.1, 1.1, 1.3, 1.3])
    p("Those are paired mean savings against Tabu in dollars per case, eight matched cases per cell, from the field validation summary. A budget study on October 1 (1,280 cases at 60 to 240 s) sharpened the fleet finding: K opt, sublist and late acceptance each saved about $6 per case on average, and each lost every one of the 40 matched 50-technician dispersed cases, by $47 to $52. A blanket switch would have hidden that trade; its run archive is not in the repository, so this document cites it and recomputes nothing from it. [12]")
    sub("The Tabu variant study")
    p(f"The question became whether Tabu's weakness on mid-size fleets and the K opt family's weakness on large ones could be traded away. The Tabu variant study of October 11 ran {len(F):,} matched cases on {daily['first']['fixtures']} fixtures (5 to 50 technicians, four workloads, budgets of 60 to 240 s, three seeds) with seven configurations: Tabu, Tabu with a shorter and a longer tabu list, Tabu with the K opt and sublist moves added, and the three late acceptance methods. [1]")
    figure("daily_fleet", "Mean saving against Tabu in modeled dollars per case by fleet size. Left: the Tabu variant study (160 cases per cell). Right: the follow-up on 40 fixtures the first study never saw (48 cases per cell), with the combined Tabu + K opt size 3. Blue saves, orange loses.")
    table(["Configuration", "Mean saving vs Tabu", "Wins / ties / losses", "5 techs", "10", "20", "50"],
          [[SOLVER_LABEL[s], dollars(saving(F, s)), wtl_text(by(F, s)), *[dollars(v) for v in per_fleet(F, s, (5, 10, 20, 50))]] for s in ("TABU_SIZE_3", "TABU_SIZE_15", "TABU_KOPT", "KOPT", "SUBLIST", "LATE_ACCEPTANCE")], [1.6, 1.2, 1.2, .75, .7, .7, .75], size=9)
    p(f"Tabu + K opt saved {dollars(saving(F, 'TABU_KOPT'))} per case with {wtl_text(tk_first)} wins, ties and losses, and it was the only configuration ahead in every fleet size, every workload ({dollars(min(saving(F, 'TABU_KOPT', workload=w) for w in ('CLUSTERED', 'DISPERSED', 'MIXED_SKILL', 'TIGHT_WINDOW')))} to {dollars(max(saving(F, 'TABU_KOPT', workload=w) for w in ('CLUSTERED', 'DISPERSED', 'MIXED_SKILL', 'TIGHT_WINDOW')))}) and every budget. The late acceptance family was the better choice on 10 and 20 technicians (about {dollars(mean(r['saving'] for r in late_family if r['fleet'] in (10, 20)))} per case) and lost about {dollars(-mean(r['saving'] for r in late_family if r['fleet'] == 50))} on 50, and the recorded search mechanics say why: on 50-technician days late acceptance found its best schedule at {mean(r['best_fraction'] for r in by(F, 'LATE_ACCEPTANCE', fleet=50)) * 100:.0f} percent of its phase on average, against {mean(r['best_fraction'] for r in by(F, 'TABU_KOPT', fleet=50)) * 100:.0f} percent for Tabu + K opt; it was still improving when its time ran out. Tabu + K opt keeps most of the mid-size gain because the K opt moves reconnect several route links at once, and avoids the large-fleet loss because Tabu's step style converges early. Its cost is balance: mean accepted variance {mean(r['variance'] for r in by(F, 'TABU_KOPT')):.4f} against Tabu's {mean(r['variance'] for r in by(F, 'TABU')):.4f}. Rerun one case at a time on {len(FC) // 2} cases it saved {dollars(saving(FC, 'TABU_KOPT'))} per case ({wtl_text(by(FC, 'TABU_KOPT'))}), so running six cases at once had not made the result. [1]")
    sub("The decision, and the follow-up that tested it")
    p(f"On October 12 the daily search became each client's own choice in the snapshot's policy (TABU_KOPT or TABU, with its own cost and fairness seconds up to 240 together), and Tabu + K opt became the default for existing clients. A study that chooses its winner on the fixtures it measured can be fooled by those fixtures, so the follow-up of October 13 ran {len(M):,} cases on {daily['follow']['fixtures']} new fixtures with new seeds, added 30 technicians, and tried the one combination the first study suggested: Tabu + K opt with the shorter tabu list of 3. [2, 13]")
    table(["Configuration", "Mean saving vs Tabu", "Wins / ties / losses", "5 techs", "10", "20", "30", "50"],
          [[SOLVER_LABEL[s], dollars(saving(M, s)), wtl_text(by(M, s)), *[dollars(v) for v in per_fleet(M, s, (5, 10, 20, 30, 50))]] for s in ("TABU_KOPT", "TABU_KOPT_SIZE_3", "TABU_SIZE_3", "KOPT")], [1.6, 1.15, 1.15, .6, .6, .6, .6, .6], size=9)
    p(f"Tabu + K opt held: {dollars(saving(M, 'TABU_KOPT'))} per case, ahead at every fleet size, {wtl_text(tk_follow)}. Tabu + K opt size 3 did better still, {dollars(saving(M, 'TABU_KOPT_SIZE_3'))} per case, ahead of Tabu + K opt itself in {sum(v > 0 for v in direct)} of {len(direct)} matched cases and at every fleet size from 10 technicians up (it trails by {dollars(saving(M, 'TABU_KOPT', fleet=5) - saving(M, 'TABU_KOPT_SIZE_3', fleet=5))} at 5). Rerun one case at a time it beat Tabu in {wtl(by(C, 'TABU_KOPT_SIZE_3'))[0]} of {len(by(C, 'TABU_KOPT_SIZE_3'))}. The long-budget arm answered the last open question: on 50 technicians late acceptance draws level with Tabu + K opt only at 600 s ({dollars(saving(L, 'LATE_ACCEPTANCE', fleet=50, budget_ms=600000))} against it; {dollars(saving(L, 'LATE_ACCEPTANCE', fleet=50, budget_ms=240000))} at 240 s), beyond the 240 s a client can set, so the large-fleet weakness is slowness, not a fault, and a fleet-size rule is not needed. Every cheaper method spreads work less evenly than Tabu when no fairness search follows ({mean(r['variance'] for r in by(M, 'TABU_KOPT_SIZE_3')):.4f} against {mean(r['variance'] for r in by(M, 'TABU')):.4f} mean variance here, cost search only); with the fairness phase on, most of that gap closes. Section 14 says what the next step is. [2]")

    heading("9 The booking search: insertion, bounded rearrangement, and the deadline")
    p("Booking finds arrival windows for a new job and holds them while the customer chooses. The search must keep every existing promise and add no overtime, and it must answer inside the client's deadline. Two methods were in play from the start: insertion, which tries the new visit in every position of every eligible route without moving anything else, and bounded search, which also tries small rearrangements (a relocation, a swap, a reversal) around the affected work, within limits on routes, depth and candidates. [14, 15]")
    sub("First iteration: the September studies")
    p(f"The booking experiment of September 24 compared the two on {b24['cases']} cases ({b24['requests']:,} requests). In all {b24['comparable']} sequential comparisons where both methods booked the same ten requests, bounded search used {b24['paid_min']:.2f} fewer paid minutes and {dollars(b24['cost'])} less modeled cost per case on average, and never more of either. The gain depended on the workload: clustered days saved {b24['by_workload']['CLUSTERED']['paid_min']:.2f} minutes and {dollars(b24['by_workload']['CLUSTERED']['cost'])} per case, dispersed days {b24['by_workload']['DISPERSED']['paid_min']:.2f} and {dollars(b24['by_workload']['DISPERSED']['cost'])}, sparse days {b24['by_workload']['SPARSE']['paid_min']:.2f} minutes and {dollars(b24['by_workload']['SPARSE']['cost'])}. The price was time: one request at a time, the search p95 rose from {b24['latency'][('INSERTION', 1)]['p95'] / 1000:.2f} s to {b24['latency'][('BOUNDED', 1)]['p95'] / 1000:.2f} s; at ten simultaneous requests from {b24['latency'][('INSERTION', 10)]['p95'] / 1000:.2f} s to {b24['latency'][('BOUNDED', 10)]['p95'] / 1000:.2f} s. [14]")
    figure("booking_study", "Left: requests served of 480 in the September 29 held-out matrix by search method, with each method's mean case p95. Right: search p95 by concurrency in the September 24 study, insertion against bounded.")
    p(f"The field validation of September 29 widened the matrix (5 to 50 technicians, concurrency 1 to 10, cold and warm caches, two held-out seeds) and added three larger searches. Bounded search served {held['BOUNDED']['served']} of {held['BOUNDED']['requests']} requests against insertion's {held['INSERTION']['served']}, with a mean case p95 of {held['BOUNDED']['meanCaseP95Ms'] / 1000:.1f} s against {held['INSERTION']['meanCaseP95Ms'] / 1000:.1f} s. The largest search, shared evaluation across windows, served {held['SHARED']['served']} at {held['SHARED']['meanCaseP95Ms'] / 1000:.1f} s and ended one run with an audit failure, so it and the expanded and ruin-and-recreate variants were set aside. The decision was to keep bounded search as the booking method and insertion as its control, and to adopt the policy the studies had been run under: a single best offer per window, zero new overtime, service and return travel inside regular availability. [15]")
    sub("What changed in the API era")
    p(f"The public API runs the same bounded search (booking.search.variant=BOUNDED, with the rollout flags that once gated it removed) as two passes inside the client's bookingDeadlineMs: an insertion pass that places the job on each day and finds the regular windows, then a refinement pass that improves cost and may move other appointments. When the insertion pass already finds more distinct windows than the client's regularWindowThreshold, the refinement is optional and capped at 250 ms. Three things about that second pass were wrong when the load tool first measured it, and each was fixed and measured in turn. [3, 16]")
    table(["Finding", "What was wrong", "What changed", "Measured"], [
        ("searchComplete (#154)", f"Reaching the refinement cap marked the search incomplete, so {b5['booking']['searchIncomplete']} of {b5['booking']['searches']} searches at five desks reported searchComplete: false while returning four offers", "The flag means what the spec says (the deadline cut a needed pass short) and OfferSet.refinement says how the second pass ended", f"{outcome['searchCompleteFalse']} of {outcome['searches']} searches incomplete in 60 s at five desks; all {outcome['refinement']['SHORTENED']} refinements SHORTENED by the cap"),
        ("Shedding (#155)", "The optional pass was spent while other customers waited for a slot", "Refinement is shed (SHED) while booking requests are queued on the replica; a client may set its own refinement budget; the dates a needed pass did not finish are reported", f"A/B at five desks on fresh addresses: {per_minute(ab_before):.0f} to {per_minute(ab_after):.0f} searches per minute, p95 {seconds(ab_before['ops']['booking offers']['ms']['p95'])} to {seconds(ab_after['ops']['booking offers']['ms']['p95'])}, every search still four offers"),
        ("Invalid refinement (#138)", "A refinement result that failed validation failed the whole search", "The insertion offers are published with refinement: UNAVAILABLE and the failure is logged and counted; an invalid insertion result still fails the search", "Decided 2026-10-13; covered by tests rather than a load run")], [1.1, 1.9, 2.1, 1.8], size=9)
    p("The thread's lesson carries into the next sections: the second pass is worth its 250 ms when the replica is idle and worth nothing when four other desks are waiting for the slot it holds, so the right policy is conditional, and the host is told which it got. Throughput at five desks doubled without changing a single offered window.")

    heading("10 Routing: from a flexible graph to a contraction hierarchy")
    p(f"Every scored route needs the directed travel time between consecutive stops, so a day of 20 technicians needs a few thousand legs and a booking search for a new address needs a few hundred that nothing has seen before. The routing service computes them on a GraphHopper road graph of the metro, and the scheduler caches them in the tiers of section 6. The question of this thread was never whether a faster graph existed; it was whether routing was the bottleneck. [17]")
    sub("First iteration: sparse pairs, and a CH experiment shelved")
    p(f"On September 23 the scheduler stopped fetching whole matrices and fetched only the directed pairs a search needs: the first sparse Omaha request took {r24['sparse']['initialSparse']:.1f} ms and a repeat {r24['sparse']['repeatedSparse']:.2f} ms, with every pair matching the full matrix. The next day a contraction hierarchy (CH) graph was prepared for the Omaha map and benchmarked against the flexible graph on 132 directed pairs: every leg matched to the metre, and an uncached 16-pair request answered in {ch_old[('CH-car-v1', 1)]['p50Ms']:.1f} ms p50 instead of {ch_old[('flexible', 1)]['p50Ms']:.1f} with one caller and {ch_old[('CH-car-v1', 10)]['p50Ms']:.1f} instead of {ch_old[('flexible', 10)]['p50Ms']:.1f} with ten. Then the same two graphs were put under a full booking run of 720 requests on real Omaha roads: the flexible graph served {omaha['flexible']['served']} requests at a pooled p95 of {omaha['flexible']['p95Ms'] / 1000:.1f} s, the CH graph {omaha['ch']['served']} at {omaha['ch']['p95Ms'] / 1000:.1f} s. A twelve-fold faster road query had bought nothing a customer would notice, because the booking search of that time spent its deadline elsewhere (reservation contention and the search itself). CH and road prewarming stayed off by default, with the note that the decision should be revisited with combined booking evidence. [17, 18]")
    sub("What the attribution showed")
    p(f"That evidence arrived on October 14. Study J0 ran the load tool against the VM-shaped rehearsal stack and, for the first time, the request log attributed each call's time to its parts. A booking search at five desks took {ms(vm_b['ops']['booking offers']['ms']['p50'])} p50 at the caller: {ms(vm_b['breakdown'][OFFERS]['queueMs']['p50'])} waiting for a slot, {ms(vm_b['breakdown'][OFFERS]['routingMs']['p50'])} routing, {ms(vm_b['breakdown'][OFFERS]['calculationMs']['p50'])} in the engine and {ms(vm_b['breakdown'][OFFERS]['dbMs']['p50'])} in the database. Each search needed about {vm_b['breakdown'][OFFERS]['routingPairs']['p50']} directed legs and fetched {vm_b['breakdown'][OFFERS]['fetched']['p50']} of them cold at about {per_leg(vm_b):.2f} ms each, because the service routed a batch's pairs one after another and the scheduler fed it through two worker threads; the routing container ran at {routing_cpu(vm_b):.0f} percent of a core for the booking replica's {[x for x in vm_b['containers'] if 'booking' in x['name']][0]['cpuPercent']['mean']:.0f}. By October the booking contention of September had been fixed (section 11), and routing was now the largest part of a search's own time. The engine was not the lever: {ms(vm_b['breakdown'][OFFERS]['calculationMs']['p50'])} of a {ms(vm_b['breakdown'][OFFERS]['elapsedMs']['p50'])} server time. [3, 19]")
    figure("attribution", "Server time of the median booking call, split into slot wait, routing, calculation engine and database, for the VM baseline (J0), the two J runs, the J2 build that shipped (section 11) and the remote solver pool (section 12). The confirm has almost no work of its own; its bar is the wait behind searches.")
    sub("The change, and what it bought")
    p("Study J (#160) changed four things the attribution pointed at. Each map version now carries a CH graph beside the flexible one, prepared with the map, listed in its manifest and required at activation, and the routing service serves it by default. The routing service routes a request's pairs on several threads at once. The scheduler trusts a routing identity it fetched within the last two seconds instead of asking /health before every lookup. And the scheduler fetches leg batches on eight workers with a queue of 64 instead of two and 16. A contraction hierarchy spends time at import ordering every junction by importance and adding shortcut edges that stand for whole chains of minor roads, so a query climbs from both ends through ever more important roads and meets near the top after touching a few hundred nodes; the answer is identical to the metre and only the work changes. The price is an import step of 27 s and about 200 MiB more on disk per map version, and a routing identity of its own, so switching starts every cache cold like a map change does. [17, 20]")
    figure("ch", "p50 of a 16-pair routing request against the routing service alone with its caches off, flexible against CH, on the 2026-09-16 map in September and the 2026-10-13 map on October 14 (the new image routes a request's pairs in parallel). Log scale.")
    table(["Run", "Searches in 60 s", "Caller p50 / p95", "Slot wait", "Routing", "Per fetched leg", "Routing CPU"],
          [[name, r["booking"]["searches"], lat(r["ops"]["booking offers"]), ms(r["breakdown"][OFFERS]["queueMs"]["p50"]), ms(r["breakdown"][OFFERS]["routingMs"]["p50"]), f"{per_leg(r):.2f} ms", f"{routing_cpu(r):.0f}%"]
           for name, r in (("J0 baseline, flexible", vm_b), ("J, flexible graph", jf_b), ("J, CH graph", jc_b))], [1.6, .9, 1.2, .8, .8, .9, .7], size=9)
    p(f"The routing service alone now answers a 16-pair request in {ch_new[('CH-car-v1', 1)]['p50Ms']:.1f} ms instead of {ch_new[('flexible', 1)]['p50Ms']:.1f} with one caller and {ch_new[('CH-car-v1', 10)]['p50Ms']:.1f} instead of {ch_new[('flexible', 10)]['p50Ms']:.1f} with ten, with exact parity on all 132 pairs. Under load the CH run answered {jc_b['booking']['searches'] / vm_b['booking']['searches']:.1f} times the baseline's searches at a quarter of the latency, and the engine and database shares did not move, so the whole gain is routing and the slot time routing was holding: a slot held {ms(vm_b['breakdown'][OFFERS]['elapsedMs']['p50'] - vm_b['breakdown'][OFFERS]['queueMs']['p50'])} per search before and {ms(jc_b['breakdown'][OFFERS]['elapsedMs']['p50'] - jc_b['breakdown'][OFFERS]['queueMs']['p50'])} after. CH is the lever and parallel flexible search is not: routing a batch's pairs in parallel on the flexible graph halved the routing time but spent {routing_cpu(jf_b) / 100:.0f} cores doing it, while CH answered in a third of that time on {routing_cpu(jc_b) / 100:.1f}. One lesson for anyone repeating the September benchmark: with the service's own cache on, both engines answer repeated pairs from memory and CH looks slower by its small overhead; the benchmark expects both instances launched with the cache limited to one entry. Why did the same graph change nothing in September and everything in October? Because the September booking search was bound by contention and its own candidate work, not by road queries, and the October one, after sections 9 and 11, was bound by routing. The order of the work was the finding. [3, 17]")

    heading("11 Admission and concurrency under load")
    p(f"The load tool shipped on October 13 with its first measurement (#151): one booking desk met the latency target with room to spare, and five or more did not. The reason was not computation. Across the next eight pull requests each measurement found one thing, one change fixed it, and the next measurement found the next. [3, 21]")
    table(["Run (#151)", "Searches / min", "Offers p50 / p95", "Confirm p50", "Holds lost", "Waited on 429"],
          [[name, f"{per_minute(r):.0f}", lat(r["ops"]["booking offers"]), ms(r["ops"]["confirm"]["ms"]["p50"]), f"{r['booking']['holdUnavailable']} of {r['booking']['selected']}", r["ops"]["booking offers"]["retriedBusy"]]
           for name, r in (("1 desk", b1), ("5 desks", b5), ("10 desks", b10))], [1.1, 1.1, 1.6, 1.0, 1.0, 1.1])
    sub("Lost holds, and the rule that keeps them (#152)")
    p(f"At five desks {pct(b5['booking']['holdUnavailable'], b5['booking']['selected'])} of the holds customers selected were lost at confirm, at ten {pct(b10['booking']['holdUnavailable'], b10['booking']['selected'])}, and the cause was reproduced deterministically. The scheduler compares a technician-day's lastModified only for equality. A request built from older records, arriving after a newer receipt was already used, looked like a fresh host edit of a route that holds depended on, and reconciliation rule 3 lost every hold on that date. Any host with several desks produces such late requests. The fix is rule 0 of section 4: each technician-day remembers the last 32 timestamps the scheduler has moved past and the route from before a confirm receipt the host has not written yet, and a snapshot carrying either is answered 409 STALE before anything is reconciled. Nothing is guessed and nothing is lost; the cost is a resend. The same pull request let concurrent bookings publish: a search had refused to publish unless every day it had read was still at the version it read, so ten desks on five dates kept invalidating each other, and now it checks only the days its offers hold. [3, 21]")
    table(["Desks", "Holds lost before / after", "Offers p95 before / after (fresh seeds)", "Searches / min before / after", "STALE resends per search after"],
          [[n, f"{b['booking']['holdUnavailable']} of {b['booking']['selected']} / {a['booking']['holdUnavailable']} of {a['booking']['selected']}", f"{seconds(b['ops']['booking offers']['ms']['p95'])} / {seconds(a['ops']['booking offers']['ms']['p95'])}",
            f"{per_minute(b):.0f} / {per_minute(a):.0f}", f"{a['booking']['stale'] / a['booking']['searches']:.2f}"] for n, b, a in ((1, b1, fresh[1]), (5, b5, fresh[5]), (10, b10, fresh[10]))], [.6, 1.7, 1.9, 1.5, 1.2], size=9)
    p(f"No hold was lost in any run after the fix and every desk count was within the 5.5 s bound. At ten desks on the same five dates {c10['booking']['stale'] / c10['booking']['searches']:.1f} resends per answered search is a high price, and it is a price of the test's shape: ten desks booking the same fifty technician-days is more contention than one metro sees. From J0 the load tool gives each desk its own dates, and the resends vanish.")
    sub("The daily side: refuse early, queue fairly, split the pools (#153, #157, #158)")
    p(f"A replica ran one daily calculation at a time, and planners behind it queued for their whole allowance and then failed: {b_daily[4]['ops']['daily proposal']['statuses'].get('503', 0)} of {b_daily[4]['daily']['proposals']} proposals at four planners were a 503 after about {seconds(b_daily[4]['ops']['daily proposal']['ms']['p50'])}, and the planner learned nothing it could act on. #153 routes the day before the calculation asks for a slot, under its own 30 s deadline, and refuses at once with 429 and the estimated wait as Retry-After any request that could not start within its reserve: at four planners {a_daily[4]['ops']['daily proposal']['statuses'].get('201', 0)} of {a_daily[4]['daily']['proposals']} were answered. What remained was fairness: retries were a lottery, and on a replica that also books, queued bookings always go first. #157 added the durable queue of section 5 with per-client turns. Measured with two clients on one replica for 90 s, client A with four planners and client B with one: synchronously B got {fair[('sync', 1)]['daily']['proposals']} proposal in 90 s because A's planners took each freed slot first; queued, B got {fair[('async', 1)]['daily']['proposals']}, one every 9 to 15 s, and the total was the same ({fair[('sync', 4)]['daily']['proposals'] + fair[('sync', 1)]['daily']['proposals']} and {fair[('async', 4)]['daily']['proposals'] + fair[('async', 1)]['daily']['proposals']}), because a replica still runs one calculation at a time; the queue shares the work out rather than adding any. [3, 21]")
    figure("daily_load", "Left: proposals answered (solid) or refused with 503 (hatched) at one and four planners, before and after #153. Middle: proposals each of two clients got in 90 s on one replica, synchronously and queued (#157). Right: proposals in 60 s from one replica doing everything or from booking and daily pools at equal capacity (#158).")
    p(f"#158 split the jar into booking replicas, which take no queued daily work, and daily replicas, which may run several calculations at once, with nginx in front. At equal total search capacity (one replica doing everything against a booking replica plus a daily replica, five desks and five queued planners for 60 s, in the order all, split, split, all), daily throughput went from {roles['all-1'][1]['daily']['proposals']} and {roles['all-2'][1]['daily']['proposals']} proposals to {roles['split-1'][1]['daily']['proposals']} and {roles['split-2'][1]['daily']['proposals']}, a queued day waited {seconds(queue_wait(roles['split-1'][1])[0])} instead of {seconds(queue_wait(roles['all-1'][1])[0])} for its turn, and the booking tail fell (offers p95 {ms(roles['all-1'][0]['ops']['booking offers']['ms']['p95'])} to {ms(roles['split-1'][0]['ops']['booking offers']['ms']['p95'])}) because no daily calculation competes for the booking replica's cores. The cost was the confirm: with two booking slots instead of four, a confirm waited behind searches, and its p50 rose from {ms(roles['all-1'][0]['ops']['confirm']['ms']['p50'])} to {ms(roles['split-1'][0]['ops']['confirm']['ms']['p50'])}. [3]")
    sub("A confirm that never waits (#162)")
    p(f"After the routing work a confirm, which does a few milliseconds of work to complete a booking the customer already chose, still spent {ms(jc_b['breakdown'][CONFIRM]['queueMs']['p50'])} of its {ms(jc_b['breakdown'][CONFIRM]['elapsedMs']['p50'])} on the server waiting for a search slot. J2 changed admission in two steps: a confirm is a ticket admission chooses before any waiting search, and a confirm-only slot beyond the search capacity means it does not wait for a running search to finish either. It then measured the JVM flags the compose file sets, with a decision rule set before the runs: a default changes only on a 5 percent latency improvement. [3]")
    figure("confirm", "Confirm (left) and offers (right) at the caller, p50 and p95, from J's CH run through the four J2 runs. The small label on each confirm bar is the median slot wait the server recorded for it.")
    p(f"A confirm now answers in {ms(final['ops']['confirm']['ms']['p50'])} p50 and {ms(final['ops']['confirm']['ms']['p95'])} p95 at the caller, from {ms(jc_b['ops']['confirm']['ms']['p50'])} and {ms(jc_b['ops']['confirm']['ms']['p95'])}; priority alone took the wait from {ms(jc_b['breakdown'][CONFIRM]['queueMs']['p50'])} to {ms(j2['j2-confirm']['breakdown'][CONFIRM]['queueMs']['p50'])} and the slot took it to {ms(final['breakdown'][CONFIRM]['queueMs']['p50'])}. Offers were unchanged within run-to-run variation and the replica served {pct(final['booking']['searches'] - jc_b['booking']['searches'], jc_b['booking']['searches'])} more searches in the minute because the desks spent less of it in confirms. The garbage collector is not a lever: ZGC on the booking replica was within 3 percent on offers and worse on confirm p95 and heap, Parallel GC on the daily replica changed nothing the load tool can see, and compact object headers were within noise; G1 and the JDK defaults stay, and both pools' flags are recorded as settings for whoever measures them on the VM itself. With confirms out of the way, the booking replica's remaining wait is searches queuing for the two slots ({ms(final['breakdown'][OFFERS]['queueMs']['p50'])} p50 of a {ms(final['breakdown'][OFFERS]['elapsedMs']['p50'])} search), which is CPU: more cores, more replicas, or a cheaper search.")
    figure("progression", "Five booking desks at each stage: searches answered per minute (left) and caller latency of offers (right). Left of the dotted line, one machine with the desks alone; right of it, the VM rehearsal stack with five daily planners working beside the desks. The two halves are not one series.")

    heading("12 Inside the engine")
    p("Two studies on October 14 looked at the calculation itself: where a daily solve spends its time, and whether moving the engine out of the scheduler into a pool is worth it. [3, 22]")
    sub("More moves per second in the same budget (study L, #163)")
    p(f"A daily proposal's latency is its queue wait plus a fixed search budget, so the daily side cannot be made faster by the load tool's measure; it can only be made better in the same seconds by evaluating more candidate moves. The benchmark harness now starts a JDK Flight Recorder recording just before each solve and stops it just after, so the samples cover the solve alone. The first profile (Tabu + K opt, 20 and 50 technicians, clustered and dispersed, a fixture seed never used for tuning, 10 s per case) showed that of {prof['before']['samples']:,} samples the engine's own timing methods held about a tenth; the rest was what scoring did around them. Every scored route looked its road legs up in an immutable map keyed by a from-to string built for each leg ({prof['before']['groups']['Immutable map probe']:.1f} percent of samples), built a HashMap of arrivals and a day plan that were then copied into a result ({prof['before']['groups']['HashMap and copies']:.1f} percent), and computed its buffered travel minutes in BigDecimal per leg ({prof['before']['groups']['BigDecimal arithmetic']:.1f} percent). [22]")
    p("The treatment keeps every number the same and removes the work around it. RoadMatrix gives each endpoint identity a number and holds the legs in two tables with their buffered minutes computed once per fact revision, built lazily on the first scoring lookup, so a plan evaluated once outside the solver (a booking candidate) pays nothing new. The timing search evaluates a route against the plan facts directly and, on the score-only path, builds no arrival map. A scoring exactness test captured a step-limited, fixed-seed solve on three fixtures before the change and asserts the same score, move count, arrangement and independent evaluation after it. Then the engine built from the commit before the change was paired against the treated one on 24 cases, one after the other on an idle workstation.")
    figure("engine", "Left: each paired case's cost-phase moves per second before (x) and after (y) the RoadMatrix treatment; the dotted line is no change and the dashed line is twice the rate. Right: offers at the caller for the embedded engine and three remote pool configurations in study K, with the daily proposals each pool answered.")
    fleet_rows = []
    for fleet in (20, 50):
        fp = [x for x in pairs if x["technicians"] == fleet]
        wins = sum(x["afterCostCents"]["after"] < x["afterCostCents"]["before"] for x in fp); ties = sum(x["afterCostCents"]["after"] == x["afterCostCents"]["before"] for x in fp)
        fleet_rows.append([f"{fleet} technicians", f"{median(x['costPhaseMovesPerSecond']['before'] for x in fp) / 1000:.0f}k to {median(x['costPhaseMovesPerSecond']['after'] for x in fp) / 1000:.0f}k ({summary['costPhaseSpeedupByFleet'][str(fleet)]['medianCostPhaseSpeedup']:.2f}x)",
                           f"{median(x['fairnessPhaseSpeedup'] for x in fp):.2f}x", f"{wins} of 12" + (f" ({ties} ties, {12 - wins - ties} dearer)" if wins < 12 else "")])
    table(["Fleet", "Cost-phase moves per second, median before to after", "Fairness-phase speedup, median", "Cases cheaper after at the same budget"], fleet_rows, [1.2, 2.3, 1.5, 1.9])
    loss = [x for x in pairs if x["afterCostCents"]["after"] > x["afterCostCents"]["before"]]
    p(f"The cost phase evaluates {summary['costPhaseSpeedup']['median']:.2f} times the moves per second (median; {summary['costPhaseSpeedup']['min']:.2f} to {summary['costPhaseSpeedup']['max']:.2f} across the 24 cases), more on 50 technicians where each scored route had more legs to look up. The fairness phase gains {summary['fairnessPhaseSpeedup']['median']:.2f} times, less, because its time goes to the fairness term and the BigDecimal score, which the treatment leaves alone. Because a solve now evaluates more moves in the same seconds, {summary['afterCostCentsWinTieLoss']['afterWins']} of 24 cases ended cheaper at the same budget, {summary['afterCostCentsWinTieLoss']['ties']} tied and {len(loss)} was dearer by {loss[0]['afterCostCents']['after'] - loss[0]['afterCostCents']['before']} cents, with no constraint violation on either side. This is not a search-policy change: the acceptor, moves and tabu size are what section 8 chose, so the promotion review that belongs to policy changes was not run for it. After the treatment the profile's largest shares are BigDecimal ({prof['after']['groups']['BigDecimal arithmetic']:.1f} percent) and the clock arithmetic of Instant and Duration ({prof['after']['groups']['Instant and Duration']:.1f} percent); section 13 names the next step. [22]")
    sub("The embedded engine against a remote pool (study K, #164)")
    p(f"The design's third scaling lever moves the calculation out of the scheduler into a solver pool (SCHEDULER_CALCULATION_MODE=REMOTE), and it had never been measured. K compared the two on the rehearsal stack under the same load as J2, with the pool at its compose shape (4 cores) beside the booking and daily replicas at 2 cores each, and an acceptance set beforehand: remote p95 at most 1.05 times embedded. Booking fails it. A remote search costs {k['k-remote5']['ops']['booking offers']['ms']['p50'] / k['k-embedded']['ops']['booking offers']['ms']['p50']:.1f} times the embedded one at the caller (p50 {k['k-remote5']['ops']['booking offers']['ms']['p50']} against {k['k-embedded']['ops']['booking offers']['ms']['p50']} ms, p95 {k['k-remote5']['ops']['booking offers']['ms']['p95']} against {k['k-embedded']['ops']['booking offers']['ms']['p95']} ms) and the replica serves two thirds of the searches, because a search makes up to three solver calls (insertion, refinement, and the retry on a moved day), each a snapshot encoded, sent, decoded and answered: the engine share of a search goes from {ms(k['k-embedded']['breakdown'][OFFERS]['calculationMs']['p50'])} to {ms(k['k-remote5']['breakdown'][OFFERS]['calculationMs']['p50'])}. So the embedded engine stays the default. [3]")
    p(f"For a daily pool, remote is the lever it was meant to be: the proposal takes the same {seconds(k_d['k-remote5']['ops']['daily proposal']['ms']['p50'])} p50 (queue wait plus a fixed 5 s search) and the daily replica's CPU falls from {daily_cpu(k['k-embedded']):.0f} to {daily_cpu(k['k-remote5']):.0f} percent of its two cores, because the solve runs in the pool. Each pool has its own calculation mode, so the daily pool can go remote while booking stays embedded. The study also found and fixed a sizing fault: with the pool's previous defaults {k_d['k-remote']['ops']['daily proposal']['statuses']['503']} of {k_d['k-remote']['daily']['proposals']} daily proposals answered 503, because the solver service admitted one daily calculation at a time whatever its capacity. The pool now opens every slot to daily work by default, and the rule for an operator is that a pool's capacity is the served replicas' search capacities summed.")

    heading("13 What comes next: performance")
    p(f"After J2 the median booking search spends {ms(final['breakdown'][OFFERS]['queueMs']['p50'])} waiting for a slot, {ms(final['breakdown'][OFFERS]['routingMs']['p50'])} routing, {ms(final['breakdown'][OFFERS]['dbMs']['p50'])} in the database and {ms(final['breakdown'][OFFERS]['calculationMs']['p50'])} in the engine, and a confirm no longer waits at all. The daily search does {summary['costPhaseSpeedup']['median']:.1f} times the work in its budget, and a remote pool can take that work off the daily replica. The evidence supports the following, in this order. [3, 22]")
    table(["Question", "Next evidence to collect", "Decision informed"], [
        ("The VM itself", "Repeat J0, J2 and K across the network on the real VM with the load generator on the host's side; every number in sections 10 to 12 is from a rehearsal stack whose CPU limits are a shape, not the VM's share.", "The latency target measured where it is defined."),
        ("Booking cores", "The slot wait is half of a search; measure five and ten desks at 2, 3 and 4 booking cores, and with a second booking replica.", "The booking pool's size on the VM."),
        ("Daily pool remote", "Run the daily pool REMOTE on the VM with the pool sized by the rule in section 12; confirm 25 of 25 proposals and the replica's CPU.", "Whether the daily pool goes remote at go-live."),
        ("Primitive time in the timing search", "BigDecimal and Instant/Duration are now the largest shares of a solve; replace them with integer seconds in the route timing search, gated by the scoring exactness test, then repeat the 24-case paired study.", "Whether a 5 s search can do still more."),
        ("Multithreaded solving", "The daily search is single threaded today: parallelism comes from several solves at once per daily replica, independent solves per core in the benchmarks, and parallel pair routing. Evaluate Timefold's moveThreadCount for the daily search (the campaign configuration currently refuses any value but NONE and the production configuration sets none), with the paired study and the exactness test as the gate, and the daily replica's core count as the constraint: a solve that uses two cores halves the calculations a replica can run at once.", "Whether one solve should use more than one core, or the replica should keep running several."),
        ("Fewer solver round trips", "Carry insertion and refinement in one solver call; only if booking is ever to use a pool.", "Whether remote booking could meet the 1.05 bound.")], [1.4, 3.6, 1.9], size=9)
    p("Two cautions carry over from every measurement. Compare only fresh-seed runs, because a seed run before leaves its roads cached and makes a one-desk search about a quarter of a second faster. And read the VM-shaped numbers as a shape, not a capacity: the real VM's share will set its own numbers, and the latency target is defined at the WaterFlex caller across the network, which has not been measured yet.")

    heading("14 What comes next: solver quality")
    p(f"The daily search studies ended with a candidate that beats the current default on every new fixture from 10 technicians up, and with three things the fixtures cannot answer. [1, 2]")
    table(["Decision", "Evidence", "Suggested next step"], [
        ("Offer Tabu + K opt size 3?", f"{dollars(mean(direct))} per case more than Tabu + K opt on 40 unseen fixtures; ahead from 10 technicians; slightly worse at 5; 16 of 16 when run alone.", "Add it as a third daily search choice beside TABU_KOPT and TABU (it exists today as an experiment configuration only), then replay real customer days before making it anyone's default."),
        ("Real schedules", "Every fixture is synthetic, with a synthetic travel matrix; the one study on real Omaha roads was a booking study.", "Replay recorded days with measured road travel through both Tabu + K opt variants; the daily proposal endpoint accepts a recorded snapshot as it is."),
        ("Balance", f"Cheaper methods are less even without a fairness search ({mean(r['variance'] for r in by(M, 'TABU_KOPT_SIZE_3')):.4f} against {mean(r['variance'] for r in by(M, 'TABU')):.4f} mean variance in the follow-up); with one, most of the gap closes.", "Measure Tabu + K opt size 3 with the fairness search on before adoption, and decide whether a client's fairness budget should carry an explicit acceptable variance."),
        ("Longer daily searches?", f"K opt passes Tabu + K opt at 50 technicians only from about 480 s ({dollars(saving(L, 'KOPT', fleet=50, budget_ms=480000))} at 480 s, {dollars(saving(L, 'KOPT', fleet=50, budget_ms=240000))} at 240 s against Tabu + K opt).", "Keep the 240 s client limit; revisit only for overnight runs, where the queue of section 5 already makes a long search harmless to other clients."),
        ("A fleet-size rule?", "The first study's hindsight rule (K opt to 20 technicians, Tabu + K opt above) held up on new fixtures, and one method matches it.", "No selector; one method per client."),
        ("Booking search quality", "The booking benchmarks of September measured the retired internal endpoints; the load tool measures speed, not offer quality.", "A booking quality study against /api/v1 on recorded days, with the served-customer denominator the September studies insisted on.")], [1.4, 2.9, 2.6], size=9)
    p("The thread that runs through both of these sections is the one that ran through the whole month: measure where the time or the money actually goes before changing anything, change one thing, and keep it only when a paired measurement on fresh inputs says so. The September CH graph and the October CH graph were the same graph; what changed between them was the evidence.")

    heading("Appendix A Studies and evidence", new_page=True)
    rows = [("2026-09-23", "Sparse routing check", "docs/evidence/sparse-routing-omaha-2026-09-23.json", "Six directed Omaha pairs; sparse against full matrix; API samples"),
            ("2026-09-24", "Flexible against CH, routing alone", "docs/evidence/routing-ch-omaha-2026-09-24.json", "132 pairs, map 2026-09-16, 1/5/10 callers"),
            ("2026-09-24", "Flexible against CH, booking", "docs/evidence/booking-omaha-routing-2026-09-24.json", "720 requests per graph on real Omaha roads"),
            ("2026-09-24", "Insertion against bounded", "docs/booking-experiment-explainer/evidence.json", f"{b24['cases']} cases, {b24['requests']:,} requests; 60 paired sequential cases"),
            ("2026-09-29", "Field validation", "docs/evidence/scheduler-field-2026-09-29/summary.json", f"{f29['held_cases']} held-out booking cases per variant; {f29['daily_cases']} daily cases, paired means by budget"),
            ("2026-10-01", "Daily budget study", "docs/daily-budget-331df1f1 (report only; run archive not in the repository)", "1,280 cases; cited, not recomputed"),
            ("2026-10-11", "Tabu variant study", "experiments/runs/20261010T040557Z-daily-tabu-variants-p6-5d15129b", f"{len(F):,} cases on {daily['first']['fixtures']} fixtures; {len(FC)} sequential confirmation cases"),
            ("2026-10-13", "Daily search follow-up", "experiments/runs/20261013T052707Z-daily-search-followup-6454fb05 (+ long budget, confirmation)", f"{len(M):,} + {len(L)} + {len(C)} cases on {daily['follow']['fixtures']} new fixtures")]
    rows += [(re.search(r"\d{4}-\d{2}-\d{2}", name).group(), f"Load study {pr}", f"docs/evidence/{name}", note) for key, name, pr, note in perf["studies"]]
    rows += [("2026-10-14", "Solver profiling (#163)", "docs/evidence/solver-profiling-2026-10-14.json", "Two solve-scoped JFR profiles; 24 paired cases"),
             ("2026-10-14", "Flexible against CH, routing alone", "docs/evidence/routing-ch-2026-10-14.json", "132 pairs, map 2026-10-13, parallel pairs")]
    table(["Date", "Study", "File", "What it holds"], rows, [.8, 1.6, 2.7, 1.8], size=9)
    p("Every load run used the load tool infra/host-conformance/load.ts on one 24-processor workstation (AMD Ryzen 9 7900X, 63 GiB) running the scheduler, the Omaha routing service, Postgres and the load generator; the solver studies ran on the same machine with the experiment runtime's daily flags. Studies #151 to #157 ran the scheduler jar alone; #158 to #164 ran the deploy/vm rehearsal stack with its CPU limits beside five queued daily planners, and the two halves are never read as one series.")

    heading("Appendix B Definitions")
    table(["Term", "Meaning in this document"], [
        ("Snapshot", "The facts of one metro's technician-days a calculating call carries: rates, policy, technicians, days with lastModified, committed appointments, locations."),
        ("Thin store", "The scheduler's own tables: holds, offer sets, proposals, receipts, request claims, tenants and tokens, the route cache; nothing the host owns."),
        ("Input revision", "A canonical hash of a snapshot; every proposal and offer records the one it was computed from."),
        ("STALE", "409 answer to a snapshot older than records the scheduler already used; the host rebuilds from current records and resends."),
        ("Slot", "One of a replica's concurrent search permits (SCHEDULER_SEARCH_CAPACITY, 2 in every run here); a confirm slot is a permit only a confirm may take."),
        ("Desk, planner", "One concurrent booking loop, or one concurrent daily loop, in the load tool."),
        ("Fresh seed", "A load run whose random addresses no earlier run used, so its roads are not cached anywhere."),
        ("Matched case", "The same starting day, seed and budget solved by two configurations, so the results differ only in the configuration."),
        ("Tabu search", "A local search that keeps a short memory of recently moved entities and will not move them back at once, and that accepts the best available step even when it is worse."),
        ("K opt move", "A move that reconnects several route links at once to explore coordinated order changes; the sublist move moves a stretch of visits together."),
        ("Late acceptance", "A local search that accepts a move if it beats the solution from a fixed number of steps ago, even when it worsens the current one."),
        ("Insertion, bounded search", "Insertion tries the new visit in every position without moving existing work; bounded search also tries limited rearrangements around the affected routes."),
        ("Contraction hierarchy", "A road graph preprocessed with shortcut edges ordered by importance, so that a query touches few nodes; identical answers, less work."),
        ("Routing identity", "A hash of the road graph and map version; every cached leg is bound to one, and a change clears the caches."),
        ("Moves per second", "Candidate moves the daily search evaluated per second of its phase."),
        ("p50, p95", "The latency that 50 or 95 percent of calls in a run stayed within, at the caller, waits included.")], [1.5, 5.4], size=9)

    heading("Appendix C Sources")
    p("[1] The Tabu variant study: docs/daily-tabu-variants-5d15129b (report, EVIDENCE.md and its validated loader, which this document imports).")
    p("[2] The daily search follow-up: docs/daily-search-followup-6454fb05 (report, EVIDENCE.md and loader), with the decision record of its driver.")
    p("[3] The performance work: docs/performance-work-2026-10-14 (report, EVIDENCE.md and loader, imported here) over the load studies of Appendix A; docs/api-load.md is the narrative kept with the code.")
    p("[4] docs/stateless-api-design.md: the design, its decisions marked Decided, the endpoint, contract change, thin store, lifetime and route cache tables read at build time.")
    p("[5] The design proposal as first written, git revision 08dae8b of the same file: the description of the scheduler before the public API, from which the comparison table of section 1 is drawn.")
    p("[6] docs/operations-runbook.md: deployment, settings, roles, health, logs, metrics, retention, maps.")
    p("[7] docs/api/openapi-v1.yaml, contract version " + docs["version"] + ", and docs/integration-guide.md: the operations, error codes, ownership table and the lastModified rules.")
    p("[8] The scheduler's booking path: scheduler-service/src/main/java/dev/waterflex/scheduler/api (BookingOffers, BookingReconciliation, BookingConfirm, DailyProposals, DailyPreparation), SearchAdmission and RoadClient in the calculation engine and scheduler.")
    p("[9] docs/scheduler-policy.md: the policy the code enforces; docs/solver-benchmarks.md: the daily solver protocol, the September 24 comparison and the profiling section.")
    p("[10] Timefold Solver documentation, local search: acceptors, foragers and move selectors, as configured in the recorded solver XML of every case.")
    p("[11] docs/scheduler-algorithm-comparison-2026-09-29.md and its PDF: the held-out booking and daily comparison; the fleet split of K opt's 15 s result is quoted from it.")
    p("[12] docs/daily-budget-331df1f1: the October 1 budget study, cited from its report; its run archive is not in the repository.")
    p("[13] Pull requests #133 (per-client daily search choice and search times) and #134 (fixture seeds, cost-only mode and Tabu + K opt size 3).")
    p("[14] docs/booking-experiment-explainer: the September 24 insertion against bounded study and its evidence.json, whose calculations this document imports.")
    p("[15] docs/scheduler-field-validation.md and docs/evidence/scheduler-field-2026-09-29: the field validation of the policy and search variants.")
    p("[16] The searchComplete, shedding and refinement fallback changes: pull requests #154, #155 and #138, with the result blocks of their evidence files.")
    p("[17] docs/scheduler-policy-progress.md (the September routing work and the Omaha CH experiment), docs/evidence/routing-ch-omaha-2026-09-24.json and docs/evidence/routing-ch-2026-10-14.json, produced by infra/benchmark-routing-ch.mjs.")
    p("[18] docs/evidence/booking-omaha-routing-2026-09-24.json and docs/evidence/sparse-routing-omaha-2026-09-23.json.")
    p("[19] The request log line (docs/operations-runbook.md, Logs): queueMs, routingMs, calculationMs, dbMs, routingPairs, cacheHits, fetched and cpuMs.")
    p("[20] GraphHopper documentation on contraction hierarchies, for the mechanism; the measurements are the project's own.")
    p("[21] The deterministic reproduction of the lost hold is recorded as staleSnapshotReproduction in the #151 and #152 evidence, with its control.")
    p("[22] docs/evidence/solver-profiling-2026-10-14.json and docs/solver-benchmarks.md, 'Solve-scoped profiling and the CR07 treatment'; the exactness gate is scheduler-service's ScoringExactnessTest.")
    p("Reproduce with docs/scheduler-explained-2026-10-14/requirements.txt, then run report_data.py, build_report.py and verify_report.py there. The builder reads every source, regenerates all charts and chart_data.json, and writes this editable Word file. EVIDENCE.md records methods and limits.")

    settings = doc.settings.element
    for el in settings.findall(qn("w:updateFields")): settings.remove(el)
    update = OxmlElement("w:updateFields"); update.set(qn("w:val"), "true"); settings.append(update)
    stream = BytesIO(); doc.save(stream)
    with ZipFile(REFERENCE) as ref, ZipFile(stream) as built, ZipFile(OUTPUT, "w", ZIP_DEFLATED) as output:
        preserve = set(PRESERVE) | {n for n in ref.namelist() if "footer" in n}
        for name in sorted(built.namelist()):
            content = ref.read(name) if name in preserve else built.read(name)
            info = ZipInfo(name, (2026, 10, 14, 0, 0, 0)); info.compress_type = ZIP_DEFLATED
            output.writestr(info, content)
    print(OUTPUT, len(figures), "figures")


if __name__ == "__main__": build()
