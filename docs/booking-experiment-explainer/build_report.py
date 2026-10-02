"""Build the Word report from reviewed evidence, prose and the retained reference.

No benchmark, database, network or service access. Run after the daily experiment.
"""
from copy import deepcopy
from datetime import datetime
from hashlib import sha256
from io import BytesIO
import json
import math
from pathlib import Path
import re
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

from report_data import SOLVERS, choose, outcomes, pairs, savings, verify

ROOT=Path(__file__).resolve().parent
REFERENCE=ROOT.parent/'scheduler-explainer/reference.docx'
OUTPUT=ROOT/'WaterFlex_Booking_Experiment_Results_Explained.docx'
REFERENCE_HASH='2d296b4ae319f04ab583a25a23df6e11e845e6d1ba3324f7dcad23d96d53df70'
COLORS=('#00959b','#e45d2a')

def num(value, places=2):
    return 'Unavailable' if value is None else f'{value:,.{places}f}'

def make_chart(path, labels, series, title, unit):
    """Simple horizontal bars with direct values and a zero origin."""
    image=Image.new('RGB',(1500,500),'white'); draw=ImageDraw.Draw(image)
    fonts=[Path('C:/Windows/Fonts/calibri.ttf'),Path('/usr/share/fonts/truetype/crosextra/Carlito-Regular.ttf')]
    font_path=next((p for p in fonts if p.exists()),None)
    if font_path is None:
        raise RuntimeError('Calibri or Carlito required for report chart typography')
    font=lambda size:ImageFont.truetype(str(font_path),size)
    draw.text((20,10),title,font=font(34),fill='black')
    x0,x1,y0,step=280,1260,110,100
    maximum=max(v for _, values in series for v in values)
    limit=100 if unit.startswith('Percent') else max(10, math.ceil(maximum/10)*10)
    for k in range(5):
        value=limit*k/4; x=x0+(x1-x0)*k/4
        draw.line((x,y0-10,x,410),fill='#d9d9d9',width=2)
        draw.text((x-18,421),f'{value:g}',font=font(24),fill='black')
    for n,label in enumerate(labels):
        y=y0+n*step
        draw.text((20,y+15),label,font=font(29),fill='black')
        for j,(name,values) in enumerate(series):
            yy=y+j*43; xx=x0+values[n]/limit*(x1-x0)
            draw.rectangle((x0,yy,xx,yy+30),fill=COLORS[j])
            draw.text((xx+12,yy-3),f'{values[n]:.2f}',font=font(26),fill='black')
    draw.text((x0,465),unit,font=font(26),fill='black')
    for j,(name,_) in enumerate(series):
        draw.rectangle((800+j*270,65,825+j*270,87),fill=COLORS[j])
        draw.text((838+j*270,59),name,font=font(26),fill='black')
    image.save(path)

def build():
    if sha256(REFERENCE.read_bytes()).hexdigest()!=REFERENCE_HASH:
        raise ValueError('Reference changed')
    e=json.loads((ROOT/'evidence.json').read_text(encoding='utf-8'))
    pp=verify(e); rows=e['cases']
    assets=ROOT/'generated'; assets.mkdir(exist_ok=True)
    # Reuse the source page system and styles; replace all body content.
    doc=Document(REFERENCE); table_properties=deepcopy(doc.tables[0]._tbl.tblPr)
    for child in list(doc.element.body):
        if child.tag!=qn('w:sectPr'): doc.element.body.remove(child)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Caption'):
        style=doc.styles[name]; style.font.name='Calibri'; style.font.color.rgb=RGBColor(0,0,0)
    doc.core_properties.title='WaterFlex Booking Experiment Results Explained'
    doc.core_properties.subject='Matched booking efficiency and response time in the September 30 archive'
    doc.core_properties.author='WaterFlex'; doc.core_properties.last_modified_by='WaterFlex'
    doc.core_properties.created=doc.core_properties.modified=datetime(2026,9,30,0,0,0)

    new_page=False
    def paragraph(text,style=None):
        nonlocal new_page
        p=doc.add_paragraph(text,style); p.paragraph_format.widow_control=True
        if new_page:
            p.paragraph_format.page_break_before=True
            new_page=False
        return p

    def table(headers,data,widths):
        t=doc.add_table(rows=1,cols=len(headers)); t.autofit=False; t.alignment=WD_TABLE_ALIGNMENT.CENTER
        t._tbl.replace(t._tbl.tblPr,deepcopy(table_properties))
        for col,w in zip(t.columns,widths): col.width=Inches(w)
        for index,values in enumerate([headers,*data]):
            row=t.rows[0] if index==0 else t.add_row()
            trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
            if index==0: trpr.append(OxmlElement('w:tblHeader'))
            for j,(cell,value,width) in enumerate(zip(row.cells,values,widths)):
                cell.width=Inches(width); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell.text=str(value); props=cell._tc.get_or_add_tcPr()
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6EF' if index==0 else ('FFFFFF' if index%2 else 'F4F6F8'));props.append(shade)
                margins=OxmlElement('w:tcMar')
                for side in ('top','left','bottom','right'):
                    n=OxmlElement('w:'+side);n.set(qn('w:w'),'65');n.set(qn('w:type'),'dxa');margins.append(n)
                props.append(margins)
                p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
                p.paragraph_format.line_spacing=1.0
                if len(headers)>=3 and j>0:p.alignment=1
                for run in p.runs:run.font.name='Calibri';run.font.size=Pt(10);run.bold=index==0
        spacer=paragraph('');spacer.paragraph_format.space_after=Pt(0)
        spacer.paragraph_format.space_before=Pt(0);spacer.paragraph_format.line_spacing=1
        spacer.add_run('').font.size=Pt(4)

    def picture(path,width=6.7,caption=None):
        p=paragraph(''); p.add_run().add_picture(str(path),width=Inches(width))
        p.paragraph_format.keep_with_next=caption is not None
        if caption:paragraph(caption,'Caption')

    def generated_table(name):
        if name=='summary':
            table(['Question','What this archive supports'],[
                ('Route efficiency','60 comparable sequential pairs; 120 concurrent pairs excluded.'),
                ('Booking outcomes','600 of 600 served per method at concurrency 1; many incomplete searches at 5 and 10.'),
                ('Search wait','Bounded search is slower, especially at higher concurrency.')],[1.5,5.4])
        elif name=='teaching':
            table(['Illustrative paid minutes','Insertion','Bounded'],[
                ('Alex',260,230),('Blair',220,230),('Fleet total',480,460),('Same new service',30,30),('Fleet saving with bounded','',20)],[3.9,1.5,1.5])
        elif name=='setup':
            table(['Dimension','Values'],[('Methods','Insertion and bounded search'),('Fleet sizes','5, 10, 20 and 50 technicians'),('Workloads','Clustered, dispersed and sparse'),('Concurrency','1, 5 and 10 simultaneous requests'),('Seeds','17, 23, 41, 59 and 83'),('Requests and cache','10 requests per case; warm scheduler cache')],[1.7,5.2])
        elif name=='units':
            table(['Measure','Unit and interpretation'],[
                ('Paid route time','Minutes within paid working segments, summed across the schedule.'),
                ('Road driving','Recorded road seconds divided by 60; excludes buffers and rounding.'),
                ('Buffers and rounding','Separate recorded seconds divided by 60; included in paid time.'),
                ('Waiting and overtime','Minutes; overtime is part of paid time, not an extra time component.'),
                ('Modeled cost','Recorded cents divided by 100; labor plus mileage.'),
                ('Search latency','Individual request observation in milliseconds, displayed here in seconds.')],[1.65,5.25])
        elif name=='efficiency':
            data=[]
            for w in ('CLUSTERED','DISPERSED','SPARSE'):
                for f in (5,10,20,50):
                    group=[p for p in pp if p['bounded']['workload']==w and p['bounded']['fleet']==f and p['comparable']]
                    data.append((w.title(),f,num(savings(group,'paid_min')),num(savings(group,'road_sec')/60),num(savings(group,'cost_cents')/100)))
            table(['Workload','Fleet','Paid min saved','Road min saved','Cost saved $'],data,[1.65,.65,1.55,1.55,1.5])
        elif name=='worked':
            p=next(p for p in pp if p['bounded']['fleet']==5 and p['bounded']['workload']=='DISPERSED' and p['bounded']['seed']==17 and p['comparable'])
            i,b=p['insertion'],p['bounded']
            table(['Whole schedule measure','Before either','Insertion after','Bounded after'],[
                ('Paid minutes',num(i['before']['paid_min'],0),num(i['after']['paid_min'],0),num(b['after']['paid_min'],0)),
                ('Modeled cost $',num(i['before']['cost_cents']/100),num(i['after']['cost_cents']/100),num(b['after']['cost_cents']/100)),
                ('Road driving minutes',num(i['before']['road_sec']/60),num(i['after']['road_sec']/60),num(b['after']['road_sec']/60)),
                ('Waiting minutes',0,0,0),('Overtime minutes',0,0,0)],[2.3,1.5,1.55,1.55])
        elif name in ('outcomes','latency'):
            data=[]
            for c in (1,5,10):
                for s in SOLVERS:
                    o=outcomes(choose(rows,concurrency=c,solver=s))
                    data.append((c,s.title(),o['served'],o['incomplete'],o['failed']) if name=='outcomes' else
                                (c,s.title(),o['latency_n'],num(o['p50']/1000),num(o['p95']/1000)))
            table(['Concurrency','Method','Served /600','Incomplete /600','Failed /600'] if name=='outcomes' else
                  ['Concurrency','Method','Samples','p50 seconds','p95 seconds'],data,[1.1,1.6,1.4,1.4,1.4])
        elif name=='glossary':
            table(['Term','Meaning in this report'],[
                ('Case','One method applied to ten requests under one fixed setup.'),
                ('Scenario group','A fleet, workload and concurrency combination with both methods and five seeds.'),
                ('Seed','A reproducible input to the fixture generator; not an independent geography.'),
                ('Matched pair','Same conditions and served request identities, with comparable audited schedules.'),
                ('Incomplete','The prescribed search did not finish; feasibility remains unresolved.'),
                ('p50 and p95','Nearest-rank median and 95th percentile of individual observations.'),
                ('Seed range','Smallest to largest observed seed-level rate; not a confidence interval.'),
                ('Modeled cost','Labor and mileage computed from configured rates, not an invoice or payroll total.')],[1.5,5.4])
        else:raise ValueError(name)

    def charts(name):
        path=assets/(name+'.png')
        if name=='efficiency_chart':
            labels=['Clustered','Dispersed','Sparse']; vals=[savings([p for p in pp if p['bounded']['workload']==w],'paid_min') for w in ('CLUSTERED','DISPERSED','SPARSE')]
            make_chart(path,labels,[('Paid minutes saved',vals)],'Average paid minutes saved per matched case','Minutes saved; 20 matched pairs per workload')
        else:
            metric='served' if name=='outcomes_chart' else 'p95'
            series=[]
            for s in SOLVERS:
                values=[outcomes(choose(rows,solver=s,concurrency=c))[metric]/(6 if metric=='served' else 1000) for c in (1,5,10)]
                series.append((s.title(),values))
            make_chart(path,['Concurrency 1','Concurrency 5','Concurrency 10'],series,
                       'Booking success' if metric=='served' else 'Search p95 by concurrency',
                       'Percent served; 600 attempts per method at each level' if metric=='served' else 'Seconds; 600 search observations per method at each level')
        picture(path,caption='Calculated from source 1. Bars start at zero; values are descriptive local results.')

    def original(name):
        _,number,side=name.split('_'); path=ROOT/'original-figures'/f'figure-{number}.png'
        im=Image.open(path); w,h=im.size
        # Original panels, including their own axes and legend. Main title appears in prose.
        box=(0,55,int(w*.5),h) if side=='left' else (int(w*.5),55,w,h)
        out=assets/(name+'.png');im.crop(box).save(out)
        picture(out,6.6,caption=f'Original figure {number}, {"outcome" if side=="left" else "latency"} panel. Source 2.')

    def appendix():
        nonlocal new_page
        for f in (5,10,20,50):
            new_page=True;paragraph(f'Appendix {f} technician scenarios','Heading 1')
            paragraph('Each scenario below has five cases and 50 requests per method. Entries show insertion / bounded search. Latency percentiles pool the 50 individual searches; times are seconds. Route savings are average insertion minus bounded values over five matched seeds. Excluded means different served customer sets, not zero savings.')
            data=[]
            for w in ('CLUSTERED','DISPERSED','SPARSE'):
                for c in (1,5,10):
                    oo=[outcomes(choose(rows,fleet=f,workload=w,concurrency=c,solver=s)) for s in SOLVERS]
                    pairgroup=[p for p in pp if p['bounded']['fleet']==f and p['bounded']['workload']==w and p['bounded']['concurrency']==c]
                    data.append((w.title(),c,f"{oo[0]['served']} / {oo[1]['served']}",
                                 f"{num(oo[0]['p50']/1000)} / {num(oo[1]['p50']/1000)}",
                                 f"{num(oo[0]['p95']/1000)} / {num(oo[1]['p95']/1000)}",
                                 num(savings(pairgroup,'paid_min')) if c==1 else 'Excluded',
                                 num(savings(pairgroup,'cost_cents')/100) if c==1 else 'Excluded'))
            table(['Workload','Conc.','Served /50','p50 I / B','p95 I / B','Paid min saved','Cost $ saved'],data,[1.1,.5,.85,1.15,1.15,1.1,1.05])
            paragraph('Incomplete and failed outcomes for the same groups','Heading 2')
            data=[]
            for w in ('CLUSTERED','DISPERSED','SPARSE'):
                for c in (1,5,10):
                    oo=[outcomes(choose(rows,fleet=f,workload=w,concurrency=c,solver=s)) for s in SOLVERS]
                    data.append((w.title(),c,f"{oo[0]['incomplete']} / {oo[1]['incomplete']}",f"{oo[0]['failed']} / {oo[1]['failed']}"))
            table(['Workload','Concurrency','Incomplete I / B','Failed I / B'],data,[1.65,1.25,2,2])
            paragraph('Counts use 50 attempts per method. Unknown completion: zero throughout. All 50 search timings per method are available in every group. Waiting and overtime: zero in all final audits. Source 1.','Caption')

    text=(ROOT/'report.md').read_text(encoding='utf-8')
    if '\u2014' in text:raise ValueError('Em dash in prose')
    text=re.sub(r'(?m)^(\[\[[a-z0-9_]+\]\])$',r'\n\1\n',text)
    for block in re.split(r'\n\s*\n',text.strip()):
        block=block.strip()
        if block=='[[page]]':new_page=True
        elif block=='[[appendix]]':appendix()
        elif block.startswith('[['):
            name=block[2:-2]
            if name.startswith('original_'):original(name)
            elif name.endswith('_chart'):charts(name)
            else:generated_table(name)
        elif block.startswith('### '):paragraph(block[4:],'Heading 2')
        elif block.startswith('## '):paragraph(block[3:],'Heading 1')
        elif block.startswith('# '):
            title,*rest=block.split('\n');paragraph(title[2:],'Title')
            if rest:paragraph(' '.join(rest),'Subtitle')
        else:paragraph(block.replace('\n',' '))
    # Keep source footer page field. Update it on opening in Word.
    settings=doc.settings.element; update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
    raw=BytesIO();doc.save(raw);raw.seek(0)
    with ZipFile(raw) as src,ZipFile(OUTPUT,'w',ZIP_DEFLATED) as dst:
        for name in sorted(src.namelist()):
            info=ZipInfo(name,(2026,9,30,0,0,0));info.compress_type=ZIP_DEFLATED;dst.writestr(info,src.read(name))
    print(OUTPUT)

if __name__=='__main__':build()
