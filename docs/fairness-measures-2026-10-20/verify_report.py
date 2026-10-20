"""Check the delivered tables, chart data and template parts against the raw evidence file.

Recomputes from docs/evidence/solver-variance-2026-10-20.json without the builder's or loader's helpers. Run normally,
without Python's -O flag. Visual pagination still requires rendering.
"""
import json
from decimal import Decimal
from hashlib import sha256
from statistics import median
from zipfile import ZipFile

from docx import Document
from build_report import HERE, OUTPUT, PRESERVE, REFERENCE, REFERENCE_SHA

EVIDENCE = HERE.parent / "evidence" / "solver-variance-2026-10-20.json"
NAMES = {"b": "b (binary64)", "c": "c (running sums)"}
MEASURES = ("variance", "maximumUtilization", "utilizationSpread", "absoluteDeviation", "loadBalance")
TITLES = {"variance": "Variance", "maximumUtilization": "Highest utilization", "utilizationSpread": "Utilization spread",
          "absoluteDeviation": "Absolute deviation", "loadBalance": "Load balance"}
OWN = {"min-max": "maximumUtilization", "spread": "utilizationSpread", "abs-dev": "absoluteDeviation", "load-balance": "loadBalance"}
SECONDS = {"15000": "15 s", "60000": "60 s"}


def signed(value):
    return f"{value:+,.1f}" if value != int(value) else f"{int(value):+,}"


def thousandths(value):
    text = f"{Decimal(value) * 1000:.2f}"
    return "0.00" if text == "-0.00" else text


def verify():
    raw = json.loads(EVIDENCE.read_text(encoding="utf-8-sig"))
    doc = Document(OUTPUT)
    assert len(doc.tables) == 22 and len(doc.inline_shapes) == 13 and len(doc.sections) == 1, (len(doc.tables), len(doc.inline_shapes))
    assert len([p for p in doc.paragraphs if p.style.name == "Heading 1"]) == 13
    captions = [p for p in doc.paragraphs if p.style.name == "Caption"]
    assert len(captions) == 13
    for i, p in enumerate(captions, 1):
        assert p.text.startswith(f"Figure {i}.") and all(str(r.font.color.rgb) == "000000" and r.italic for r in p.runs)
    for shape in doc.inline_shapes:
        assert shape._inline.docPr.get("descr", "").startswith("Figure ")

    def rows(index):
        return [[c.text for c in r.cells] for r in doc.tables[index].rows[1:]]

    def check(index, expected):
        actual = rows(index)
        assert actual == [[str(c) for c in r] for r in expected], (index, actual[:2], expected[:2])

    budgets = raw["budgets"]
    runs = {key: {(r["arm"], r["technicians"], r["workload"], r["seed"], r["rep"]): r for r in budgets[key]["runs"]} for key in budgets}
    keys = sorted({(r["technicians"], r["workload"], r["seed"]) for r in budgets["15000"]["runs"]})
    assert len(keys) == 24 and all(len(runs[key]) == 336 for key in runs)

    def mean(key, arm, case, field):
        values = [Decimal(runs[key][(arm, *case, rep)]["measures"][field]) if field in MEASURES else Decimal(runs[key][(arm, *case, rep)][field]) for rep in (1, 2)]
        return sum(values) / 2

    # Table 7: b and c against the bars, recomputed from the runs.
    expected = []
    for key in ("15000", "60000"):
        for arm in ("b", "c"):
            pairs = [runs[key][(arm, *c, rep)]["fairnessMovesPerSecond"] / runs[key][("baseline", *c, rep)]["fairnessMovesPerSecond"] for c in keys for rep in (1, 2)]
            reference = [runs[key][(arm, *c, rep)]["referenceMovesPerSecond"] / runs[key][("baseline", *c, rep)]["referenceMovesPerSecond"] for c in keys for rep in (1, 2)]
            fairer = same = less = beyond = 0
            for c in keys:
                base = [Decimal(runs[key][("baseline", *c, rep)]["measures"]["variance"]) for rep in (1, 2)]
                gap = mean(key, arm, c, "variance") - sum(base) / 2
                fairer += gap < 0; same += gap == 0; less += gap > 0; beyond += gap > max(base) - min(base)
            diffs = [float(mean(key, arm, c, "costCents") - mean(key, "baseline", c, "costCents")) for c in keys]
            within = all(runs[key][(arm, *c, rep)]["costCents"] <= runs[key][(arm, *c, rep)]["costCeilingCents"] for c in keys for rep in (1, 2))
            passes = median(pairs) >= 1.25 and beyond == 0 and median(diffs) <= 0 and within
            assert passes == budgets[key]["summary"][arm]["passes"]
            expected.append([f"{arm}, {SECONDS[key]}", f"{round(median(pairs), 3):.2f}x ({round(min(pairs), 3):.2f} to {round(max(pairs), 3):.2f})", f"{round(median(reference), 3):.3f}x",
                             f"{fairer} / {same} / {less} / {beyond}",
                             f"{signed(median(diffs))} c ({sum(d < 0 for d in diffs)} / {sum(d == 0 for d in diffs)} / {sum(d > 0 for d in diffs)})", "Yes" if within else "No", "Yes" if passes else "No"])
    check(7, expected)
    assert raw["decision"]["b"]["promoted"] and not raw["decision"]["c"]["promoted"]

    # Tables 9 and 10: every arm on every measure.
    arms = ("b", "c", "min-max", "spread", "abs-dev", "load-balance")
    for index, key in ((9, "15000"), (10, "60000")):
        expected = []
        for arm in arms:
            cells = []
            for name in MEASURES:
                values = [(mean(key, arm, c, name), mean(key, "baseline", c, name)) for c in keys]
                ratio = round(median([float(a / b) for a, b in values if b > 0]), 3)
                cells.append(f"{ratio:.2f} ({sum(a < b for a, b in values)}/{sum(a > b for a, b in values)})")
            diffs = [float(mean(key, arm, c, "costCents") - mean(key, "baseline", c, "costCents")) for c in keys]
            pairs = [runs[key][(arm, *c, rep)]["fairnessMovesPerSecond"] / runs[key][("baseline", *c, rep)]["fairnessMovesPerSecond"] for c in keys for rep in (1, 2)]
            expected.append([NAMES.get(arm, arm), TITLES[OWN.get(arm, "variance")], *cells, f"{signed(median(diffs))} c", f"{round(median(pairs), 3):.2f}x"])
        check(index, expected)

    # Table 11: allowance spent, own-measure reduction and acceptance reasons.
    expected = []
    for arm in ("baseline",) + arms:
        cells = []
        for field in ("used", "own"):
            parts = []
            for key in ("15000", "60000"):
                values = []
                for r in budgets[key]["runs"]:
                    if r["arm"] != arm:
                        continue
                    if field == "used":
                        room = r["costCeilingCents"] - r["referenceCostCents"]
                        if room > 0: values.append((r["costCents"] - r["referenceCostCents"]) / room)
                    else:
                        name = OWN.get(arm, "variance"); start = Decimal(r["referenceMeasures"][name])
                        if start != 0: values.append(float(Decimal(r["measures"][name]) / start))
                parts.append(f"{median(values) * 100:.0f}%")
            cells.append(" / ".join(parts))
        reasons = []
        for key in ("15000", "60000"):
            got = [r["acceptanceReason"] for r in budgets[key]["runs"] if r["arm"] == arm]
            reasons.append(f"{got.count('FAIRNESS_IMPROVEMENT')} / {got.count('COST_REDUCTION')}")
        expected.append([NAMES.get(arm, arm), *cells, *reasons])
    check(11, expected)

    # Tables 15 and 16: every paired case for b and c.
    for index, key in ((15, "15000"), (16, "60000")):
        expected = []
        for arm in ("b", "c"):
            for c in keys:
                base = [Decimal(runs[key][("baseline", *c, rep)]["measures"]["variance"]) for rep in (1, 2)]
                own = [Decimal(runs[key][(arm, *c, rep)]["measures"]["variance"]) for rep in (1, 2)]
                expected.append([arm, c[0], c[1].replace("_", " ").title(), c[2], " / ".join(thousandths(v) for v in base), thousandths(max(base) - min(base)),
                                 " / ".join(thousandths(v) for v in own), thousandths(sum(own) / 2 - sum(base) / 2),
                                 signed(float(mean(key, arm, c, "costCents") - mean(key, "baseline", c, "costCents")))])
        check(index, expected)

    # Tables 17 and 18: every policy case.
    for index, key in ((17, "15000"), (18, "60000")):
        expected = []
        for c in keys:
            cells = []
            for arm in ("min-max", "spread", "abs-dev", "load-balance"):
                base = mean(key, "baseline", c, OWN[arm]); cost = signed(float(mean(key, arm, c, "costCents") - mean(key, "baseline", c, "costCents")))
                cells.append(f"{float(mean(key, arm, c, OWN[arm]) / base):.2f} / {cost}" if base > 0 else f"n/a / {cost}")
            expected.append([c[0], c[1].replace("_", " ").title(), c[2], *cells])
        check(index, expected)

    # Table 6: the profile shares.
    shares = lambda name, group: f"{raw['profiles'][name]['samplesPercent'].get(group, 0):.1f}%"
    for row in rows(6):
        assert row[1:] == [shares(n, row[0]) for n in ("baseline cost", "baseline fairness", "b fairness", "c fairness")], row

    # Chart data agrees with the evidence.
    chart = json.loads((HERE / "generated" / "chart_data.json").read_text(encoding="utf-8"))
    for key in ("15000", "60000"):
        for arm in arms:
            pairs = [runs[key][(arm, *c, rep)]["fairnessMovesPerSecond"] / runs[key][("baseline", *c, rep)]["fairnessMovesPerSecond"] for c in keys for rep in (1, 2)]
            assert [round(v, 4) for v in pairs] == chart["speed"][key][arm]
            assert {m: budgets[key]["comparison"][arm]["measures"][m]["medianRatioToBaseline"] for m in MEASURES} == chart["heatmap"][key][arm]
    for name in ("profile",):
        for row in chart[name]:
            assert row["shares"] == raw["profiles"][f"{row['arm']} {row['phase']}"]["samplesPercent"]

    # Template parts preserved byte for byte, and no em or en dash anywhere in the document XML.
    assert sha256(REFERENCE.read_bytes()).hexdigest() == REFERENCE_SHA
    with ZipFile(REFERENCE) as ref, ZipFile(OUTPUT) as out:
        for name in set(PRESERVE) | {n for n in ref.namelist() if "footer" in n}:
            assert ref.read(name) == out.read(name), name
        for name in out.namelist():
            if name.endswith(".xml"):
                text = out.read(name).decode("utf-8")
                assert "\u2014" not in text and "\u2013" not in text, name
        body = out.read("word/document.xml").decode("utf-8")
        assert body.count("<w:cantSplit/>") == sum(len(t.rows) for t in doc.tables)
    print("verified", OUTPUT.name)


if __name__ == "__main__":
    verify()
