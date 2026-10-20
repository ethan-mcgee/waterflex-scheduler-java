"""Build the fairness measure study report from the validated evidence and the reference explainer's template.

Install requirements.txt first. No network, services or benchmark execution. The two diagrams in diagrams/ are
rendered to generated/ separately (see EVIDENCE.md); the charts are regenerated here.
"""
from copy import deepcopy
from datetime import datetime
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from statistics import median
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from charts import LABEL, case_label, make_charts
from report_data import (ARMS, BUDGETS, MEASURES, POLICIES, REPS, allowance_used, case_mean, cases, index, load, own_measure,
                         own_reduction, speedups)

HERE = Path(__file__).resolve().parent
REFERENCE = HERE.parent / "WaterFlex_Scheduler_Algorithm_Comparison_Explained.docx"
REFERENCE_SHA = "5366f3017d6f18c15de295ba5fb9c1a014b6222911040f1c5e044845e57f832d"
OUTPUT = HERE / "Fairness_Measures_2026-10-20_Report.docx"
PRESERVE = ("word/styles.xml", "word/stylesWithEffects.xml", "word/numbering.xml", "word/theme/theme1.xml", "word/fontTable.xml")
DIAGRAMS = ("fairness-computation", "fairness-acceptance")
MEASURE_NAME = {"variance": "Variance", "maximumUtilization": "Highest utilization", "utilizationSpread": "Utilization spread",
                "absoluteDeviation": "Absolute deviation", "loadBalance": "Load balance"}
GROUPS = ("decimal: variance", "variance (non-decimal)", "decimal: score impacts", "decimal: money", "collections", "engine", "scheduler code",
          "score engine", "time", "other")
SECONDS = {"15000": "15 s", "60000": "60 s"}


def speed(entry):
    return f"{entry['median']:.2f}x ({entry['min']:.2f} to {entry['max']:.2f})"


def cents(value):
    return f"{value:+,.1f}" if value != int(value) else f"{int(value):+,}"


def milli(value):
    """A variance as thousandths, the unit every variance in the report is shown in."""
    text = f"{Decimal(value) * 1000:.2f}"
    return "0.00" if text == "-0.00" else text


def pct(value):
    return f"{value * 100:.0f}%"


def build():
    assert sha256(REFERENCE.read_bytes()).hexdigest() == REFERENCE_SHA
    document = load()
    assets = HERE / "generated"
    for name in DIAGRAMS:
        assert (assets / f"{name}.png").exists(), f"render diagrams/{name}.html first (EVIDENCE.md)"
    make_charts(document, assets)
    doc = Document(REFERENCE)
    table_properties = deepcopy(doc.tables[0]._tbl.tblPr)
    for child in list(doc.element.body):
        if child.tag != qn("w:sectPr"): doc.element.body.remove(child)
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype.endswith(("/image", "/hyperlink")): del doc.part.rels[rid]
    doc.core_properties.title = "WaterFlex scheduler fairness measures explained"
    doc.core_properties.subject = "Two faster formulas for the fairness variance and four alternative fairness policies, measured against production, 2026-10-20"
    doc.core_properties.author = doc.core_properties.last_modified_by = "WaterFlex"
    doc.core_properties.created = doc.core_properties.modified = datetime(2026, 10, 20)
    figures = []

    def p(text, style=None):
        assert "\u2014" not in text and "\u2013" not in text
        para = doc.add_paragraph(text, style); para.paragraph_format.widow_control = True
        return para

    def heading(text):
        h = p(text, "Heading 1"); h.paragraph_format.page_break_before = True

    def sub(text): p(text, "Heading 2")

    def table(headers, data, widths, size=10):
        assert abs(sum(widths) - 6.9) < .01
        t = doc.add_table(rows=1, cols=len(headers)); t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t._tbl.replace(t._tbl.tblPr, deepcopy(table_properties))
        for col, w in zip(t.columns, widths): col.width = Inches(w)
        for i, values in enumerate([headers, *data]):
            assert len(values) == len(headers)
            row = t.rows[0] if i == 0 else t.add_row()
            trpr = row._tr.get_or_add_trPr(); trpr.append(OxmlElement("w:cantSplit"))
            if i == 0: trpr.append(OxmlElement("w:tblHeader"))
            for j, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
                text = str(value)
                assert "\u2014" not in text and "\u2013" not in text
                cell.width = Inches(width); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell.text = text
                props = cell._tc.get_or_add_tcPr()
                shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "DCE6EF" if i == 0 else ("FFFFFF" if i % 2 else "F4F6F8")); props.append(shade)
                margins = OxmlElement("w:tcMar")
                for side in ("top", "left", "bottom", "right"):
                    n = OxmlElement("w:" + side); n.set(qn("w:w"), "65"); n.set(qn("w:type"), "dxa"); margins.append(n)
                props.append(margins)
                para = cell.paragraphs[0]; para.paragraph_format.space_after = Pt(2); para.paragraph_format.space_before = Pt(2); para.paragraph_format.line_spacing = 1.0
                if len(headers) > 2 and j > 0: para.alignment = 1
                if len(data) <= 12 and i < len(data): para.paragraph_format.keep_with_next = True  # Short tables stay on one page.
                for run in para.runs: run.font.name = "Calibri"; run.font.size = Pt(size); run.bold = i == 0
        spacer = p(""); spacer.paragraph_format.space_after = Pt(0); spacer.paragraph_format.line_spacing = Pt(3)

    def figure(name, caption, width=6.85):
        figures.append(name)
        caption = f"Figure {len(figures)}. {caption}"
        para = p(""); shape = para.add_run().add_picture(str(assets / f"{name}.png"), width=Inches(width))
        shape._inline.docPr.set("descr", caption); para.paragraph_format.keep_with_next = True; para.alignment = 1
        cap = p(caption, "Caption"); cap.paragraph_format.space_after = Pt(7)
        for run in cap.runs:
            run.font.color.rgb = RGBColor(0, 0, 0); run.bold = False; run.italic = True

    # The numbers the narrative refers to by name.
    budgets = document["budgets"]
    by = {key: index(budgets[key]) for key in BUDGETS}
    all_cases = cases(budgets["15000"])
    summary = {key: budgets[key]["summary"] for key in BUDGETS}
    comparison = {key: budgets[key]["comparison"] for key in BUDGETS}
    profiles = document["profiles"]
    sb, sc = summary["15000"]["b"], summary["15000"]["c"]
    lb, lc = summary["60000"]["b"], summary["60000"]["c"]

    def fleet_speed(key, arm, size, phase="fairnessMovesPerSecond"):
        return median([by[key][(arm, c, r)][phase] / by[key][("baseline", c, r)][phase] for c in all_cases if c[0] == size for r in REPS])

    def rate(key, arm, phase):
        return median([r[phase] for r in budgets[key]["runs"] if r["arm"] == arm])

    def ratio(key, arm, name):
        return comparison[key][arm]["measures"][name]["medianRatioToBaseline"]

    def own_ratio(key, arm):
        return ratio(key, arm, own_measure(arm))

    def allowance(key, arm):
        return median([u for u in (allowance_used(r) for r in budgets[key]["runs"] if r["arm"] == arm) if u is not None])

    def reduction(key, arm):
        return median([o for o in (own_reduction(r) for r in budgets[key]["runs"] if r["arm"] == arm) if o is not None])

    b15_beats_base60 = sum(1 for c in all_cases if case_mean(by["15000"], "b", c, "variance") < case_mean(by["60000"], "baseline", c, "variance"))
    variance_share = profiles["baseline fairness"]["samplesPercent"]["decimal: variance"]
    variance_total = variance_share + profiles["baseline fairness"]["samplesPercent"]["variance (non-decimal)"]
    c_fail_60 = [x for x in budgets["60000"]["pairs"] if x["arm"] == "c" and Decimal(x["meanVarianceDifference"]) > Decimal(x["baselineVarianceSpread"])]
    assert len(c_fail_60) == 1
    c_fail = c_fail_60[0]
    leaves_b = profiles["b fairness"]["leafFramesPercent"]
    probe = leaves_b["java.util.ImmutableCollections$MapN.probe"]
    put = leaves_b["java.util.HashMap.putVal"]
    total_runs = sum(len(budgets[key]["runs"]) for key in BUDGETS)

    p("WaterFlex scheduler fairness measures explained", "Title")
    p("Two faster formulas for the fairness variance, and four other ways to measure fairness, measured against production", "Subtitle")
    p("October 20 2026  |  Fairness measure study of pull request #187, run 2026-10-20 from 04:53 to 12:24 UTC")
    p(f"The daily search ends with a fairness phase, which spends up to two percent more cost to even out how busy each technician is. The representation study before this one found that this phase evaluated moves about four times slower than the cost phase, and that the slowdown was the formula for fairness itself: a DECIMAL128 variance recomputed over the whole fleet for every candidate move. This study tested two faster ways to compute that same variance (options b and c), and four other fairness policies, in {total_runs} solves over 24 fleets at two budgets. Option b, the same variance in binary64, makes the fairness phase {sb['fairnessPhaseSpeedup']['median']:.1f} times faster at the 15 s production budget and {lb['fairnessPhaseSpeedup']['median']:.1f} times faster at 60 s. It ends fairer than production in {sb['fairness']['fairer']} and {lb['fairness']['fairer']} of the 24 cases, at a median cost difference of {cents(sb['cost']['medianDifferenceCents'])} and {cents(lb['cost']['medianDifferenceCents'])} cents. It passes every bar the owner set before the run, so it is the recommended production formula. Option c passes on speed and misses two bars narrowly. None of the four other policies is better than the variance it would replace.")
    sub("What the measurements tell us")
    table(["Arm", "What it changes", "15 s: fairness-phase speed, variance, cost", "60 s: the same", "Verdict"], [
        [LABEL[arm], what,
         *[f"{comparison[key][arm]['fairnessPhaseSpeedup']['median']:.2f}x, {ratio(key, arm, 'variance'):.2f}, {cents(comparison[key][arm]['cost']['medianDifferenceCents'])} c" for key in BUDGETS],
         verdict]
        for arm, what, verdict in (
            ("b", "Variance in binary64", "Passes every bar; promote"),
            ("c", "Variance from running sums", "Misses cost at 15 s and fairness at 60 s"),
            ("min-max", "Busiest technician", "Less fair on every measure, its own included"),
            ("spread", "Busiest minus least busy", "Less fair on every measure, its own included"),
            ("abs-dev", "Mean absolute deviation", "Fairer on its own measure, no better on the busiest day"),
            ("load-balance", "Timefold load balance", "Within 0.01 of b or behind it on every measure"))],
        [1.0, 1.55, 1.6, 1.6, 1.15])
    p("Each speed is the median over 48 paired solves of the fairness phase's moves per second divided by production's in the same case and repetition. Each variance is the median over the 24 cases of the final variance divided by production's; lower is fairer. Each cost is the median over the cases of the final cost minus production's, in cents. Every one of the 672 solves ended with no constraint violation and inside its cost ceiling.")
    p("The report explains each result through what the solver recorded: moves per second in each phase, the five fairness measures of every final plan, how much of its fairness allowance each run spent, why each result was accepted, and where the CPU went in a flight recording of each arm. All of it comes from one evidence file, docs/evidence/solver-variance-2026-10-20.json. None of it is customer data; the fleets are synthetic.")
    sub("How to read this report")
    table(["Sections", "What you will learn"], [
        ("1", "What was measured: the seven arms, the 24 fleets, the five fairness measures and the bars."),
        ("2", "Why production's fairness phase is slow, and what each formula does differently after a move."),
        ("3 and 4", "Options b and c against each bar, and what their extra moves buy in fairness."),
        ("5", "The four other fairness policies, compared on every measure, with what each spent and why each was accepted."),
        ("6", "Where the fairness phase spends its time once the variance is cheap."),
        ("7 and 8", "Conclusions, and what should be tested or implemented next."),
        ("Appendices", "Every paired case for b and c, every policy case, every profile, definitions and sources.")], [1.45, 5.45])

    heading("1 What was measured")
    p("The daily search runs in two phases. The cost phase looks for the cheapest plan that breaks no rule; its result is the reference plan. The fairness phase then starts from the reference and looks for a plan that spreads the work more evenly, but it may not cost more than the reference plus the client's fairness allowance (two percent here). The fairness phase's objective is the fairness measure; its second objective is cost. Production measures fairness as the capacity-weighted variance of each technician's utilization, which is the share of the technician's regular hours the plan fills. [1, 2]")
    p("This study changed only the fairness phase's objective. Each arm is a constraint provider that replaces one constraint and keeps every other score level as production has it. Two arms compute production's own variance faster and are judged against bars; four change what fairness means and are compared on every measure, without bars, because changing the policy is the owner's decision.")
    table(["Arm", "Variant", "Fairness objective", "Judged by"], [
        ("baseline", "TABU_KOPT", "Production: DECIMAL128 variance over the whole fleet, on every move", "The reference"),
        ("b", "TABU_KOPT_BINARY_VARIANCE", "The same variance in binary64, carried as rounded units of 1e-9", "Bars"),
        ("c", "TABU_KOPT_INCREMENTAL_VARIANCE", "The same variance from three exact running sums updated per move", "Bars"),
        ("min-max", "TABU_KOPT_MIN_MAX_UTILIZATION", "The highest technician utilization", "Comparison"),
        ("spread", "TABU_KOPT_UTILIZATION_SPREAD", "The highest minus the lowest utilization", "Comparison"),
        ("abs-dev", "TABU_KOPT_ABSOLUTE_DEVIATION", "The capacity-weighted mean absolute deviation of utilization", "Comparison"),
        ("load-balance", "TABU_KOPT_LOAD_BALANCE", "Timefold's loadBalance over utilization, unweighted", "Comparison")], [.95, 2.25, 2.75, .95], size=9)
    manifest = document["manifest"]
    table(["Dimension", "Values in this study"], [
        ("Fleets", f"{', '.join(manifest['sizes'])} technicians; workloads {', '.join(w.title().replace('_', ' ') for w in manifest['workloads'])}; seeds {', '.join(manifest['seeds'])}: 24 cases on fixture seed {manifest['fixtureSeed']}"),
        ("Budgets", "15 s (the production default: 10 s cost, 5 s fairness) and 60 s (40 s and 20 s), to see whether a gain holds when the fairness search has room"),
        ("Repetitions", "Two of every case at every budget; all seven arms in one JVM per case, their order rotating from case to case"),
        ("Search", "TABU_KOPT at NO_ASSERT, the production move selectors; fairness allowance 2 percent"),
        ("Machine", f"One {manifest['processors']}-processor workstation; each solve limited to 2 processors, 1.5 GiB heap and the serial collector ({' '.join(manifest['jvm'])})"),
        ("Tree", f"One frozen build of revision {manifest['revision'][:7]}, carrying the nanosecond timing of #186"),
        ("Profiles", "A flight recording of every arm's two phases at 15 s on seed 17, eight recordings per arm and phase")], [1.3, 5.6])
    sub("The five fairness measures")
    p("Every final plan was read on all five measures, whatever its arm searched on, so that each policy can be judged on the measures it did not optimize. Each counts the technicians the production variance counts: those qualified for the day's demand, with regular capacity.")
    table(["Measure", "How to interpret it"], [
        ("Variance", "Production's measure: the capacity-weighted mean of the squared distance between each technician's utilization and the fleet's. Shown here in thousandths."),
        ("Highest utilization", "The busiest technician's share of regular hours. Lower means the busiest day is lighter."),
        ("Utilization spread", "The busiest technician's utilization minus the least busy one's."),
        ("Absolute deviation", "The capacity-weighted mean distance of each utilization from the fleet's, without squaring, so one large gap counts less than in the variance."),
        ("Load balance", "Timefold's loadBalance unfairness over each utilization in millionths, not weighted by capacity.")], [1.55, 5.35])
    sub("The bars for options b and c")
    p("The owner set four bars before the run. Each option is promoted only if every bar holds at both budgets.")
    table(["Bar", "Test"], [
        ("Speed", "The median of the 48 paired fairness-phase speedups is at least 1.25x."),
        ("Fairness", "No case is less fair than the baseline by more than that case's own spread between the baseline's two repetitions."),
        ("Cost", "The median per-case cost difference is not above zero. A single case may be dearer, because the fairness phase may spend its allowance."),
        ("Ceiling", "Every accepted result is inside its cost ceiling, with zero violations.")], [1.2, 5.7])

    heading("2 Why the fairness phase was slow")
    p(f"A local search evaluates a candidate move by changing the plan, rescoring it, and undoing the change. Timefold rescores incrementally: only the routes the move touched are walked again. Production's fairness constraint cannot be incremental in that way, because a variance depends on the fleet's mean, and the mean changes whenever any route's paid minutes change. So for every candidate move, production collects every route's paid minutes and capacity and recomputes the variance with DECIMAL128 divides. In the flight recording of production's fairness phase that work held {variance_total:.0f} percent of the CPU samples: {variance_share:.0f} percent in decimal arithmetic, mostly long division (MutableBigInteger.divideMagnitude and BigDecimal.divideAndRound lead the leaf frames), and {profiles['baseline fairness']['samplesPercent']['variance (non-decimal)']:.0f} percent in the surrounding collection. [3]")
    figure("fairness-computation", "What each arm does after a candidate move. The routes the move changed are walked again in every arm. Production then recomputes the variance over the whole fleet in DECIMAL128; option b recomputes the same variance in binary64; option c updates three exact running sums for the changed routes only.")
    p("Option b keeps production's structure and changes only the arithmetic. The variance is computed in binary64, which a modern CPU does in a few cycles, and rounded to an integer count of 1e-9 units, the resolution the numeric contract in scheduler-policy.md already declares. Two plans whose variances differ by less than a unit tie. Option c changes the structure as well: it keeps three running sums over the fleet (capacity, paid minutes, and paid minutes squared over capacity at a scale of 1e8, all exact longs) and updates them only for the routes a move changes; the variance is one binary64 subtraction of the sums.")
    figure("profile", "Share of CPU samples in each phase by kind of work, from eight flight recordings per arm and phase at 15 s. Red is production's variance. With any other formula, the fairness phase spends its time the way the cost phase does.")
    table(["Kind of work", "Baseline, cost phase", "Baseline, fairness phase", "b, fairness phase", "c, fairness phase"],
          [[group, *[f"{profiles[name]['samplesPercent'].get(group, 0):.1f}%" for name in ("baseline cost", "baseline fairness", "b fairness", "c fairness")]]
           for group in GROUPS if any(profiles[name]["samplesPercent"].get(group, 0) > 0 for name in ("baseline cost", "baseline fairness", "b fairness", "c fairness"))],
          [1.9, 1.25, 1.3, 1.2, 1.25])
    p(f"The profiles say the variance was the whole difference between the two phases. Once it is cheap, the fairness phase's profile matches the cost phase's within a few points for every arm, and the absolute rates agree: production's fairness phase evaluated a median {rate('15000', 'baseline', 'fairnessMovesPerSecond'):,.0f} moves per second at 15 s against {rate('15000', 'baseline', 'referenceMovesPerSecond'):,.0f} in its cost phase, and option b's fairness phase evaluated {rate('15000', 'b', 'fairnessMovesPerSecond'):,.0f}. The fairness phase is now as fast as the cost phase.")
    sub("Why the gain grows with the fleet")
    p(f"Production's per-move fairness cost grows with the number of technicians, because the whole fleet is recomputed, while the rest of a move's cost depends mostly on the routes it touches. So the larger the fleet, the larger option b's gain: at 15 s its median fairness-phase speedup is {fleet_speed('15000', 'b', 20):.2f}x at 20 technicians and {fleet_speed('15000', 'b', 50):.2f}x at 50, and at 60 s {fleet_speed('60000', 'b', 20):.2f}x and {fleet_speed('60000', 'b', 50):.2f}x. The speed figure in the next section shows the two fleet sizes as two bands. A metro larger than 50 technicians should expect more.")

    heading("3 Options b and c against the bars")
    table(["Option, budget", "Fairness-phase speed", "Cost-phase speed", "Fairer / same / less fair / beyond noise", "Median cost difference (cheaper / ties / dearer)", "Inside ceiling", "Passes"],
          [[f"{arm}, {SECONDS[key]}", speed(s["fairnessPhaseSpeedup"]), f"{s['referencePhaseSpeedup']['median']:.3f}x",
            f"{s['fairness']['fairer']} / {s['fairness']['same']} / {s['fairness']['lessFair']} / {s['fairness']['lessFairBeyondNoise']}",
            f"{cents(s['cost']['medianDifferenceCents'])} c ({s['cost']['cheaper']} / {s['cost']['ties']} / {s['cost']['dearer']})",
            "Yes" if s["bars"]["ceiling"] else "No", "Yes" if s["passes"] else "No"]
           for key in BUDGETS for arm, s in (("b", summary[key]["b"]), ("c", summary[key]["c"]))], [1.0, 1.25, .8, 1.3, 1.35, .6, .6], size=9)
    figure("speed", f"Fairness-phase moves per second of each arm relative to production in the same case and repetition, all 48 pairs per budget. Triangles are 50-technician fleets, circles 20. The black bar is the median; the dashed line is the 1.25x bar. Every pair of every arm clears it.")
    p(f"Speed is not close for either option: the slowest single pair of b ran {sb['fairnessPhaseSpeedup']['min']:.2f}x production's rate and the slowest of c {sc['fairnessPhaseSpeedup']['min']:.2f}x. The four policy arms are as fast, because none of them uses decimal division either; the measure's arithmetic, not its definition, was the cost. The cost phase is a different matter. b's cost phase runs at production's rate ({sb['referencePhaseSpeedup']['median']:.3f}x), as it should, since the cost phase does not score fairness. c's runs at {sc['referencePhaseSpeedup']['median']:.3f}x and {lc['referencePhaseSpeedup']['median']:.3f}x, because its running sums are maintained on every move of both phases even though only the fairness phase reads them. Avoiding the whole-fleet walk did not make c faster than b in the fairness phase either: once the arithmetic is binary64, collecting 20 or 50 numbers is cheap next to walking the changed routes.")
    figure("pairs", "Final variance of each case relative to production's (the mean of two repetitions on each side). The grey bar reaches production's value plus its own spread between repetitions; a marker above its bar fails the fairness bar. Points at 0 are cases where the option found a perfectly even day.")
    p(f"Option b was fairer in {sb['fairness']['fairer']} of 24 cases at 15 s and the same in {sb['fairness']['same']}; at 60 s fairer in {lb['fairness']['fairer']} and the same in {lb['fairness']['same']}. It was never less fair. Its median variance is {sb['fairness']['medianVarianceRatio']:.2f} and {lb['fairness']['medianVarianceRatio']:.2f} of production's. The cases where it ties are small clustered fleets where both reach the same plan.")
    figure("cost", "Final cost of each case minus production's, in dollars, for b and c at both budgets. The black bar is the median, which the cost bar holds at or below zero. Every point is inside its ceiling.")
    p(f"The cost bar is where the options differ. At 15 s b's median difference is {cents(sb['cost']['medianDifferenceCents'])} cents with {sb['cost']['cheaper']} cases cheaper and {sb['cost']['dearer']} dearer, and at 60 s {cents(lb['cost']['medianDifferenceCents'])} cents with {lb['cost']['cheaper']} cheaper. c's median at 15 s is {cents(sc['cost']['medianDifferenceCents'])} cents with {sc['cost']['dearer']} cases dearer, so it misses the bar by three cents. The reason is the allowance: every variance arm spends almost all of its two percent (section 5), so a fairer final plan sits close to the same ceiling, and a few cents either way are which plan the search ended on. b is the one that does not end dearer in the median.")
    sub("The two cases where option c was less fair")
    figure("c_cases", "Final variance in each repetition for the two cases where option c ended less fair than production. Pale bars are repetition 1, solid repetition 2.")
    table(["Case, budget", "Baseline (two repetitions)", "Baseline spread", "c (two repetitions)", "b (two repetitions)", "Beyond the spread"],
          [[f"{case_label(case)}, {SECONDS[key]}", " / ".join(milli(by[key][("baseline", case, r)]["measures"]["variance"]) for r in REPS),
            milli(x["baselineVarianceSpread"]), " / ".join(milli(v) for v in x["armVariance"]),
            " / ".join(milli(by[key][("b", case, r)]["measures"]["variance"]) for r in REPS), "Yes" if Decimal(x["meanVarianceDifference"]) > Decimal(x["baselineVarianceSpread"]) else "No"]
           for key in BUDGETS for x in budgets[key]["pairs"] if x["arm"] == "c" and Decimal(x["meanVarianceDifference"]) > 0
           for case in [(x["technicians"], x["workload"], x["seed"])]], [1.55, 1.25, .9, 1.15, 1.15, .9], size=9)
    p(f"At 15 s, 20 ABSENCE 17 was less fair under c but inside production's own spread, so it passes the bar. At 60 s, {c_fail['technicians']} {c_fail['workload']} {c_fail['seed']} fails it: production reached the same plan in both repetitions, so its spread is zero and any loss counts, and c's second repetition ended on a plan with a variance of {milli(c_fail['armVariance'][1])} thousandths against production's {milli(c_fail['baselineVariance'][0])} in both. Its first repetition was fairer than production. One unlucky repetition is enough to fail a bar measured against a zero spread, and the bar was set that way before the run. The report does not reargue it: c is slower than b in the cost phase and no faster in the fairness phase, so nothing would be gained by promoting it instead.")

    heading("4 What the extra moves buy")
    p(f"The fairness phase does not run longer under option b. It evaluates more moves in the same seconds, and the extra moves find fairer plans within the same allowance. A useful way to read the size of the gain is to compare budgets: at 15 s, b's median variance is {milli(comparison['15000']['b']['measures']['variance']['median'])} thousandths, below production's {milli(comparison['60000']['baseline']['measures']['variance']['median'])} at 60 s, four times the time. Case by case, b at 15 s was fairer than production at 60 s in {b15_beats_base60} of 24 cases.")
    figure("budget", "Left: the median final variance of each variance-reducing arm at each budget. Right: each case's final variance under option b at 15 s against production at 60 s; points below the dashed line are cases where b in a quarter of the time ended fairer.")
    p(f"Production gains from the longer budget (its median falls from {milli(comparison['15000']['baseline']['measures']['variance']['median'])} to {milli(comparison['60000']['baseline']['measures']['variance']['median'])} thousandths), but b gains more ({milli(comparison['15000']['b']['measures']['variance']['median'])} to {milli(comparison['60000']['b']['measures']['variance']['median'])}), because at 60 s its fairness phase has 20 s at four times the rate. The two policies that also reduce variance, abs-dev and load-balance, do not improve with the longer budget: they optimize a different measure and reach their own plateau, and the variance they leave is whatever that plateau implies.")

    heading("5 The four other fairness policies")
    p("Each policy searched on its own measure and was accepted on it: a result counts as a fairness improvement if it is fairer on that measure than the host's plan, as a cost reduction if it is no less fair and cheaper, and is rejected otherwise. The rule is production's own, with the policy's measure in place of the variance. Every final plan was then read on all five measures.")
    figure("fairness-acceptance", "How a fairness-phase result is accepted. The search runs under the cost ceiling; its result is then compared with the host's plan on the arm's own measure M.", width=5.4)
    for key in BUDGETS:
        sub(f"Every arm on every measure at {SECONDS[key]}")
        table(["Arm", "Accepted on", *[MEASURE_NAME[m] for m in MEASURES], "Median cost difference", "Fairness-phase speed"],
              [[LABEL[arm], MEASURE_NAME[comparison[key][arm]["acceptedOn"]],
                *[f"{comparison[key][arm]['measures'][m]['medianRatioToBaseline']:.2f} ({comparison[key][arm]['measures'][m]['lowerThanBaseline']}/{comparison[key][arm]['measures'][m]['higherThanBaseline']})" for m in MEASURES],
                f"{cents(comparison[key][arm]['cost']['medianDifferenceCents'])} c", f"{comparison[key][arm]['fairnessPhaseSpeedup']['median']:.2f}x"]
               for arm in ARMS[1:]], [1.05, .9, .7, .7, .7, .7, .7, .75, .7], size=9)
        p(f"Each measure cell is the median over the 24 cases of the arm's value divided by production's, then in brackets the cases lower and higher than production; lower is fairer on every measure. Production's own medians at {SECONDS[key]}: variance {milli(comparison[key]['baseline']['measures']['variance']['median'])} thousandths, highest utilization {float(comparison[key]['baseline']['measures']['maximumUtilization']['median']):.3f}, spread {float(comparison[key]['baseline']['measures']['utilizationSpread']['median']):.3f}, absolute deviation {float(comparison[key]['baseline']['measures']['absoluteDeviation']['median']):.4f}, load balance {float(comparison[key]['baseline']['measures']['loadBalance']['median']):,.0f}.")
    figure("heatmap", "The two comparison tables as one picture: each arm's median ratio to production on each measure. Teal is fairer than production, red less fair; bold is the measure the arm searched on.")
    p(f"Three things stand out. First, the variance does well on measures it does not optimize: option b is fairer than production on all five, and on highest utilization and spread it is better than min-max and spread themselves. Second, min-max and spread are less fair than production on every measure, including their own: the busiest technician is about {(own_ratio('15000', 'min-max') - 1) * 100:.0f} percent busier under min-max than under production, and the variance is about three times production's. Third, abs-dev wins on its own measure ({own_ratio('60000', 'abs-dev'):.2f} at 60 s, against b's {ratio('60000', 'b', 'absoluteDeviation'):.2f}) but leaves the busiest technician busier than production in {comparison['15000']['abs-dev']['measures']['maximumUtilization']['higherThanBaseline']} cases at 15 s, because an absolute deviation does not care whether a gap is one large one or many small ones. Load-balance tracks b closely: on every measure at both budgets it is within 0.01 of b or behind it.")
    sub("Why min-max and spread fail on their own measure")
    p("A maximum is flat almost everywhere. Moving a visit between two technicians who are not the busiest leaves the highest utilization unchanged, so nearly every candidate move scores the same on the policy's level and the search falls through to its second objective, cost. The search therefore spends the fairness phase reducing cost, not unfairness. The figure below shows the result: these two arms spend almost none of their allowance, and end within a few percent of the reference plan's value on their own measure, while every other arm spends almost all of it.")
    figure("allowance", "Left: the share of the 2 percent allowance each run spent above its own reference plan. Right: each run's final value on its own measure as a percent of its reference plan's. Left of each pair is 15 s, right 60 s; the black bar is the median.")
    table(["Arm", "Allowance spent, 15 s / 60 s", "Own measure, final over reference, 15 s / 60 s", "Fairness improvement / cost reduction, 15 s", "60 s"],
          [[LABEL[arm], f"{pct(allowance('15000', arm))} / {pct(allowance('60000', arm))}", f"{pct(reduction('15000', arm))} / {pct(reduction('60000', arm))}",
            *[f"{comparison[key][arm]['acceptance'].get('FAIRNESS_IMPROVEMENT', 0)} / {comparison[key][arm]['acceptance'].get('COST_REDUCTION', 0)}" for key in BUDGETS]]
           for arm in ARMS], [1.2, 1.4, 1.9, 1.5, .9])
    p(f"Their cost reduction is the other side of this. Measured against production, min-max is {-comparison['60000']['min-max']['cost']['medianDifferenceCents'] / 100:.0f} dollars cheaper at the median and spread {-comparison['60000']['spread']['cost']['medianDifferenceCents'] / 100:.0f} dollars at 60 s; that is roughly the allowance production spends on fairness, kept instead. A host that wanted the cheapest plan would get it by setting the allowance to zero, not by switching policy.")
    figure("tradeoff", "Every case at 60 s for b and the four policies: final cost minus production's against final variance relative to production's (log scale). Min-max and spread buy their savings with large losses of fairness; b, abs-dev and load-balance sit at production's cost.")
    figure("acceptance", "How each arm's 48 results per budget were accepted. Spread never improved on the host's plan on its own measure; min-max did in a quarter of runs. Abs-dev improved in every run.")
    reductions = lambda arm: " and ".join(str(comparison[key][arm]["acceptance"].get("COST_REDUCTION", 0)) for key in BUDGETS)
    p(f"The acceptance reasons agree. Production itself is accepted as a cost reduction in {reductions('baseline')} of 48 runs at 15 s and 60 s, where its fairness phase could not beat the host's plan on variance within the ceiling; b only {reductions('b')} times and c {reductions('c')}, because their extra moves find a fairer plan. Spread is a cost reduction in every run.")

    heading("6 Where the time goes now")
    p(f"With the variance cheap, the fairness phase's profile is the cost phase's, and its leaders are the leaders for both phases. Collections hold about {profiles['b fairness']['samplesPercent']['collections']:.0f} percent of option b's fairness-phase samples: an immutable map lookup (ImmutableCollections$MapN.probe) alone is {probe:.1f} percent of leaf frames and HashMap.putVal {put:.1f} percent. BigDecimal score impacts and money are {profiles['b fairness']['samplesPercent']['decimal: score impacts'] + profiles['b fairness']['samplesPercent']['decimal: money']:.0f} percent together, Timefold's own engine {profiles['b fairness']['samplesPercent']['engine']:.0f} percent, and the scheduler's route timing {profiles['b fairness']['samplesPercent']['scheduler code']:.0f} percent.")
    figure("leaves", "The twelve leaf frames holding the most of option b's fairness-phase samples at 15 s. A leaf frame is the method running when a sample was taken.")
    table(["Leaf frame", "Share of samples", "What it most likely is"], [
        ("ImmutableCollections$MapN.probe", f"{probe:.1f}%", "RoadMatrix's identity-to-number lookup, once per visit per move (identified in the CR07 write-up)"),
        ("HashMap.putVal, LinkedHashMap.afterNodeInsertion, HashMap.getNode", f"{put + leaves_b['java.util.LinkedHashMap.afterNodeInsertion'] + leaves_b['java.util.HashMap.getNode']:.1f}%", "Maps built and read during a move"),
        ("BigDecimal.add, valueOf, multiply", f"{leaves_b['java.math.BigDecimal.add'] + leaves_b['java.math.BigDecimal.valueOf'] + leaves_b['java.math.BigDecimal.multiply']:.1f}%", "Score impacts and money in BendableBigDecimalScore"),
        ("RouteTimeline.evaluate", f"{leaves_b['dev.waterflex.scheduler.optimizer.RouteTimeline.evaluate']:.1f}%", "The route re-walk itself"),
        ("Timefold undo, getValue, setScore", f"{leaves_b['ai.timefold.solver.core.impl.score.stream.common.inliner.BendableBigDecimalScoreContext$SingleHardImpact.undo'] + leaves_b['ai.timefold.solver.core.impl.domain.variable.descriptor.VariableDescriptor.getValue'] + leaves_b['ai.timefold.solver.core.impl.domain.score.descriptor.ScoreDescriptor.setScore']:.1f}%", "The engine's move and score bookkeeping")], [2.6, 1.0, 3.3], size=9)
    p(f"The map lookup is the largest single target, and it is a known one. The CR07 treatment replaced string-keyed leg lookups with integer indexes and raised the cost phase's rate 1.54x at 20 technicians and 2.00x at 50; its write-up identified the MapN.probe that remained as RoadMatrix's identity-to-number probe, one per visit per move (RoadMatrix.number itself is {leaves_b['dev.waterflex.scheduler.optimizer.RoadMatrix.number']:.1f} percent of leaves here). Carrying each visit's road number on the visit, so the probe is done once per plan instead of once per move, is the next treatment. The HashMap work is the maps the timing search builds per route; a profile with stack traces should confirm their callers before they are changed. [3]")

    heading("7 Conclusions")
    table(["Question", "Answer from the evidence"], [
        ("Can production's fairness phase be made faster without changing fairness?", f"Yes. Option b makes it {sb['fairnessPhaseSpeedup']['median']:.1f}x faster at 15 s and {lb['fairnessPhaseSpeedup']['median']:.1f}x at 60 s, passes every bar at both, and was never less fair than production in any case."),
        ("Was the slowness the score type or the formula?", f"The formula. The decimal variance held {variance_total:.0f} percent of the fairness phase; with it gone the fairness phase runs at the cost phase's rate, still on the decimal score."),
        ("Do running sums (option c) help further?", "No. They are no faster than b in the fairness phase, they slow the cost phase by about 4 percent, and they missed the cost bar at 15 s and the fairness bar at 60 s."),
        ("What do the extra moves buy?", f"Fairer days at the same cost: a median variance about two thirds of production's, and fairer at 15 s than production at 60 s in {b15_beats_base60} of 24 cases."),
        ("Is a different fairness policy better?", "No. Min-max and spread are less fair than production on every measure, their own included. Abs-dev wins only on its own measure. Load-balance is within 0.01 of b or behind it on every measure."),
        ("Does the gain depend on the fleet?", f"It grows with it: {fleet_speed('15000', 'b', 20):.1f}x at 20 technicians and {fleet_speed('15000', 'b', 50):.1f}x at 50, because production's cost per move grows with the fleet."),
        ("What limits the search now?", "RoadMatrix's per-move identity lookup and the timing search's maps, then BigDecimal score impacts, in both phases alike.")], [2.4, 4.5])
    p("Two cautions carry over. The fleets are synthetic and at most 50 technicians, and each solve was limited to two processors on a workstation, so the absolute rates are the harness's, not the VM's; the paired ratios are what transfers. And a fairer plan is fairer by the policy's definition: whether planners see the difference between a variance of 19 and 12 thousandths is a question for the pilot, not for this harness.")

    heading("8 What should happen next")
    table(["Next step", "What to do", "What it decides"], [
        ("Ship option b", "Replace DayConstraintProvider.fairness with b's binary64 units in production. Keep SchedulingPolicy.measure (DECIMAL128) for acceptance and reporting, as the study did. Run the exactness test and the FULL_ASSERT campaign, and record the tie tolerance of one unit in the numeric contract's change log.", "The production fairness formula."),
        ("Keep the variance as the policy", "Make no policy change. Record in scheduler-policy.md that min-max, spread, abs-dev and load-balance were measured and why each was not chosen.", "The owner's policy decision, on this evidence."),
        ("Remove the study arms", "Once b is in main, remove c and the four policy providers from main with file edits, keeping them on the study branch.", "A smaller solver surface."),
        ("Re-split the 15 s budget", "With b, the fairness phase reaches in 5 s what production did not reach in 20 s. Measure 12 s and 3 s, and 11 s and 4 s, against 10 s and 5 s on the same 24 cases: cost, variance and acceptance.", "Whether time moves from fairness to cost."),
        ("Number the visits once", "Carry each visit's RoadMatrix number on the visit so the identity-to-number probe runs once per plan, not once per move; then profile with stack traces for the per-route HashMaps. Gate with ScoringExactnessTest and repeat the 24-case paired study.", "The next speedup, in both phases."),
        ("Daily load with b", "Repeat the VM rehearsal stack's daily load (five queued planners) with b, to confirm the fairness gain at the service and see whether proposals per minute change.", "Whether the gain shows at the API.")], [1.45, 3.7, 1.75])
    p("The long score (the S treatment of the representation study) is not on this list. With the variance gone its share of the fairness phase is the same as in the cost phase, where it gained 1.04x; it should be revisited only after the map lookups, if BigDecimal then leads the profile.")
    sub("How the calculations were checked")
    p("Every table and chart value is computed from the evidence file at build time; nothing is typed in. report_data.py checks the evidence before use: the run is the full study on the pinned revision; every arm, case and repetition is present exactly once at both budgets; no run has a violation or ends over its ceiling; and each policy row was accepted on its own measure. It then recomputes from the individual runs the medians the decision rests on (the fairness-phase speedups, the cost differences, the count of cases beyond the noise, and every comparison ratio and acceptance count) and requires them to equal the recorded summary. verify_report.py recomputes the tables a second time from the raw JSON without the builder's helpers, and checks the figure captions, the template parts and the absence of dashes. [4]")

    for key, letter in (("15000", "A"), ("60000", "B")):
        heading(f"Appendix {letter} Every paired case for b and c at {SECONDS[key]}")
        p(f"Each row is one case: production's final variance in its two repetitions, its spread, the option's two final variances, all in thousandths, and the mean cost difference in cents. The fairness bar fails where the option's mean exceeds production's by more than the spread. [4]")
        table(["Arm", "Techs", "Workload", "Seed", "Baseline variance", "Spread", "Option variance", "Mean difference", "Cost difference"],
              [[x["arm"], x["technicians"], x["workload"].replace("_", " ").title(), x["seed"], " / ".join(milli(v) for v in x["baselineVariance"]), milli(x["baselineVarianceSpread"]),
                " / ".join(milli(v) for v in x["armVariance"]), milli(x["meanVarianceDifference"]), cents(x["meanCostDifferenceCents"])]
               for x in sorted(budgets[key]["pairs"], key=lambda x: (x["arm"], x["technicians"], x["workload"], x["seed"]))],
              [.4, .45, 1.0, .45, 1.15, .65, 1.15, .8, .85], size=9)

    heading("Appendix C Every policy case")
    p("Each cell is the policy's final value on its own measure divided by production's, then the mean cost difference in cents, both over the two repetitions. A ratio above 1 is less fair than production on the measure the policy searched on. [4]")
    for key in BUDGETS:
        sub(f"At {SECONDS[key]}")
        table(["Techs", "Workload", "Seed", *[LABEL[arm] for arm in POLICIES]],
              [[c[0], c[1].replace("_", " ").title(), c[2],
                *[f"{float(case_mean(by[key], arm, c, own_measure(arm)) / case_mean(by[key], 'baseline', c, own_measure(arm))):.2f} / {cents(float(case_mean(by[key], arm, c, 'costCents') - case_mean(by[key], 'baseline', c, 'costCents')))}"
                  if case_mean(by[key], "baseline", c, own_measure(arm)) > 0 else f"n/a / {cents(float(case_mean(by[key], arm, c, 'costCents') - case_mean(by[key], 'baseline', c, 'costCents')))}" for arm in POLICIES]]
               for c in all_cases], [.5, 1.1, .5, 1.2, 1.2, 1.2, 1.2], size=9)

    heading("Appendix D Every profile")
    p("Each column is one arm and phase at 15 s on seed 17: the share of its CPU samples by kind of work, over eight recordings. A blank is a kind with no samples. [3]")
    names = [f"{arm} {phase}" for arm in ARMS for phase in ("cost", "fairness")]
    for chunk in (names[:7], names[7:]):
        table(["Kind of work", *[n.replace(" cost", ", cost").replace(" fairness", ", fair.") for n in chunk]],
              [[group, *[f"{profiles[n]['samplesPercent'][group]:.1f}" if group in profiles[n]["samplesPercent"] else "" for n in chunk]] for group in GROUPS] +
              [["Samples", *[f"{sum(profiles[n]['samples'].values()):,}" for n in chunk]]], [1.65, *([5.25 / len(chunk)] * len(chunk))], size=9)

    heading("Appendix E Definitions and sources")
    table(["Term", "Meaning in this report"], [
        ("Arm", "One fairness objective under test, run as its own solver variant in every case."),
        ("Case", "One fleet size, workload and seed; each was solved twice by every arm at each budget."),
        ("Reference plan", "The plan the cost phase ends with; the fairness phase starts from it."),
        ("Allowance and ceiling", "The fairness phase may spend up to 2 percent above the reference plan's cost; the ceiling is that cost times 1.02, rounded down to a cent."),
        ("Utilization", "A technician's paid minutes divided by regular capacity minutes for the day."),
        ("Moves per second", "Candidate moves a phase evaluated per second; more moves in the same budget is a more thorough search, not a faster answer."),
        ("Speedup", "An arm's moves per second divided by production's in the same case and repetition; the report gives the median of these ratios."),
        ("Spread between repetitions", "The difference between production's two final variances in a case; the fairness bar's noise floor for that case."),
        ("Unit of 1e-9", "The resolution the numeric contract declares for the fairness variance; option b ties plans closer than one unit."),
        ("Leaf frame", "The method at the top of a sampled stack in a flight recording: where the CPU was at that instant."),
        ("Thousandths", "Variances in this report are shown multiplied by 1000.")], [1.6, 5.3])
    p("[1] docs/scheduler-policy.md: the fairness phase, the allowance, the acceptance rule and the numeric contract.")
    p("[2] The constraint providers: calculation-engine/src/main/java/dev/waterflex/scheduler/optimizer, DayConstraintProvider (production) and FairnessMeasures (the six study arms); the score levels in DayScores; the acceptance rule in SchedulingPolicy.compare.")
    p("[3] Profiles: the profiles block of the evidence, from one flight recording per solve of every arm at 15 s on seed 17, attributed by infra/representation_study.py's frame groups. The representation study's profile and its conclusion that the variance formula, not the score type, was the decimal cost: docs/solver-benchmarks.md, 'Score and time representation'.")
    p("[4] Evidence: docs/evidence/solver-variance-2026-10-20.json, written by infra/variance_study.py from the run infra/run-variance-study.ps1 made (manifest, every run, every pair, both budgets' summaries and comparisons, the decision and the profiles). docs/solver-benchmarks.md, 'Fairness measures (PR 187)', is the narrative kept with the code.")
    p("Reproduce with docs/fairness-measures-2026-10-20/requirements.txt: render the two diagrams as EVIDENCE.md says, then run report_data.py, build_report.py and verify_report.py there. The builder regenerates the charts and chart_data.json and writes this editable Word file.")

    settings = doc.settings.element
    for el in settings.findall(qn("w:updateFields")): settings.remove(el)
    update = OxmlElement("w:updateFields"); update.set(qn("w:val"), "true"); settings.append(update)
    stream = BytesIO(); doc.save(stream)
    with ZipFile(REFERENCE) as ref, ZipFile(stream) as built, ZipFile(OUTPUT, "w", ZIP_DEFLATED) as output:
        preserve = set(PRESERVE) | {n for n in ref.namelist() if "footer" in n}
        for name in sorted(built.namelist()):
            content = ref.read(name) if name in preserve else built.read(name)
            info = ZipInfo(name, (2026, 10, 20, 0, 0, 0)); info.compress_type = ZIP_DEFLATED
            output.writestr(info, content)
    print(OUTPUT)


if __name__ == "__main__": build()
