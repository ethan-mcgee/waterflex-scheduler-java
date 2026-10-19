"""Build the booking configuration study report from the validated evidence and the selected explainer template.

Install requirements.txt first. No network, services or study runs.
"""
from copy import deepcopy
from datetime import datetime
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from charts import GUARD_MS, make_charts
from report_data import CAPS, FIRST_FILE, FLEETS, LAYOUTS, RERUN_FILE, THRESHOLDS, load

HERE = Path(__file__).resolve().parent
REFERENCE = HERE.parent / "WaterFlex_Scheduler_Algorithm_Comparison_Explained.docx"
REFERENCE_SHA = "5366f3017d6f18c15de295ba5fb9c1a014b6222911040f1c5e044845e57f832d"
OUTPUT = HERE / "Booking_Configuration_Study_2026-10-19_Report.docx"
PRESERVE = ("word/styles.xml", "word/stylesWithEffects.xml", "word/numbering.xml", "word/theme/theme1.xml", "word/fontTable.xml")
FIGURES = ("pipeline", "gap", "left", "routing", "parts", "noise", "horizons", "desk", "layouts", "reserve", "budget")
FIG = {name: i + 1 for i, name in enumerate(FIGURES)}
WIDTH = 6.9


def money(cents, sign=True):
    if cents is None:
        return "n/a"
    value = cents / 100
    if not sign:
        return f"${value:.2f}"
    return f"{'-' if value < 0 else '+'}${abs(value):.2f}"


def num(value, digits=0):
    return "n/a" if value is None else f"{value:,.{digits}f}"


def secs(value_ms, digits=1):
    return "n/a" if value_ms is None else f"{value_ms / 1000:.{digits}f} s"


def stops(row):
    return ", ".join(f"{k} {v}" for k, v in sorted(row["stops"].items())) or "none"


def span(values, fmt):
    low, high = min(values), max(values)
    return fmt(low) if fmt(low) == fmt(high) else f"{fmt(low)} to {fmt(high)}"


def build():
    assert sha256(REFERENCE.read_bytes()).hexdigest() == REFERENCE_SHA
    data = load()
    assets = HERE / "generated"
    make_charts(data, assets)
    doc = Document(REFERENCE)
    table_properties = deepcopy(doc.tables[0]._tbl.tblPr)
    for child in list(doc.element.body):
        if child.tag != qn("w:sectPr"): doc.element.body.remove(child)
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype.endswith(("/image", "/hyperlink")): del doc.part.rels[rid]
    doc.core_properties.title = "WaterFlex booking configuration study explained"
    doc.core_properties.subject = "Which refinement cap, window threshold, booking pool layout and publish reserve the booking search should ship with, 2026-10-17 to 2026-10-19"
    doc.core_properties.author = doc.core_properties.last_modified_by = "WaterFlex"
    doc.core_properties.created = doc.core_properties.modified = datetime(2026, 10, 19)
    figures = []

    def p(text, style=None):
        assert "—" not in text and "–" not in text
        para = doc.add_paragraph(text, style); para.paragraph_format.widow_control = True
        return para

    def lead(label, text):
        para = p("")
        run = para.add_run(label + " "); run.bold = True
        assert "—" not in text and "–" not in text
        para.add_run(text)

    def heading(text):
        h = p(text, "Heading 1"); h.paragraph_format.page_break_before = True

    def sub(text): p(text, "Heading 2")

    def table(headers, rows, weights, size=10):
        widths = [w * WIDTH / sum(weights) for w in weights]
        t = doc.add_table(rows=1, cols=len(headers)); t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t._tbl.replace(t._tbl.tblPr, deepcopy(table_properties))
        for col, w in zip(t.columns, widths): col.width = Inches(w)
        for i, values in enumerate([headers, *rows]):
            assert len(values) == len(headers), values
            row = t.rows[0] if i == 0 else t.add_row()
            trpr = row._tr.get_or_add_trPr(); trpr.append(OxmlElement("w:cantSplit"))
            if i == 0: trpr.append(OxmlElement("w:tblHeader"))
            for j, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
                text = str(value)
                assert "—" not in text and "–" not in text
                cell.width = Inches(width); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell.text = text
                props = cell._tc.get_or_add_tcPr()
                shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "DCE6EF" if i == 0 else ("FFFFFF" if i % 2 else "F4F6F8")); props.append(shade)
                margins = OxmlElement("w:tcMar")
                for side in ("top", "left", "bottom", "right"):
                    n = OxmlElement("w:" + side); n.set(qn("w:w"), "65"); n.set(qn("w:type"), "dxa"); margins.append(n)
                props.append(margins)
                para = cell.paragraphs[0]; para.paragraph_format.space_after = Pt(2); para.paragraph_format.space_before = Pt(2); para.paragraph_format.line_spacing = 1.0
                if len(headers) > 2 and j > 0: para.alignment = 1
                if len(rows) <= 12 and i < len(rows): para.paragraph_format.keep_with_next = True  # Short tables stay on one page.
                for run in para.runs: run.font.name = "Calibri"; run.font.size = Pt(size); run.bold = i == 0
        spacer = p(""); spacer.paragraph_format.space_after = Pt(0); spacer.paragraph_format.line_spacing = Pt(3)

    def figure(name, caption):
        figures.append(name)
        assert FIG[name] == len(figures), name
        caption = f"Figure {len(figures)}. {caption}"
        para = p(""); shape = para.add_run().add_picture(str(assets / f"{name}.png"), width=Inches(6.85))
        shape._inline.docPr.set("descr", caption); para.paragraph_format.keep_with_next = True
        cap = p(caption, "Caption"); cap.paragraph_format.space_after = Pt(7)
        for run in cap.runs:
            run.font.color.rgb = RGBColor(0, 0, 0); run.bold = False; run.italic = True

    # The figures the narrative refers to by name.
    fq, fq1 = data["rerun"]["quality"], data["first"]["quality"]
    load_, cap, rl, rq = data["first"]["load"], data["first"]["capacity"], data["rerun"]["reserveLoad"], data["rerun"]["reserveQuality"]
    first_rl, shape = data["first"]["reserveLoad"], data["rerun"]["shape"]
    q = lambda fleet, variant: fq[f"{fleet}/full"][variant]
    L = lambda fleet, variant, profile="full": load_[f"{fleet}/{profile}/{variant}"]
    C = lambda variant, layout, clients, fleet=10, profile="full": cap[f"{fleet}/{profile}/{variant}/{layout}/{clients}"]
    gap = {v: [q(f, v)["meanCents"] for f in FLEETS] for v in ("2:0", "0:0", "2:250", "2:1000", "2:2000", "2:3000")}
    spread = lambda fleet, c: max(q(fleet, f"{t}:{c}")["meanCents"] for t in THRESHOLDS) - min(q(fleet, f"{t}:{c}")["meanCents"] for t in THRESHOLDS)
    noise = max(spread(f, 3000) for f in FLEETS)
    insertion_spread = max(abs(q(f, "0:0")["meanCents"] - q(f, "2:0")["meanCents"]) for f in FLEETS)
    short_gap = max(fq1[f"{f}/{p_}"]["2:250"]["meanCents"] for f in FLEETS for p_ in ("empty", "light"))
    windows_short = min(fq1[f"{f}/{p_}"][v]["windowsMin"] for f in FLEETS for p_ in ("empty", "light") for v in fq1[f"{f}/{p_}"] if fq1[f"{f}/{p_}"][v]["windowsMin"] is not None)
    windows_full = min(q(f, v)["windowsMin"] for f in FLEETS for v in fq[f"{f}/full"] if q(f, v)["windowsMin"] is not None)
    capped = ("2:250", "2:1000", "2:3000")
    shed10 = [C(v, l, 10)["shed"] for v in capped for l in LAYOUTS]
    grid_required = [C("1000", l, c) for l in LAYOUTS for c in (5, 10)]
    required_rates, required_503s = [r["perMinuteAll"] for r in grid_required], [r["unavailable"] for r in grid_required]
    capped_503 = max(C(v, l, c)["unavailable"] for v in ("2:250", "2:3000") for l in LAYOUTS for c in (5, 10))
    commit = {f: max(rl[f"{f}/full/{v}/800"]["commitP95"] for v in ("1000", "2:3000")) for f in FLEETS}
    commit_left = 800 - GUARD_MS
    closed2000 = sum(q(f, "2:250")["meanCents"] - q(f, "2:2000")["meanCents"] for f in FLEETS)
    closed3000 = sum(q(f, "2:250")["meanCents"] - q(f, "2:3000")["meanCents"] for f in FLEETS)
    deadline = lambda row: row["stops"].get("DEADLINE", 0)
    extra_routing = {f: L(f, "2:250")["routingP50"] - L(f, "2:0")["routingP50"] for f in FLEETS}
    unavailable20 = rl["20/full/2:3000/800"]["refinements"].get("UNAVAILABLE", 0)
    rate_kept = {(f, v): rl[f"{f}/full/{v}/800"]["perMinute"] / L(f, v)["perMinute"] * 100 for f in FLEETS for v in ("2:3000", "1000")}
    p50_3000, p50_250 = [q(f, "2:3000")["p50ms"] for f in FLEETS], [q(f, "2:250")["p50ms"] for f in FLEETS]
    reports = data["first"]["reports"] + data["rerun"]["reports"]

    p("WaterFlex booking configuration study explained", "Title")
    p("Which refinement cap, window threshold, booking pool layout and publish reserve the booking search should ship with, and why", "Subtitle")
    p("October 19 2026  |  Booking configuration study, 2026-10-17 to 2026-10-19, with the full-schedule rerun after 530d356")
    p(f"Over three days the booking search was run against every setting that could change its default: the window threshold crossed with the optional "
      f"refinement cap up to the new 3000 ms limit, empty to booked-out schedules under the dynamic horizon, 5 to 30 technicians, seven booking pool "
      f"layouts with 5 and 10 dealerships at once, and the publish reserve. Quality was the judge: each variant's best offer was priced against an "
      f"uncapped search on the same case. Across {data['first']['cells'] + data['rerun']['cells']:,} cells and {reports:,} kit reports, today's 250 ms cap "
      f"left a booked-out schedule {span(gap['2:250'], lambda c: money(c, False))} a booking dearer than the uncapped search, no better than running no "
      f"refinement; a 3000 ms cap brought the gap to between {money(min(gap['2:3000']))} and {money(max(gap['2:3000']))}, at a search p50 of "
      f"{span(p50_3000, secs)}; under 10 clients the scheduler shed {num(min(shed10))} to {num(max(shed10))} percent of optional refinements on every "
      f"layout, so a bigger cap costs little when the pool is busy; and the publish reserve floor went to 800 ms once the study showed that the database "
      f"admission guard, not the commit, refused searches below it. [1]")
    sub("What the measurements tell us")
    table(["Setting", "Today", "Recommended", "What the study found"], [
        ["Optional refinement cap", "250 ms", "3000 ms (2000 ms if 3 s is too slow)",
         f"On full schedules 250 ms is {span(gap['2:250'], lambda c: money(c, False))} a booking dearer than uncapped; 3000 ms is within {money(max(gap['2:3000']), False)}"],
        ["Window threshold", "2", "2", f"Never fired: every search found at least {windows_short} windows. A required refinement turns load into 503s"],
        ["Booking pool layout", "1 replica, 2 slots, 2 vCPU", "1 replica, slots equal to vCPU (3 or 4)",
         f"1x4 answered {num(C('2:3000', '1x4', 10)['perMinuteAll'])} searches a minute at 10 clients, the most of any layout"],
        ["Publish reserve", "1000 ms (floor 800 ms since 530d356)", "Keep 1000 ms",
         f"800 ms is safe for the contract, but at 20 technicians {unavailable20} of {rl['20/full/2:3000/800']['searches']} refinements were lost"]], [1.3, 1.1, 1.5, 3.0])
    p(f"The quality results on booked-out schedules come from the rerun, which fixed a fault that had made the first study search 2 dates where the "
      f"schedule called for 15 weekdays (section 3); everything else comes from the first study. Differences under about {money(noise, False)} a booking "
      f"are noise: four thresholds that ran the identical search landed that far apart (section 6). The conclusions rest on gaps of $1 to $4 or on "
      f"mechanisms read from the scheduler's code, so they hold at that noise level.")
    p("The report explains each result through what the scheduler recorded: how the refinement ended (complete, shortened by its cap or by the deadline, "
      "shed, unavailable), how many regular windows the first pass found, and the time each search spent waiting for a slot, routing, in the engine, in "
      "the database and committing. Latencies are measured at the caller. None of it is customer traffic: the addresses are random points inside Omaha "
      "and the schedules are synthetic. [1, 2]")
    sub("How to read this report")
    table(["Sections", "What you will learn"], [
        ("1 to 3", "What was measured, how a booking search spends its time, and what the rerun fixed."),
        ("4 to 7", "Offer quality: the gap each cap leaves on a booked-out schedule, why it closes only past 2000 ms, the noise floor, and shorter horizons."),
        ("8 and 9", "The price: one desk's latency and throughput, and several clients on seven booking pool layouts."),
        ("10", "The publish reserve: why its floor is 800 ms and why the default should stay at 1000 ms."),
        ("11 to 13", "The hypotheses recorded before the run, the conclusions, and what should happen next."),
        ("Appendices", "Every quality cell, every one-desk run, every capacity run, every reserve run, definitions and sources.")], [1.45, 5.45])

    heading("1 What was measured")
    p("The scheduler is reached only through its public API, so the host conformance kit plays the host. Two tools ran. quality.ts builds a schedule, "
      "optimizes every date overnight, then searches the same job addresses once per variant on that held schedule and prices each variant's best offer "
      "on a copy of the day with routes/evaluate. load.ts is a booking desk: it searches, selects an offer and confirms it (30 percent of searches) or "
      "releases it, so the schedule fills as confirms land. The runner, infra/run-booking-configuration-study.ps1, set the booking pool's shape for each "
      "part, restarted only that pool, and recorded every kit report. [1, 2]")
    table(["Dimension", "Values in this study"], [
        ("Variants", "threshold:cap. Thresholds 0, 2, 5 and 10 crossed with caps 0, 250, 1000, 2000 and 3000 ms, plus 1000: a threshold no search reaches, "
                     "so the refinement is required and never capped. 21 variants; 1000 is the reference, 2:250 is today, x:0 is insertion only"),
        ("Schedules", "empty (no visits, 2 dates), light (1 visit a day, 2 dates), half (7 weekdays booked out, 9 dates), full (15 weekdays booked out, the "
                      "15-weekday cap). The dynamic horizon (lead 2 weekdays, full at 85 percent) follows the schedule"),
        ("Fleets", "5, 10, 15, 20 and 30 technicians"),
        ("Policy", "bookingDeadlineMs 5000, offerLimit 4, windowHours 2; cold roads on every run, as for a real new address"),
        ("Stack", f"deploy/vm rehearsal: nginx proxy, booking pool of {shape['bookingReplicas']} replica with {shape['bookingSlots']} slots on "
                  f"{shape['bookingCpus']} vCPU ({shape['bookingMemory']}, database pool {shape['bookingPool']}), daily pool on {shape['dailyCpus']} vCPU, "
                  f"Omaha routing capped at {shape['routingCpus']} vCPU; shipped routing and admission defaults"),
        ("Machine", "One AMD Ryzen 9 7900X (12 cores, 24 hardware threads). Docker cpus count threads, as a VM vCPU does, so figures are per vCPU")], [1.4, 5.5])
    table(["Part", "Tool", "File", "Cells", "What it judged"], [
        ["Quality", "quality.ts", "first study", 40, "5 fleets x 4 schedules x 2 schedule seeds, 20 cases each, all 21 variants"],
        ["Load", "load.ts", "first study", 840, "One desk, 5 fleets x 4 schedules x 21 variants x 2 repeats, 120 s each"],
        ["Capacity", "load.ts", "first study", 224, "5 or 10 tenants with one desk each, 7 layouts, variants 2:250, 2:1000, 2:3000 and 1000, 10 and 5 technicians, half and full"],
        ["Reserve", "both", "first study", 80, "Reserve 500 and 250 ms, 5 to 30 technicians, half and full, 2:3000 and 1000"],
        ["Quality", "quality.ts", "rerun", 10, "Full schedules only, after the horizon fix, all 21 variants"],
        ["Reserve", "both", "rerun", 25, "Reserve 800 ms, full schedules, 5 to 30 technicians"]], [.8, .8, .8, .5, 4.0], size=9)
    p("A run's numbers come from three places: the kit's own clock at the caller (each search's time, and p50, p95 and p99 per run), the OfferSet each "
      "search returned (refinement, refinementStop, regularWindowsFound, searchComplete), and the scheduler's request log line for each call, which "
      "attributes its time to slot wait, routing, the engine, the database and the commit. [2, 5]")
    sub("The measures")
    table(["Measure", "How to interpret it"], [
        ("Gap", "A variant's best offer's added fleet cost minus the reference's on the same case, averaged per booking. Positive means a dearer slot."),
        ("Cheaper / dearer", "Cases whose best offer cost less or more than the reference's; the rest cost the same."),
        ("Same offer set", "Cases that returned exactly the reference's offers. A different set can still cost the same."),
        ("Search p50", "The median time of the variant's searches at the caller. In quality runs, one search at a time on a held schedule."),
        ("Stopped by", "How SHORTENED refinements ended: CAP (their own cap) or DEADLINE (the 5 s deadline less the reserve)."),
        ("Searches per minute", "Searches one desk, or all clients together, had answered per minute of the run."),
        ("Shed", "Share of answered searches whose optional refinement was skipped because a booking was waiting for a slot."),
        ("503", "Searches answered unavailable: the search could not finish inside its deadline."),
        ("BUSY resends", "Searches answered 429 BUSY and resent by the kit after Retry-After."),
        ("Commit", "Scheduler time from the start of the commit (validation, offer set, publish) to the end of the search.")], [1.55, 5.35])

    heading("2 How a booking search spends its time")
    p("A booking search is a short pipeline inside a 5 s deadline. Knowing its stages makes the rest of the report readable, because every result below "
      "is one stage running out of time or being skipped. [3]")
    table(["Stage", "What happens", "What limits it"], [
        ("Admission", "The search takes one of the replica's slots; confirms go first and have a slot of their own.", "Other searches on the replica"),
        ("First pass", "The job is inserted into every technician-day on the horizon; the regular windows are counted.", "Days on the horizon and fleet size"),
        ("Shed check", "An optional refinement is skipped when any booking is waiting in this replica's queue.", "queuedBookings() > 0"),
        ("Refinement", "A local search that may move existing visits to make room more cheaply. Its first step fetches routes for the neighborhoods around the "
                       "insertion candidates.", "Its cap (optional only) and the deadline"),
        ("Commit", "Every candidate is validated, the offer set prepared and the holds published.", "The publish reserve and the database admission guard")], [1.2, 3.8, 1.9])
    figure("pipeline", "The five stages of one booking search, left to right, over its 5 s deadline: exploration, then the publish reserve with the database "
                       "admission guard inside it. The dashed path is the shed under load; the refinement is the stage the cap governs. Read from the code.")
    p("Whether the refinement is optional is the window threshold's only job. When the first pass finds no more distinct regular windows than the "
      "threshold, the refinement is required: it is never shed and never capped, and it runs until the deadline less the reserve. Otherwise it is "
      "optional: it is shed under load and stops at its cap. The cap's clock starts before the neighborhood fetch, so on cold roads that fetch is paid "
      "out of the cap. [3]")
    sub("Why quality is judged against the uncapped search")
    p("The reference variant, 1000, makes every refinement required and uncapped, so it is the best answer the 5 s deadline allows. It is not a true "
      "optimum: from 15 technicians its own refinement runs into the deadline (section 4). A variant level with it is as good as this deadline can buy, "
      "and its gap measures what a cap leaves on the table.")

    heading("3 What changed between the first study and the rerun")
    p(f"The first study ({data['first']['cells']:,} cells, runner {data['first']['revision'][:7]}) finished with no contract defects but had two faults that "
      f"made its booked-out results unusable. Both were fixed in 530d356 before the rerun ({data['rerun']['cells']} cells, runner "
      f"{data['rerun']['revision'][:7]}). [1, 6]")
    lead("Quality horizon.", f"quality.ts read the dynamic horizon after its overnight optimization. Optimized days fall below the 85 percent full mark, so "
         f"every quality case searched 2 dates and {windows_short} windows instead of the 15 weekdays the schedule called for. The fix reads the horizon "
         f"from the booked schedule before optimizing; on full schedules the rerun then searched 21 calendar days and at least {windows_full} windows on "
         f"every case. Empty and light schedules have a 2-date horizon either way, so their first-study results stand. Half schedules were not rerun.")
    lead("Reserve floor.", f"At reserves of 500 and 250 ms every uncapped search on a full schedule was answered 429 BUSY, and the kit resent it for about "
         f"three minutes ({len(data['first']['unanswered'])} runs never got an answer). SearchDatabaseTelemetry refuses database admission when less than "
         f"the pool's connect timeout plus validation timeout ({GUARD_MS} ms) is left. The floor is now 800 ms, enforced at startup and by a test that reads "
         f"both timeouts from application.properties. Section 10 explains the arithmetic.")
    table(["Part", "First study", "Rerun"], [
        ["Quality, empty and light", "Valid (2-date horizon either way)", "Not rerun"],
        ["Quality, half", "2-date searches; shown as an indication only", "Not rerun"],
        ["Quality, full", "Not valid (2 dates instead of 15 weekdays)", "Valid, all 21 variants"],
        ["Load, one desk", "Valid (load.ts read the horizon correctly)", "Not rerun"],
        ["Capacity", "Valid", "Not rerun"],
        ["Reserve", "500 and 250 ms: refused at publish", "800 ms: valid"]], [1.8, 2.6, 2.5])

    heading("4 Offer quality on booked-out schedules")
    p(f"Each case is one job address searched once per variant on the same optimized full schedule; 40 cases per cell (2 schedules of 20). Figure {FIG['gap']} shows "
      "the gap each cap leaves at threshold 2, fleet by fleet; the table beneath it adds how many of the 40 bookings paid more than they would have "
      "with the uncapped search. [1]")
    figure("gap", "Full schedules, threshold 2: mean added cost per booking against the uncapped search, by fleet size and refinement cap. Darker is "
                  "dearer; the red cell is today's 250 ms cap at the fleet where it costs most.")
    rows = []
    for f in FLEETS:
        for v in ("2:0", "2:250", "2:1000", "2:2000", "2:3000"):
            r = q(f, v)
            rows.append([f, v, money(r["meanCents"]), money(r["medianCents"]), f"{r['cheaper']} / {r['dearer']}", r["same"], num(r["p50ms"]), stops(r)])
        r = q(f, "1000")
        rows.append([f, "1000 (ref)", "", "", "", "", num(r["p50ms"]), f"{stops(r)}; incomplete {r['incomplete']}"])
    table(["Techs", "Variant", "Mean", "Median", "Cheaper / dearer", "Same offer set", "Search p50 ms", "Stopped by"], rows, [.5, .8, .65, .65, .9, .7, .8, 1.9], size=9)
    p(f"Cap 250 is no better than cap 0 on a booked-out schedule: 2:250 is {span(gap['2:250'], lambda c: money(c, False))} a booking dearer than the "
      f"reference and 2:0 is {span(gap['2:0'], lambda c: money(c, False))}. On 2-date horizons the same 250 ms lands within {money(short_gap, False)} "
      f"(section 7). The gap shrinks with the cap and reaches {span(gap['2:3000'], money)} at 3000 ms, median $0.00 at every fleet. 10 and 20 technicians "
      f"are the hardest cases ({money(q(10, '2:250')['meanCents'], False)} and {money(q(20, '2:250')['meanCents'], False)} at 250 ms).")
    p(f"At 250 and 1000 ms the median gap is close to the mean, so most bookings pay a little extra. At 2000 and 3000 ms the median is $0.00 and the "
      f"mean comes from a minority of dearer cases; at 20 and 30 technicians that minority is larger ({q(20, '2:3000')['dearer']} and "
      f"{q(30, '2:3000')['dearer']} of 40), which is where 3000 ms falls short of the reference. The reference itself stopped on the deadline in "
      f"{deadline(q(15, '1000'))}, {deadline(q(20, '1000'))} and {deadline(q(30, '1000'))} of 40 cases at 15, 20 and 30 technicians, so from 15 "
      f"technicians 3000 ms and the reference are both spending most of the deadline.")
    figure("left", "The gap left at today's 250 ms cap (hollow), at 2000 ms (tick) and at 3000 ms (red), one row per fleet on one dollar scale; the "
                   "line spans today's cap to 3000 ms.")
    sub("2000 or 3000 ms")
    rows = []
    for f in FLEETS:
        base = q(f, "2:250")["meanCents"]
        rows.append([f, money(base), money(base - q(f, "2:2000")["meanCents"], False), money(base - q(f, "2:3000")["meanCents"], False),
                     secs(q(f, "2:2000")["p50ms"]), secs(q(f, "2:3000")["p50ms"])])
    table(["Techs", "Gap at 2:250", "Closed by 2000", "Closed by 3000", "p50 at 2000", "p50 at 3000"], rows, [.6, 1.1, 1.1, 1.1, 1.0, 1.0])
    p(f"Summed over the five fleets, 2000 ms closes {money(closed2000, False)} of the {money(closed3000, False)} that 3000 ms closes "
      f"({100 * closed2000 / closed3000:.0f} percent) for about 1 s less per search. The shortfall sits at 10 and 30 technicians. If the desk can wait "
      f"3 s, 3000 is the better default; 2000 is the fallback if it cannot.")

    heading("5 Why the gap closes only past 2000 ms")
    p("The refinement's cap is measured from the moment the refinement starts, and its first step fetches routes for the neighborhoods around the "
      "insertion candidates (SnapshotRouting.neighborhoods inside the cap's checkpoint in BookingOffers). On a full schedule those neighborhoods span 21 "
      "calendar days and 168 windows, not 2 dates and 16. Routing time at one desk on full schedules shows how much of a 250 ms cap that fetch takes: [3, 5]")
    rows = []
    for f in FLEETS:
        extra = extra_routing[f]
        rows.append([f, f"{num(L(f, '2:0')['routingP50'])} ms", f"{num(L(f, '2:250')['routingP50'])} ms", f"{num(extra)} ms" if extra > 10 else "none measurable",
                     "all of it" if extra >= 250 else (f"about {100 * extra / 250:.0f}%" if extra > 10 else "small")])
    table(["Techs", "Routing p50, 2:0", "Routing p50, 2:250", "Extra for the refinement", "Share of a 250 ms cap"], rows, [.6, 1.1, 1.1, 1.6, 1.6])
    figure("routing", "The scheduler's routing time per search at one desk on full schedules, by fleet, for insertion only, today's cap, 3000 ms and the "
                      "uncapped search.")
    figure("parts", "The scheduler's p50 time by part at 30 technicians on a full schedule, for five variants, against the 5 s deadline; the parts are "
                    "separate medians, so they add up only roughly.")
    p(f"These are p50s of separate runs, so the difference is an estimate, not a per-search measurement. At 20 and 30 technicians a 250 ms refinement is "
      f"mostly fetching legs ({num(extra_routing[20])} and {num(extra_routing[30])} ms more than insertion only). At 5 and 10 the fetch is cheap, and the "
      f"gap there says the engine itself needs more than a quarter second to make progress over 21 days; on a 2-date horizon it does not (the engine's "
      f"progress is not logged, so this part is an inference).")
    p(f"The ceiling comes from the other end of the deadline. With the default 1000 ms reserve about 4000 ms is left for the first pass and the "
      f"refinement together. At 30 technicians the first pass alone takes about {secs(q(30, '2:0')['p50ms'])}, so a 3000 ms refinement brings the search "
      f"to {secs(q(30, '2:3000')['p50ms'])} p50 and the deadline already cuts {deadline(q(30, '2:3000'))} of 40 searches before the cap. A larger cap "
      f"would be cut by the deadline more often at large fleets and buy nothing at small ones, where 3000 ms already matches. 3000 ms is about the most "
      f"this deadline can use.")
    sub("Where the uncapped search's time goes")
    ref = {f: L(f, "1000") for f in FLEETS}
    p(f"The reference's routing p50 rises from {num(ref[5]['routingP50'])} ms at 5 technicians to {num(ref[10]['routingP50'])} ms at 10, "
      f"{secs(ref[15]['routingP50'], 2)} at 15 and about {secs(ref[20]['routingP50'])} at 20 and 30, while its engine time falls from "
      f"{secs(ref[5]['engineP50'])} to about {secs(ref[30]['engineP50'])} (Figure {FIG['parts']}). Its p50 stays between {secs(min(r['p50'] for r in ref.values()))} "
      f"and {secs(max(r['p50'] for r in ref.values()))} because the deadline ends it. From 20 technicians the reference spends most of its deadline fetching "
      f"legs and has little left to improve the offer; the code comment on the BUSY path in BookingOffers describes the same thing. [3]")
    p(f"Two things follow. The reference is weakest exactly where 2:3000 looks furthest from it, so the 20 and 30 technician gaps say less about 2:3000 "
      f"than the small fleets do. And 2:3000 at the same fleets spends only {num(L(20, '2:3000')['routingP50'])} to {num(L(30, '2:3000')['routingP50'])} ms "
      f"routing. Why the required refinement fetches so much more than the capped one is not established; the default never runs uncapped, but it would "
      f"matter for any client whose threshold makes refinements required.")

    heading("6 The window threshold, and how much noise there is")
    p(f"The threshold is compared with the first pass's distinct regular windows (regularWindowsFound). The fewest any search found was {windows_short}, on "
      f"2-date horizons, and {windows_full} on every full-schedule search, so no threshold up to 10 ever made a refinement required, and thresholds 0, 2, 5 "
      f"and 10 at one cap ran the same search. That makes them a free repeatability test. [1]")
    rows = [[f] + [money(q(f, f"{t}:250")["meanCents"]) for t in THRESHOLDS] + [money(spread(f, 250), False)]
            + [money(q(f, f"{t}:3000")["meanCents"]) for t in THRESHOLDS] + [money(spread(f, 3000), False)] for f in FLEETS]
    table(["Techs", "0:250", "2:250", "5:250", "10:250", "Spread", "0:3000", "2:3000", "5:3000", "10:3000", "Spread"], rows, [.45] + [.6] * 10, size=9)
    figure("noise", "Four thresholds that ran the identical search, at cap 250 ms (hollow) and at cap 3000 ms (solid), one row per fleet; the width "
                    "of each cluster is run-to-run noise.")
    p(f"The same configuration lands up to {money(noise, False)} a booking apart over 40 cases. The cause is not pinned down: the two insertion-only "
      f"variants, 0:0 and 2:0, which run no refinement at all, also differ by up to {money(insertion_spread, False)}, so it is not only the refinement's "
      f"clock. Read differences smaller than about {money(noise, False)} as noise. By that measure 2:3000 is level with the reference at every fleet, and "
      f"2:250 is clearly behind it at 10, 20 and 30 technicians and just outside noise at 5 and 15.")
    p("The threshold would matter for a client whose horizon shows very few open windows, such as a single technician booked out for weeks. None of the "
      "profiles here came close, and section 9 shows what a required refinement does under load.")

    heading("7 Shorter horizons")
    p("Empty and light schedules have a 2-date horizon, so the first study's quality results are valid for them. Half-schedule figures are from 2-date "
      "searches because of the horizon fault (a real half schedule searches 9 dates) and are shown only as an indication. Mean gap against the uncapped "
      "search, 40 cases per cell. [1]")
    rows = [[profile + (" (2 dates, fault)" if profile == "half" else ""), f] + [money(fq1[f"{f}/{profile}"][v]["meanCents"]) for v in ("2:0", "2:250", "2:1000", "2:3000")]
            for profile in ("empty", "light", "half") for f in FLEETS]
    table(["Schedule", "Techs", "2:0", "2:250", "2:1000", "2:3000"], rows, [1.6, .6, .9, .9, .9, .9])
    figure("horizons", "Today's 250 ms cap on a light schedule with a 2-date horizon (hollow) and on a full schedule with 15 weekdays (red), one row "
                       "per fleet.")
    light = lambda v: [fq1[f"{f}/light"][v]["meanCents"] for f in FLEETS]
    p(f"On an empty schedule every variant gave the uncapped search's offers; on a free day insertion is already optimal. On a light schedule no refinement "
      f"costs {span(light('2:0'), lambda c: money(c, False))} a booking, and 250 ms recovers nearly all of it ({span(light('2:250'), money)}). On 2-date "
      f"searches 250 ms is enough. Figure {FIG['horizons']} puts today's cap on the two horizons side by side: it was tuned where it works, and with the dynamic horizon a "
      f"booked-out client gets 15 weekdays, where it does not.")

    heading("8 One desk: latency and throughput")
    p("One desk per tenant, cold roads, the schedule filling as confirms land, 2 repeats pooled per cell. The load part read the dynamic horizon "
      "correctly, so its full-schedule figures are valid. Engine, routing and commit are the scheduler's own p50 from its request log; latency is the "
      "caller's. [1, 5]")
    rows = []
    for f in FLEETS:
        for v in ("2:0", "2:250", "2:1000", "2:3000", "1000"):
            r = L(f, v)
            rows.append([f, v, num(r["perMinute"], 1), num(r["p50"]), num(r["p95"]), num(r["p99"]), num(r["engineP50"]), num(r["routingP50"]),
                         f"{num(r['commitP50'])} / {num(r['commitP95'])}", f"{r['incomplete']} of {r['searches']}"])
    table(["Techs", "Variant", "Searches / min", "p50", "p95", "p99", "Engine", "Routing", "Commit p50 / p95", "Incomplete"], rows,
          [.45, .6, .65, .55, .55, .55, .6, .6, .85, .8], size=9)
    figure("desk", "Offers p95 at the caller for one desk on full schedules, by fleet and variant; the dashed line is the 5 s deadline.")
    p(f"Engine time is the cap: every capped variant spends almost exactly its cap in the engine on top of the first pass, and the first pass alone grows "
      f"from {num(L(5, '2:0')['engineP50'])} ms at 5 technicians to {num(L(30, '2:0')['engineP50'])} ms at 30. Commit time grows with the fleet, not the "
      f"cap. Slot wait was {num(max(L(f, v)['queueP50'] for f in FLEETS for v in ('2:0', '2:250', '2:3000', '1000')))} ms in every cell: at one desk nothing "
      f"queues.")
    assert all(L(f, "2:3000")["incomplete"] == 0 for f in FLEETS)
    assert all(L(f, "1000")["incomplete"] == L(f, "1000")["searches"] for f in FLEETS if f >= 10)
    p(f"The uncapped search reported searchComplete false on every search from 10 technicians up, because it always runs into the deadline. 2:3000 stays "
      f"complete at every fleet, with p99 under {secs(max(L(f, '2:3000')['p99'] for f in FLEETS))} against the 5 s deadline. A desk books far fewer than 15 "
      f"jobs a minute, so even 2:3000 keeps up with one desk; throughput matters when several clients share the pool (section 9).")
    sub("The price of quality")
    p("What each step up from 250 ms buys per second of extra search time. Gaps come from the quality rerun; p50 and throughput from the one-desk runs.")
    rows = []
    for f in FLEETS:
        g250, g2000, g3000 = (q(f, f"2:{c}")["meanCents"] for c in (250, 2000, 3000))
        p250, p2000, p3000 = (L(f, f"2:{c}")["p50"] / 1000 for c in (250, 2000, 3000))
        step = g2000 - g3000
        rows.append([f, f"{money(g250 - g2000, False)} / {p2000 - p250:.1f} s", money((g250 - g2000) / (p2000 - p250), False),
                     f"{money(step)} / {p3000 - p2000:.1f} s", money(step / (p3000 - p2000), False) if step > 0 else "none",
                     f"{100 * L(f, '2:3000')['perMinute'] / L(f, '2:250')['perMinute']:.0f}%"])
    table(["Techs", "250 to 2000: closed / extra p50", "Per second", "2000 to 3000: closed / extra p50", "Per second", "Throughput kept, 3000 vs 250"], rows,
          [.5, 1.6, .8, 1.6, .8, 1.2])
    p("The last second, from 2000 to 3000 ms, pays at 10 and 30 technicians and is within noise at 5, 15 and 20. The throughput column is a desk's own "
      "ceiling, not the client's booking rate: people take far longer than 4 s between searches, so a single desk never waits on the scheduler.")

    heading("9 Several clients: booking pool layouts")
    p(f"5 or 10 dealerships at once, each its own tenant with one desk, against seven booking pool layouts inside 4 vCPU (replicas x slots, one vCPU a "
      f"slot). The table is 10 technicians on a full schedule at 10 clients, for today's default, 3000 ms and the uncapped variant; searches a minute are "
      f"summed over the clients. [1]")
    rows = []
    for l in LAYOUTS:
        a, b, c = C("2:250", l, 10), C("2:3000", l, 10), C("1000", l, 10)
        rows.append([l, a["threads"], num(a["perMinuteAll"]), secs(a["p95"]), f"{num(a['shed'])}%", num(b["perMinuteAll"]), secs(b["p95"]), f"{num(b['shed'])}%",
                     num(c["perMinuteAll"]), c["unavailable"]])
    table(["Layout", "vCPU", "2:250 / min", "p95", "Shed", "2:3000 / min", "p95", "Shed", "1000 / min", "503s"], rows, [.6, .5, .75, .6, .55, .8, .6, .55, .75, .5])
    figure("layouts", "10 technicians, full schedule, 10 clients: searches per minute over all clients by booking pool layout, for today's cap "
                      "(hollow), 3000 ms (red) and the uncapped search (tick).")
    p(f"At 3 and 4 vCPU one replica wins. 1x4 gives the most searches a minute at 10 clients for every variant ({num(C('2:3000', '1x4', 10)['perMinuteAll'])} "
      f"with 2:3000) and 1x3 beats 3x1. At 2 vCPU the order reverses for 2:250: 2x1 answered {num(C('2:250', '2x1', 10)['perMinuteAll'])} searches a minute "
      f"against {num(C('2:250', '1x2', 10)['perMinuteAll'])} on 1x2. One admission queue never leaves a slot idle while a call waits elsewhere, and the "
      f"proxy's round-robin is blind to load; the reversal at 2 vCPU is measured but not explained. The confirm-only slot is added on top of the search "
      f"slots, so it is not the cause.")
    sub("How shedding works")
    p("Each booking replica admits searches into its slots through one queue. When a search's first pass is done and its refinement is optional, the "
      "scheduler checks that replica's queue: if any booking call is waiting, the refinement is skipped and the search publishes its first-pass offers at "
      "once (queuedBookings() > 0 in BookingOffers). A required refinement is never skipped this way. That one rule explains the grid. [3, 7]")
    p(f"With one replica and many clients some call is nearly always waiting, so almost every refinement is shed ({num(C('2:3000', '1x4', 10)['shed'])} "
      f"percent on 1x4). Searches are short, slots turn over fast, and throughput and p95 are the best of any layout. With several replicas the proxy spreads "
      f"calls round-robin, so each replica of 4x1 sees about a quarter of them; its queue is empty more often, and {num(C('2:3000', '4x1', 10)['shed'])} to "
      f"{num(C('2:250', '4x1', 10)['shed'])} percent are shed instead of 100. Each refinement that does run holds its slot for its whole cap: 250 ms is a "
      f"small cost and 3000 ms is not, which is why 2:3000 on 4x1 answered {num(C('2:3000', '4x1', 10)['perMinuteAll'])} searches a minute against "
      f"{num(C('2:250', '4x1', 10)['perMinuteAll'])} with p95 {secs(C('2:3000', '4x1', 10)['p95'])}.")
    p(f"A shed search gets the insertion-only offer, which on a full schedule is {span(gap['2:0'], lambda c: money(c, False))} dearer than the reference "
      f"(section 4). During busy periods that is what most bookings get, whatever the cap. The design choice behind the rule is to serve the next customer "
      f"rather than polish an offer when the pool is busy, and the study supports it: the capped variants answered 503 at most {capped_503} times in any run, "
      f"while the uncapped variant, whose refinements are required, answered 503 in every cell, up to {max(required_503s)} times, and held only "
      f"{num(min(required_rates))} to {num(max(required_rates))} searches a minute.")
    sub("Five clients")
    rows = []
    for l in LAYOUTS:
        a, b, c = C("2:250", l, 5), C("2:3000", l, 5), C("1000", l, 5)
        rows.append([l, a["threads"], num(a["perMinuteAll"]), secs(a["p95"]), f"{num(a['shed'])}%", num(b["perMinuteAll"]), secs(b["p95"]), f"{num(b['shed'])}%",
                     num(c["perMinuteAll"]), c["unavailable"]])
    table(["Layout", "vCPU", "2:250 / min", "p95", "Shed", "2:3000 / min", "p95", "Shed", "1000 / min", "503s"], rows, [.6, .5, .75, .6, .55, .8, .6, .55, .75, .5])
    p(f"At 5 clients the queue is empty more often, so more refinements run and the cap shows. 3x1 falls from {num(C('2:250', '3x1', 5)['perMinuteAll'])} "
      f"searches a minute with 2:250 to {num(C('2:3000', '3x1', 5)['perMinuteAll'])} with 2:3000, and its p95 rises to {secs(C('2:3000', '3x1', 5)['p95'])}. "
      f"So a bigger cap trades throughput for quality only where refinements still run: on 1x3 and 1x4 the 3000 ms default changes little under load. "
      f"Keeping quality during busy periods needs more slots, not a bigger cap or more replicas. Appendix C lists every capacity run, including 5 "
      f"technicians and half schedules, which follow the same order.")

    heading("10 The publish reserve")
    p("The reserve is the end of the 5 s deadline kept back from exploration so the search can validate its candidates, prepare the offer set and publish "
      "its holds. The rerun ran the reserve arm at the new 800 ms floor on full schedules: 2 load repeats per variant and 20 quality cases per fleet. [1, 6]")
    rows = []
    for f in FLEETS:
        for v in ("1000", "2:3000"):
            new = rl[f"{f}/full/{v}/800"]
            rows.append([f, v, f"{num(new['commitP50'])} / {num(new['commitP95'])}", num(new["perMinute"], 1), num(L(f, v)["perMinute"], 1), new["busy"],
                         new["unavailable"], new["incomplete"]])
    table(["Techs", "Variant", "Commit p50 / p95 ms", "Searches / min at 800", "At 1000", "BUSY resends", "503s", "Incomplete"], rows, [.5, .7, 1.2, 1.0, .7, .8, .5, .9])
    figure("reserve", f"What has to fit inside the reserve: the {GUARD_MS} ms the database admission guard demands plus the worse commit p95 at the "
                      "800 ms floor, per fleet, against the 800 ms floor and the 1000 ms default.")
    p(f"No search was refused BUSY or answered 503 at 800 ms at any fleet, and commit p95 never exceeded {num(max(commit.values()))} ms. 2:3000 ran "
      f"{span([rate_kept[(f, '2:3000')] for f in FLEETS if f != 20], lambda x: f'{x:.0f}')} percent of its 1000 ms rate at every fleet but 20, where it ran "
      f"{rate_kept[(20, '2:3000')]:.0f} percent; the uncapped search, which always runs to the deadline, ran "
      f"{span([rate_kept[(f, '1000')] for f in FLEETS], lambda x: f'{x:.0f}')} percent, because it spends the extra 200 ms exploring. On quality, 2:3000 "
      f"against the reference at 800 ms landed between {money(min(rq[f'{f}/full']['2:3000']['meanCents'] for f in FLEETS))} and "
      f"{money(max(rq[f'{f}/full']['2:3000']['meanCents'] for f in FLEETS))} on 20 cases per fleet, on both sides of zero. With half the cases these "
      f"figures are noisier than the 40-case ones, and they show no quality gain. The lower reserve is safe; it is not a quality lever.")
    sub("How the reserve, the admission guard and the commit fit together")
    p(f"Exploration must stop when the reserve begins; inside it the search validates every candidate, prepares the offer set and publishes its holds, which "
      f"needs the database. Every database call goes through an admission guard that refuses with BUSY when less than the pool's connect timeout plus "
      f"validation timeout is left (250 + 250 = {GUARD_MS} ms), so the call cannot hang past the deadline. The reserve therefore has to cover the guard and "
      f"the commit work together. [3]")
    c30 = commit[30]
    table(["Reserve", "Guard", "Left for the commit", "Commit p95 at 30 techs", "Outcome"], [
        ["250 ms", f"{GUARD_MS} ms", "none", f"{num(c30)} ms", "every long search refused BUSY"],
        ["500 ms", f"{GUARD_MS} ms", "none", f"{num(c30)} ms", "every long search refused BUSY"],
        ["800 ms", f"{GUARD_MS} ms", f"{commit_left} ms", f"{num(c30)} ms", f"passes with {num(commit_left - c30)} ms to spare"],
        ["1000 ms", f"{GUARD_MS} ms", f"{1000 - GUARD_MS} ms", f"{num(c30)} ms", f"passes with {num(1000 - GUARD_MS - c30)} ms to spare"]], [.8, .7, 1.2, 1.4, 2.6])
    figure("budget", f"The {GUARD_MS} ms guard plus the commit p95 at 30 technicians ({num(c30)} ms) inside a dashed frame of each reserve the studies ran; "
                     f"at 250 and 500 ms the stack overflows the frame, which is why every long search was refused there.")
    p("Short searches were never affected, because they finish exploring well before the reserve and reach the database with time to spare. That is why "
      "the half-schedule runs at 5 to 15 technicians passed even at 250 ms: what failed was any search that used its whole exploration time, which on full "
      "schedules is every uncapped one.")
    sub("Reserve detail by fleet")
    rows = []
    for f in FLEETS:
        r = rl[f"{f}/full/2:3000/800"]
        parts = [f"{k.lower()} {v}" for k, v in sorted(r["stops"].items())] + [f"unavailable {v}" for k, v in r["refinements"].items() if k == "UNAVAILABLE"]
        rows.append([f, f"{num(commit[f])} ms", f"{num(commit_left - commit[f])} ms", ", ".join(parts) + f", of {r['searches']}"])
    table(["Techs", "Worst commit p95", "Margin at 800 ms", "2:3000 refinement at 800 ms"], rows, [.6, 1.2, 1.1, 4.0])
    p(f"One result needs a caveat. At 20 technicians {unavailable20} of {rl['20/full/2:3000/800']['searches']} 2:3000 searches ended with refinement "
      f"UNAVAILABLE, which the API reports when the refinement could not get routing or database capacity in time (the REFINEMENT_BUSY path); those "
      f"searches returned their first-pass offers. At the 1000 ms default the same cell finished all {L(20, '2:3000')['searches']} searches on its cap. In "
      f"the same cell {deadline(rl['20/full/2:3000/800'])} searches stopped on the deadline at 800 ms, although a smaller reserve leaves more exploration "
      f"time, not less. Both point at that one cell being slower in the rerun, which needs a repeat to confirm. Either way the reading is the same: 800 ms "
      f"keeps the contract, 1000 ms keeps the refinements.")
    p(f"In the first study's 500 and 250 ms runs the uncapped variant on full schedules managed 2 to 4 searches a run, each refused BUSY and resent for "
      f"about three minutes. 2:3000 got through at 5 to 15 technicians but met BUSY at 30 ({first_rl['30/full/2:3000/250']['busy']} at 250 ms, "
      f"{first_rl['30/full/2:3000/500']['busy']} at 500 ms). That pattern, long searches refused and short ones not, is what pointed at the admission "
      f"guard rather than commit time (Appendix D).")

    heading("11 Hypotheses recorded before the run")
    p("The study file records eight hypotheses written before the first run. Each verdict below is computed from the tables above. [1]")
    finish_cap = min(q(f, f"2:{c}")["answered"] - deadline(q(f, f"2:{c}")) for f in (20, 30) for c in (2000, 3000))
    zero = lambda profile, source: [source[f"{f}/{profile}"]["0:0"]["meanCents"] for f in FLEETS]
    shed_grid = lambda layouts, clients: [C("2:1000", l, clients)["shed"] for l in layouts]
    gain = {c: C("2:1000", "1x4", c)["perMinuteAll"] / C("2:1000", "1x2", c)["perMinuteAll"] for c in (5, 10)}
    queue = [C("2:1000", l, c)["queueP50"] for l in LAYOUTS for c in (5, 10)]
    small_commit = [commit[f] for f in (5, 10, 15)]
    verdicts = [
        ("1", "Cap needed grows with the fleet; at 20 and 30 caps above 2000 stop on the deadline and none matches 1000.",
         f"Partly supported. 2:1000 does not match at 5 technicians ({money(q(5, '2:1000')['meanCents'])}) but 2:2000 does ({money(q(5, '2:2000')['meanCents'])}); "
         f"10 needs 3000 (2:2000 leaves {money(q(10, '2:2000')['meanCents'])}). At 20 and 30, 2000 and 3000 finished or stopped on their cap in {finish_cap} or "
         f"more of 40 cases, and 2:3000 leaves {money(q(30, '2:3000')['meanCents'])} to {money(q(20, '2:3000')['meanCents'])}. It is the uncapped search "
         f"that hits the deadline."),
        ("2", "x:0 is $1 to $3 dearer on half and full at 10 to 20 technicians, within $0.50 on empty and light.",
         f"Supported with an outlier. Full: 0:0 is {money(q(10, '0:0')['meanCents'])} at 10 and {money(q(20, '0:0')['meanCents'])} at 20, but "
         f"{money(q(15, '0:0')['meanCents'])} at 15. Empty {span(zero('empty', fq1), money)}; light {span(zero('light', fq1), money)}; half (2-date searches) "
         f"{span(zero('half', fq1), money)}."),
        ("3", "Thresholds 5 and 10 make the refinement required on a share of empty and light searches; inert on half and full.",
         f"Refuted for empty and light, supported for full. The first pass found at least {windows_short} windows on every 2-date search and {windows_full} on "
         f"every full one, so no threshold up to 10 ever fired."),
        ("4", "calculationMs grows with cap and fleet; routingMs with fleet and dates; queueMs near zero at one desk.",
         f"Supported. Engine time tracks the cap; insertion-only routing grows from {num(L(5, '2:0')['routingP50'])} to {num(L(30, '2:0')['routingP50'])} ms "
         f"with the fleet on full schedules; slot wait is 0 at one desk and {span(queue, num)} ms p50 under 5 and 10 clients (10 "
         f"technicians, full, 2:1000)."),
        ("5", "Throughput falls with the cap at one desk; 80 percent shed on 2 slots, under 30 percent on 4; 4 cores at least 1.8x of 2; 10 clients shed everywhere.",
         f"Mostly supported (10 technicians, full, 2:1000). One-desk throughput falls with the cap; 2 vCPU at 5 clients shed "
         f"{span(shed_grid(('1x2', '2x1'), 5), num)} percent and 10 clients shed {span(shed_grid(LAYOUTS, 10), num)} percent everywhere. Refuted on 4 vCPU at 5 "
         f"clients: {span(shed_grid(('1x4', '2x2', '4x1'), 5), num)} percent shed, not under 30. 1x4 against 1x2 is {gain[5]:.1f}x at 5 clients and "
         f"{gain[10]:.1f}x at 10."),
        ("6", "At equal cores one replica with more slots beats several replicas by 10 to 20 percent on p50 and routing.",
         f"Partly supported, and larger than predicted at 3 and 4 vCPU: 1x3 against 3x1 at 5 clients is {num(C('2:1000', '1x3', 5)['p50'])} against "
         f"{num(C('2:1000', '3x1', 5)['p50'])} ms p50 with {C('2:1000', '1x3', 5)['perMinuteAll'] / C('2:1000', '3x1', 5)['perMinuteAll']:.1f} times the "
         f"throughput. Reversed at 2 vCPU. Confounded: the multi-replica layouts also refine more, because they shed less."),
        ("7", "Offer limit 8.", "Dropped before the run: the API allows at most 4 offers per set."),
        ("8", "Commit under 150 ms p99 at 5 to 15 and 400 to 900 ms at 30; 500 ms reserve safe to 15; 250 fails at 20 and above.",
         f"Refuted. Commit p95 is {span(small_commit, num)} ms at 5 to 15 and {num(commit[30])} ms at 30, far under the guess. But 500 ms was not safe at any "
         f"fleet: the database admission guard, not commit time, sets the floor, which is why it is now 800 ms.")]
    table(["#", "Hypothesis", "Verdict"], [list(v) for v in verdicts], [.3, 2.4, 4.2], size=9)

    heading("12 Conclusions")
    p("Each conclusion says what was measured, why the scheduler behaves that way, and what follows. Where the cause is read from the code it says so; "
      "where it is an inference it says that too.")
    sub("1. Today's 250 ms cap does not work on a booked-out schedule")
    lead("Measured.", f"On full schedules 2:250 is {span(gap['2:250'], lambda c: money(c, False))} a booking dearer than the uncapped search, no better than "
         f"no refinement ({span(gap['2:0'], lambda c: money(c, False))}). On 2-date horizons the same 250 ms lands within {money(short_gap, False)}.")
    lead("Why.", f"The cap's clock starts before the refinement improves anything, because its first step fetches routes for the neighborhoods around the "
         f"insertion candidates (read from the code). At 20 and 30 technicians that fetch costs {num(extra_routing[20])} and {num(extra_routing[30])} ms more "
         f"than insertion only, most or all of 250 ms. At 5 to 15 technicians the fetch is cheaper, but a quarter second moves almost nothing in a search "
         f"space about ten times the 2-date one (an inference: the engine's progress is not logged).")
    lead("So.", "The 250 ms default was tuned on 2-date horizons. With the dynamic horizon a booked-out client gets 15 weekdays, and the cap has to grow with "
         "the horizon, not the fleet.")
    same3000 = [q(f, "2:3000")["same"] for f in (5, 10, 15)]
    sub("2. 3000 ms is enough, and about the most the deadline can use")
    lead("Measured.", f"2:3000's gap to the reference is between {money(min(gap['2:3000']))} and {money(max(gap['2:3000']))}, median $0.00 at every fleet, and returns the reference's exact offer "
         f"set in {span(same3000, str)} of 40 cases at 5 to 15 technicians. It finished or stopped on its own cap in "
         f"{min(q(f, '2:3000')['answered'] - deadline(q(f, '2:3000')) for f in FLEETS)} or more of 40 cases at every fleet.")
    lead("Why.", f"Given about 3 s the refinement settles on the same best offer as an unlimited one in most cases. The reference is bounded by the deadline "
         f"less the reserve, about 4 s, and from 15 technicians it runs into that bound ({deadline(q(15, '1000'))}, {deadline(q(20, '1000'))} and "
         f"{deadline(q(30, '1000'))} of 40 cases). At 30 technicians a 3000 ms refinement already brings the search to {secs(q(30, '2:3000')['p50ms'])}.")
    lead("So.", "A cap above 3000 ms would mostly turn into the deadline at large fleets and buy nothing at small ones. Going further needs a longer deadline, "
         "which was out of scope.")
    sub("3. The threshold does not matter for quality, but it must stay low")
    lead("Measured.", f"The first pass found at least {windows_short} regular windows on every 2-date search and {windows_full} on every full one, so "
         f"thresholds 0 to 10 never made a refinement required. Separately, the uncapped variant answered 503 on every layout under 5 and 10 clients.")
    lead("Why.", "A required refinement skips both safety valves: the shed check and the cap's checkpoint apply only to optional refinements (read from the "
         "code). Under load each such search holds its slot for up to 4 s, the queue grows, and searches that cannot start before their deadline are refused.")
    lead("So.", "The threshold is a guard for clients with almost no open windows, and today it never fires. Raising it would not improve offers on these "
         "profiles; it would only remove the shed and cap protections. Keep it at 2.")
    sub("4. Under load, shedding decides quality and throughput, not the cap")
    lead("Measured.", f"At 10 clients {num(min(shed10))} to {num(max(shed10))} percent of optional refinements were shed on every layout up to 4 vCPU. On "
         f"one-replica layouts 2:250 and 2:3000 behave alike; on multi-replica layouts 2:3000 sheds less and runs slower.")
    lead("Why.", "An optional refinement is shed whenever any booking is waiting in that replica's queue (read from the code). One replica sees every client's "
         "call in one queue, which is rarely empty under 5 or more clients. Several replicas split the queue, so each is empty more often, more refinements "
         "run, and each one holds a slot for its whole cap.")
    lead("So.", f"The bigger cap is safe: under load it degrades to insertion only instead of refusing bookings. The price is that busy-period bookings on full "
         f"schedules cost what insertion only costs ({span(gap['2:0'], lambda c: money(c, False))} more). Keeping quality in busy periods needs more slots.")
    sub("5. One replica with slots equal to vCPU is the best layout at 3 and 4 vCPU")
    lead("Measured.", f"1x4 gave the most searches a minute at 10 clients ({num(C('2:3000', '1x4', 10)['perMinuteAll'])} with 2:3000), the lowest p95 "
         f"({secs(C('2:3000', '1x4', 10)['p95'])}) and the fewest 503s for required refinements ({C('1000', '1x4', 10)['unavailable']}). 1x3 beat 3x1. At 2 "
         f"vCPU the order reversed for 2:250.")
    lead("Why.", "Partly the expected reasons: one admission queue never leaves a slot idle while a call waits elsewhere, and round-robin is blind to load. "
         "Partly conclusion 4: the single replica sheds more, so part of its lead is bought with quality in busy periods. The 2 vCPU reversal is not explained.")
    lead("So.", "Run the booking pool as one replica with 3 or 4 slots on as many vCPU. Add replicas for isolation and rolling restarts, knowing they shed less "
         "and answer slower with the 3000 ms cap.")
    sub("6. The reserve floor is set by database admission, and 1000 ms should stay the default")
    lead("Measured.", f"At 500 and 250 ms every long uncapped search was refused BUSY; at 800 ms nothing was refused, but {unavailable20} of "
         f"{rl['20/full/2:3000/800']['searches']} 2:3000 searches at 20 technicians lost their refinement. Commit p95 never exceeded {num(max(commit.values()))} ms.")
    lead("Why.", f"A search needs {GUARD_MS} ms left at the database admission guard before publishing (read from the code), so a reserve under {GUARD_MS} ms "
         f"guarantees a refusal for any search that uses its whole exploration time. Why refinements went UNAVAILABLE only at 20 technicians and only at "
         f"800 ms is not established.")
    lead("So.", "800 ms is a correct floor for the contract. Lowering the default buys no measurable quality and may cost refinements on large fleets. Keep "
         "1000 ms.")
    sub("7. How far to trust the numbers")
    p(f"The judge is modeled fleet cost on the kit's synthetic records, priced on a copy of the day. Identical searches spread by up to {money(noise, False)} "
      f"over 40 cases and the insertion-only variants by up to {money(insertion_spread, False)}, so the noise is not only the refinement's clock and its "
      f"cause is open. The conclusions above rest on gaps of $1 to $4 or on mechanisms read from the code. The exact gap at a given fleet does not: the "
      f"fleets are separate synthetic schedules, so the gap does not rise steadily with fleet size (at 15 technicians even insertion only came within "
      f"{money(q(15, '2:0')['meanCents'], False)}, against {money(q(10, '2:0')['meanCents'], False)} at 10).")

    heading("13 What should happen next")
    p(f"The study supports one change: raise the default optional refinement cap from 250 ms to 3000 ms (or 2000 ms), keep the threshold at 2 and the reserve "
      f"at 1000 ms, and lay the booking pool out as one replica with slots equal to vCPU. The cap is a policy change and the owner's call. 3000 ms closes "
      f"about {100 * closed3000 / closed2000 - 100:.0f} percent more of the gap than 2000 ms for about 1 s more per search.")
    table(["Question", "Next evidence to collect", "Decision informed"], [
        ("The cap", "None: the choice between 3000 and 2000 ms is the desk's tolerance for a 3 s search.", "The shipped default."),
        ("Half schedules", "Rerun the quality part on half schedules with the horizon fix (9 dates).", "Whether 9-date horizons need the full cap."),
        ("The 20-technician cell at 800 ms", "Repeat it at 800 and 1000 ms.", "Whether a lower reserve costs refinements."),
        ("Uncapped routing", "Log which legs a required refinement fetches at 20 and 30 technicians.", "Whether high-threshold clients are safe."),
        ("Above 30 technicians", "Measure commitMs and the admission margin before such a client goes live.", "Onboarding a large fleet."),
        ("After rollout", "Compare refinement outcomes per tenant (COMPLETE, CAP, DEADLINE, SHED, UNAVAILABLE) with this study's mix.", "Whether the study holds live.")],
          [1.5, 3.3, 2.1])
    sub("Risks and what to watch")
    table(["Risk", "How large", "What to watch after the change"], [
        ["Slower answers at the desk", f"Search p50 on a booked-out schedule goes from {span(p50_250, secs)} to {span(p50_3000, secs)}; 2-date horizons stay fast",
         "Booking search p95 in the request log; searches abandoned before an offer is picked"],
        ["Multi-replica pools", f"Unshed refinements now run 3 s; 3x1 at 5 clients fell from {num(C('2:250', '3x1', 5)['perMinuteAll'])} to "
                                f"{num(C('2:3000', '3x1', 5)['perMinuteAll'])} searches a minute", "Slot wait and scheduler_booking_refinement_shed_total"],
        ["Fleets above 30", f"Commit p95 {num(commit[30])} ms at 30 leaves {num(commit_left - commit[30])} ms at the 800 ms floor", "commitMs p95 per tenant"],
        ["Refinements lost at the floor", f"{unavailable20} of {rl['20/full/2:3000/800']['searches']} at 20 technicians at 800 ms", "The share of refinement UNAVAILABLE"],
        ["The judge itself", f"Identical searches differ by up to {money(noise, False)}", "Treat sub-dollar gaps as ties"],
        ["A threshold set high by a client", "A required refinement is never shed and answered 503 on every layout under load",
         "Threshold values in client policies"]], [1.5, 2.8, 2.6])
    p("Two cautions carry over. Compare only like with like: the first study's full-schedule quality figures are not valid, and its half-schedule ones are "
      "2-date searches. And read the rehearsal's numbers as a shape: the VM's clock speed will set its own absolute times.")
    sub("How the calculations were checked")
    p("Every table and chart value is computed from the two evidence files; nothing is typed in. report_data.py checks every kit report before use: no "
      "contract breach or host anomaly, statuses that add up to their calls, ordered percentiles, one entry per variant per case, and the cell counts the "
      "protocol promised. Figures are pooled with the summarizer's own functions, so this report and the tables in docs/api-load.md are one computation. "
      "verify_report.py rebuilds the main tables and every chart dataset from the raw evidence without report_data.py's pooling, and checks the template "
      "parts, the captions and that no dash characters slipped in. [1, 4]")

    heading("Appendix A Every full-schedule quality cell")
    p("Every variant of the rerun at every fleet: 40 cases each, the gap against the uncapped search, the cases cheaper, dearer and with the reference's "
      "exact offers, the search p50 at the caller, and how shortened refinements stopped. [1]")
    variants = sorted({v for f in FLEETS for v in fq[f"{f}/full"]}, key=lambda v: (v != "1000", [int(x) for x in v.split(":")]))
    rows = [[f, v, q(f, v)["priced"] if v != "1000" else "", money(q(f, v)["meanCents"]) if v != "1000" else "", money(q(f, v)["medianCents"]) if v != "1000" else "",
             f"{q(f, v)['cheaper']} / {q(f, v)['dearer']}" if v != "1000" else "", q(f, v)["same"] if v != "1000" else "", num(q(f, v)["p50ms"]), stops(q(f, v))]
            for f in FLEETS for v in variants]
    table(["Techs", "Variant", "Priced", "Mean", "Median", "Cheaper / dearer", "Same set", "p50 ms", "Stopped by"], rows, [.45, .6, .5, .65, .65, .8, .55, .6, 1.4], size=9)

    heading("Appendix B Every one-desk run on half and full schedules")
    p("The load part's cells with refinement and the reference, pooled over 2 repeats: searches answered, the rate, caller latency in milliseconds, the "
      "scheduler's engine, routing and commit p50, refinements shed and searches incomplete. Empty and light schedules are in the evidence file. [1, 5]")
    load_variants = ("2:0", "2:250", "2:1000", "2:2000", "2:3000", "0:3000", "10:3000", "1000")
    rows = [[profile, f, v, L(f, v, profile)["searches"], num(L(f, v, profile)["perMinute"], 1),
             f"{num(L(f, v, profile)['p50'])} / {num(L(f, v, profile)['p95'])} / {num(L(f, v, profile)['p99'])}", num(L(f, v, profile)["engineP50"]),
             num(L(f, v, profile)["routingP50"]), num(L(f, v, profile)["commitP50"]), num(L(f, v, profile)["shed"]), L(f, v, profile)["incomplete"]]
            for profile in ("half", "full") for f in FLEETS for v in load_variants]
    table(["Schedule", "Techs", "Variant", "Searches", "/ min", "p50 / p95 / p99", "Engine", "Routing", "Commit", "Shed %", "Incomplete"], rows,
          [.6, .45, .6, .6, .5, 1.25, .55, .6, .55, .5, .7], size=9)

    heading("Appendix C Every capacity run")
    p("Every layout, client count, variant, fleet and schedule of the capacity part: searches a minute summed over the clients, the clients' mean p50 and "
      "p95 at the caller, the scheduler's median slot wait, refinements shed and searches answered 503. [1]")
    rows = []
    for f in (5, 10):
        for profile in ("half", "full"):
            for v in ("2:250", "2:1000", "2:3000", "1000"):
                for l in LAYOUTS:
                    for c in (5, 10):
                        r = C(v, l, c, f, profile)
                        rows.append([f, profile, v, l, c, num(r["perMinuteAll"]), num(r["p50"]), num(r["p95"]), num(r["queueP50"]), num(r["shed"]), r["unavailable"]])
    table(["Techs", "Schedule", "Variant", "Layout", "Clients", "/ min", "p50", "p95", "Slot wait", "Shed %", "503s"], rows,
          [.45, .6, .6, .55, .55, .55, .6, .6, .65, .55, .45], size=9)

    heading("Appendix D Every reserve run")
    p("The first study's 500 and 250 ms arms and the rerun's 800 ms arm at one desk: searches, the rate, commit p50 and p95, BUSY resends, 503s and "
      "searches incomplete. A run whose searches were all refused shows n/a where the scheduler logged no completed search. [1]")
    rows = []
    for source, label in ((first_rl, "first"), (rl, "rerun")):
        for key in sorted(source, key=lambda k: (k.split("/")[1], int(k.split("/")[0]), k.split("/")[2], -int(k.split("/")[3]))):
            f, profile, v, reserve = key.split("/")
            r = source[key]
            rows.append([label, f, profile, v, reserve, r["searches"], num(r["perMinute"], 1), f"{num(r['commitP50'])} / {num(r['commitP95'])}", r["busy"],
                         r["unavailable"], r["incomplete"]])
    table(["Study", "Techs", "Schedule", "Variant", "Reserve", "Searches", "/ min", "Commit p50 / p95", "BUSY", "503s", "Incomplete"], rows,
          [.55, .45, .6, .6, .6, .65, .5, 1.0, .5, .45, .7], size=9)

    heading("Appendix E Definitions and evidence")
    table(["Term", "Meaning in this report"], [
        ("First pass", "The insertion search: every technician-day on the horizon is tried for the new job. It always runs."),
        ("Refinement", "A local search after the first pass that may move existing visits to make room more cheaply."),
        ("Window threshold", "If the first pass finds no more distinct regular windows than this, the refinement is required: never shed, never capped."),
        ("Refinement cap", "The most time an optional refinement may run (bookingRefinementMs). 250 ms today; the API now allows up to 3000 ms."),
        ("threshold:cap", "How a variant is written. 2:250 is today's default; 2:0 is insertion only; 1000 makes every refinement required and uncapped."),
        ("CAP and DEADLINE", "What cut a SHORTENED refinement short: its own cap, or the deadline less the publish reserve (refinementStop)."),
        ("Shed", "An optional refinement skipped because a booking was waiting in the replica's queue; the search answers with the first pass's offers."),
        ("UNAVAILABLE", "A refinement that could not get routing or database capacity in time; the search answers with the first pass's offers."),
        ("Publish reserve", "The end of the deadline kept back to validate, prepare and publish the offer set; 1000 ms by default, 800 ms floor."),
        ("Admission guard", f"The check that refuses a database call with BUSY when less than {GUARD_MS} ms (connect plus validation timeout) is left."),
        ("Layout RxS", "R booking replicas with S search slots each, one vCPU per slot: 1x4 is one replica with 4 slots."),
        ("Dynamic horizon", "The dates searched: 2 lead weekdays past the last day over 85 percent full, at most 15 weekdays."),
        ("p50, p95, p99", "The latency that 50, 95 or 99 percent of searches in the run stayed within.")], [1.5, 5.4])
    p(f"[1] Study evidence: docs/evidence/{FIRST_FILE} (first study, runner {data['first']['revision'][:7]}, {data['first']['cells']:,} cells) and "
      f"docs/evidence/{RERUN_FILE} (full-schedule rerun, runner {data['rerun']['revision'][:7]}, {data['rerun']['cells']} cells). Each records the stack, "
      f"the booking pool shape, the protocol, the hypotheses and every kit report. docs/api-load.md, 'Booking configuration (C)', is the narrative kept "
      f"with the code.")
    p("[2] The study runner infra/run-booking-configuration-study.ps1 and the host conformance kit in infra/host-conformance (quality.ts, load.ts and "
      "src/load/serverlog.ts, which reads the scheduler's request log beside a run).")
    p("[3] The booking path: scheduler-service/src/main/java/dev/waterflex/scheduler/api/BookingOffers.java (the shed check, the cap's checkpoint and the "
      "neighborhood fetch inside it, the REFINEMENT_BUSY path), SearchDatabaseTelemetry and DatabasePool in scheduler-service, and SearchAdmission and "
      "SearchDeadline in calculation-engine.")
    p("[4] infra/summarize_booking_configuration.py and infra/test_summarize_booking_configuration.py: the pooling and pairing every figure uses.")
    p("[5] The request log line: scheduler-service RequestLogging, fields queueMs, routingMs, calculationMs, dbMs and commitMs; docs/operations-runbook.md, "
      "'Logs'.")
    p("[6] Commit 530d356: the publish reserve floor of 800 ms and the quality horizon read before optimizing.")
    p("[7] Refinement shedding was introduced in #155; docs/performance-work-2026-10-14 explains it and measures its first effect.")
    p("Reproduce with docs/booking-configuration-study-2026-10-19/requirements.txt, then run report_data.py, build_report.py and verify_report.py there. The "
      "builder reads the two evidence files, regenerates all eleven figures and chart_data.json, and writes this editable Word file. EVIDENCE.md records "
      "methods.")

    assert figures == list(FIGURES)
    settings = doc.settings.element
    for el in settings.findall(qn("w:updateFields")): settings.remove(el)
    update = OxmlElement("w:updateFields"); update.set(qn("w:val"), "true"); settings.append(update)
    stream = BytesIO(); doc.save(stream)
    with ZipFile(REFERENCE) as ref, ZipFile(stream) as built, ZipFile(OUTPUT, "w", ZIP_DEFLATED) as output:
        preserve = set(PRESERVE) | {n for n in ref.namelist() if "footer" in n}
        for name in sorted(built.namelist()):
            content = ref.read(name) if name in preserve else built.read(name)
            info = ZipInfo(name, (2026, 10, 19, 0, 0, 0)); info.compress_type = ZIP_DEFLATED
            output.writestr(info, content)
    print(OUTPUT)


if __name__ == "__main__": build()
