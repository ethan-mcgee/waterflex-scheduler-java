"""Reject malformed evidence and document drift at the report boundary."""
import tempfile
from pathlib import Path
import unittest
import copy
from docx import Document
from collect_retained_evidence import read
from verify_evidence import verify_hash, verify_parity, verify_campaign
from build_report import inline
from campaign_spec import verify_configuration
class EvidenceBoundaryTest(unittest.TestCase):
    def test_duplicate_fields_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'input.json';path.write_text('{"cost":1,"cost":2}')
            with self.assertRaises(ValueError):read(path)
    def test_nonfinite_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'input.json'
            for value in ('NaN','Infinity','-Infinity'):
                path.write_text('{"cost":'+value+'}')
                with self.assertRaises(ValueError):read(path)
    def test_changed_evidence_rejected(self):
        with self.assertRaises(AssertionError):verify_hash(b'changed',{'path':'input','sha256':'0'*64})
    def test_word_omission_and_reordering_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'report.docx';doc=Document();doc.add_paragraph('First');doc.add_paragraph('Second');doc.save(path)
            verify_parity('First\n\nSecond',path)
            for source in ('Second\nFirst','First\nSecond\nMissing'):
                with self.assertRaises(AssertionError):verify_parity(source,path)
    def test_word_link_target_mismatch_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'report.docx';doc=Document()
            inline(doc.add_paragraph(), '[Guide](https://example.com/wrong)');doc.save(path)
            with self.assertRaises(AssertionError):verify_parity('[Guide](https://example.com/right)',path)
    def test_campaign_arithmetic_rejects_wrong_counts(self):
        root=Path(__file__).resolve().parent
        design=read(root/'evidence/fourth-review/campaign-design.json')
        report=(root/'Timefold_Audit.md').read_text(encoding='utf-8')
        verify_campaign(design,report)
        design['stages'][0]['expected_cases']+=1
        with self.assertRaises(AssertionError):verify_campaign(design,report)
    def test_campaign_null_and_nonfinite_inputs_rejected(self):
        root=Path(__file__).resolve().parent
        valid=read(root/'evidence/fourth-review/campaign-design.json')
        report=(root/'Timefold_Audit.md').read_text(encoding='utf-8')
        for field,value in [('factors',None),('factors',[True]),('budgets_seconds',[None]),
                            ('budgets_seconds',[float('nan')]),('warmup_seconds',None),
                            ('parallel_cases',True),('fresh_jvms',1.5),('overhead_seconds_per_case',-1)]:
            design=copy.deepcopy(valid);design['stages'][0][field]=value
            with self.subTest(field=field,value=value):
                with self.assertRaises(ValueError):verify_campaign(design,report)
    def test_parameter_inventory_missing_null_unknown_and_nonfinite_fields_rejected(self):
        root=Path(__file__).resolve().parent
        valid=read(root/'evidence/fourth-review/campaign-example.json')
        self.assertEqual(verify_configuration(valid),[2,3,4,2,2])
        def objects(value, path=()):
            if isinstance(value,dict):
                yield path,value
                for key,item in value.items():yield from objects(item,path+(key,))
            elif isinstance(value,list):
                for index,item in enumerate(value):yield from objects(item,path+(index,))
        for path,obj in objects(valid):
            for key in obj:
                if key=='campaignCutoffUtc':continue  # Intentional nullable dispatch boundary.
                broken=copy.deepcopy(valid);target=broken
                for part in path:target=target[part]
                target[key]=None
                with self.subTest(path=path,key=key):
                    with self.assertRaises(ValueError):verify_configuration(broken)
            broken=copy.deepcopy(valid);target=broken
            for part in path:target=target[part]
            target['unknown']=1
            with self.assertRaises(ValueError):verify_configuration(broken)
            broken=copy.deepcopy(valid);target=broken
            for part in path:target=target[part]
            del target[next(iter(obj))]
            with self.assertRaises(ValueError):verify_configuration(broken)
        for value in (float('nan'),float('inf'),True,0):
            broken=copy.deepcopy(valid);broken['budgetCohorts'][0]['seconds']=value
            with self.assertRaises(ValueError):verify_configuration(broken)
    def test_community_and_resource_contract_rejects_incompatible_design(self):
        root=Path(__file__).resolve().parent
        valid=read(root/'evidence/fourth-review/campaign-example.json')
        for path,value in [(('edition',),'ENTERPRISE'),(('runtime','moveThreads'),'AUTO'),
                           (('runtime','nativeParallelBenchmarkCount'),2),
                           (('runtime','heapMinMiB'),2048),
                           (('resources','parallelCases'),3),
                           (('analysis','confidenceLevel'),1),
                           (('execution','campaignCutoffUtc'),'2026-10-06T20:00:00'),
                           (('solverSeeds',),[17,17])]:
            broken=copy.deepcopy(valid);target=broken
            for part in path[:-1]:target=target[part]
            target[path[-1]]=value
            with self.subTest(path=path):
                with self.assertRaises(ValueError):verify_configuration(broken)
    def test_worked_arithmetic_is_bound_to_json_parameters(self):
        root=Path(__file__).resolve().parent
        design=read(root/'evidence/fourth-review/campaign-design.json')
        config=read(root/'evidence/fourth-review/campaign-example.json')
        report=(root/'Timefold_Audit.md').read_text(encoding='utf-8')
        verify_campaign(design,report,config)
        config['processRepetitions']=3
        with self.assertRaises(AssertionError):verify_campaign(design,report,config)
        for field in ('expected_warmup_seconds','expected_parallel_seconds','expected_overhead_seconds'):
            broken=copy.deepcopy(design);broken['stages'][0][field]+=1
            with self.assertRaises(AssertionError):verify_campaign(broken,report)
if __name__=='__main__':unittest.main()
