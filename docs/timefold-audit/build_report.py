"""Build the Word companion from the canonical Markdown audit. Requires python-docx."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent


def inline(paragraph, text, size=None):
    lead = re.match(r'^(Current behavior|Current gap|Consequence|Guidance and change|Why this helps|Acceptance|Proposed integration|Configuration example|Verification|Proposed corpus|Proposed pilot|Worked example|Proposed analysis|Correctness gate|Quality gate|Performance gate|Implementation sequence|Example resolved campaign|Evidence|Impact|Recommendation|Priority|Classification|Status|Reference|References): ', text)
    if lead:
        paragraph.add_run(lead[0]).bold = True
        text = text[lead.end():]
    cursor = 0
    for match in re.finditer(r'\[([^\]]+)\]\(([^)]+)\)', text):
        paragraph.add_run(text[cursor:match.start()])
        link = OxmlElement('w:hyperlink')
        link.set(qn('r:id'), paragraph.part.relate_to(match[2], RT.HYPERLINK, is_external=True))
        run = OxmlElement('w:r')
        props = OxmlElement('w:rPr')
        color = OxmlElement('w:color'); color.set(qn('w:val'), '205781'); props.append(color)
        if size is not None:
            font_size = OxmlElement('w:sz'); font_size.set(qn('w:val'), str(int(size * 2))); props.append(font_size)
        run.append(props)
        value = OxmlElement('w:t'); value.text = match[1]; run.append(value)
        link.append(run); paragraph._p.append(link)
        cursor = match.end()
    paragraph.add_run(text[cursor:])


def table(doc, lines):
    rows = [[cell.strip() for cell in line.strip().strip('|').split('|')] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r':?-+:?', cell) for cell in row)]
    grid = doc.add_table(rows=0, cols=len(rows[0]))
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid.autofit = False
    # Allocate width by the actual table schema, not just its column count.
    widths_by_header = {
        'ID': [0.55, 0.7, 3.75, 2.0],
        'Workflow or variant': [1.65, 5.35],
        'Disposition': [1.9, 5.1],
        'Mechanism': [1.6, 5.4],
        'Priority and evidence': [1.65, 5.35],
        'Evidence': [1.75, 5.25],
        'Fleet and variant': [1.6, 2.7, 2.7],
        'Finding': [0.8, 6.2],
        'Concern': [1.25, 2.875, 2.875],
        'Check': [2.1, 4.9],
        'Source': [0.7, 6.3],
        'Decision': [1.5, 5.5],
        'Layer': [1.25, 2.75, 3.0],
        'Measurement': [1.7, 5.3],
        'Contract': [1.1, 5.9],
        'Stage': [1.4, 3.2, 2.4],
        'Term': [1.55, 5.45],
        'Local chapter and official source': [3.0, 4.0],
    }
    widths = widths_by_header.get(rows[0][0], [7 / len(rows[0])] * len(rows[0]))
    if len(widths) != len(rows[0]) or any(len(row) != len(widths) for row in rows):
        raise ValueError('Table schema does not match column widths')
    compact = rows[0][0] == 'ID'
    font_size = 9 if compact else 10
    for column, width in zip(grid.columns, widths): column.width = Inches(width)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        node = OxmlElement('w:' + edge)
        for name, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]: node.set(qn('w:' + name), value)
        borders.append(node)
    grid._tbl.tblPr.append(borders)
    for index, row in enumerate(rows):
        cells = grid.add_row().cells
        for column_index, (cell, value, width) in enumerate(zip(cells, row, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                padding = 70 if compact and edge in ('top', 'bottom') else 100
                item = OxmlElement('w:' + edge); item.set(qn('w:w'), str(padding)); item.set(qn('w:type'), 'dxa'); margins.append(item)
            props.append(margins)
            fill = OxmlElement('w:shd')
            fill.set(qn('w:fill'), 'DDEBF7' if index == 0 else ('F4F8FB' if index % 2 == 0 else 'FFFFFF'))
            props.append(fill)
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.line_spacing = 1.04
            p.paragraph_format.widow_control = True
            if (compact and column_index < 2) or (rows[0][0] == 'Finding' and column_index == 0):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if index == 0:
                p.paragraph_format.keep_with_next = True
            inline(p, value, size=font_size)
            for run in p.runs:
                run.font.size = Pt(font_size)
                run.bold = index == 0 or column_index == 0
                run.font.color.rgb = RGBColor(0, 0, 0)
        trpr = grid.rows[index]._tr.get_or_add_trPr()
        avoid_split = OxmlElement('w:cantSplit'); trpr.append(avoid_split)
        if index == 0:
            repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(4)
    gap.paragraph_format.line_spacing = Pt(3)


def main():
    source = (ROOT / 'Timefold_Audit.md').read_text(encoding='utf-8')
    assert '\u2014' not in source
    doc = Document()
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = Inches(.65); section.bottom_margin = Inches(.65)
    section.left_margin = Inches(.75); section.right_margin = Inches(.75)
    section.header_distance = Inches(.25); section.footer_distance = Inches(.3)
    section.different_first_page_header_footer = True
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'; normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.widow_control = True
    normal.font.color.rgb = RGBColor.from_string('222222')
    for name, size in [('Title', 26), ('Heading 1', 17), ('Heading 2', 13)]:
        style = doc.styles[name]; style.font.name = 'Calibri'; style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(16)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.keep_with_next = True
    for style_name in ('Header', 'Footer'):
        style = doc.styles[style_name]
        style.font.name = 'Calibri'; style.font.size = Pt(8)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.tab_stops.clear_all()
        style.paragraph_format.space_after = Pt(0)
    header = section.header.paragraphs[0]
    header.add_run('WATERFLEX').bold = True
    header.add_run('\tTimefold audit and benchmark design')
    header.paragraph_format.tab_stops.add_tab_stop(Inches(7), WD_TAB_ALIGNMENT.RIGHT)
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0, 0, 0)
    for footer_part in (section.footer, section.first_page_footer):
        footer = footer_part.paragraphs[0]
        footer.paragraph_format.tab_stops.add_tab_stop(Inches(7), WD_TAB_ALIGNMENT.RIGHT)
        footer.add_run('WaterFlex  |  Technical review\t')
        for instruction in ('PAGE', 'NUMPAGES'):
            if instruction == 'NUMPAGES': footer.add_run(' / ')
            field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), instruction); footer._p.append(field)
        for run in footer.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string('555555')
    doc.core_properties.title = 'WaterFlex Timefold Audit and Benchmark Design'
    doc.core_properties.subject = 'Explained architecture findings and benchmark redesign specification'
    doc.core_properties.author = 'WaterFlex'
    lines = source.splitlines(); i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith('|'):
            group = []
            while i < len(lines) and lines[i].startswith('|'): group.append(lines[i]); i += 1
            table(doc, group); continue
        if line.startswith('# '):
            title = doc.add_paragraph(style='Title')
            first, second = line[2:].split(' and ', 1)
            title.add_run(first + ' ').add_break()
            title.add_run('and ' + second)
            title.paragraph_format.line_spacing = 1.02
            title.paragraph_format.space_before = Pt(0)
            title.paragraph_format.space_after = Pt(10)
        elif line.startswith('## '):
            heading = doc.add_paragraph(line[3:], 'Heading 1')
            if line == '## Benchmark redesign':
                heading.paragraph_format.page_break_before = True
        elif line.startswith('### '): doc.add_paragraph(line[4:], 'Heading 2')
        elif line.strip():
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.keep_together = False
            inline(paragraph, line)
            if line.startswith('Priority '):
                paragraph.paragraph_format.keep_with_next = True
            if line.startswith('Prepared '):
                paragraph.paragraph_format.space_after = Pt(10)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor.from_string('555555')
        i += 1
    output = ROOT / 'Timefold_Audit.docx'
    doc.save(output)
    # Verify every Markdown line, including table cells and hyperlink labels, survived authoring.
    from lxml import etree
    from zipfile import ZipFile
    with ZipFile(output) as archive:
        xml = etree.fromstring(archive.read('word/document.xml'))
    text = ''.join(xml.itertext())
    for line in lines:
        if not line or re.fullmatch(r'[| :\-]+', line): continue
        cells = line.strip('|').split('|') if line.startswith('|') else [re.sub(r'^#+ ', '', line)]
        for cell in cells:
            plain = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', cell.strip())
            assert plain in text, plain
    print(f'Created {output}; Markdown content verified')


if __name__ == '__main__': main()
