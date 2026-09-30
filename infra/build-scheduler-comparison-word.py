"""Create the comparison DOCX from its verified PDF content and retained explainer.

Requires python-docx 1.2.0, pypdf and Pillow. The source PDF, Markdown,
reference DOCX and PNG figures are read-only. Render the output separately.
"""
from copy import deepcopy
from datetime import datetime
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import re
from zipfile import ZipFile, ZIP_DEFLATED, ZipInfo

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT / "docs/scheduler-explainer/reference.docx"
REFERENCE_SHA = "2d296b4ae319f04ab583a25a23df6e11e845e6d1ba3324f7dcad23d96d53df70"
SOURCE = ROOT / "docs/scheduler-algorithm-comparison-2026-09-29.md"
PDF = SOURCE.with_suffix(".pdf")
OUTPUT = ROOT / "docs/WaterFlex_Scheduler_Algorithm_Comparison_Explained.docx"
REVISION = "001ea09d85a489f90e1304d41f9a8b43410ad512"
BASE_URL = f"https://github.com/ethan-mcgee/waterflex-scheduler-java/blob/{REVISION}/docs/"

CHAPTERS = {
    "Experiment overview": "1 What was tested",
    "Mechanisms in plain language": "2 How the search methods work",
    "Booking comparison": "3 What booking results establish",
    "Held-out service by operating condition": "4 How operating conditions change booking",
    "Latency and matched-customer cost": "5 How to read latency and cost",
    "Search quality and booking implications": "6 What visible options tell us",
    "Daily-solver comparison": "7 How daily solvers compare",
    "Fleet size changes the daily conclusion": "8 Why fleet size and time matter",
    "Individual seeds and paired wins": "9 What the individual seeds show",
    "Daily implications": "10 What the daily results mean",
    "Recommendations and follow-up": "11 What should happen next",
    "Technical appendix: metric glossary": "Appendix A Measurement definitions",
    "Technical appendix: complete daily fleet comparison": "Appendix B Daily savings by fleet size",
    "Technical appendix: accepted schedule diagnostics": "Appendix C Fairness waiting and schedule changes",
    "Technical appendix: exploratory intervals": "Appendix D Exploratory uncertainty intervals",
    "Technical appendix: evidence and reproducibility": "Appendix E Evidence and sources",
}

INTRODUCTIONS = {
    "Mechanisms in plain language": [
        "Imagine rearranging furniture in a crowded room. Moving one chair can make space for a table, even if the first move alone does not improve the room. Scheduling search has a similar challenge: a useful rearrangement may require several coordinated changes. This analogy explains search behavior; it is not an additional benchmark result.",
        "A move describes a possible change to the schedule. A search strategy decides which changes to explore and how long to keep looking. Relocation and swap are moves. Tabu and late acceptance are strategies for exploring them. A larger collection of moves gives the search more possibilities, but evaluating those possibilities also takes time.",
    ],
    "Booking comparison": [
        "The first booking question is whether the attempt ended with confirmed service. The second is how long the customer waited for the search. Cost is then compared only where the variants served the same customers. This order keeps a cheaper but less successful search from appearing automatically better.",
    ],
    "Individual seeds and paired wins": [
        "A seed fixes the experiment's repeatable random choices. Showing each retained seed separately helps reveal whether an average gain appears in both repeats or is carried mainly by one. A win counts a cheaper accepted case; the size of the saving is a separate question. Several small wins can be outweighed by a few large losses.",
    ],
}


def compact(text):
    return re.sub(r"\s+", "", text).replace("\u200b", "")


def read_blocks():
    """Parse only the small, explicit format emitted by the PDF report builder."""
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    blocks = []
    index = 0
    while index < len(lines):
        line = lines[index]
        index += 1
        if not line or line.startswith("# "):
            continue
        if line.startswith("## "):
            blocks.append(("heading", line[3:]))
        elif line.startswith("| "):
            rows = [line.strip("| ").split(" | ")]
            while index < len(lines) and lines[index].startswith("| "):
                rows.append(lines[index].strip("| ").split(" | "))
                index += 1
            assert all(cell == "---" for cell in rows[1])
            assert all(len(row) == len(rows[0]) for row in rows)
            blocks.append(("table", rows[0], rows[2:]))
        elif line.startswith("!["):
            match = re.fullmatch(r"!\[(.*)\]\(([^)]+)\)", line)
            assert match is not None
            caption, target = match.groups()
            path = SOURCE.parent / target
            assert path.is_file()
            blocks.append(("figure", path, caption))
            # The source repeats the caption as accessible Markdown prose.
            while index < len(lines) and not lines[index]:
                index += 1
            assert lines[index] == caption
            index += 1
        elif line.startswith("- ["):
            match = re.fullmatch(r"- \[([^]]+)\]\(([^)]+)\)", line)
            assert match is not None
            title, target = match.groups()
            assert (SOURCE.parent / target).is_file()
            blocks.append(("link", title, target))
        else:
            blocks.append(("paragraph", line))
    pdf_text = compact(" ".join(page.extract_text() for page in PdfReader(PDF).pages))
    for block in blocks:
        if block[0] in ("heading", "paragraph"):
            assert compact(block[1]) in pdf_text, f"PDF content mismatch: {block[1]}"
        elif block[0] == "table":
            for row in block[2]:
                assert compact("".join(row)) in pdf_text, f"PDF table mismatch: {row}"
        elif block[0] == "figure":
            assert compact(block[2]) in pdf_text
    # Keep the compact worked fixture beside the experiment overview.
    start = next(i for i, b in enumerate(blocks) if b[0] == "heading" and b[1].startswith("Illustrative fixture evidence"))
    end = next(i for i in range(start + 1, len(blocks)) if blocks[i][0] == "heading")
    fixture = blocks[start:end]
    del blocks[start:end]
    target = next(i for i, b in enumerate(blocks) if b[0] == "heading" and b[1] == "Mechanisms in plain language")
    blocks[target:target] = fixture
    sequential = next(i for i, b in enumerate(blocks) if b[0] == "paragraph" and b[1].startswith("At concurrency 1,"))
    note = blocks.pop(sequential)
    target = next(i for i, b in enumerate(blocks) if b[0] == "heading" and b[1] == "Held-out service by operating condition")
    blocks.insert(target, note)
    return blocks


def table_widths(headers):
    if len(headers) == 2:
        return [1.85, 5.05]
    if len(headers) == 3:
        return [1.65, 2.55, 2.70]
    if len(headers) == 4:
        return [2.05, 1.62, 1.62, 1.61]
    if len(headers) == 5 and headers[0] == "Stage":
        return [.65, 1.60, 1.70, 1.05, 1.90]
    if len(headers) == 5 and headers[0] in ("Concurrency", "Fleet size", "Cache condition"):
        return [1.05, 1.85, 1.30, 1.30, 1.40]
    if len(headers) == 5:
        return [2.05, 1.20, 1.20, 1.20, 1.25]
    if len(headers) == 6:
        if headers[0] == "Held variant":
            return [1.20, 1.23, 1.10, 1.05, 1.20, 1.12]
        return [.65, 1.55, .75, 1.32, 1.32, 1.31]
    if len(headers) == 7:
        return [1.98, .85, .72, .93, .85, .80, .77]
    raise ValueError(f"Unsupported table: {headers}")


def build(blocks):
    assert sha256(REFERENCE.read_bytes()).hexdigest() == REFERENCE_SHA
    doc = Document(REFERENCE)
    component = deepcopy(doc.tables[0]._tbl)
    for child in list(doc.element.body):
        if child.tag != qn("w:sectPr"):
            doc.element.body.remove(child)

    def paragraph(text, style="Normal"):
        p = doc.add_paragraph(text, style)
        p.paragraph_format.widow_control = True
        return p

    def bookmark(p, name, number):
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:name"), name)
        start.set(qn("w:id"), str(number))
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(number))
        p._p.insert(1, start)
        p._p.append(end)

    def hyperlink(p, title, url=None, anchor=None):
        link = OxmlElement("w:hyperlink")
        if anchor:
            link.set(qn("w:anchor"), anchor)
        else:
            rel = doc.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
            link.set(qn("r:id"), rel)
        run = OxmlElement("w:r")
        props = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "245A81")
        props.append(color)
        run.append(props)
        text = OxmlElement("w:t")
        text.text = title
        run.append(text)
        link.append(run)
        p._p.append(link)

    def table(headers, rows, widths=None):
        widths = widths or table_widths(headers)
        assert abs(sum(widths) - 6.9) < .001
        result = doc.add_table(rows=1, cols=len(headers))
        result._tbl.replace(result._tbl.tblPr, deepcopy(component.find(qn("w:tblPr"))))
        result.autofit = False
        for col, width in zip(result.columns, widths):
            col.width = Inches(width)
        for _ in rows:
            result.add_row()
        source_rows = component.findall(qn("w:tr"))
        for index, (row, values) in enumerate(zip(result.rows, [headers] + rows)):
            sample = source_rows[0 if index == 0 else 1 if index % 2 else 2].find(qn("w:tc"))
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            if index == 0:
                row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for col, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
                cell._tc.replace(cell._tc.get_or_add_tcPr(), deepcopy(sample.find(qn("w:tcPr"))))
                cell.width = Inches(width)
                p = cell.paragraphs[0]
                p._p.insert(0, deepcopy(sample.find(qn("w:p")).find(qn("w:pPr"))))
                p.paragraph_format.keep_with_next = index == 0
                run = p.add_run(str(value).replace("_", "_\u200b"))
                run._r.insert(0, deepcopy(sample.find(qn("w:p")).find(qn("w:r")).find(qn("w:rPr"))))
                if len(headers) >= 4:
                    run.font.size = Pt(9)
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.05
                    end_props = OxmlElement("w:rPr")
                    end_size = OxmlElement("w:sz")
                    end_size.set(qn("w:val"), "18")
                    end_props.append(end_size)
                    p._p.get_or_add_pPr().append(end_props)
                    margins = cell._tc.get_or_add_tcPr().find(qn("w:tcMar"))
                    for edge in ("top", "bottom"):
                        margins.find(qn("w:" + edge)).set(qn("w:w"), "40")
                if col > 0 and re.fullmatch(r"[\d$%+., /()\[\]-]+", str(value)):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        spacer = paragraph("")
        spacer.paragraph_format.space_after = Pt(2)
        spacer.paragraph_format.space_before = Pt(0)
        spacer.paragraph_format.line_spacing = Pt(3)
        return result

    paragraph("WaterFlex scheduler algorithm comparison explained", "Title")
    paragraph("A practical guide to the results and the next decisions", "Subtitle")
    paragraph("September 29 2026")
    paragraph("WaterFlex uses different searches to fit a new customer into existing commitments and to improve an already booked day. This report explains what the retained comparison says about those choices, why the results differ, and which follow-up tests could justify a change.")

    chapter_index = 0
    current = None
    for block in blocks:
        kind = block[0]
        if kind == "heading":
            current = block[1]
            if current == "Decision summary":
                paragraph("What the comparison tells us", "Heading 2")
            elif current in CHAPTERS:
                if chapter_index == 0:
                    paragraph("How to read this report", "Heading 2")
                    nav = table(["Sections", "What you will learn"], [
                        ["1 and 2", "What was tested and how the search methods work"],
                        ["3 to 6", "Booking service, operating conditions, latency and comparable cost"],
                        ["7 to 10", "Daily solvers, fleet size, time budgets and individual seed results"],
                        ["11", "The decisions and additional tests to prioritize"],
                        ["Appendices", "Metric definitions, full measurements, uncertainty and sources"],
                    ], [1.05, 5.85])
                    for row, anchor in zip(nav.rows[1:], (1, 3, 7, 11, 12)):
                        p = row.cells[1].paragraphs[0]
                        title = p.text
                        p.clear()
                        hyperlink(p, title, anchor=f"chapter_{anchor}")
                chapter_index += 1
                p = paragraph(CHAPTERS[current], "Heading 1")
                p.paragraph_format.page_break_before = True
                bookmark(p, f"chapter_{chapter_index}", chapter_index)
                for text in INTRODUCTIONS.get(current, []):
                    paragraph(text)
            else:
                title = re.sub(r"[^A-Za-z0-9 ]", " ", current)
                paragraph(re.sub(r" +", " ", title), "Heading 2")
        elif kind == "paragraph":
            text = block[1]
            p = paragraph(text)
            if current in ("Daily implications", "Search quality and booking implications") and ":" in text:
                lead, rest = text.split(":", 1)
                if re.fullmatch(r"[A-Z_]+", lead):
                    p.clear()
                    p.add_run(lead + ":").bold = True
                    p.add_run(rest)
        elif kind == "table":
            table(block[1], block[2])
        elif kind == "figure":
            p = paragraph("")
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(4)
            width = 6.0 if block[1].name == "booking-quality.png" else 6.9
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            shape = p.add_run().add_picture(str(block[1]), width=Inches(width))
            shape._inline.docPr.set("descr", block[2])
            caption = paragraph(block[2])
            caption.paragraph_format.keep_together = True
            for run in caption.runs:
                run.font.size = Pt(10)
                run.font.italic = True
                run.font.color.rgb = RGBColor(0, 0, 0)
        elif kind == "link":
            hyperlink(paragraph(""), block[1], BASE_URL + block[2])
        else:
            raise ValueError(kind)

    paragraph("Comparison report", "Heading 2")
    paragraph("The PDF and its companion Markdown provide the retained numerical comparison. The Word document carries the same measurements, figures and limitations with the explanatory structure used in the WaterFlex scheduler guide. Linked project sources are pinned to the comparison revision.")
    hyperlink(paragraph(""), "Scheduler algorithm comparison PDF", BASE_URL + PDF.name)
    hyperlink(paragraph(""), "Scheduler algorithm comparison Markdown", BASE_URL + SOURCE.name)
    hyperlink(paragraph(""), "WaterFlex scheduler and route optimization explained", BASE_URL + "scheduler-explainer/WaterFlex_Scheduler_and_Route_Optimization_Explained_Expanded.docx")
    paragraph("To rebuild this Word edition, run infra/build-scheduler-comparison-word.py with Python, python-docx 1.2.0 and pypdf installed. Add --check to verify the existing document against the retained PDF, tables, figures and reference template.")
    doc.core_properties.title = "WaterFlex scheduler algorithm comparison explained"
    doc.core_properties.subject = "Booking and daily optimization evidence from September 29 2026"
    doc.core_properties.author = "WaterFlex"
    doc.core_properties.last_modified_by = "WaterFlex"
    doc.core_properties.created = datetime(2026, 9, 29)
    doc.core_properties.modified = datetime(2026, 9, 29)
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")
    buffer = BytesIO()
    doc.save(buffer)
    editable = {"word/document.xml", "word/_rels/document.xml.rels", "[Content_Types].xml", "docProps/core.xml", "word/settings.xml"}
    with ZipFile(REFERENCE) as original, ZipFile(buffer) as generated, ZipFile(OUTPUT, "w", ZIP_DEFLATED) as output:
        for name in sorted(set(original.namelist()) | set(generated.namelist())):
            source = generated if name in editable or name not in original.namelist() else original
            entry = ZipInfo(name, (2026, 9, 29, 0, 0, 0))
            entry.compress_type = ZIP_DEFLATED
            output.writestr(entry, source.read(name))
    verify(blocks)


def verify(blocks):
    assert sha256(REFERENCE.read_bytes()).hexdigest() == REFERENCE_SHA
    doc = Document(OUTPUT)
    text = " ".join(doc.element.itertext())
    assert "\u2014" not in text
    narrative = compact(" ".join(p.text for p in doc.paragraphs))
    for block in blocks:
        if block[0] == "paragraph":
            assert compact(block[1]) in narrative, f"Missing source paragraph: {block[1]}"
        elif block[0] == "figure":
            assert compact(block[2]) in narrative, f"Missing caption: {block[2]}"
    # Every source numeric table remains editable Word content in original order.
    expected = [block for block in blocks if block[0] == "table"]
    assert len(doc.tables) == len(expected) + 1
    for table, block in zip(doc.tables[1:], expected):
        actual = [[compact(cell.text) for cell in row.cells] for row in table.rows]
        wanted = [[compact(cell) for cell in row] for row in [block[1]] + block[2]]
        assert actual == wanted, block[1]
    assert len(doc.inline_shapes) == 4
    assert len(doc.sections) == 1
    with ZipFile(REFERENCE) as reference, ZipFile(OUTPUT) as result:
        editable = {"word/document.xml", "word/_rels/document.xml.rels", "[Content_Types].xml", "docProps/core.xml", "word/settings.xml"}
        for name in reference.namelist():
            if name not in editable:
                assert reference.read(name) == result.read(name), name
        for block in blocks:
            if block[0] == "figure":
                assert block[1].read_bytes() in [result.read(n) for n in result.namelist() if n.startswith("word/media/")]
    section_before = Document(REFERENCE).sections[0]._sectPr
    assert doc.sections[0]._sectPr.xml == section_before.xml
    names = {node.get(qn("w:name")) for node in doc.element.iter(qn("w:bookmarkStart"))}
    for link in doc.element.iter(qn("w:hyperlink")):
        anchor = link.get(qn("w:anchor"))
        assert anchor is None or anchor in names
    print(f"Verified {len(expected)} editable tables, 4 original figures, template parts, navigation and source PDF content")


if __name__ == "__main__":
    import sys
    blocks = read_blocks()
    if "--check" in sys.argv:
        verify(blocks)
    else:
        build(blocks)
        print(OUTPUT)
